"""A material or equipment name without the protocols.io list labels (lane FX, decision 4).

Human decision of 2026-10-07: in the report's materials and equipment
tables a name is shown without what the protocols.io list writes around it
-- "NAME …", "TYPE …", "BRAND …", "SKU" and a long number code at its end.
Display only: the name as listed stays as it was. The steps an item is used
in are looked for by the name so shown.

Measured on the in-gel report (lane N): "Eppendorf Thermomixer C Model 5382
NAME Thermomixer C TYPE Eppendorf BRAND 5382000023 SKU", used in "—".
"""

from __future__ import annotations

import io
import unittest
from pathlib import Path
from types import SimpleNamespace

from docx import Document

from tests.test_lab_report import _ReportCase, _step
from tests.test_lab_report_polish import items_double
from voiney_lab import experiment_reports as er

ROOT = Path(__file__).resolve().parents[1]
INGEL_PDF = ROOT / "data" / "runtime" / "candidate-a-source" / "in-gel-digestion.pdf"
#: As the in-gel analysis lists it (data/fixtures/development_protocols).
THERMOMIXER = (
    "Eppendorf Thermomixer C Model 5382\nNAME\nThermomixer C TYPE\nEppendorf\nBRAND\n"
    "5382000023\nSKU")
THERMOMIXER_LINE = " ".join(THERMOMIXER.split())


class DisplayNameTests(unittest.TestCase):
    def test_the_list_labels_and_what_follows_them_are_left_out(self) -> None:
        for listed, shown in (
            (THERMOMIXER, "Eppendorf Thermomixer C Model 5382"),
            (THERMOMIXER_LINE, "Eppendorf Thermomixer C Model 5382"),
            ("Thermomixer C Model 5382\nNAME\nThermomixer C", "Thermomixer C Model 5382"),
            ("Centrifuge 5424 R TYPE Eppendorf BRAND 5424000410 SKU", "Centrifuge 5424 R"),
            ("Vortex-Genie 2 BRAND Scientific Industries", "Vortex-Genie 2"),
            ("Falcon 50 mL conical tube 352070", "Falcon 50 mL conical tube"),
            ("Heat sealer SKU 1234567", "Heat sealer"),
        ):
            with self.subTest(listed=listed):
                self.assertEqual(er.item_display_name(listed), shown)

    def test_a_name_without_them_is_shown_as_listed(self) -> None:
        for listed in (
            "Acetonitrile LC-MS grade B&J Brand VWR International (Avantor) Catalog #BJLC015-2.5",
            "Formic acid, LC-MS grade Thermo Fisher Scientific Catalog #28905",
            "Promega trypsin Promega Catalog #V5113",
            "Lysogeny Broth (LB) (Sigma-Aldrich)",
            "Centrifuge 5424 R",
            "Eppendorf Thermomixer C Model 5382",
            "TYPE II collagenase",
            "modified heating oven",
        ):
            with self.subTest(listed=listed):
                self.assertEqual(er.item_display_name(listed), " ".join(listed.split()))


def tagged_double() -> SimpleNamespace:
    double = items_double()
    double.draft.protocol.sections = (SimpleNamespace(title_source_text="All", steps=(
        _step("t-step-01", "1", "Place the gel pieces in Eppendorf tubes."),
        _step("t-step-02", "2", "Shake at 800 rpm in the Eppendorf Thermomixer C Model 5382."),
        _step("t-step-03", "3", "Spin in the Centrifuge 5424 R."),
    )),)
    double.draft.protocol.materials = (SimpleNamespace(name_source_text="Eppendorf tubes"),)
    double.draft.protocol.equipment = tuple(SimpleNamespace(name_source_text=name) for name in (
        THERMOMIXER_LINE,
        "Centrifuge 5424 R\nNAME\nCentrifuge TYPE\nEppendorf\nBRAND\n5424000410\nSKU",
    ))
    return double


class ReportTableTests(_ReportCase):
    def setUp(self) -> None:
        super().setUp()
        self.fixture = tagged_double()
        self.event("00:30", "session_started", "1")
        self.event("00:36", "step_completed", "1")

    def facts(self) -> er.ReportFacts:
        return er.ReportWriterBrain.facts_for(self.doc(), fixture=self.fixture)

    def test_the_item_keeps_the_name_as_listed(self) -> None:
        facts = self.facts()
        thermomixer = facts.items[1]
        self.assertEqual(thermomixer.name, "Eppendorf Thermomixer C Model 5382")
        self.assertEqual(thermomixer.listed, THERMOMIXER_LINE)
        self.assertEqual(facts.equipment[0], THERMOMIXER_LINE)
        self.assertEqual(self.fixture.draft.protocol.equipment[0].name_source_text, THERMOMIXER_LINE)

    def test_the_steps_are_found_by_the_name_shown(self) -> None:
        # With the labels the name is cut down to "Eppendorf", which the
        # tubes also carry, and nothing was found.
        steps = {item.name: er._label_runs(item.steps) or "—" for item in self.facts().items}
        self.assertEqual(steps, {
            "Eppendorf tubes": "1",
            "Eppendorf Thermomixer C Model 5382": "2",
            "Centrifuge 5424 R": "3",
        })

    def test_markdown_and_word_show_the_name_and_say_what_was_left_out(self) -> None:
        text = self.store.export_markdown(self.report_id, fixture=self.fixture).decode()
        block = text.split("### 3-1. 재료와 장비")[1].split("### 3-2.")[0]
        self.assertIn("| Eppendorf Thermomixer C Model 5382 | 2 |", block)
        self.assertIn("| Centrifuge 5424 R | 3 |", block)
        for label in ("NAME", "TYPE", "BRAND", "SKU", "5382000023"):
            self.assertNotIn(label, block.split("> ")[0])
        self.assertIn(er.ITEM_LABELS_NOTE, block)
        document = Document(io.BytesIO(self.store.export_docx(self.report_id, fixture=self.fixture)))
        tables = [[[cell.text for cell in row.cells] for row in table.rows]
                  for table in document.tables]
        rows = [row for table in tables if tuple(table[0]) == er.ITEM_HEADER for row in table[1:]]
        self.assertIn(["Eppendorf Thermomixer C Model 5382", "2"], rows)
        self.assertIn(er.ITEM_LABELS_NOTE, [p.text for p in document.paragraphs])

    def test_no_note_when_no_name_carried_labels(self) -> None:
        self.fixture.draft.protocol.equipment = (SimpleNamespace(name_source_text="Centrifuge 5424 R"),)
        text = self.store.export_markdown(self.report_id, fixture=self.fixture).decode()
        self.assertNotIn(er.ITEM_LABELS_NOTE, text)


@unittest.skipUnless(INGEL_PDF.is_file(), "licensed in-gel PDF absent (condition B)")
class InGelThermomixerTests(unittest.TestCase):
    """The in-gel fixture: the name is cleaned; its use is still not found."""

    def test_the_thermomixer(self) -> None:
        from voiney_lab.curated_protocol import load_curated_protocol_fixture

        fixtures = ROOT / "data" / "fixtures" / "development_protocols"
        fixture = load_curated_protocol_fixture(
            fixtures / "candidate_a_curated_analysis.json",
            fixtures / "candidate_a_curated_analysis.provenance.json",
            INGEL_PDF)
        protocol = fixture.draft.protocol
        self.assertEqual(protocol.equipment[0].name_source_text, THERMOMIXER)
        items = er.item_step_labels(
            [("재료", str(m.name_source_text)) for m in protocol.materials]
            + [("장비", str(e.name_source_text)) for e in protocol.equipment],
            er._source_steps_from_fixture(fixture))
        steps = {item.name: er._label_runs(item.steps) or "—" for item in items}
        self.assertEqual(steps, {
            "Acetonitrile LC-MS grade B&J Brand VWR International (Avantor) Catalog #BJLC015-2.5":
                "2, 8–9, 19–20, 24",
            "Ammonium bicarbonate Merck MilliporeSigma (Sigma-Aldrich) Catalog #A6141":
                "1–2, 5, 10, 17, 21",
            "Formic acid, LC-MS grade Thermo Fisher Scientific Catalog #28905": "24",
            "Promega trypsin Promega Catalog #V5113": "21–22",
            "DTT Merck MilliporeSigma (Sigma-Aldrich) Catalog #D0632": "10–11",
            "Iodoacetamide Merck MilliporeSigma (Sigma-Aldrich) Catalog #I1149-5G": "10, 15",
            # Only a note under step 3 names it ("You can use an Eppendorf
            # Thermomixer for the incubation steps at 37C"); the lookup reads
            # the steps' text, not their notes (lane RP2: "—" was right).
            "Eppendorf Thermomixer C Model 5382": "—",
        })


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
