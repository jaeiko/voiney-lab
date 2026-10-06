"""The experiment report, polished (lane RP2, decisions of 2026-10-06).

A person read a headspace report in Word and found long gaps between lines,
materials and equipment as bare names in one list, and nothing on a later
start or a return within a repeat. Built on lane RP's protocol double and
ledger helpers (tests/test_lab_report.py); provider calls are a fake client.
"""

from __future__ import annotations

import asyncio
import io
import json
from types import SimpleNamespace
from unittest import mock

from docx import Document
from docx.shared import Pt

from tests.test_lab_report import GOOD_REPLY, _Client, _ReportCase, _step, protocol_double
from voiney_lab import experiment_reports as er


class WordSpacingTests(_ReportCase):
    """Decision 1: tighter body, list and table spacing; room around headings."""

    def document(self):
        self.run_steps_one_to_four_then_stop()
        return Document(io.BytesIO(self.store.export_docx(self.report_id, fixture=self.fixture)))

    def test_the_decided_values(self) -> None:
        # Line spacing about 1.15, 3-4 pt after a paragraph, 0-2 pt between
        # list items; headings keep more room before them than after.
        self.assertEqual(er.DOCX_SPACING["p"], (0, 4, 1.15))
        self.assertEqual(er.DOCX_SPACING["list"], (0, 1, 1.15))
        self.assertEqual(er.DOCX_SPACING["table"], (0, 0, 1.0))
        self.assertEqual(er.DOCX_SPACING["note"], (2, 4, 1.15))
        self.assertEqual(er.DOCX_SPACING["h1"], (12, 4, 1.0))
        self.assertEqual(er.DOCX_SPACING["h2"], (8, 3, 1.0))
        self.assertEqual(er.DOCX_SPACING["title"], (0, 8, 1.0))

    def test_word_paragraphs_carry_them(self) -> None:
        document = self.document()
        normal = document.styles["Normal"].paragraph_format
        self.assertEqual((normal.space_after, normal.line_spacing), (Pt(4), 1.15))
        paragraphs = document.paragraphs
        texts = [p.text for p in paragraphs]

        def spacing(paragraph):
            form = paragraph.paragraph_format
            return form.space_before, form.space_after, form.line_spacing

        self.assertEqual(spacing(paragraphs[texts.index("1. 실험 목적")]), (Pt(12), Pt(4), 1.0))
        self.assertEqual(spacing(paragraphs[texts.index("3-2. 수행한 단계")]), (Pt(8), Pt(3), 1.0))
        body = paragraphs[texts.index("1. 실험 목적") + 1]
        self.assertEqual(spacing(body), (Pt(0), Pt(4), 1.15))
        # 5. 고찰 (나): one list; its items are close, its last item is not.
        start = texts.index("(나) 확인이 필요한 점") + 1
        items = []
        while texts[start + len(items)].startswith("• "):
            items.append(paragraphs[start + len(items)])
        self.assertGreaterEqual(len(items), 2)
        self.assertEqual({spacing(p) for p in items[:-1]}, {(Pt(0), Pt(1), 1.15)})
        self.assertEqual(spacing(items[-1]), (Pt(0), Pt(4), 1.15))
        cells = {spacing(p) for table in document.tables for row in table.rows
                 for cell in row.cells for p in cell.paragraphs}
        self.assertEqual(cells, {(Pt(0), Pt(0), 1.0)})

    def test_no_paragraph_falls_back_to_the_templates_ten_points(self) -> None:
        document = self.document()
        paragraphs = list(document.paragraphs) + [
            p for table in document.tables for row in table.rows for cell in row.cells
            for p in cell.paragraphs]
        for paragraph in paragraphs:
            self.assertIsNotNone(paragraph.paragraph_format.space_after, paragraph.text)
            self.assertLessEqual(paragraph.paragraph_format.space_after, Pt(8), paragraph.text)


def items_double() -> SimpleNamespace:
    """Names as protocols list them, steps that call them by name or short name."""

    steps = (
        _step("i-step-01", "1", "Place the gel pieces in 200 µL 25mM AMBIC."),
        _step("i-step-02", "2", "Make solution A from 25mM ammonium bicarbonate (AMBIC) and acetonitrile."),
        _step("i-step-03", "3", "Add trypsin made up in AMBIC."),
        _step("i-step-04", "4", "Add formic acid (FA, 10% v/v) to reach 1%."),
        _step("i-step-05", "5", "Make 200 mL of LB broth."),
        _step("i-step-06", "6", "Pour the LB agar into the plates."),
        _step("i-step-07", "7", "Fill two tubes with autoclaved LB."),
        _step("i-step-08", "8", "Use a glass Pasteur pipette and a small ferrule."),
        _step("i-step-09", "9", "Dry in the modified heating oven, then in the modified heating block."),
        _step("i-step-10", "10", "Mix with a pipette."),
    )
    double = protocol_double()
    double.draft.protocol.sections = (SimpleNamespace(title_source_text="All", steps=steps),)
    double.draft.protocol.materials = tuple(SimpleNamespace(name_source_text=name) for name in (
        "Acetonitrile LC-MS grade B&J Brand VWR International (Avantor) Catalog #BJLC015-2.5",
        "Ammonium bicarbonate Merck MilliporeSigma (Sigma-Aldrich) Catalog #A6141",
        "Promega trypsin Promega Catalog #V5113",
        "Formic acid, LC-MS grade Thermo Fisher Scientific Catalog #28905",
        "Lysogeny Broth (LB) (Sigma-Aldrich)",
        "Lysogeny Agar (Sigma-Aldrich)",
        "Glass pipettes (10 mL)",
        "Glass Pasteur pipette (150 mm)",
    ))
    double.draft.protocol.equipment = tuple(SimpleNamespace(name_source_text=name) for name in (
        "Thermomixer C Model 5382\nNAME\nThermomixer C",
        "modified heating oven",
        "modified heating block",
        "Ferrules",
    ))
    double.timer_manifest = {}
    double.localized_fact = lambda step_id, fact_id: None
    return double


class MaterialsAndEquipmentTablesTests(_ReportCase):
    """Decisions 2, 3 and 5: two tables, name as listed and the steps that name it."""

    def setUp(self) -> None:
        super().setUp()
        self.fixture = items_double()
        self.event("00:30", "session_started", "1")
        self.event("00:36", "step_completed", "1")

    def facts(self) -> er.ReportFacts:
        return er.ReportWriterBrain.facts_for(self.doc(), fixture=self.fixture)

    def test_each_item_lists_the_source_steps_that_name_it(self) -> None:
        steps = {item.name: er._label_runs(item.steps) or "—" for item in self.facts().items}
        self.assertEqual(steps, {
            # The name before its company and grade words.
            "Acetonitrile LC-MS grade B&J Brand VWR International (Avantor) Catalog #BJLC015-2.5": "2",
            # The short name the source gives it, wherever it is used.
            "Ammonium bicarbonate Merck MilliporeSigma (Sigma-Aldrich) Catalog #A6141": "1–3",
            "Promega trypsin Promega Catalog #V5113": "3",
            "Formic acid, LC-MS grade Thermo Fisher Scientific Catalog #28905": "4",
            # "LB" is the broth's short name; "LB agar" is the agar.
            "Lysogeny Broth (LB) (Sigma-Aldrich)": "5, 7",
            "Lysogeny Agar (Sigma-Aldrich)": "6",
            # Not "pipette" alone: other listed items are pipettes too.
            "Glass pipettes (10 mL)": "—",
            "Glass Pasteur pipette (150 mm)": "8",
            "Thermomixer C Model 5382 NAME Thermomixer C": "—",
            "modified heating oven": "9",
            "modified heating block": "9",
            "Ferrules": "8",
        })

    def test_two_tables_in_markdown_and_word_without_a_prefix(self) -> None:
        text = self.store.export_markdown(self.report_id, fixture=self.fixture).decode()
        block = text.split("### 3-1. 재료와 장비")[1].split("### 3-2.")[0]
        self.assertNotIn("장비:", text)
        self.assertEqual([line for line in block.splitlines() if line.startswith("#")],
                         ["#### 재료", "#### 장비"])
        self.assertEqual(block.count("| 이름 (원문 그대로) | 사용 단계 |"), 2)
        self.assertIn("| Promega trypsin Promega Catalog #V5113 | 3 |", block)
        self.assertIn("| Thermomixer C Model 5382 NAME Thermomixer C | — |", block)
        self.assertIn("사용 단계는 서버가 원문 단계 글에서 그 이름(또는 원문이 쓰는 줄임말)을 찾아 적었다.", block)

        document = Document(io.BytesIO(self.store.export_docx(self.report_id, fixture=self.fixture)))
        tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
        item_tables = [rows for rows in tables if tuple(rows[0]) == er.ITEM_HEADER]
        self.assertEqual(len(item_tables), 2)
        self.assertEqual(len(item_tables[0]) - 1, 8)
        self.assertEqual(item_tables[1][1:], [
            ["Thermomixer C Model 5382 NAME Thermomixer C", "—"], ["modified heating oven", "9"],
            ["modified heating block", "9"], ["Ferrules", "8"]])
        # Markdown carries the same rows.
        for row in item_tables[0][1:] + item_tables[1][1:]:
            self.assertIn("| " + " | ".join(row) + " |", block)
        # A caption stands between the two tables, so Word keeps them apart.
        body = document.element.body
        kinds = [child.tag.rsplit("}", 1)[1] for child in body.iterchildren()]
        texts = [p.text for p in document.paragraphs]
        self.assertEqual(texts[texts.index("재료") + 1], "장비")
        self.assertNotIn(["tbl", "tbl"], [kinds[i:i + 2] for i in range(len(kinds) - 1)])

    def test_the_steps_column_is_narrow(self) -> None:
        name, steps = er.TABLE_WIDTHS[er.ITEM_HEADER]
        self.assertGreater(name, steps)
        self.assertAlmostEqual(name + steps, 1.0)

    def test_a_protocol_without_a_list_says_so(self) -> None:
        self.fixture.draft.protocol.materials = ()
        self.fixture.draft.protocol.equipment = ()
        text = self.store.export_markdown(self.report_id, fixture=self.fixture).decode()
        self.assertIn("> 원문에서 재료 목록을 불러오지 못했다.", text)
        self.assertNotIn("이름 (원문 그대로)", text)


USES = [
    {"번호": 1, "용도": "세척·탈수 용매"},
    {"번호": 2, "용도": "완충 용액 성분"},
    {"번호": 3, "용도": "단백질 분해 효소"},
    {"번호": 5, "용도": "세균 배양 배지"},
    {"번호": 10, "용도": "건조 가열 장치"},
]


class ItemUseTests(_ReportCase):
    """Decision 4: a short use per item from the same call, each checked."""

    def setUp(self) -> None:
        super().setUp()
        self.fixture = items_double()
        self.event("00:30", "session_started", "1")
        self.event("00:36", "step_completed", "1")
        self.event("00:40", "session_stopped", "2", payload={"stop_reason": "stopped_by_user"})
        self.store.finalize(self.report_id, status="stopped", event_key="turn-final")

    def write(self, reply) -> tuple[er.ReportNarrative, _Client]:
        client = _Client(reply)
        brain = er.ReportWriterBrain(client=client, model="fake-model", timeout_seconds=5)
        doc = self.doc()
        return asyncio.run(brain.generate_narrative(doc, list(doc["events"]), fixture=self.fixture)), client

    def test_uses_from_the_one_call_fill_a_use_column(self) -> None:
        narrative, client = self.write(dict(GOOD_REPLY, item_uses=USES))
        self.assertEqual(len(client.calls), 1)
        text = er.render_markdown(narrative)
        block = text.split("### 3-1. 재료와 장비")[1].split("### 3-2.")[0]
        self.assertEqual(block.count("| 이름 (원문 그대로) | 용도 | 사용 단계 |"), 2)
        self.assertIn("| Promega trypsin Promega Catalog #V5113 | 단백질 분해 효소 | 3 |", block)
        self.assertIn("| Lysogeny Agar (Sigma-Aldrich) |   | 6 |", block)
        self.assertIn("| modified heating oven | 건조 가열 장치 | 9 |", block)
        self.assertIn("> 용도는 AI 가 원문 단계를 바탕으로 정리했다.", block)
        document = Document(io.BytesIO(er.render_docx(narrative)))
        rows = [[cell.text for cell in row.cells] for table in document.tables
                for row in table.rows]
        self.assertIn(["Lysogeny Broth (LB) (Sigma-Aldrich)", "세균 배양 배지", "5, 7"], rows)
        self.assertIn("용도는 AI 가 원문 단계를 바탕으로 정리했다.", [p.text for p in document.paragraphs])

    def test_the_model_reads_each_item_with_its_steps(self) -> None:
        _, client = self.write(dict(GOOD_REPLY, item_uses=USES))
        sent = json.loads(client.calls[0]["messages"][1]["content"])
        self.assertEqual(sent["재료·장비"][2], {
            "번호": 3, "종류": "재료", "이름": "Promega trypsin Promega Catalog #V5113", "사용 단계": ["3"]})
        self.assertEqual(sent["재료·장비가 나오는 원문 단계"]["3"], "Add trypsin made up in AMBIC.")
        self.assertNotIn("10", sent["재료·장비가 나오는 원문 단계"])  # no listed item is named there
        self.assertIn("item_uses", client.calls[0]["messages"][0]["content"])

    def test_each_use_is_checked_and_a_failing_one_is_left_blank(self) -> None:
        narrative, _ = self.write(dict(GOOD_REPLY, item_uses=[
            {"번호": 1, "용도": "단백질 펩타이드를 젤 조각에서 추출하고 탈수하는 데 쓰는 유기 용매"},
            {"번호": 2, "용도": "25 mM 완충 용액"},
            {"번호": 3, "용도": "µL 단위 효소 용액"},
            {"번호": 4, "용도": "피부 접촉 주의 산"},
            {"번호": 5, "용도": "배지를 먼저 데우세요"},
            {"번호": 6, "용도": "고체 배지를 만든다."},
            {"번호": 99, "용도": "원문에 없는 장비"},
            {"번호": 7, "이름": "Centrifuge", "용도": "원심 분리"},
            {"번호": 8, "용도": "용액 옮기기"},
            {"번호": 12, "용도": "작은 고정 부품"},
        ]))
        self.assertEqual(dict(narrative.item_uses), {
            "Glass Pasteur pipette (150 mm)": "용액 옮기기", "Ferrules": "작은 고정 부품"})
        self.assertEqual(list(narrative.item_uses_rejected), [
            ("Acetonitrile LC-MS grade B&J Brand VWR International (Avantor) Catalog #BJLC015-2.5", "25자 넘음"),
            ("Ammonium bicarbonate Merck MilliporeSigma (Sigma-Aldrich) Catalog #A6141", "숫자"),
            ("Promega trypsin Promega Catalog #V5113", "단위 “µL”"),
            ("Formic acid, LC-MS grade Thermo Fisher Scientific Catalog #28905", "안전 지시 “주의”"),
            ("Lysogeny Broth (LB) (Sigma-Aldrich)", "절차 지시 “먼저”"),
            ("Lysogeny Agar (Sigma-Aldrich)", "절차 지시 “다.”"),
            ("번호 99", "원문에 없는 항목"),
            ("Centrifuge", "원문에 없는 항목"),
        ])
        text = er.render_markdown(narrative)
        self.assertIn("| Promega trypsin Promega Catalog #V5113 |   | 3 |", text)
        self.assertNotIn("원심 분리", text)

    def test_no_model_a_failed_call_or_no_passing_use_leaves_no_use_column(self) -> None:
        server = er.ReportWriterBrain().build_deterministic_narrative(
            self.doc(), list(self.doc()["events"]), fixture=self.fixture)
        failed, _ = self.write("not json")
        refused, _ = self.write(dict(GOOD_REPLY, item_uses=[{"번호": 1, "용도": "10% 용액"}]))
        for narrative in (server, failed, refused):
            text = er.render_markdown(narrative)
            self.assertIn("| 이름 (원문 그대로) | 사용 단계 |", text)
            self.assertNotIn("| 용도 |", text)
            self.assertNotIn("용도는 AI", text)

    def test_uses_alone_still_say_ai_wrote_part_of_the_report(self) -> None:
        narrative, _ = self.write({"item_uses": USES})
        self.assertNotIn("모델", {narrative.section_origin[key] for key in er.MODEL_SECTIONS})
        self.assertIn("문장 일부는 AI(fake-model)", er.authorship_line(narrative))

    def test_the_use_column_is_the_widest_and_the_steps_the_narrowest(self) -> None:
        name, use, steps = er.TABLE_WIDTHS[er.ITEM_USE_HEADER]
        self.assertGreater(use, name)
        self.assertGreater(name, steps)
        self.assertAlmostEqual(name + use + steps, 1.0)

    def test_a_download_uses_the_kept_reply_and_checks_its_uses_again(self) -> None:
        client = _Client(dict(GOOD_REPLY, item_uses=USES))
        preparer = er.ReportProsePreparer()
        with mock.patch.object(er, "report_protocol_fixture", return_value=(self.fixture, "")):
            preparer.start(self.store, self.report_id, lambda: er.ReportWriterBrain(
                client=client, model="fake-model", timeout_seconds=5))
            self.assertTrue(preparer.wait(self.store, self.report_id, 10))
            text = self.store.export_markdown(self.report_id).decode()
            self.store.export_docx(self.report_id)
        self.assertEqual(len(client.calls), 1)
        self.assertIn("| Promega trypsin Promega Catalog #V5113 | 단백질 분해 효소 | 3 |", text)
