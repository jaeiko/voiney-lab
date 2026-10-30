"""NOTE2: append-only, evidence-audited binders of Voiney experiment notebooks."""

from __future__ import annotations

import base64
import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from tests.test_lane_note1_notebook import battery_protocol
from tests.lane_rp3_support import cell_protocol, coating_protocol
from tests.test_lane_own_reports import records  # noqa: F401
from voiney_lab import experiment_reports as er
from voiney_lab import server as srv
from voiney_lab.workspace_store import initialize_workspace_store


FIXTURES = {
    "protocol-biology": cell_protocol,
    "protocol-chemistry": coating_protocol,
    "protocol-materials": battery_protocol,
}


def _report(
    store: er.ExperimentReportStore, *, field: str, day: str, status: str,
    observation: str, continuation: dict[str, str] | None = None,
    lab_revision: dict[str, object] | None = None,
) -> str:
    protocol_id = f"protocol-{field}"
    started = {"run_changes": []}
    if lab_revision is not None:
        started["lab_adapted_revision"] = lab_revision
    with pytest.MonkeyPatch.context() as patch:
        clock = f"{day}T09:00:00+00:00"
        patch.setattr(er, "_now", lambda: clock)
        report = store.open_report(
            session_id=f"session-{field}-{day}", protocol_id=protocol_id,
            protocol_title=f"{field} protocol", protocol_revision="revision-1",
            protocol_sha256={"biology": "1", "chemistry": "2", "materials": "3"}[field] * 64,
            readiness_status="analysis_passed", development_only=False,
            organization_id="org-note2", owner_principal_id="owner-note2",
        )
        report_id = report["report_id"]
        store.append_event(
            report_id, event_key=f"start-{field}", event_type="session_started",
            step_id="p-1", step_label="1", payload=started,
        )
        if continuation is not None:
            store.append_event(
                report_id, event_key=f"continuation-{field}",
                event_type="checkpoint_continuation", payload=continuation,
            )
        store.append_event(
            report_id, event_key=f"observation-{field}", event_type="observation",
            step_id="p-1", step_label="1", user_wording=observation,
            category="appearance", confirmation_state="read_back",
        )
        store.append_event(
            report_id, event_key=f"end-{field}",
            event_type="workflow_completed" if status == "completed" else "session_stopped",
            step_id="p-1", step_label="1",
        )
        store.finalize(report_id, status=status, event_key=f"final-{field}")
    return report_id


def _text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def test_binder_keeps_three_fields_isolated_and_renders_one_projection(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    store = er.ExperimentReportStore(tmp_path / "reports.sqlite")
    biology = _report(
        store, field="biology", day="2026-10-27", status="completed",
        observation="BIOLOGY-ONLY cells clear",
    )
    chemistry = _report(
        store, field="chemistry", day="2026-10-28", status="stopped",
        observation="CHEMISTRY-ONLY coating cloudy",
    )
    lab_revision = {
        "revision_id": "lab-adapted-revision-one", "version": 1,
        "label_ko": "연구실판 v1", "source_revision_id": "revision-1",
        "recorded_at": "2026-10-29T08:50:00+00:00",
        "changes": [{
            "step_label": "1", "source_written": "80 µL", "shown_ko": "90 µL",
        }],
    }
    materials = _report(
        store, field="materials", day="2026-10-29", status="completed",
        observation="MATERIALS-ONLY cell sealed",
        continuation={
            "source_report_id": chemistry,
            "source_session_id": "session-chemistry-2026-10-28",
        },
        lab_revision=lab_revision,
    )
    monkeypatch.setattr(
        er, "report_protocol_fixture",
        lambda report: (FIXTURES[report["protocol_id"]](), ""),
    )

    created = store.create_notebook_binder(
        organization_id="org-note2", owner_principal_id="owner-note2",
        topic_name="전달용 조건 검토", date_from="2026-10-27", date_to="2026-10-29",
        protocol_id="", report_ids=[materials, biology, chemistry],
    )
    binder = store.notebook_binder(
        created.binder_id, organization_id="org-note2",
        owner_principal_id="owner-note2",
    )

    assert binder.topic_name == "전달용 조건 검토"
    assert [entry.report_id for entry in binder.entries] == [biology, chemistry, materials]
    assert [entry.status for entry in binder.entries] == ["completed", "stopped", "completed"]
    assert binder.entries[2].continuation is not None
    assert binder.entries[2].continuation.source_report_id == chemistry
    assert binder.entries[2].condition_revisions[0].label == "연구실판 v1"
    assert binder.entries[2].condition_revisions[0].changes == ("1단계 · 80 µL → 90 µL",)
    assert er.audit_notebook_binder_isolation(binder) == ()
    sentinels = ["BIOLOGY-ONLY", "CHEMISTRY-ONLY", "MATERIALS-ONLY"]
    for entry, own in zip(binder.entries, sentinels):
        words = "\n".join(item.text for item in entry.notebook.observations)
        assert own in words
        assert all(other not in words for other in sentinels if other != own)

    public = binder.public_dict()
    document = Document(io.BytesIO(er.render_notebook_binder_docx(binder)))
    word = _text(document)
    assert [entry["report_id"] for entry in public["entries"]] == [biology, chemistry, materials]
    assert "목차" in word
    assert "전달용 조건 검토" in word
    assert all(value in word for value in sentinels)
    assert "연구실판 v1" in word
    assert chemistry in word

    anchor = store.get_report(biology)
    events = [event for event in anchor["events"] if event["event_type"] == "notebook_binder_created"]
    assert len(events) == 1
    assert events[0]["payload"]["topic_name"] == "전달용 조건 검토"
    assert events[0]["payload"]["report_ids"] == [biology, chemistry, materials]


def test_binder_filters_and_tenant_fences_are_checked_before_append(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    store = er.ExperimentReportStore(tmp_path / "reports.sqlite")
    first = _report(
        store, field="biology", day="2026-10-27", status="completed",
        observation="first",
    )
    second = _report(
        store, field="chemistry", day="2026-10-28", status="completed",
        observation="second",
    )
    monkeypatch.setattr(
        er, "report_protocol_fixture",
        lambda report: (FIXTURES[report["protocol_id"]](), ""),
    )
    before = {report_id: store.get_report(report_id)["events"] for report_id in (first, second)}

    with pytest.raises(ValueError, match="filter"):
        store.create_notebook_binder(
            organization_id="org-note2", owner_principal_id="owner-note2",
            topic_name="wrong protocol", date_from="2026-10-27", date_to="2026-10-28",
            protocol_id="protocol-biology", report_ids=[first, second],
        )
    with pytest.raises(KeyError):
        store.create_notebook_binder(
            organization_id="org-note2", owner_principal_id="another-owner",
            topic_name="other owner", date_from="", date_to="",
            protocol_id="", report_ids=[first, second],
        )
    monkeypatch.setattr(
        er, "audit_notebook_number_provenance", lambda *args, **kwargs: ("99",),
    )
    with pytest.raises(RuntimeError, match="ungrounded number"):
        store.create_notebook_binder(
            organization_id="org-note2", owner_principal_id="owner-note2",
            topic_name="failed audit", date_from="", date_to="",
            protocol_id="", report_ids=[first, second],
        )
    assert {report_id: store.get_report(report_id)["events"] for report_id in (first, second)} == before


def test_owner_builds_and_exports_the_same_binder_projection_through_http(
    records, monkeypatch: pytest.MonkeyPatch,
) -> None:
    devices, store, first = records
    # Resolve the legacy fixture's owner before the binder query.
    assert devices["owner"].get(f"/api/experiment-reports/{first}/notebook").status_code == 200
    owner = srv._account_for_cookie(devices["owner"].session_cookie).principal()
    workspace = initialize_workspace_store(srv._workspace_settings())
    selected = []
    for suffix, observation in (
        ("second", "second-only observation"),
        ("third", "third-only observation"),
    ):
        session_id = f"session-owned-{suffix}"
        report_id = store.open_report(
            session_id=session_id, protocol_id="protocol-shared",
            protocol_title="Fictional source", protocol_revision="revision-one",
            protocol_sha256="a" * 64, readiness_status="ready", development_only=True,
            organization_id=owner.organization_id, owner_principal_id=owner.principal_id,
        )["report_id"]
        workspace.start_experiment(
            owner, session_id=session_id, protocol_id="protocol-shared",
            protocol_revision_id="revision-one",
        )
        workspace.bind_resource(owner, "experiment_report", report_id)
        store.append_event(
            report_id, event_key=f"{suffix}-start", event_type="session_started",
            step_id="p-1", step_label="1", payload={"run_changes": []},
        )
        store.append_event(
            report_id, event_key=f"{suffix}-step", event_type="step_completed",
            step_id="p-1", step_label="1",
        )
        store.append_event(
            report_id, event_key=f"{suffix}-observation", event_type="observation",
            step_id="p-1", step_label="1", user_wording=observation,
            category="appearance", confirmation_state="read_back",
        )
        store.append_event(
            report_id, event_key=f"{suffix}-complete", event_type="workflow_completed",
            step_id="p-1", step_label="1",
        )
        store.finalize(report_id, status="completed", event_key=f"{suffix}-final")
        selected.append(report_id)
    workspace.close()
    figure = SimpleNamespace(
        asset_id="source-crop-one", source_document_id="a" * 64,
        source_page=2, normalized_bounding_box=(0.1, 0.2, 0.3, 0.4),
        caption_primary="shared protocol figure", caption_source="Figure 1.",
        mime_type="image/png", sha256="b" * 64,
    )
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )

    def fixture_for_report(report):
        fixture = cell_protocol()
        fixture.steps = tuple(
            step for section in fixture.draft.protocol.sections for step in section.steps
        )
        fixture.source_pdf_sha256 = "a" * 64
        fixture.visual_for_step = lambda index: figure if index == 0 else None
        fixture.visual_content = lambda index: (figure, png)
        return fixture, ""

    monkeypatch.setattr(er, "report_protocol_fixture", fixture_for_report)
    monkeypatch.setattr(srv, "report_protocol_fixture", fixture_for_report)
    monkeypatch.setattr(er, "audit_notebook_number_provenance", lambda *args, **kwargs: ())

    candidates = devices["owner"].get("/api/notebook-binders/candidates")
    assert candidates.status_code == 200, candidates.text
    assert {item["report_id"] for item in candidates.json()["reports"]} == {first, *selected}
    invalid = devices["owner"].post("/api/notebook-binders", json={
        "topic_name": "   ",
        "date_from": "", "date_to": "", "protocol_id": "",
        "report_ids": selected,
    })
    assert invalid.status_code == 422, invalid.text
    created = devices["owner"].post("/api/notebook-binders", json={
        "topic_name": "HTTP 묶음", "date_from": "", "date_to": "",
        "protocol_id": "protocol-shared", "report_ids": list(reversed(selected)),
    })
    assert created.status_code == 201, created.text
    binder_id = created.json()["binder_id"]
    shown = devices["owner"].get(f"/api/notebook-binders/{binder_id}")
    assert shown.status_code == 200
    assert shown.json() == created.json()
    assert shown.json()["audit"] == {
        "ungrounded_numbers": 0, "cross_experiment_mixing": 0,
    }
    word = devices["owner"].get(f"/api/notebook-binders/{binder_id}.docx")
    assert word.status_code == 200
    document = Document(io.BytesIO(word.content))
    assert "HTTP 묶음" in _text(document)
    assert "second-only observation" in _text(document)
    assert "third-only observation" in _text(document)
    assert len(document.inline_shapes) == 2

    rendered = []
    monkeypatch.setattr(
        srv, "render_notebook_binder_pdf",
        lambda binder, **kwargs: rendered.append(binder.public_dict()) or b"%PDF-1.4 synthetic",
    )
    pdf = devices["owner"].get(f"/api/notebook-binders/{binder_id}.pdf")
    assert pdf.status_code == 200
    assert rendered == [shown.json()]
    for who in ("peer", "admin", "outside"):
        assert devices[who].get(f"/api/notebook-binders/{binder_id}").status_code == 404


def test_report_bridge_keeps_checkpoint_and_lab_revision_without_actor_ids(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    chosen = SimpleNamespace(event_type="protocol_run_changes_chosen", payload={
        "change_set_id": "set-one", "actor_principal_id": "private-person-id",
        "lab_adapted_revision": {
            "revision_id": "lab-adapted-one", "version": 1,
            "label_ko": "연구실판 v1", "source_revision_id": "revision-one",
            "recorded_at": "2026-10-29T01:00:00+00:00",
            "actor_principal_id": "private-person-id", "target": {"revision": 7},
            "changes": [{"step_label": "2", "source_written": "80 µL", "shown_ko": "90 µL"}],
        },
    })
    catalog_store = SimpleNamespace(list_events=lambda protocol_id: [chosen], close=lambda: None)
    monkeypatch.setattr(
        srv, "_open_protocol_catalog",
        lambda: (SimpleNamespace(store=catalog_store), catalog_store),
    )
    workspace = SimpleNamespace(
        get_experiment=lambda principal, session_id: {"events": [{
            "event_type": "session_restarted_from_checkpoint",
            "payload": {"source_session_id": "session-before"},
        }]},
        close=lambda: None,
    )
    monkeypatch.setattr(srv, "_commercial_workspace", lambda: (SimpleNamespace(), workspace))

    revision = srv._report_lab_revision("protocol-one", "set-one")
    continuation = srv._report_checkpoint_continuation(
        "session-after", "protocol-one", "revision-one",
    )

    assert revision == {
        "revision_id": "lab-adapted-one", "version": 1,
        "label_ko": "연구실판 v1", "source_revision_id": "revision-one",
        "recorded_at": "2026-10-29T01:00:00+00:00",
        "changes": [{"step_label": "2", "source_written": "80 µL", "shown_ko": "90 µL"}],
    }
    assert "actor_principal_id" not in revision
    assert "target" not in revision
    assert continuation["source_session_id"] == "session-before"
    assert continuation["source_report_id"].startswith("ER-")
