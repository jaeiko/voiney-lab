"""The report as PDF (lane RP2, decision 2, 2026-10-19).

Human decision: the PDF is the same report, the Word file printed to PDF by
the server's LibreOffice; the export route takes one more format. The
screen's button is added later by the screen lane (this lane does not touch
static/). LibreOffice is faked in all but one test, which runs only where
LibreOffice is installed and makes no external call.
"""

from __future__ import annotations

import asyncio
import io
import os
import unittest
from pathlib import Path
from unittest import mock

import httpx

from voiney_lab import experiment_reports as er
from voiney_lab import server
from voiney_lab.experiment_protocol_pdf import find_office_converter

from tests.lane_rp2_support import START_CHECK, deviation_run, environment, protocol_double, temporary


class _FakeOffice:
    """Stands in for ``soffice``: keeps the Word file it was given and prints a PDF."""

    def __init__(self, *, returncode: int = 0, write: bool = True) -> None:
        self.returncode = returncode
        self.write = write
        self.document: bytes | None = None
        self.calls: list[list[str]] = []

    def __call__(self, command, **_kwargs):
        self.calls.append(list(command))
        if "--version" in command:
            return mock.Mock(returncode=0, stdout=b"LibreOffice 26.2.6.3 test\n")
        source = Path(command[-1])
        self.document = source.read_bytes()
        if self.write:
            out = Path(command[command.index("--outdir") + 1])
            out.mkdir(parents=True, exist_ok=True)
            (out / (source.stem + ".pdf")).write_bytes(b"%PDF-1.7\n% fake\n")
        return mock.Mock(returncode=self.returncode, stdout=b"", stderr=b"")


class ExportPdfTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = temporary()
        self.env = environment()
        self.env.start()
        er._NARRATIVE_CACHE.clear()
        self.record = deviation_run(self.temporary.name)

    def tearDown(self) -> None:
        self.env.stop()
        self.temporary.cleanup()

    def export(self, office: _FakeOffice, converter: str | None = "/usr/bin/soffice") -> bytes:
        return self.record.store.export_pdf(
            self.record.report_id, fixture=protocol_double(), start_check=START_CHECK,
            converter=converter, run=office)

    def test_the_pdf_is_the_word_report_printed_by_libreoffice(self) -> None:
        from docx import Document

        office = _FakeOffice()
        pdf = self.export(office)
        self.assertEqual(pdf, b"%PDF-1.7\n% fake\n")
        convert = next(call for call in office.calls if "--convert-to" in call)
        self.assertEqual(convert[convert.index("--convert-to") + 1], "pdf")
        self.assertIn("--headless", convert)
        printed = [paragraph.text for paragraph in Document(io.BytesIO(office.document)).paragraphs]
        word = self.record.store.export_docx(
            self.record.report_id, fixture=protocol_double(), start_check=START_CHECK)
        self.assertEqual(printed, [paragraph.text for paragraph in Document(io.BytesIO(word)).paragraphs])

    def test_without_libreoffice_the_pdf_is_refused_by_name(self) -> None:
        with mock.patch("voiney_lab.experiment_protocol_pdf.find_office_converter", return_value=None):
            with self.assertRaises(er.ReportPdfError) as raised:
                self.export(_FakeOffice(), converter=None)
        self.assertEqual(raised.exception.code, "report_pdf_unavailable")

    def test_a_failed_print_is_refused_by_name(self) -> None:
        with self.assertRaises(er.ReportPdfError) as raised:
            self.export(_FakeOffice(returncode=1, write=False))
        self.assertEqual(raised.exception.code, "report_pdf_failed")


class ExportRouteTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = temporary()
        self.addCleanup(self.temporary.cleanup)
        self.record = deviation_run(self.temporary.name)
        patches = [
            mock.patch.dict(os.environ, {
                "VOINEY_LAB_EXPERIMENT_REPORTS_ENABLED": "true",
                "VOINEY_LAB_EXPERIMENT_REPORT_DB": str(self.record.store.path),
                "VOINEY_LAB_REPORT_WRITER_ENABLED": "false"}),
            mock.patch.object(server, "_scope_tenant_resource"),
            mock.patch.object(server, "_pilot_mark"),
        ]
        for patcher in patches:
            patcher.start()
            self.addCleanup(patcher.stop)

    def get(self, format_name: str) -> httpx.Response:
        async def run() -> httpx.Response:
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app),
                                         base_url="http://testserver") as http:
                return await http.get(f"/api/experiment-reports/{self.record.report_id}.{format_name}")
        return asyncio.run(run())

    def test_pdf_is_one_more_format_of_the_export_route(self) -> None:
        with mock.patch.object(er.ExperimentReportStore, "export_pdf", return_value=b"%PDF-1.7 fake") as export:
            response = self.get("pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/pdf")
        self.assertEqual(response.headers["content-disposition"],
                         f'attachment; filename="{self.record.report_id}.pdf"')
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertEqual(response.content, b"%PDF-1.7 fake")
        export.assert_called_once_with(self.record.report_id)

    def test_a_server_without_libreoffice_says_so(self) -> None:
        refusal = er.ReportPdfError("LibreOffice is not installed on this server.", code="report_pdf_unavailable")
        with mock.patch.object(er.ExperimentReportStore, "export_pdf", side_effect=refusal):
            response = self.get("pdf")
        self.assertEqual(response.status_code, 503)
        self.assertEqual(response.json()["detail"], "report_pdf_unavailable")

    def test_the_other_formats_are_as_they_were(self) -> None:
        self.assertEqual(self.get("docx").status_code, 200)
        self.assertEqual(self.get("md").status_code, 200)
        self.assertEqual(self.get("xls").status_code, 404)


@unittest.skipUnless(find_office_converter(), "LibreOffice is not installed here")
class LibreOfficePdfTests(unittest.TestCase):
    """The PDF this server's LibreOffice prints (local; no external call)."""

    def test_an_a4_pdf_whose_korean_text_reads_as_written(self) -> None:
        import fitz

        with temporary() as directory, environment():
            record = deviation_run(directory)
            pdf = record.store.export_pdf(record.report_id, fixture=protocol_double(), start_check=START_CHECK)
        document = fitz.open(stream=pdf, filetype="pdf")
        first = document[0]
        self.assertAlmostEqual(first.rect.width, 595.3, delta=1.0)
        self.assertAlmostEqual(first.rect.height, 841.9, delta=1.0)
        text = first.get_text()
        for words in ("Test digestion protocol 실험 보고서", "한눈 요약", "검토자", "김검토", "37°C → 40°C"):
            self.assertIn(words, text)
        whole = "".join(page.get_text() for page in document)
        self.assertIn("튜브를 쏟았어", whole)
        self.assertIn("1. 실험 목적", whole)
        fonts = {font[3] for page in document for font in page.get_fonts()}
        self.assertTrue(any("Nanum" in name for name in fonts), fonts)


if __name__ == "__main__":
    unittest.main()
