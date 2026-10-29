"""Lane CF, decision 5 (field interviews, 2026-10-07/08): which file the run follows.

"내가 쓰는 프로토콜의 특정 버전 기준으로 정확히 답해 달라." The screen's top
rail and the before-start check say "지금 기준: {파일 이름} · {올린 날짜}";
"어느 프로토콜 기준이야?" is answered with the same words; the report's run
information carries them.
"""

from __future__ import annotations

import io
import unittest
from pathlib import Path
from unittest import mock

from tests.lane_cf_support import RecordedSession, Session
from voiney_lab import experiment_reports as er
from voiney_lab import server as server_module
from voiney_lab.curated_protocol import protocol_basis_question

INDEX_HTML = Path(__file__).resolve().parents[1] / "src/voiney_lab/static/index.html"
BASIS = {"filename": "usingdynamicheadspacecollections.pdf", "uploaded_at": "2026-10-07T09:00:00+00:00"}
DISCLOSED_BASIS = {
    **BASIS,
    "source_disclosure": {
        "recorded": True,
        "status": "yes",
        "status_ko": "예 — 공개·시판 원문",
        "note": "시판 제품 설명서",
        "recorded_at": "2026-10-07T08:59:00+00:00",
    },
}
WORDS = "지금 기준: usingdynamicheadspacecollections.pdf · 2026년 10월 7일 올림"


class _CaptureOffice:
    """Capture the exact Word document handed to the local PDF converter."""

    def __init__(self) -> None:
        self.document: bytes | None = None

    def __call__(self, command, **_kwargs):
        if "--version" in command:
            return mock.Mock(returncode=0, stdout=b"LibreOffice test\n")
        source = Path(command[-1])
        self.document = source.read_bytes()
        out = Path(command[command.index("--outdir") + 1])
        out.mkdir(parents=True, exist_ok=True)
        (out / (source.stem + ".pdf")).write_bytes(b"%PDF-1.7\n% fake\n")
        return mock.Mock(returncode=0, stdout=b"", stderr=b"")


def _word_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*(paragraph.text for paragraph in document.paragraphs), *cells])


class BasisWordsTests(unittest.TestCase):

    def test_the_questions_that_ask_for_it(self) -> None:
        for said in ("어느 프로토콜 기준이야?", "어떤 프로토콜 기준이야", "어느 파일 기준이야?",
                     "지금 기준 뭐야?", "기준 파일 알려줘", "무슨 프로토콜로 하고 있어?",
                     "어느 버전 기준이야?"):
            with self.subTest(said=said):
                self.assertTrue(protocol_basis_question(said))

    def test_other_questions_are_left_as_they_were(self) -> None:
        # "현재 프로토콜 버전 알려줘" keeps its own answer (the replay's turn 2).
        for said in ("현재 프로토콜 버전 알려줘.", "프로토콜 시작해줘", "이 단계 기준이 뭐야?",
                     "pH 기준이 뭐야?"):
            with self.subTest(said=said):
                self.assertFalse(protocol_basis_question(said))

    def test_the_words_from_the_file_and_its_upload(self) -> None:
        self.assertEqual(server_module._source_basis_from(
            "usingdynamicheadspacecollections.pdf",
            {"source_filename": "headspace-upload.pdf", "created_at": "2026-10-07T09:00:00+00:00"},
        ), {"filename": "headspace-upload.pdf", "uploaded_at": "2026-10-07T09:00:00+00:00"})
        self.assertEqual(server_module._source_basis_from("a.pdf", None),
                         {"filename": "a.pdf", "uploaded_at": None})


class BasisRuleTests(Session, unittest.TestCase):

    def test_said_back_before_and_during_the_run(self) -> None:
        self.open_with(None)
        self.session.set_source_basis(BASIS)
        self.assertEqual(self.session.source_basis_words(), WORDS)
        plan = self.say("어느 프로토콜 기준이야?")
        self.assertEqual(plan.display_text, f"{WORDS}.")
        self.assertEqual(self.session.last_front_rule, "protocol_basis")
        self.assertFalse(plan.state_changed)
        self.say("프로토콜 시작해줘")
        before = self.projection()
        plan = self.say("지금 기준 뭐야?")
        self.assertEqual(plan.display_text, f"{WORDS}.")
        self.assertEqual(self.projection(), before)

    def test_without_an_upload_date_it_says_so(self) -> None:
        self.open_with(None)
        self.session.set_source_basis({"filename": "a.pdf", "uploaded_at": None})
        self.assertEqual(self.session.source_basis_words(), "지금 기준: a.pdf · 올린 날짜 기록 없음")

    def test_the_screen_fields_carry_it(self) -> None:
        self.open_with(1)
        self.session.set_source_basis(BASIS)
        fields = server_module.curated_screen_fields(self.session)
        self.assertEqual(fields["source_basis"], {**BASIS, "words": WORDS})


class BasisReportTests(RecordedSession, unittest.TestCase):

    def setUp(self) -> None:
        self.start_recording()

    def test_the_run_information_names_the_file(self) -> None:
        self.open_with(None)
        self.session.set_source_basis(BASIS)
        self.record("프로토콜 시작해줘")
        kinds = [kind for kind, _, _ in self.events()]
        self.assertIn("protocol_basis_recorded", kinds)
        facts = er.build_report_facts(self.report(), fixture=self.session.fixture)
        self.assertIn(("기준 파일", "usingdynamicheadspacecollections.pdf · 2026년 10월 7일 올림"),
                      facts.run_rows)

    def test_the_notebook_carries_the_source_disclosure_and_note(self) -> None:
        self.open_with(None)
        self.session.set_source_basis(DISCLOSED_BASIS)
        self.record("프로토콜 시작해줘")
        facts = er.build_report_facts(self.report(), fixture=self.session.fixture)
        self.assertIn(("원문 공개 범위", "예 — 공개·시판 원문"), facts.notebook.header_rows)
        self.assertIn(("출처 메모", "시판 제품 설명서"), facts.notebook.header_rows)

    def test_person_source_note_numbers_survive_audit_word_and_pdf_without_admitting_others(self) -> None:
        basis = {
            **DISCLOSED_BASIS,
            "source_disclosure": {
                **DISCLOSED_BASIS["source_disclosure"],
                "note": "시판 제품 설명서 문서번호 918273645",
            },
        }
        self.open_with(None)
        self.session.set_source_basis(basis)
        self.record("프로토콜 시작해줘")
        report = self.report()
        facts = er.build_report_facts(report, fixture=self.session.fixture)
        assert facts.notebook is not None
        source_steps = er._source_steps_from_fixture(self.session.fixture)
        self.assertEqual(
            er.audit_notebook_number_provenance(
                facts.notebook, source_steps=source_steps, events=report["events"],
            ),
            (),
        )
        from dataclasses import replace

        tampered = replace(facts.notebook, interpretation_words="기록에 없는 번호 918273646")
        self.assertIn(
            "notebook.interpretation_words: 918273646",
            er.audit_notebook_number_provenance(
                tampered, source_steps=source_steps, events=report["events"],
            ),
        )

        store = self.listener.experiment_report_store
        report_id = self.listener.experiment_report_id
        word = store.export_docx(
            report_id, fixture=self.session.fixture, start_check=None, previous=None,
        )
        for words in ("실험 정보", "원문 공개 범위", "예 — 공개·시판 원문", "출처 메모",
                      "시판 제품 설명서 문서번호 918273645"):
            self.assertIn(words, _word_text(word))
        office = _CaptureOffice()
        self.assertTrue(store.export_pdf(
            report_id, fixture=self.session.fixture, start_check=None, previous=None,
            converter="/usr/bin/soffice", run=office,
        ).startswith(b"%PDF"))
        assert office.document is not None
        self.assertIn("시판 제품 설명서 문서번호 918273645", _word_text(office.document))


class BasisScreenTests(unittest.TestCase):

    def test_the_rail_and_the_start_check_show_it(self) -> None:
        html = INDEX_HTML.read_text(encoding="utf-8")
        self.assertIn('id="rail-basis"', html)
        self.assertIn("source_basis", html)
        summary = html.split("function renderStartSummary", 1)[1].split("\nfunction ", 1)[0]
        self.assertIn("basisWords(", summary)
        words = html.split("function basisWords", 1)[1].split("\nfunction ", 1)[0]
        self.assertIn("지금 기준: ${name} · ${when}", words)
        self.assertIn('timeZone:"Asia/Seoul"', words)
        self.assertNotIn("innerHTML", summary)


if __name__ == "__main__":
    unittest.main()
