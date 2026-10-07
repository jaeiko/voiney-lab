"""The report's text: Korean steps, an AI background, and rounds (lane N, decisions 6-8).

6. The "수행한 단계" table and what the model reads for the methods use a
   step's stored translation when there is one -- marked "(자동 번역)", the
   source under it in the same cell in small type. A translation that fails
   its check is not stored as usable, and the source is used.
7. With no background or principle in the source, the report model may give
   a short one from general knowledge (3-4 sentences, no numbers, procedures
   or safety directions), labelled "AI 일반 지식 — 출처 없음, 확인 필요"; one
   that fails the server's check leaves the server's sentence. A source
   background is used alone.
8. A model sentence that gives a number of rounds or repetitions must say
   "말로 확인한 돌아가기 기준"; otherwise the section is the server's.
"""

from __future__ import annotations

import dataclasses
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from tests.lane_n_support import notes_fixture
from voiney_lab import experiment_reports as er
from voiney_lab.protocol_translation import translation_revision_key, translation_units
from voiney_lab.workspace_store import (
    FactTranslationRecord,
    WorkspaceSettings,
    initialize_workspace_store,
)

KOREAN_STEP_1 = "1 염색된 밴드를 조각으로 잘라 1.5 mL 튜브에 넣는다."


def _report(*events: dict, status: str = "stopped") -> dict:
    return {
        "protocol_id": "lane-n-wash", "protocol_revision": "lane-n-wash-v1",
        "protocol_title": "Fictional wash", "status": status,
        "started_at": "2026-10-07T00:00:00+00:00", "ended_at": "2026-10-07T01:00:00+00:00",
        "events": list(events),
    }


def _done(label: str, key: str, **payload) -> dict:
    return {"event_key": key, "event_type": "step_completed", "step_label": label,
            "step_id": f"step-{label}", "user_wording": None, "category": None,
            "payload": payload, "created_at": "2026-10-07T00:10:00+00:00"}


def _translated_fixture():
    return dataclasses.replace(
        notes_fixture(), machine_localizations={"step-1/current_step": KOREAN_STEP_1})


class KoreanStepTableTests(unittest.TestCase):
    """Decision 6: the stored Korean, marked, with the source under it."""

    def facts(self) -> er.ReportFacts:
        return er.build_report_facts(
            _report(_done("1", "d1"), _done("2", "d2")), fixture=_translated_fixture())

    def test_the_markdown_cell_holds_the_korean_and_the_source_under_it(self) -> None:
        markdown = er.render_markdown(er.narrative_from_sections(self.facts()))
        self.assertIn(
            f"| 1 | {KOREAN_STEP_1} (자동 번역)<br><small>원문: 1 Cut the stained band into "
            "pieces and place them in a 1.5 mL tube.</small> |",
            markdown)
        self.assertIn("| 2 | 2 Wash the pieces with 500 µL of Solution A for 10 min. (원문 영어) |",
                      markdown)

    def test_the_word_cell_has_the_source_below_in_small_type(self) -> None:
        document = Document(io.BytesIO(er.render_docx(er.narrative_from_sections(self.facts()))))
        cell = next(
            row.cells[1] for table in document.tables for row in table.rows
            if row.cells[0].text == "1" and "자동 번역" in row.cells[1].text
        )
        self.assertEqual(cell.paragraphs[0].text, f"{KOREAN_STEP_1} (자동 번역)")
        below = cell.paragraphs[1].runs[0]
        self.assertTrue(below.text.startswith("원문: 1 Cut the stained band"))
        self.assertLess(below.font.size.pt, 9)

    def test_the_model_reads_the_korean_marked_beside_the_source(self) -> None:
        steps = {step["단계"]: step for step in er._writer_facts(self.facts())["수행한 단계"]}
        self.assertEqual(steps["1"]["한국어(자동 번역)"], KOREAN_STEP_1)
        self.assertTrue(steps["1"]["원문"].startswith("1 Cut the stained band"))
        self.assertNotIn("한국어(자동 번역)", steps["2"])

    def test_the_report_reads_stored_translations_without_starting_one(self) -> None:
        import voiney_lab.server as server

        fixture = notes_fixture()
        unit = next(u for u in translation_units(fixture) if u.fact_key == "step-1/current_step")
        with tempfile.TemporaryDirectory() as directory:
            workspace = Path(directory) / "workspace"
            store = initialize_workspace_store(WorkspaceSettings(True, workspace))
            try:
                store.record_fact_translations([
                    FactTranslationRecord(
                        revision_id=translation_revision_key(fixture), fact_key=unit.fact_key,
                        language="ko", source_sha256=unit.source_sha256,
                        translated_text=KOREAN_STEP_1, status="machine", check_result="passed",
                        model="fake-translator", model_version="v1",
                        created_at="2026-10-07T00:00:00+00:00",
                    ),
                    FactTranslationRecord(
                        revision_id=translation_revision_key(fixture), fact_key="step-2/current_step",
                        language="ko", source_sha256="0" * 64,
                        translated_text="2 조각을 씻는다.", status="machine",
                        check_result="quantities_changed", model="fake-translator",
                        model_version="v1", created_at="2026-10-07T00:00:00+00:00",
                    ),
                ])
            finally:
                store.close()
            with patch.dict("os.environ", {
                "VOINEY_LAB_WORKSPACE_ENABLED": "true",
                "VOINEY_LAB_WORKSPACE_DATA_DIR": str(workspace),
            }), patch.object(server, "server_config", return_value=None), \
                    patch.object(server, "_configured_candidate_fixture", return_value=fixture), \
                    patch.object(server, "_revision_translation_runner",
                                 side_effect=AssertionError("a report never starts a translation")):
                loaded, problem = er.report_protocol_fixture(
                    {"protocol_id": "lane-n-wash", "protocol_revision": "lane-n-wash-v1"})
        self.assertEqual(problem, "")
        self.assertEqual(loaded.localized_fact("step-1", "current_step"), KOREAN_STEP_1)
        self.assertEqual(loaded.localization_source("step-1", "current_step"), "machine")
        self.assertIsNone(loaded.localized_fact("step-2", "current_step"))


class GeneralBackgroundTests(unittest.TestCase):
    """Decision 7: a labelled general-knowledge background, checked by the server."""

    GENERAL = ("젤 속 단백질은 염색약과 결합해 밴드로 보인다. 세척 용액은 결합하지 않은 염색약을 "
               "씻어 낸다. 밴드가 투명해지면 다음 처리에 방해가 되는 염색약이 빠진 것이다.")

    def narrative(self, reply: dict) -> er.ReportNarrative:
        facts = er.build_report_facts(_report(_done("1", "d1")), fixture=notes_fixture())
        return er.narrative_from_reply(facts, reply, writer="fake-writer")

    def test_with_no_source_background_the_general_one_is_used_and_labelled(self) -> None:
        narrative = self.narrative({"background": "", "background_general": self.GENERAL})
        self.assertTrue(narrative.background_general)
        self.assertEqual(narrative.background, self.GENERAL)
        markdown = er.render_markdown(narrative)
        self.assertIn(f"{self.GENERAL}\n\n> AI 일반 지식 — 출처 없음, 확인 필요", markdown)
        self.assertIn("AI 가 일반 지식으로 썼으며 출처가 없다", markdown)

    def test_a_source_that_says_it_has_none_is_replaced_too(self) -> None:
        narrative = self.narrative({
            "background": "원문에는 원리에 대한 설명이 기재되어 있지 않다[1].",
            "background_general": self.GENERAL,
        })
        self.assertTrue(narrative.background_general)

    def test_numbers_procedures_or_safety_leave_the_servers_sentence(self) -> None:
        for general in (
            "젤 조각을 37도에서 세척한다. 염색약이 빠진다. 밴드가 투명해진다.",
            "젤 조각을 세척해야 한다. 염색약이 빠진다. 밴드가 투명해진다.",
            "염색약은 독성이 있어 장갑을 낀다. 염색약이 빠진다. 밴드가 투명해진다.",
            "염색약이 빠진다[1]. 밴드가 투명해진다. 다음 처리에 쓴다.",
        ):
            with self.subTest(general=general):
                narrative = self.narrative({"background": "", "background_general": general})
                self.assertFalse(narrative.background_general)
                self.assertIn("background_general", narrative.rejected)
                self.assertIn("외부 자료는 쓰지 않았다. 이 칸은 프로토콜 원문만으로 썼다.",
                              er.render_markdown(narrative))

    def test_a_source_background_is_used_alone(self) -> None:
        narrative = self.narrative({
            "background": "원문은 염색약을 씻어 낸 뒤 다음 단계로 간다고 적고 있다[1].",
            "background_general": self.GENERAL,
        })
        self.assertFalse(narrative.background_general)
        self.assertEqual(narrative.background, "원문은 염색약을 씻어 낸 뒤 다음 단계로 간다고 적고 있다[1].")


class RoundBasisTests(unittest.TestCase):
    """Decision 8: a number of rounds is said with its basis, or not at all."""

    def facts(self) -> er.ReportFacts:
        events = [
            _done("2", "d2a"), _done("3", "d3a"), _done("4", "d4a"),
            {"event_key": "back", "event_type": "repeat_returned", "step_label": "2",
             "step_id": "step-2", "user_wording": None, "category": None,
             "payload": {"step_record": {"kind": "repeat_return", "from_step": "5",
                                         "to_step": "2", "round": 2}},
             "created_at": "2026-10-07T00:20:00+00:00"},
            _done("2", "d2b", step_record={"kind": "repeat_round_completion", "round": 2}),
        ]
        return er.build_report_facts(_report(*events), fixture=notes_fixture())

    def test_a_round_count_without_its_basis_is_refused(self) -> None:
        for methods in ("2–4단계 세척을 총 2회 반복 수행했다.", "2단계를 2회차까지 수행했다.",
                        "세척을 두 번 반복했다."):
            with self.subTest(methods=methods):
                refused = er.check_report_sections({"methods_summary": methods}, self.facts())
                self.assertIn("methods_summary", refused)
                self.assertTrue(any("말로 확인한 돌아가기 기준" in why for why in refused["methods_summary"]))

    def test_a_round_count_with_its_basis_passes(self) -> None:
        methods = "5단계에서 2단계로 돌아가 2회차(말로 확인한 돌아가기 기준)를 수행했다."
        self.assertEqual(er.check_report_sections({"methods_summary": methods}, self.facts()), {})

    def test_the_sources_own_count_is_not_a_round(self) -> None:
        self.assertEqual(
            er.check_report_sections({"methods_summary": "2–4단계를 수행했다."}, self.facts()), {})
        self.assertIsNone(er._ROUND_COUNT.search("원문은 2회 세척 후 투명해진다고 적는다."))


if __name__ == "__main__":
    unittest.main()
