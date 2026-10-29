"""Idempotent, event-driven experiment records owned by the server.

The service stores only explicit workflow events and user-confirmed observations.
It never infers completion, approval, or a laboratory result from model output.
"""

from __future__ import annotations

import asyncio
import hashlib
import csv
import io
import json
import logging
import os
import re
import secrets
import sqlite3
import threading
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field, fields as _dataclass_fields, is_dataclass, replace as _replace_dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Callable
from urllib.parse import quote
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from voiney_lab.experiment_protocol_analysis import program_address_shape
from voiney_lab.report_projection import project_protocol_for_report, project_step_for_report

log = logging.getLogger(__name__)


_TRUE = frozenset({"1", "true", "yes", "on"})
#: Lane RP2, decision 1: an export left to read the check before the start
#: from the protocol's ledger itself (the server's path); a test gives it.
_FROM_LEDGER = object()
_FALSE = frozenset({"0", "false", "no", "off", ""})
_SAFE_ID = re.compile(r"^[A-Za-z0-9._:-]{1,160}$")
_STATUSES = frozenset({"in_progress", "completed", "blocked", "stopped", "incomplete"})


@dataclass(frozen=True)
class ExperimentReportSettings:
    enabled: bool
    database_path: Path | None = None

    @classmethod
    def from_environment(cls) -> "ExperimentReportSettings":
        raw = os.environ.get(
            "VOINEY_LAB_EXPERIMENT_REPORTS_ENABLED", ""
        ).strip().casefold()
        if raw in _FALSE:
            return cls(False)
        if raw not in _TRUE:
            raise ValueError(
                "VOINEY_LAB_EXPERIMENT_REPORTS_ENABLED must be a boolean"
            )
        path = os.environ.get(
            "VOINEY_LAB_EXPERIMENT_REPORT_DB",
            os.environ.get("VOINEY_LAB_EXPERIMENT_REPORTS_DATABASE", ""),
        ).strip()
        if not path:
            raise ValueError(
                "VOINEY_LAB_EXPERIMENT_REPORT_DB is required"
            )
        return cls(True, Path(path))


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean_identifier(value: str, field: str) -> str:
    if not isinstance(value, str) or _SAFE_ID.fullmatch(value) is None:
        raise ValueError(f"{field} is invalid")
    return value


def _clean_text(value: str, *, maximum: int = 1600) -> str:
    if not isinstance(value, str):
        raise ValueError("report text is invalid")
    cleaned = " ".join(value.split())
    if not cleaned or len(cleaned) > maximum:
        raise ValueError("report text is invalid")
    return cleaned


#: The experiment record's events by their Korean names, as the screen's
#: event list shows them; the report's timeline (CSV) carries them beside the
#: event type. Lane VT, decision 8: the timer notices the server gave first.
#: Lane FX4, decision 2: every event type the server writes to the record has
#: its name (tests/test_lane_fx4_event_names.py reads them from server.py).
EVENT_NAMES_KO: dict[str, str] = {
    "session_started": "세션 시작", "step_completed": "단계 완료", "step_advanced": "단계 이동",
    "observation": "관찰 기록", "blocked": "진행 차단", "anomaly": "이상 사항",
    "session_stopped": "세션 종료", "timer_started": "타이머 시작",
    "workflow_paused": "일시정지", "workflow_resumed": "재개", "step_presented": "단계 안내",
    "source_consulted": "참고 자료 확인", "system_anomaly": "시스템 이상",
    "safety_notices_acknowledged": "안전 주의 확인", "workflow_completed": "실험 완료",
    "report_finalized": "기록 마감", "record_corrected": "기록 정정",
    "record_retracted": "기록 취소", "report_value_confirmed": "값 확인",
    "report_review_deferred": "값 확인 화면으로",
    "timer_end_notice": "타이머 끝 알림", "timer_last_minute_notice": "타이머 1분 전 알림",
    # Lane HEAR2, decision 3: the timer's voice commands.
    "timer_paused": "타이머 일시정지", "timer_resumed": "타이머 다시 시작",
    "timer_cancelled": "타이머 취소", "timer_restarted": "타이머 처음부터 다시",
    "connection_restored": "연결 끊김 뒤 다시 연결", "photo_attached": "사진 첨부",
    "web_reference": "웹 참고 링크", "experimenter_recorded": "실험자 기록",
    # Lane HELP: the sources' help after a problem said or asked about.
    "problem_help": "문제 도움",
    "protocol_basis_recorded": "따른 원문 파일 기록", "steps_skipped": "단계 건너뜀",
    "step_reverted": "이전 단계로 되돌림", "repeat_returned": "반복 구간으로 돌아감",
    "repeat_registered": "반복 횟수 등록", "repeat_rounds_declined": "반복 더 하지 않음",
    "repeat_closed": "반복 마침", "branch_answered": "조건 답변",
    "branch_steps_skipped": "조건에 따라 단계 건너뜀",
    # Lane RP3 (decisions ②, ④): what the experimenter says at the start and
    # the end of a run, and a PI's confirmation of the notebook.
    "purpose_recorded": "목적·시료 기록", "result_recorded": "결과·다음 할 일 기록",
    "pi_confirmed": "PI 확인", "notebook_domain_selected": "실험 노트 분야 양식 변경",
    "notebook_inventory_recorded": "Methods 정보 기록",
    "notebook_share_changed": "실험 노트 비교 공유 변경",
}


class ExperimentReportStore:
    """Small SQLite store with one report per procedure session."""

    SCHEMA_VERSION = 3

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    @property
    def database_path(self) -> Path:
        return self.path

    def _connect(self) -> sqlite3.Connection:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(self.path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA busy_timeout = 5000")
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS experiment_report_metadata (
              schema_version INTEGER NOT NULL
            );
            CREATE TABLE IF NOT EXISTS experiment_reports (
              report_id TEXT PRIMARY KEY,
              session_id TEXT NOT NULL UNIQUE,
              protocol_id TEXT NOT NULL,
              protocol_title TEXT NOT NULL,
              protocol_revision TEXT NOT NULL,
              protocol_sha256 TEXT NOT NULL,
              readiness_status TEXT NOT NULL,
              development_only INTEGER NOT NULL,
              status TEXT NOT NULL,
              started_at TEXT NOT NULL,
              ended_at TEXT,
              timezone TEXT NOT NULL,
              anomaly_count INTEGER NOT NULL DEFAULT 0,
              blocker_count INTEGER NOT NULL DEFAULT 0,
              finalization_version INTEGER NOT NULL DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS experiment_report_events (
              id INTEGER PRIMARY KEY,
              report_id TEXT NOT NULL REFERENCES experiment_reports(report_id),
              event_key TEXT NOT NULL,
              event_type TEXT NOT NULL,
              step_id TEXT,
              step_label TEXT,
              user_wording TEXT,
              category TEXT,
              severity TEXT,
              confirmation_state TEXT,
              source_tier TEXT,
              citation_identities TEXT NOT NULL,
              payload TEXT NOT NULL,
              created_at TEXT NOT NULL,
              UNIQUE(report_id, event_key)
            );
            -- Lane RP, decision 9: the report's prose, written once when the
            -- experiment ends. Derived output, not part of the event ledger.
            CREATE TABLE IF NOT EXISTS experiment_report_prose (
              report_id TEXT PRIMARY KEY REFERENCES experiment_reports(report_id),
              writer TEXT NOT NULL,
              reply TEXT,
              failure TEXT NOT NULL DEFAULT '',
              event_count INTEGER NOT NULL,
              prepared_at TEXT NOT NULL
            );
            -- Organisation notebook headings/order are append-only versions.
            -- They change projections only, never experiment state.
            CREATE TABLE IF NOT EXISTS experiment_notebook_templates (
              id INTEGER PRIMARY KEY,
              organization_id TEXT NOT NULL,
              event_key TEXT NOT NULL,
              sections TEXT NOT NULL,
              actor_principal_id TEXT NOT NULL,
              created_at TEXT NOT NULL,
              UNIQUE(organization_id,event_key)
            );
            """
        )
        connection.execute("BEGIN IMMEDIATE")
        rows = connection.execute(
            "SELECT schema_version FROM experiment_report_metadata"
        ).fetchall()
        if not rows:
            connection.execute(
                "INSERT INTO experiment_report_metadata(schema_version) VALUES (?)",
                (self.SCHEMA_VERSION,),
            )
        elif len(rows) == 1 and rows[0][0] == 1:
            # Ownership is metadata; historical ledger events are never rewritten.
            for column in ("organization_id TEXT", "owner_principal_id TEXT",
                           "ownership_migrated INTEGER NOT NULL DEFAULT 0"):
                connection.execute(f"ALTER TABLE experiment_reports ADD COLUMN {column}")
            connection.execute("UPDATE experiment_report_metadata SET schema_version=?", (self.SCHEMA_VERSION,))
        elif len(rows) == 1 and rows[0][0] == 2:
            # Version 3 only adds append-only notebook presentation settings.
            connection.execute("UPDATE experiment_report_metadata SET schema_version=?", (self.SCHEMA_VERSION,))
        elif len(rows) != 1 or rows[0][0] != self.SCHEMA_VERSION:
            connection.close()
            raise RuntimeError("experiment report schema is unsupported")
        columns = {row[1] for row in connection.execute("PRAGMA table_info(experiment_reports)")}
        if "organization_id" not in columns:
            connection.execute("ALTER TABLE experiment_reports ADD COLUMN organization_id TEXT")
            connection.execute("ALTER TABLE experiment_reports ADD COLUMN owner_principal_id TEXT")
            connection.execute("ALTER TABLE experiment_reports ADD COLUMN ownership_migrated INTEGER NOT NULL DEFAULT 0")
        connection.execute("CREATE INDEX IF NOT EXISTS experiment_report_owner_history "
                           "ON experiment_reports(organization_id,owner_principal_id,protocol_id,started_at,report_id)")
        connection.execute("CREATE INDEX IF NOT EXISTS experiment_notebook_template_history "
                           "ON experiment_notebook_templates(organization_id,id)")
        connection.commit()
        return connection

    def backfill_ownership(
        self, resolve: Callable[[str, str, str], tuple[str | None, str | None]],
    ) -> dict[str, int]:
        """Fill legacy metadata from workspace sessions, never from the caller."""

        counts = {"owned": 0, "admin_only": 0, "unscoped": 0}
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT report_id,session_id,protocol_id FROM experiment_reports WHERE ownership_migrated=0"
            ).fetchall()
            for row in rows:
                organization, owner = resolve(row["report_id"], row["session_id"], row["protocol_id"])
                cursor = connection.execute(
                    "UPDATE experiment_reports SET organization_id=?,owner_principal_id=?,ownership_migrated=1 "
                    "WHERE report_id=? AND ownership_migrated=0",
                    (organization, owner, row["report_id"]),
                )
                counts["owned" if owner else "admin_only" if organization else "unscoped"] += cursor.rowcount
        return counts

    def require_access(
        self, report_id: str, *, organization_id: str, owner_principal_id: str,
        every_owner: bool = False,
    ) -> None:
        with self._connect() as connection:
            row = connection.execute(
                "SELECT organization_id,owner_principal_id FROM experiment_reports WHERE report_id=?", (report_id,)
            ).fetchone()
        if (row is None or row["organization_id"] != organization_id
                or (not every_owner and row["owner_principal_id"] != owner_principal_id)):
            raise KeyError("experiment report not found")

    @staticmethod
    def notebook_sharing(report: Mapping[str, Any]) -> dict[str, Any]:
        """Current opt-in state, derived from the latest append-only event."""

        shared = False
        changed_at = None
        changes = 0
        if report.get("owner_principal_id") is None:
            return {"shared": False, "changed_at": None, "changes": 0}
        for event in report.get("events") or ():
            if str(event.get("event_type") or "") != "notebook_share_changed":
                continue
            payload = event.get("payload") if isinstance(event.get("payload"), Mapping) else {}
            if not isinstance(payload.get("shared"), bool):
                continue
            shared = bool(payload["shared"])
            changed_at = event.get("created_at")
            changes += 1
        return {"shared": shared, "changed_at": changed_at, "changes": changes}

    def notebook_comparison_reports(
        self, report_id: str, *, requester_principal_id: str, limit: int = 20,
    ) -> tuple[dict[str, Any], ...]:
        """Own and explicitly shared records for one lab/protocol.

        This method does not grant report access.  Its caller must first prove
        ownership of the anchor and must expose only a bounded read projection.
        """

        anchor = self.get_report(report_id)
        requester_principal_id = _clean_identifier(requester_principal_id, "requester_principal_id")
        if (anchor.get("owner_principal_id") != requester_principal_id
                or anchor.get("organization_id") is None):
            raise KeyError("experiment report not found")
        bounded = max(1, min(int(limit), 20))
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT candidate.report_id FROM experiment_reports AS candidate
                 WHERE candidate.organization_id=? AND candidate.protocol_id=?
                   AND candidate.owner_principal_id IS NOT NULL AND candidate.ended_at IS NOT NULL
                   AND (candidate.owner_principal_id=? OR COALESCE((
                       SELECT json_extract(event.payload,'$.shared')
                         FROM experiment_report_events AS event
                        WHERE event.report_id=candidate.report_id
                          AND event.event_type='notebook_share_changed'
                        ORDER BY event.id DESC LIMIT 1
                   ),0)=1)
                 ORDER BY candidate.started_at,candidate.report_id
                 LIMIT ?
                """,
                (anchor["organization_id"], anchor["protocol_id"], requester_principal_id, bounded),
            ).fetchall()
        visible = []
        for row in rows:
            report = self.get_report(str(row["report_id"]))
            if (report.get("owner_principal_id") == requester_principal_id
                    or self.notebook_sharing(report)["shared"]):
                visible.append(report)
        return tuple(visible)

    def notebook_template(self, organization_id: str | None) -> dict[str, Any]:
        """Latest lab notebook presentation version, or the product default."""

        if organization_id is None:
            return {
                "version": 0,
                "sections": [{"id": item.id, "name": item.name} for item in DEFAULT_NOTEBOOK_TEMPLATE],
                "created_at": None,
            }
        _clean_identifier(organization_id, "organization_id")
        with self._connect() as connection:
            row = connection.execute(
                "SELECT id,sections,created_at FROM experiment_notebook_templates "
                "WHERE organization_id=? ORDER BY id DESC LIMIT 1",
                (organization_id,),
            ).fetchone()
        if row is None:
            return {
                "version": 0,
                "sections": [{"id": item.id, "name": item.name} for item in DEFAULT_NOTEBOOK_TEMPLATE],
                "created_at": None,
            }
        return {"version": int(row["id"]), "sections": json.loads(row["sections"]),
                "created_at": row["created_at"]}

    def save_notebook_template(
        self, *, organization_id: str, actor_principal_id: str,
        sections: Sequence[Mapping[str, Any]], event_key: str,
    ) -> dict[str, Any]:
        """Append one validated heading/order version for a lab."""

        organization_id = _clean_identifier(organization_id, "organization_id")
        actor_principal_id = _clean_identifier(actor_principal_id, "actor_principal_id")
        event_key = _clean_identifier(event_key, "event_key")
        normalized = []
        for item in sections:
            section_id = str(item.get("id") or "") if isinstance(item, Mapping) else ""
            name = " ".join(str(item.get("name") or "").split()) if isinstance(item, Mapping) else ""
            if section_id not in NOTEBOOK_SECTION_IDS or not name or len(name) > 60:
                raise ValueError("notebook template section is invalid")
            normalized.append({"id": section_id, "name": name})
        if len(normalized) != len(NOTEBOOK_SECTION_IDS) or {item["id"] for item in normalized} != set(
                NOTEBOOK_SECTION_IDS):
            raise ValueError("notebook template must contain every section once")
        encoded = json.dumps(normalized, ensure_ascii=False, separators=(",", ":"))
        with self._connect() as connection:
            connection.execute(
                "INSERT OR IGNORE INTO experiment_notebook_templates("
                "organization_id,event_key,sections,actor_principal_id,created_at) VALUES(?,?,?,?,?)",
                (organization_id, event_key, encoded, actor_principal_id, _now()),
            )
        return self.notebook_template(organization_id)

    def open_report(
        self,
        *,
        session_id: str,
        protocol_id: str,
        protocol_title: str,
        protocol_revision: str,
        protocol_sha256: str,
        readiness_status: str,
        development_only: bool,
        organization_id: str | None = None,
        owner_principal_id: str | None = None,
    ) -> dict[str, Any]:
        session_id = _clean_identifier(session_id, "session_id")
        protocol_id = _clean_identifier(protocol_id, "protocol_id")
        protocol_revision = _clean_identifier(protocol_revision, "protocol_revision")
        if not re.fullmatch(r"[0-9a-f]{64}", protocol_sha256):
            raise ValueError("protocol_sha256 is invalid")
        title = _clean_text(protocol_title, maximum=400)
        readiness = _clean_identifier(readiness_status, "readiness_status")
        if owner_principal_id is not None and organization_id is None:
            raise ValueError("report owner requires an organization")
        for value, name in ((organization_id, "organization_id"), (owner_principal_id, "owner_principal_id")):
            if value is not None:
                _clean_identifier(value, name)
        report_id = "ER-" + hashlib.sha256(
            f"{session_id}\x1f{protocol_id}\x1f{protocol_revision}".encode()
        ).hexdigest()[:20].upper()
        now = _now()
        with self._connect() as connection:
            connection.execute(
                """
                INSERT OR IGNORE INTO experiment_reports(
                  report_id,session_id,protocol_id,protocol_title,
                  protocol_revision,protocol_sha256,readiness_status,
                  development_only,status,started_at,timezone,organization_id,owner_principal_id,ownership_migrated
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    report_id, session_id, protocol_id, title,
                    protocol_revision, protocol_sha256, readiness,
                    int(bool(development_only)), "in_progress", now, "UTC",
                    organization_id, owner_principal_id, int(organization_id is not None),
                ),
            )
            stored = connection.execute(
                "SELECT organization_id,owner_principal_id FROM experiment_reports WHERE report_id=?", (report_id,)
            ).fetchone()
            if organization_id is not None and tuple(stored) != (organization_id, owner_principal_id):
                raise ValueError("report ownership cannot change")
        return self.get_report(report_id)

    def append_event(
        self,
        report_id: str,
        *,
        event_key: str,
        event_type: str,
        step_id: str | None = None,
        step_label: str | None = None,
        user_wording: str | None = None,
        category: str | None = None,
        severity: str | None = None,
        confirmation_state: str | None = None,
        source_tier: str | None = None,
        citation_identities: tuple[str, ...] = (),
        payload: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        report_id = _clean_identifier(report_id, "report_id")
        event_key = _clean_identifier(event_key, "event_key")
        event_type = _clean_identifier(event_type, "event_type")
        for value, field in (
            (step_id, "step_id"), (step_label, "step_label"),
            (category, "category"), (severity, "severity"),
            (confirmation_state, "confirmation_state"),
            (source_tier, "source_tier"),
        ):
            if value is not None:
                _clean_identifier(value, field)
        wording = (
            _clean_text(user_wording, maximum=800)
            if user_wording is not None else None
        )
        citations = tuple(
            _clean_identifier(item, "citation_identity")
            for item in citation_identities[:20]
        )
        safe_payload = payload or {}
        encoded = json.dumps(
            safe_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        )
        if len(encoded) > 12000:
            raise ValueError("report payload is too large")
        with self._connect() as connection:
            cursor = connection.execute(
                """
                INSERT OR IGNORE INTO experiment_report_events(
                  report_id,event_key,event_type,step_id,step_label,user_wording,
                  category,severity,confirmation_state,source_tier,
                  citation_identities,payload,created_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    report_id, event_key, event_type, step_id, step_label,
                    wording, category, severity, confirmation_state, source_tier,
                    json.dumps(citations), encoded, _now(),
                ),
            )
            if cursor.rowcount and event_type == "anomaly":
                connection.execute(
                    "UPDATE experiment_reports SET anomaly_count=anomaly_count+1 "
                    "WHERE report_id=?",
                    (report_id,),
                )
            if cursor.rowcount and event_type == "blocked":
                connection.execute(
                    "UPDATE experiment_reports SET blocker_count=blocker_count+1 "
                    "WHERE report_id=?",
                    (report_id,),
                )
        result = self.get_report(report_id)
        result["event_inserted"] = bool(cursor.rowcount)
        return result

    def finalize(
        self, report_id: str, *, status: str, event_key: str
    ) -> dict[str, Any]:
        report_id = _clean_identifier(report_id, "report_id")
        event_key = _clean_identifier(event_key, "event_key")
        if status not in _STATUSES - {"in_progress"}:
            raise ValueError("report status is invalid")
        self.append_event(
            report_id, event_key=event_key, event_type="report_finalized",
            payload={"status": status},
        )
        with self._connect() as connection:
            connection.execute(
                """
                UPDATE experiment_reports
                   SET status=?, ended_at=COALESCE(ended_at,?),
                       finalization_version=CASE
                         WHEN ended_at IS NULL THEN finalization_version+1
                         ELSE finalization_version END
                 WHERE report_id=?
                """,
                (status, _now(), report_id),
            )
        return self.get_report(report_id)

    def get_report(self, report_id: str) -> dict[str, Any]:
        report_id = _clean_identifier(report_id, "report_id")
        with self._connect() as connection:
            report = connection.execute(
                "SELECT * FROM experiment_reports WHERE report_id=?", (report_id,)
            ).fetchone()
            if report is None:
                raise KeyError("experiment report not found")
            events = connection.execute(
                """
                SELECT event_key,event_type,step_id,step_label,user_wording,
                       category,severity,confirmation_state,source_tier,
                       citation_identities,payload,created_at
                  FROM experiment_report_events
                 WHERE report_id=? ORDER BY id
                """,
                (report_id,),
            ).fetchall()
        result = dict(report)
        result["development_only"] = bool(result["development_only"])
        result.update(self._day_sequence(result))
        result["events"] = [
            {
                **{key: row[key] for key in (
                    "event_key", "event_type", "step_id", "step_label",
                    "user_wording", "category", "severity",
                    "confirmation_state", "source_tier", "created_at",
                )},
                "citation_identities": json.loads(row["citation_identities"]),
                "payload": json.loads(row["payload"]),
            }
            for row in events
        ]
        return result

    def list_reports(self, *, session_id: str | None = None) -> list[dict[str, Any]]:
        """List report summaries in stable start order for one development session."""

        parameters: tuple[Any, ...] = ()
        where = ""
        if session_id is not None:
            where = "WHERE session_id=?"
            parameters = (_clean_identifier(session_id, "session_id"),)
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT report_id,session_id,protocol_id,protocol_title,status,started_at,ended_at,
                       anomaly_count,blocker_count,finalization_version,
                       development_only,organization_id,owner_principal_id
                  FROM experiment_reports {where}
                 ORDER BY started_at,report_id
                """,
                parameters,
            ).fetchall()
        return [
            {
                **dict(row), "development_only": bool(row["development_only"]),
                **self._day_sequence(dict(row)),
            }
            for row in rows
        ]

    def _day_sequence(self, report: dict[str, Any]) -> dict[str, Any]:
        """The record's place among those started the same day (lane R6, decision 5).

        Records keep their start in UTC (``timezone``), so the day is the UTC
        date of ``started_at``; 1 is the first record started that day.
        """

        day = str(report["started_at"])[:10]
        if report.get("organization_id") is not None and report.get("owner_principal_id") is None:
            return {"day_sequence": None, "day_sequence_date": day}
        with self._connect() as connection:
            sequence = connection.execute(
                """
                SELECT COUNT(*) FROM experiment_reports
                 WHERE substr(started_at,1,10)=?
                   AND protocol_id=? AND organization_id IS ? AND owner_principal_id IS ?
                   AND (started_at<? OR (started_at=? AND report_id<=?))
                """,
                (day, report["protocol_id"], report.get("organization_id"), report.get("owner_principal_id"),
                 report["started_at"], report["started_at"], report["report_id"]),
            ).fetchone()[0]
        return {"day_sequence": sequence, "day_sequence_date": day}

    def aggregate_metrics(self) -> dict[str, Any]:
        """Return privacy-minimized operational aggregates for administrators.

        This projection intentionally omits report/session identifiers, protocol
        titles, free-form user wording, event payloads, and citation identities.
        """

        tracked_events = (
            "session_started",
            "step_presented",
            "step_completed",
            "timer_started",
            "observation",
            "anomaly",
            "blocked",
            "workflow_paused",
            "workflow_resumed",
            "source_consulted",
            "workflow_completed",
        )
        with self._connect() as connection:
            report_row = connection.execute(
                """
                SELECT COUNT(*) AS total,
                       SUM(CASE WHEN status='completed' THEN 1 ELSE 0 END) AS completed,
                       SUM(anomaly_count) AS anomalies,
                       SUM(blocker_count) AS blockers
                  FROM experiment_reports
                """
            ).fetchone()
            status_rows = connection.execute(
                """
                SELECT status,COUNT(*) AS count
                  FROM experiment_reports
                 GROUP BY status ORDER BY status
                """
            ).fetchall()
            placeholders = ",".join("?" for _ in tracked_events)
            event_rows = connection.execute(
                f"""
                SELECT event_type,COUNT(*) AS count
                  FROM experiment_report_events
                 WHERE event_type IN ({placeholders})
                 GROUP BY event_type ORDER BY event_type
                """,
                tracked_events,
            ).fetchall()
            blocked_rows = connection.execute(
                """
                SELECT COALESCE(step_label,'unlabeled') AS step_label,
                       COUNT(*) AS count
                  FROM experiment_report_events
                 WHERE event_type='blocked'
                 GROUP BY COALESCE(step_label,'unlabeled')
                 ORDER BY count DESC,step_label
                 LIMIT 10
                """
            ).fetchall()
        total = int(report_row["total"] or 0)
        completed = int(report_row["completed"] or 0)
        event_counts = {item: 0 for item in tracked_events}
        event_counts.update(
            {str(row["event_type"]): int(row["count"]) for row in event_rows}
        )
        return {
            "reports": {
                "total": total,
                "completed": completed,
                "completion_rate": round(completed / total, 4) if total else None,
                "by_status": {
                    str(row["status"]): int(row["count"])
                    for row in status_rows
                },
            },
            "workflow_events": event_counts,
            "quality": {
                "anomalies": int(report_row["anomalies"] or 0),
                "blockers": int(report_row["blockers"] or 0),
                "common_blocked_steps": [
                    {
                        "step_label": str(row["step_label"]),
                        "count": int(row["count"]),
                    }
                    for row in blocked_rows
                ],
            },
            "privacy": {
                "raw_audio_included": False,
                "transcripts_included": False,
                "free_text_included": False,
                "report_identifiers_included": False,
            },
        }

    def summary_rows(self) -> list[dict[str, Any]]:
        """One content-free row per experiment record, for the pilot metrics (lane LOG).

        Ids, codes, times and counts only: the steps passed, the records,
        deviations and anomalies kept, the timer notices given, the
        reconnections and their gaps (lane NET), and whether the report's
        prose was written. No wording, payload or step label.
        Reads with a plain connection, so pass a copy of a live file: this
        never creates a table or a missing file.
        """

        if not self.path.is_file():
            return []
        connection = sqlite3.connect(self.path)
        connection.row_factory = sqlite3.Row
        try:
            reports = connection.execute(
                """
                SELECT report_id,session_id,protocol_id,development_only,status,
                       started_at,ended_at,organization_id
                  FROM experiment_reports ORDER BY started_at,report_id
                """
            ).fetchall()
            events = connection.execute(
                """
                SELECT report_id,event_type,COALESCE(category,'') AS category,
                       COUNT(*) AS count
                  FROM experiment_report_events
                 GROUP BY report_id,event_type,COALESCE(category,'')
                """
            ).fetchall()
            # Lane NET, decision 5: one connection_restored per reconnection,
            # its gap in seconds (a number, or null when unknown).
            gaps = {
                str(row["report_id"]): (int(row["count"]), row["gap"])
                for row in connection.execute(
                    """
                    SELECT report_id,COUNT(*) AS count,
                           SUM(CASE WHEN json_type(payload,'$.gap_seconds') IN ('integer','real')
                                    THEN json_extract(payload,'$.gap_seconds') END) AS gap
                      FROM experiment_report_events
                     WHERE event_type='connection_restored'
                     GROUP BY report_id
                    """
                )
            }
            try:
                prose = {
                    str(row["report_id"])
                    for row in connection.execute(
                        "SELECT report_id FROM experiment_report_prose WHERE reply IS NOT NULL"
                    )
                }
            except sqlite3.OperationalError:  # a record from before lane RP's prose
                prose = set()
        finally:
            connection.close()
        counts: dict[str, dict[tuple[str, str], int]] = {}
        for row in events:
            counts.setdefault(str(row["report_id"]), {})[
                (str(row["event_type"]), str(row["category"]))
            ] = int(row["count"])

        def total(report_id: str, *types: str, category: str | None = None,
                  except_category: str | None = None) -> int:
            return sum(
                count for (kind, kind_category), count in counts.get(report_id, {}).items()
                if kind in types
                and (category is None or kind_category == category)
                and (except_category is None or kind_category != except_category)
            )

        return [
            {
                "report_id": str(report["report_id"]),
                "session_id": str(report["session_id"]),
                "protocol_id": str(report["protocol_id"]),
                "organization_id": report["organization_id"],
                "development_only": bool(report["development_only"]),
                "status": str(report["status"]),
                "started_at": str(report["started_at"]),
                "ended_at": report["ended_at"],
                "steps": total(report["report_id"], "step_completed", "step_advanced"),
                "reverts": total(report["report_id"], "step_reverted"),
                "records": total(report["report_id"], "observation", except_category="deviation"),
                "deviations": total(report["report_id"], "observation", category="deviation"),
                "anomalies": total(report["report_id"], "anomaly"),
                "system_anomalies": total(report["report_id"], "system_anomaly"),
                "corrections": total(report["report_id"], "record_corrected", "record_retracted"),
                "timer_last_minute_alerts": total(report["report_id"], "timer_last_minute_notice"),
                "timer_end_alerts": total(report["report_id"], "timer_end_notice"),
                "report_prose": str(report["report_id"]) in prose,
                "reconnections": gaps.get(str(report["report_id"]), (0, None))[0],
                "reconnection_gap_seconds": gaps.get(str(report["report_id"]), (0, None))[1],
            }
            for report in reports
        ]

    def export_json(self, report_id: str) -> bytes:
        return (
            json.dumps(
                self.get_report(report_id), ensure_ascii=False,
                sort_keys=True, indent=2,
            ) + "\n"
        ).encode()

    def export_markdown(
        self, report_id: str, narrative: ReportNarrative | None = None,
        *, fixture: Any = None, start_check: Any = _FROM_LEDGER, previous: Any = _FROM_LEDGER,
    ) -> bytes:
        """The researcher's notebook as Markdown (same structure as the .docx).

        Without a narrative the server's own sentences fill every section.
        ``fixture`` is the protocol the session ran; left out, it is looked up
        the way the server loads it, and so are the answers given in the check
        before the start (``start_check``, lane RP2: None when not read) and
        the run of the same protocol before this one (``previous``, lane RP3:
        None when there is none).
        """

        return render_markdown(self._narrative(report_id, narrative, fixture, start_check, previous)).encode()

    def _narrative(
        self, report_id: str, narrative: ReportNarrative | None, fixture: Any,
        start_check: Any = _FROM_LEDGER, previous: Any = _FROM_LEDGER,
    ) -> ReportNarrative:
        """The given narrative, else the prepared prose, else the server's sentences.

        Prepared prose is checked again against the record as it is now, so a
        photo or note added after it was written still appears in the tables.
        """

        report = self.get_report(report_id)
        if narrative is not None and narrative.facts is not None:
            prepared = narrative
            return self._with_notebook_template(prepared, report)
        if start_check is _FROM_LEDGER:
            start_check = (
                report_start_check_answers(report, since=self._previous_start(report))
                if fixture is None else None)
        if previous is _FROM_LEDGER:
            # Lane RP3, decision ① (6): the finished run of the same protocol
            # before this one, from this store; its facts are built the same
            # way (with the same protocol when it ran the same revision).
            earlier = self.previous_finished_report(report)
            previous = None
            if earlier is not None:
                same = fixture is not None and str(earlier.get("protocol_revision")) == str(
                    report.get("protocol_revision"))
                previous = ReportWriterBrain.facts_for(
                    earlier, list(earlier["events"]), fixture=fixture if same else _AUTO)
        facts = ReportWriterBrain.facts_for(
            report, list(report["events"]), fixture=_AUTO if fixture is None else fixture,
            start_check=start_check, previous=previous)
        prose = self.get_prose(report_id)
        if prose is None or prose["writer"] == SERVER_WRITER:
            prepared = narrative_from_sections(facts)
        else:
            prepared = narrative_from_reply(
                facts, prose["reply"], failure=prose["failure"], writer=prose["writer"],
                written_at=prose["prepared_at"],
            )
        return self._with_notebook_template(prepared, report)

    def _with_notebook_template(
        self, narrative: ReportNarrative, report: Mapping[str, Any],
    ) -> ReportNarrative:
        facts = narrative.facts
        if facts is None or facts.notebook is None:
            return narrative
        stored = self.notebook_template(
            str(report.get("organization_id")) if report.get("organization_id") else None)
        template = tuple(NotebookSection(str(item["id"]), str(item["name"]))
                         for item in stored["sections"])
        book = _replace_dataclass(facts.notebook, section_template=template)
        return _replace_dataclass(narrative, facts=_replace_dataclass(facts, notebook=book))

    def _previous_start(self, report: Mapping[str, Any]) -> str | None:
        """When the run of the same protocol before this one started, or None.

        Lane RP2, decision 1: the check before the start that belongs to a run
        is the one answered after the run before it started.
        """

        if report.get("organization_id") is not None and report.get("owner_principal_id") is None:
            return None
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT MAX(started_at) FROM experiment_reports
                 WHERE protocol_id=? AND started_at<?
                   AND organization_id IS ? AND owner_principal_id IS ?
                """,
                (report["protocol_id"], report["started_at"], report.get("organization_id"), report.get("owner_principal_id")),
            ).fetchone()
        return row[0] if row and row[0] else None

    def last_run_values(
        self, *, protocol_id: str, organization_id: str | None, owner_principal_id: str | None,
    ) -> dict[str, Any] | None:
        """The values the owner's latest run of this protocol used in place of
        the source's (lane EASY, decision 3: "지난번처럼"), or None.

        The record's own owner only (the ownership rule of the records): an
        unknown owner gets nothing; any revision, finished or not.
        """

        if owner_principal_id is None:
            return None
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT report_id FROM experiment_reports
                 WHERE protocol_id=? AND organization_id IS ? AND owner_principal_id=?
                 ORDER BY started_at DESC, report_id DESC LIMIT 1
                """,
                (protocol_id, organization_id, owner_principal_id),
            ).fetchone()
        if row is None:
            return None
        report = self.get_report(row[0])
        return {"report_id": report["report_id"], "started_at": report.get("started_at"),
                "protocol_revision": report.get("protocol_revision"),
                "changes": run_values(report.get("events") or ())}

    def previous_finished_report(self, report: Mapping[str, Any]) -> dict[str, Any] | None:
        """The finished run of the same protocol that started last before this
        one, with its events, or None (lane RP3, decision ① (6): "지난번").

        Same protocol, any revision; a run still in progress is not "지난번".
        """

        if report.get("organization_id") is not None and report.get("owner_principal_id") is None:
            return None
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT report_id FROM experiment_reports
                 WHERE protocol_id=? AND started_at<? AND ended_at IS NOT NULL AND report_id<>?
                   AND organization_id IS ? AND owner_principal_id IS ?
                 ORDER BY started_at DESC, report_id DESC LIMIT 1
                """,
                (report["protocol_id"], report["started_at"], report["report_id"],
                 report.get("organization_id"), report.get("owner_principal_id")),
            ).fetchone()
        return self.get_report(row[0]) if row else None

    def earlier_reports(
        self, report_id: str, *, limit: int = 20,
    ) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        """This record and the runs of the same protocol that started before
        it, newest first, with their events (lane HELP, decision 2 ②).

        The scope is ``previous_finished_report``'s: the same protocol (any
        revision), organisation and owner; a run still going counts too. A
        lab-only record (an organisation and no owner) has none.
        """

        report = self.get_report(report_id)
        if report.get("organization_id") is not None and report.get("owner_principal_id") is None:
            return report, []
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT report_id FROM experiment_reports
                 WHERE protocol_id=? AND started_at<? AND report_id<>?
                   AND organization_id IS ? AND owner_principal_id IS ?
                 ORDER BY started_at DESC, report_id DESC LIMIT ?
                """,
                (report["protocol_id"], report["started_at"], report["report_id"],
                 report.get("organization_id"), report.get("owner_principal_id"), int(limit)),
            ).fetchall()
        return report, [self.get_report(row[0]) for row in rows]

    def save_prose(
        self, report_id: str, *, writer: str, reply: Mapping[str, Any] | None,
        failure: str, event_count: int,
    ) -> None:
        """Keep one report's prose (replacing any earlier one)."""

        report_id = _clean_identifier(report_id, "report_id")
        with self._connect() as connection:
            connection.execute(
                """
                INSERT OR REPLACE INTO experiment_report_prose(
                  report_id,writer,reply,failure,event_count,prepared_at
                ) VALUES(?,?,?,?,?,?)
                """,
                (report_id, writer[:200],
                 None if reply is None else json.dumps(reply, ensure_ascii=False),
                 failure[:200], int(event_count), _now()),
            )

    def get_prose(self, report_id: str) -> dict[str, Any] | None:
        report_id = _clean_identifier(report_id, "report_id")
        with self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM experiment_report_prose WHERE report_id=?", (report_id,)
            ).fetchone()
        if row is None:
            return None
        prose = dict(row)
        prose["reply"] = json.loads(prose["reply"]) if prose["reply"] is not None else None
        return prose

    def export_csv(self, report_id: str) -> bytes:
        """Export the stable event timeline as UTF-8 CSV."""

        report = self.get_report(report_id)
        output = io.StringIO(newline="")
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow((
            "report_id", "event_key", "event_type", "step_id", "step_label",
            "category", "severity", "confirmation_state", "source_tier",
            "user_wording", "created_at",
            "source_duration_seconds", "timer_started_at", "elapsed_seconds",
            "remaining_seconds", "completion_state", "demo_bypassed",
            "event_name_ko",
        ))
        for event in report["events"]:
            payload = event.get("payload") or {}
            timer = payload.get("timer") if isinstance(payload.get("timer"), dict) else {}
            writer.writerow((
                report["report_id"], event["event_key"], event["event_type"],
                event["step_id"] or "", event["step_label"] or "",
                event["category"] or "", event["severity"] or "",
                event["confirmation_state"] or "", event["source_tier"] or "",
                event["user_wording"] or "", event["created_at"],
                timer.get("source_duration_seconds", ""),
                timer.get("started_at", ""),
                timer.get("elapsed_seconds", ""),
                timer.get("remaining_seconds", ""),
                timer.get("completion_state", ""),
                timer.get("demo_bypassed", ""),
                EVENT_NAMES_KO.get(str(event["event_type"]), ""),
            ))
        return output.getvalue().encode("utf-8-sig")

    def docx_bytes(self, report_id: str) -> bytes:
        return self.export_docx(report_id)

    def export_docx(
        self, report_id: str, narrative: ReportNarrative | None = None,
        *, fixture: Any = None, start_check: Any = _FROM_LEDGER, previous: Any = _FROM_LEDGER,
        image_resolver: Callable[["NotebookImage"], tuple[bytes, str] | None] | None = None,
    ) -> bytes:
        """The researcher's notebook as .docx (same structure as the Markdown).

        The SQLite event ledger remains the authority; this document is a
        derived, human-readable projection of it.
        """

        prepared = self._narrative(report_id, narrative, fixture, start_check, previous)
        source_fixture = fixture
        if source_fixture is None:
            source_fixture, _ = report_protocol_fixture(self.get_report(report_id))

        def resolve(image: NotebookImage) -> tuple[bytes, str] | None:
            if image_resolver is not None:
                found = image_resolver(image)
                if found is not None:
                    return found
            if image.kind != "source_figure" or source_fixture is None:
                return None
            from voiney_lab import source_figures
            found = source_figures.figure_content(
                getattr(source_fixture, "source_pdf_path", None),
                getattr(source_fixture, "source_pdf_sha256", None), image.asset_id,
            )
            return None if found is None else (found[1], found[0].mime_type)

        return render_docx(prepared, image_resolver=resolve)

    def export_pdf(
        self, report_id: str, narrative: ReportNarrative | None = None,
        *, fixture: Any = None, start_check: Any = _FROM_LEDGER, previous: Any = _FROM_LEDGER,
        converter: str | None = None, run: Any = None,
        image_resolver: Callable[["NotebookImage"], tuple[bytes, str] | None] | None = None,
    ) -> bytes:
        """The researcher's notebook as PDF: the .docx above, printed by the
        server's LibreOffice (lane RP2, decision 2). ``ReportPdfError`` when
        it cannot be printed."""

        return report_docx_to_pdf(
            self.export_docx(
                report_id, narrative, fixture=fixture, start_check=start_check, previous=previous,
                image_resolver=image_resolver,
            ),
            converter=converter, run=run)

    def export_meeting_docx(
        self, report_id: str, narrative: ReportNarrative | None = None,
        *, fixture: Any = None, start_check: Any = _FROM_LEDGER, previous: Any = _FROM_LEDGER,
    ) -> bytes:
        """The one-page lab-meeting summary as .docx (lane RP3, decision ③)."""

        return render_meeting_docx(self._narrative(report_id, narrative, fixture, start_check, previous))

    def export_meeting_pdf(
        self, report_id: str, narrative: ReportNarrative | None = None,
        *, fixture: Any = None, start_check: Any = _FROM_LEDGER, previous: Any = _FROM_LEDGER,
        converter: str | None = None, run: Any = None,
    ) -> bytes:
        """The lab-meeting summary printed by the server's LibreOffice."""

        return report_docx_to_pdf(
            self.export_meeting_docx(report_id, narrative, fixture=fixture, start_check=start_check,
                                     previous=previous),
            converter=converter, run=run)

    def export_xlsx(
        self, report_id: str, narrative: ReportNarrative | None = None,
        *, fixture: Any = None, start_check: Any = _FROM_LEDGER, previous: Any = _FROM_LEDGER,
    ) -> bytes:
        """The results table as .xlsx (lane RP3, decision ③)."""

        return render_xlsx(self._narrative(report_id, narrative, fixture, start_check, previous))


@dataclass(frozen=True)
class StepExecutionContext:
    """Detailed factual context for one protocol step and its execution status."""

    step_id: str
    step_label: str
    section_title: str
    instruction_source_text: str
    sub_actions: tuple[str, ...] = ()
    quantities: tuple[str, ...] = ()
    conditions: tuple[str, ...] = ()
    expected_results: tuple[str, ...] = ()
    warnings: tuple[str, ...] = ()
    notes: tuple[str, ...] = ()
    tips: tuple[str, ...] = ()
    source_page: int = 1
    evidence_ids: tuple[str, ...] = ()
    entered_at: str | None = None
    completed_at: str | None = None
    completion_state: str = "not_started"  # "completed", "in_progress", "not_started"
    timer_configuration: str | None = None
    timer_actuals: str | None = None
    user_confirmed_observations: tuple[str, ...] = ()
    anomalies_deviations: tuple[str, ...] = ()
    applicable_safety_references: tuple[str, ...] = ()

    @property
    def instruction(self) -> str:
        return self.instruction_source_text

    def public_dict(self) -> dict[str, Any]:
        return {
            "step_id": self.step_id,
            "step_label": self.step_label,
            "section_title": self.section_title,
            "instruction_source_text": self.instruction_source_text,
            "sub_actions": list(self.sub_actions),
            "quantities": list(self.quantities),
            "conditions": list(self.conditions),
            "expected_results": list(self.expected_results),
            "warnings": list(self.warnings),
            "notes": list(self.notes),
            "tips": list(self.tips),
            "source_page": self.source_page,
            "evidence_ids": list(self.evidence_ids),
            "entered_at": self.entered_at,
            "completed_at": self.completed_at,
            "completion_state": self.completion_state,
            "timer_configuration": self.timer_configuration,
            "timer_actuals": self.timer_actuals,
            "user_confirmed_observations": list(self.user_confirmed_observations),
            "anomalies_deviations": list(self.anomalies_deviations),
            "applicable_safety_references": list(self.applicable_safety_references),
        }


@dataclass(frozen=True)
class GroundedReportContext:
    """Rich factual context constructed from the ledger and verified protocol store."""

    report_metadata: dict[str, Any]
    protocol_metadata: dict[str, Any]
    experiment_objective: str
    materials: tuple[str, ...]
    equipment: tuple[str, ...]
    prerequisites: tuple[str, ...]
    executed_steps: tuple[StepExecutionContext, ...]
    all_protocol_steps: tuple[StepExecutionContext, ...]
    observations: tuple[dict[str, Any], ...]
    timers: tuple[dict[str, Any], ...]
    anomalies: tuple[dict[str, Any], ...]
    deviations: tuple[dict[str, Any], ...]
    safety_pack_summary: dict[str, Any] | None
    source_references: tuple[str, ...]
    session_timing: dict[str, Any]
    event_ledger: tuple[dict[str, Any], ...]

    @property
    def protocol_id(self) -> str:
        return str(self.protocol_metadata.get("protocol_id") or "")

    @property
    def report_id(self) -> str:
        return str(self.report_metadata.get("report_id") or "")

    @property
    def all_steps(self) -> tuple[StepExecutionContext, ...]:
        return self.all_protocol_steps

    def public_dict(self) -> dict[str, Any]:
        return {
            "report_metadata": self.report_metadata,
            "protocol_metadata": self.protocol_metadata,
            "experiment_objective": self.experiment_objective,
            "materials": list(self.materials),
            "equipment": list(self.equipment),
            "prerequisites": list(self.prerequisites),
            "executed_steps": [s.public_dict() for s in self.executed_steps],
            "observations": list(self.observations),
            "timers": list(self.timers),
            "anomalies": list(self.anomalies),
            "deviations": list(self.deviations),
            "safety_pack_summary": self.safety_pack_summary,
            "source_references": list(self.source_references),
            "session_timing": self.session_timing,
            "total_event_count": len(self.event_ledger),
        }


def build_grounded_report_context(
    report_data: dict[str, Any],
    events: list[dict[str, Any]] | None = None,
    *,
    fixture: Any = None,
    lookup: bool = True,
) -> GroundedReportContext:
    """Rehydrate protocol steps and merge execution event timestamps, observations, and timers."""
    if events is None:
        events = list(report_data.get("events") or [])
    protocol_id = str(report_data.get("protocol_id") or "")
    protocol_rev = str(report_data.get("protocol_revision") or "")
    protocol_title = str(report_data.get("protocol_title") or "Experiment Protocol")

    step_events: dict[str, list[dict[str, Any]]] = {}
    obs_list: list[dict[str, Any]] = []
    timer_list: list[dict[str, Any]] = []
    anomaly_list: list[dict[str, Any]] = []
    deviation_list: list[dict[str, Any]] = []
    step_snapshots: dict[str, dict[str, Any]] = {}

    for e in events:
        s_lbl = str(e.get("step_label") or "")
        if s_lbl and s_lbl != "—":
            step_events.setdefault(s_lbl, []).append(e)
        ev_type = str(e.get("event_type") or "")
        payload = e.get("payload") if isinstance(e.get("payload"), dict) else {}

        if "step_snapshot" in payload and isinstance(payload["step_snapshot"], dict):
            snap = payload["step_snapshot"]
            snap_lbl = str(snap.get("step_label") or s_lbl)
            if snap_lbl:
                step_snapshots[snap_lbl] = snap

        if ev_type == "observation":
            obs_list.append(e)
        elif ev_type in ("timer_started", "timer_expired") or "timer" in payload:
            timer_list.append(e)
        elif ev_type in ("anomaly", "system_anomaly") or e.get("anomaly_category"):
            anomaly_list.append(e)
        elif ev_type in ("workflow_paused", "workflow_blocked"):
            deviation_list.append(e)

    # The protocol the session ran, looked up the way the server loads it
    # (report_protocol_fixture) unless the caller passes it or asks not to.
    if fixture is None and lookup:
        fixture, _ = report_protocol_fixture(report_data)
    stored_protocol = fixture.draft.protocol if fixture is not None else None

    all_steps: list[StepExecutionContext] = []
    executed_steps: list[StepExecutionContext] = []
    all_materials: set[str] = set()
    all_equipment: set[str] = set()
    all_prereqs: set[str] = set()
    objective = f"본 실험은 '{protocol_title}' 지침에 따라 표준화된 실험 절차를 수행하고 검증된 실험 데이터를 기록하는 것을 목적으로 한다."

    if stored_protocol is not None:
        projected = project_protocol_for_report(stored_protocol)
        if projected.objective:
            objective = projected.objective
        all_materials = set(projected.materials)
        all_equipment = set(projected.equipment)
        all_prereqs = set(projected.prerequisites)

        for p_step in projected.steps:
            lbl = p_step.step_label
            evs = step_events.get(lbl, [])
            entered = next((_human_event_clock(ev.get("created_at")) for ev in evs if ev.get("event_type") in ("step_entered", "step_presented", "session_started")), None)
            completed = next((_human_event_clock(ev.get("created_at")) for ev in evs if ev.get("event_type") == "step_completed"), None)
            state = "completed" if completed else ("in_progress" if entered else "not_started")

            # Timers
            t_cfg = None
            t_act = None
            t_ev = next((ev for ev in evs if ev.get("event_type") == "timer_started" or "timer" in (ev.get("payload") or {})), None)
            if t_ev:
                t_pay = (t_ev.get("payload") or {}).get("timer", {})
                dur = t_pay.get("source_duration_seconds", t_pay.get("duration_seconds"))
                elap = t_pay.get("elapsed_seconds")
                if dur:
                    t_cfg = f"{dur}초 ({_format_elapsed_clock(dur)})"
                if elap:
                    t_act = f"{elap}초 경과 ({_format_elapsed_clock(elap)})"

            # Observations
            obs_collected = []
            for ev in evs:
                if ev.get("event_type") == "observation":
                    text = str(ev.get("user_wording") or (ev.get("payload") or {}).get("text") or "관찰 기록")
                    if text:
                        obs_collected.append(text)
                payload = ev.get("payload") if isinstance(ev.get("payload"), dict) else {}
                if "observations" in payload and isinstance(payload["observations"], (list, tuple)):
                    for o in payload["observations"]:
                        if o and str(o).strip():
                            obs_collected.append(str(o).strip())
            step_obs = tuple(obs_collected)

            # Anomalies
            step_anom = tuple(
                str(ev.get("user_wording") or (ev.get("payload") or {}).get("text") or "이상 보고")
                for ev in evs if ev.get("event_type") in ("anomaly", "system_anomaly") or ev.get("anomaly_category")
            )

            ctx_step = StepExecutionContext(
                step_id=p_step.step_id,
                step_label=lbl,
                section_title=p_step.section_title,
                instruction_source_text=p_step.instruction_source_text,
                sub_actions=p_step.sub_actions,
                quantities=p_step.quantities,
                conditions=p_step.conditions,
                expected_results=p_step.expected_results,
                warnings=p_step.warnings,
                notes=p_step.notes,
                tips=p_step.tips,
                source_page=p_step.source_page,
                evidence_ids=p_step.evidence_ids,
                entered_at=entered,
                completed_at=completed,
                completion_state=state,
                timer_configuration=t_cfg,
                timer_actuals=t_act,
                user_confirmed_observations=step_obs,
                anomalies_deviations=step_anom,
            )
            all_steps.append(ctx_step)
            if state in ("completed", "in_progress") or step_obs or step_anom:
                executed_steps.append(ctx_step)

    else:
        # Fall back to step_snapshots or step_events
        for lbl, evs in sorted(step_events.items(), key=lambda x: str(x[0])):
            snap = step_snapshots.get(lbl, {})
            entered = next((_human_event_clock(ev.get("created_at")) for ev in evs if ev.get("event_type") in ("step_entered", "step_presented", "session_started")), None)
            completed = next((_human_event_clock(ev.get("created_at")) for ev in evs if ev.get("event_type") == "step_completed"), None)
            state = "completed" if completed else "in_progress"

            # Lane RP: no placeholder text. A step whose words are unknown
            # stays empty, and the report says the source could not be read.
            inst = snap.get("instruction_source_text") or snap.get("instruction") or ""
            sec_title = snap.get("section_title") or "Experiment Execution"
            quantities = tuple(snap.get("quantities") or ())
            conditions = tuple(snap.get("conditions") or ())
            exp_res = tuple(snap.get("expected_results") or ())
            warnings = tuple(snap.get("warnings") or ())
            page = int(snap.get("source_page") or 1)

            obs_collected = []
            for ev in evs:
                if ev.get("event_type") == "observation":
                    text = str(ev.get("user_wording") or (ev.get("payload") or {}).get("text") or "관찰 기록")
                    if text:
                        obs_collected.append(text)
                payload = ev.get("payload") if isinstance(ev.get("payload"), dict) else {}
                if "observations" in payload and isinstance(payload["observations"], (list, tuple)):
                    for o in payload["observations"]:
                        if o and str(o).strip():
                            obs_collected.append(str(o).strip())
            step_obs = tuple(obs_collected)

            step_anom = tuple(
                str(ev.get("user_wording") or (ev.get("payload") or {}).get("text") or "이상 보고")
                for ev in evs if ev.get("event_type") in ("anomaly", "system_anomaly") or ev.get("anomaly_category")
            )

            ctx_step = StepExecutionContext(
                step_id=snap.get("step_id") or f"step-{lbl}",
                step_label=lbl,
                section_title=sec_title,
                instruction_source_text=inst,
                quantities=quantities,
                conditions=conditions,
                expected_results=exp_res,
                warnings=warnings,
                source_page=page,
                entered_at=entered,
                completed_at=completed,
                completion_state=state,
                user_confirmed_observations=step_obs,
                anomalies_deviations=step_anom,
            )
            all_steps.append(ctx_step)
            executed_steps.append(ctx_step)

    refs = [f"실험 PDF: {report_data.get('protocol_title') or protocol_title} (ID: {protocol_id}, Rev: {protocol_rev})"]
    if report_data.get("protocol_sha256"):
        refs.append(f"PDF SHA-256: {report_data.get('protocol_sha256')}")

    return GroundedReportContext(
        report_metadata={
            "report_id": report_data.get("report_id"),
            "session_id": report_data.get("session_id"),
            "status": report_data.get("status"),
            "anomaly_count": report_data.get("anomaly_count", len(anomaly_list)),
            "blocker_count": report_data.get("blocker_count", 0),
            "finalization_version": report_data.get("finalization_version", 0),
        },
        protocol_metadata={
            "protocol_id": protocol_id,
            "protocol_title": protocol_title,
            "protocol_revision": protocol_rev,
            "protocol_sha256": report_data.get("protocol_sha256"),
            "total_steps": len(all_steps),
        },
        experiment_objective=objective,
        materials=tuple(sorted(all_materials)),
        equipment=tuple(sorted(all_equipment)),
        prerequisites=tuple(sorted(all_prereqs)),
        executed_steps=tuple(executed_steps),
        all_protocol_steps=tuple(all_steps),
        observations=tuple(obs_list),
        timers=tuple(timer_list),
        anomalies=tuple(anomaly_list),
        deviations=tuple(deviation_list),
        safety_pack_summary=None,
        source_references=tuple(refs),
        session_timing={
            "started_at": report_data.get("started_at"),
            "ended_at": report_data.get("ended_at"),
            "timezone": report_data.get("timezone", "UTC"),
        },
        event_ledger=tuple(events),
    )


# --- The report a researcher reads (lane RP, decisions of 2026-10-06) -------------
#
# The report is an experiment report, not a log of the system: purpose,
# background, materials and methods, results, discussion, conclusion,
# and references. The record's identifiers, hashes and raw event list stay in
# the ledger and the JSON export; the researcher's report leaves them out.
#
# The server builds every fact from the ledger and the protocol it ran
# (``build_report_facts``): the step tables, the results table, what was done
# differently from the source, and what still needs checking. The model writes
# only the prose around those facts, from a projection with no identifiers,
# hashes, status values or command names in it, and each of its sections is
# checked before it is used (``check_report_sections``); a section that fails
# is replaced by the server's own sentence for it. The report uses no outside
# knowledge: Google Search grounding was taken out (decision of 2026-10-06,
# option (a)) because the service terms forbid caching, storing, or rewriting
# grounded results, and a report is a stored file other people read.

REPORT_TIMEZONE_SETTING = "VOINEY_LAB_REPORT_TIMEZONE"
DEFAULT_REPORT_TIMEZONE = "Asia/Seoul"

_WEEKDAYS = "월화수목금토일"


def report_timezone(environment: Mapping[str, str] | None = None) -> ZoneInfo:
    """The researcher's time zone for the report's clock times."""

    env = os.environ if environment is None else environment
    name = env.get(REPORT_TIMEZONE_SETTING, "").strip() or DEFAULT_REPORT_TIMEZONE
    try:
        return ZoneInfo(name)
    except (ZoneInfoNotFoundError, ValueError):
        log.warning("report timezone setting is not a known zone; using %s", DEFAULT_REPORT_TIMEZONE)
        return ZoneInfo(DEFAULT_REPORT_TIMEZONE)


@dataclass(frozen=True)
class ReportSource:
    """One numbered reference: [1] is the protocol's own PDF."""

    number: int
    kind: str  # "protocol"
    title: str
    url: str = ""
    note: str = ""


@dataclass(frozen=True)
class ReportRecord:
    """One thing the researcher recorded, in their own words."""

    number: int
    step_label: str
    kind: str  # "측정" | "편차" | "관찰" | "메모" | "이상" | "사진"
    text: str  # the latest words: a correction's, when one was made
    at: str  # local HH:MM
    #: Lane N, decision 2: "정정됨" or "취소됨" once amended, and the words
    #: first recorded; the record's event key, which a correction names.
    status: str = ""
    original: str = ""
    key: str = ""
    #: Lane N, decision 3: the researcher confirmed this value for the report.
    confirmed: bool = False

    @property
    def shown(self) -> str:
        """The record as the report's table writes it."""

        if self.status == "정정됨":
            shown = f"{self.text} (정정됨 — 처음 기록 “{self.original}”)"
        elif self.status == "취소됨":
            return f"{self.text} (취소됨)"
        else:
            shown = self.text
        return f"{shown} (실험자 확인)" if self.confirmed else shown


#: A note's category -> its kind in the report (lane N, decision 1). Any other
#: observation -- an endpoint answer, an appearance -- is an observation.
RECORD_KINDS = {"measurement": "측정", "deviation": "편차", "note": "메모"}
#: The kinds counted in the report, in the order they are listed. 측정, 편차
#: and 메모 are listed only when there is one, so a record without notes reads
#: as it did before lane N.
_COUNTED_KINDS = ("측정", "편차", "관찰", "메모", "이상", "사진")
_ALWAYS_COUNTED = frozenset({"관찰", "이상", "사진"})


def record_amendments(events: Sequence[Mapping[str, Any]]) -> dict[str, dict[str, str]]:
    """{event key of a record: {"status": "정정됨"|"취소됨", "text": latest words}}.

    Lane N, decision 2: a correction or a withdrawal is its own event naming
    the record it amends; the record itself is never changed. A withdrawal
    keeps the latest words ("" when it was never corrected).
    """

    amended: dict[str, dict[str, str]] = {}
    for event in events:
        kind = str(event.get("event_type") or "")
        if kind not in {"record_corrected", "record_retracted"}:
            continue
        payload = event.get("payload") if isinstance(event.get("payload"), dict) else {}
        fix = payload.get("record_fix") if isinstance(payload.get("record_fix"), dict) else {}
        target = str(fix.get("target_event_key") or "")
        if not target:
            continue
        if kind == "record_retracted":
            amended[target] = {"status": "취소됨", "text": amended.get(target, {}).get("text", "")}
        else:
            amended[target] = {"status": "정정됨", "text": " ".join(str(fix.get("text_after") or "").split())}
    return amended


# --- The values the report holds, confirmed before they go in (lane N, decision 3)
#
# When the experiment ends the server lists the record's important values:
# each measurement, each observation with a number in it, each point done
# differently from the source (a deviation note, a later start, a confirmed
# return within a repeat, a timer ended early) and each anomaly. The
# researcher confirms them one by one by voice ("네"), corrects one (a
# correction appended as in decision 2) or leaves the rest for the screen's
# checklist. Each confirmation is its own event. A value confirmed carries
# "실험자 확인" in the report; one left unconfirmed is listed apart under
# "확인되지 않은 값", and model prose may not state it.
#
# Lane VX, decision 3: a value said "네" to when it was recorded, in the 바로
# 확인 way, is confirmed already -- its observation event says so
# (``experimenter_confirmed``) -- and is not asked again; a correction made
# after it takes the confirmation back, as it does a reviewed value's.

REVIEW_CONFIRMED = "report_value_confirmed"
REVIEW_DEFERRED = "report_review_deferred"
#: Events about the record rather than new things recorded: they do not count
#: as records added after the prose was written.
REVIEW_EVENT_TYPES = frozenset({REVIEW_CONFIRMED, REVIEW_DEFERRED})
_NUMBER_IN_TEXT = re.compile(r"\d")


def report_review_items(events: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """The values the report holds, each with whether the researcher confirmed it.

    Read from the record alone, in the order recorded. ``item_id`` is stable:
    it names the event the value came from, so a correction or a
    confirmation made later still finds it. A withdrawn record is not listed.
    """

    amended = record_amendments(events)
    confirmed: dict[str, str] = {}
    for event in events:
        if event.get("event_type") == REVIEW_CONFIRMED:
            payload = event.get("payload") if isinstance(event.get("payload"), dict) else {}
            confirmed[str(payload.get("item_id") or "")] = " ".join(str(payload.get("text") or "").split())
    items: list[dict[str, Any]] = []

    def add(
        prefix: str, event: Mapping[str, Any], kind: str, text: str, *,
        confirmed_when_recorded: bool = False, **extra: Any,
    ) -> None:
        key = str(event.get("event_key") or "")
        change = amended.get(key, {})
        if change.get("status") == "취소됨":
            return
        latest = change.get("text") or " ".join(text.split())
        item_id = f"{prefix}.{key}"
        items.append({
            "item_id": item_id,
            "kind": kind,
            "step_label": str(event.get("step_label") or ""),
            "text": latest,
            "source_key": key,
            "corrected": change.get("status") == "정정됨",
            "original": " ".join(text.split()) if change.get("status") == "정정됨" else "",
            "confirmed": confirmed.get(item_id) == latest or (
                confirmed_when_recorded and not change
            ),
            **extra,
        })

    for event in events:
        kind = str(event.get("event_type") or "")
        payload = event.get("payload") if isinstance(event.get("payload"), dict) else {}
        wording = str(event.get("user_wording") or "").strip()
        record = payload.get("step_record") if isinstance(payload.get("step_record"), Mapping) else {}
        if kind == "observation" and wording:
            category = str(event.get("category") or "")
            extra = {
                "workspace_event_key": payload.get("workspace_event_key"),
                "confirmed_when_recorded": payload.get("experimenter_confirmed") is True,
            }
            if category == "measurement":
                add("r", event, "측정", wording, **extra)
            elif category == "deviation":
                # Lane FX3, decision 5: the source values the note differs
                # from, as the event recorded them (none for a note the
                # experimenter called a deviation without a source value).
                deviations = [
                    dict(item) for item in payload.get("deviations") or ()
                    if isinstance(item, Mapping)
                ]
                add("r", event, "편차", wording, deviations=deviations, **extra)
            elif category != "note" and _NUMBER_IN_TEXT.search(
                amended.get(str(event.get("event_key") or ""), {}).get("text") or wording
            ):
                add("r", event, "관찰", _readable_record_text(wording), **extra)
        elif kind == "anomaly" and wording:
            add("r", event, "이상", wording)
        elif kind == "steps_skipped":
            skipped = [str(item) for item in record.get("skipped_step_labels") or ()]
            if skipped:
                start = str(record.get("start_step") or event.get("step_label") or "")
                add("s", event, "건너뜀", f"{start}단계부터 시작({_ranges(skipped)} 건너뜀)")
        elif kind == "repeat_returned":
            round_ = record.get("round")
            add("b", event, "돌아감",
                f"{record.get('from_step') or ''}단계에서 {record.get('to_step') or event.get('step_label') or ''}"
                f"단계로 돌아감"
                + (f" — {round_}회차(말로 확인한 돌아가기 기준)" if round_ else ""))
        elif kind == "step_completed":
            timer = payload.get("timer") if isinstance(payload.get("timer"), dict) else {}
            if timer.get("step_exited_before_timer_elapsed") or (
                timer.get("completion_state") == "step_exited_before_timer_elapsed"
            ):
                try:
                    elapsed = int(timer.get("elapsed_seconds") or 0)
                    duration = int(timer.get("source_duration_seconds") or timer.get("duration_seconds") or 0)
                except (TypeError, ValueError):
                    continue
                add("t", event, "타이머",
                    f"원문 시간 {_duration_words(duration)} 중 {_duration_words(elapsed)}에 단계를 끝냄")
    return items


def _range_pair(record: Mapping[str, Any]) -> tuple[str, str]:
    """The first and last labels of a repeat's range, as its record names them."""

    labels = [str(item) for item in record.get("repeated_step_labels") or ()]
    if not labels:
        return "", ""
    return labels[0], labels[-1]


def record_counts(records: Sequence["ReportRecord"], *, separator: str = " · ") -> str:
    """"관찰 2건 · 이상 0건 · 사진 0건", with 측정·편차·메모 when there are any."""

    live = [record for record in records if record.status != "취소됨"]
    parts = []
    for kind in _COUNTED_KINDS:
        count = sum(1 for record in live if record.kind == kind)
        if count or kind in _ALWAYS_COUNTED:
            parts.append(f"{kind} {count}건")
    return separator.join(parts)


@dataclass(frozen=True)
class ReportReturn:
    """A confirmed return within a repeat (lane R7's repeat_returned event)."""

    from_label: str
    to_label: str
    round: int | None  # counted from returns confirmed in words, not rounds done
    at: str  # local HH:MM
    key: str = ""


@dataclass(frozen=True)
class ReportRevert:
    """A confirmed revert (lane CF's step_reverted event): the completions it took back."""

    from_label: str
    to_label: str
    reverted_labels: tuple[str, ...]
    in_repeat: tuple[str, str] | None  # the repeat's first and last step, when inside one
    at: str  # local HH:MM
    key: str = ""


#: Lane CF, decision 4: an answer given before the start, in the report's words.
BEFORE_START_WORDS = "시작 전에 답함"

#: Lane CB: where a value the run acted on came from, in the report's words.
VALUE_SOURCE_WORDS = {"source": "원문", "operator": "사람이 답함", "lab_default": "연구실 기본값"}


def value_source_words(value: Any) -> str:
    return VALUE_SOURCE_WORDS.get(str(value or ""), str(value or ""))


@dataclass(frozen=True)
class ReportBranchAnswer:
    """A source condition answered on the run (lane CB's branch_answered event)."""

    step_label: str
    condition: str
    answer: str  # "예" | "아니요"
    value_source: str  # in the report's words
    at: str  # local HH:MM
    skipped_labels: tuple[str, ...] = ()
    key: str = ""


@dataclass(frozen=True)
class ReportRepetition:
    """A repeat's count given by a person (lane CB's repeat_registered / repeat_closed)."""

    first: str
    last: str
    asked_at: str  # the step the count was asked at
    count: int | None
    value_source: str  # in the report's words
    per_round: bool  # asked after each round ("아직 몰라")
    at: str  # local HH:MM
    key: str = ""


@dataclass(frozen=True)
class ReportReviewChange:
    """One item a reviewer changed or confirmed in the review version the run
    followed (lane RV, decision g)."""

    step_label: str
    kind: str  # the item's kind in Korean ("단계", "단계 타이머", ...)
    original: str  # the source's or the AI's reading; "—" when none
    #: the version's value; "(비움 · 원문에 없음)" when left empty, and
    #: "프로그램에 하는 말이라 뺌" when the words read as written to the program
    #: (lane KO2, decision 7)
    reviewed: str
    origin: str  # "검토자 입력" or "AI 제안 · 검토자 확인"
    reviewer: str
    changed: bool  # fixed, entered or left empty (not only confirmed)


#: An item a reviewer took out because its words read as written to the
#: program (lane TIDY's card, [빼 주세요]); lane KO2, decision 7.
PROGRAM_ADDRESS_CLEARED = "프로그램에 하는 말이라 뺌"

#: The review version table's heading row (lane RV, decision g).
REVIEW_CHANGE_HEADER = ("단계", "항목", "원문 · AI 제안", "검토본", "출처", "검토자")


@dataclass(frozen=True)
class ReportItem:
    """A material or a piece of equipment, as the protocol lists it."""

    kind: str  # "재료" | "장비"
    #: As listed, with its whitespace made single and without the
    #: protocols.io list labels (lane FX, decision 4: ``item_display_name``).
    name: str
    steps: tuple[str, ...]  # source steps whose text names it (decision 3)
    listed: str = ""  # as listed, with its whitespace made single


@dataclass(frozen=True)
class ReportStepFacts:
    label: str
    step_id: str
    section: str
    text: str
    source_text: str
    translated: bool
    expected: tuple[str, ...]
    source_timer_seconds: int | None
    completed: bool
    completed_at: str
    timer_note: str
    records: tuple[ReportRecord, ...]
    #: "machine" for a stored automatic translation, "reviewed" for a
    #: reviewed one, "" for the source (lane N, decision 6).
    translation: str = ""


@dataclass(frozen=True)
class StackedCell:
    """A table cell with a line below in small type: the source under its Korean."""

    main: str
    below: str

    def __str__(self) -> str:
        return f"{self.main} — 원문: {self.below}"


@dataclass(frozen=True)
class ReportFacts:
    """Everything the report states, built by the server from the record and the PDF."""

    protocol_title: str
    protocol_loaded: bool
    total_steps: int
    run_rows: tuple[tuple[str, str], ...]
    completed_labels: tuple[str, ...]
    outcome: str  # "completed" | "stopped" | "in_progress"
    stop_label: str | None
    steps: tuple[ReportStepFacts, ...]
    records: tuple[ReportRecord, ...]
    purpose_from_pdf: str
    keywords: str
    sections: tuple[str, ...]
    materials: tuple[str, ...]
    equipment: tuple[str, ...]
    deviations: tuple[str, ...]
    confirmed: tuple[str, ...]
    to_check: tuple[str, ...]
    protocol_reference: ReportSource
    zone: ZoneInfo
    items: tuple[ReportItem, ...] = ()
    #: (label, source text) of every step an item is used in, for the model's
    #: uses (decision 4).
    item_step_texts: tuple[tuple[str, str], ...] = ()
    returns: tuple[ReportReturn, ...] = ()
    reverts: tuple[ReportRevert, ...] = ()
    #: Lane N, decision 3: the values the report holds and whether each was
    #: confirmed; and the texts of those not confirmed, which model prose may
    #: not state.
    review: tuple[dict[str, Any], ...] = ()
    unconfirmed_texts: tuple[str, ...] = ()
    #: Lane CB, decision 5: the source conditions answered and the counts a
    #: person gave, each with where its value came from.
    branch_answers: tuple[ReportBranchAnswer, ...] = ()
    repetitions: tuple[ReportRepetition, ...] = ()
    #: Lane RV, decision g: the review version the run followed ("검토본 v2 ·
    #: 김재준 · 2026-10-14"; empty for an analysis as it came) and every item
    #: its reviewer changed or confirmed, against the source.
    review_version_words: str = ""
    review_changes: tuple[ReportReviewChange, ...] = ()
    #: Lane RP2, decision 1: the report's first pages, from the record.
    glance: "ReportGlance | None" = None
    #: Lane RP3, decision ①: the lab notebook the exports render, from the
    #: record, the source and the run before.
    notebook: "LabNotebook | None" = None

    @property
    def record_texts(self) -> tuple[str, ...]:
        return tuple(record.text for record in self.records)


def _review_version_facts(fixture: Any) -> tuple[str, tuple[ReportReviewChange, ...]]:
    """The review version a run followed and its changes, or ("", ()) (lane RV, g)."""

    version = getattr(getattr(fixture, "draft", None), "review_version", None)
    if version is None:
        return "", ()
    when = str(getattr(version, "registered_on", "") or version.registered_at or "")[:10]
    words = f"검토본 v{version.version} · {version.reviewer_name}" + (f" · {when}" if when else "")
    rows = []
    for change in version.changes:
        if change.action == "confirm" and change.origin != "ai_suggestion":
            continue
        reviewed = (
            # Lane KO2, decision 7: words taken out as written to the program
            # (lane TIDY's card) are in the source; they were left out for that.
            PROGRAM_ADDRESS_CLEARED
            if change.action == "clear" and change.kind != "timer"
            and program_address_shape(change.original) is not None
            else "(비움 · 원문에 없음)" if change.action == "clear" and change.kind != "timer"
            else "(타이머 없음)" if change.action == "clear"
            else change.reviewed or "—")
        rows.append(ReportReviewChange(
            step_label=change.step_label or "",
            kind=change.kind_ko,
            original=change.original or "—",
            reviewed=reviewed,
            origin="검토자 입력" if change.origin == "reviewer_input" else "AI 제안 · 검토자 확인",
            reviewer=change.decided_by or version.reviewer_name,
            changed=change.action in {"fix", "enter", "clear"},
        ))
    return words, tuple(rows)


def _source_steps_from_fixture(fixture: Any) -> list[dict[str, Any]]:
    steps: list[dict[str, Any]] = []
    timers = getattr(fixture, "timer_manifest", None) or {}
    for section in fixture.draft.protocol.sections:
        title = " ".join(str(getattr(section, "title_source_text", "") or "").split())
        for step in section.steps:
            source = " ".join(str(step.instruction_source_text or "").split())
            korean = None
            localized = getattr(fixture, "localized_fact", None)
            if callable(localized):
                korean = localized(step.step_id, "current_step")
            if korean:
                korean = re.sub(r"^\s*\d+\s*단계\s*[:：]\s*", "", " ".join(korean.split()))
            origin = getattr(fixture, "localization_source", None)
            # Lane N, decision 6: a stored machine translation is shown as one.
            translation = (origin(step.step_id, "current_step") or "") if korean and callable(origin) else ""
            expected = [" ".join(str(item.source_text).split()) for item in step.expected_results]
            whole = [source]
            for action in getattr(step, "sub_actions", ()) or ():
                whole.append(" ".join(str(getattr(action, "instruction_source_text", "") or "").split()))
                for item in getattr(action, "expected_results", ()) or ():
                    text = " ".join(str(item.source_text).split())
                    if text not in expected:
                        expected.append(text)
            steps.append({
                "label": str(step.source_label), "step_id": str(step.step_id),
                "section": title, "source_text": source,
                "text": korean or source, "translated": bool(korean),
                "translation": translation,
                "expected": tuple(expected),
                "timer": timers.get(step.step_id),
                # Lane RP2, decision 1: the source page the step is on.
                "page": getattr(getattr(step, "evidence", None), "source_page_number", None),
                # The step's whole source text, its sub-actions too: where a
                # material or piece of equipment is looked for (decision 3).
                "whole_source_text": " ".join(part for part in whole if part),
            })
    return steps


def _source_steps_from_context(context: "GroundedReportContext") -> list[dict[str, Any]]:
    return [
        {
            "label": step.step_label, "step_id": step.step_id,
            "section": step.section_title if step.section_title != "Experiment Execution" else "",
            "source_text": step.instruction_source_text, "text": step.instruction_source_text,
            "translated": False, "expected": tuple(step.expected_results), "timer": None,
        }
        for step in context.all_protocol_steps
    ]


def report_protocol_fixture(report_data: Mapping[str, Any]) -> tuple[Any, str]:
    """The protocol a report's session ran, as sessions load it, or (None, why).

    The configured development fixture when the record names it, otherwise the
    catalog's executable fixture for the protocol, accepted only when its
    revision is the one the record names. Lane RP: this lookup used to call
    ``load_executable_fixture`` with a revision argument it does not take, and
    any failure was swallowed, so the report printed placeholder steps.
    """

    protocol_id = str(report_data.get("protocol_id") or "")
    revision = str(report_data.get("protocol_revision") or "")
    try:
        from voiney_lab.server import (
            _configured_candidate_fixture,
            _open_protocol_catalog,
            _with_revision_translations,
            server_config,
        )

        candidate = _configured_candidate_fixture(server_config())
        if candidate is not None and candidate.protocol_id == protocol_id:
            fixture = candidate
        else:
            catalog, store = _open_protocol_catalog()
            try:
                # Lane RV, decision g: the revision the session ran -- an
                # earlier review version or an analysis analysed again since
                # is still read -- else what runs now.
                fixture = (
                    catalog.load_fixture_for_revision(protocol_id, revision)
                    if revision and not revision.startswith("fixture-")
                    else catalog.load_executable_fixture(protocol_id))
            finally:
                store.close()
    except Exception as exc:  # noqa: BLE001 -- reported on the report, never raised
        log.warning("report protocol lookup failed error=%s", type(exc).__name__)
        return None, "프로토콜 원문을 불러오지 못했다"
    if revision and getattr(fixture, "revision_id", revision) != revision:
        return None, "기록의 프로토콜 버전과 지금 프로토콜 버전이 달라 원문을 싣지 않았다"
    # Lane N, decision 6: the Korean a session shows for the steps -- stored
    # translations that pass their check again -- read, never started here.
    try:
        fixture = _with_revision_translations(fixture, start_missing=False)
    except Exception as exc:  # noqa: BLE001 -- the source is used instead
        log.warning("report translations not read error=%s", type(exc).__name__)
    return fixture, ""


def _local(value: Any, zone: ZoneInfo) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(zone)


def _duration_words(seconds: float | int | None) -> str:
    if seconds is None:
        return ""
    minutes = int(round(float(seconds) / 60))
    hours, minutes = divmod(max(minutes, 0), 60)
    if hours and minutes:
        return f"{hours}시간 {minutes}분"
    if hours:
        return f"{hours}시간"
    return f"{minutes}분"


def _label_runs(labels: Sequence[str]) -> str:
    """'1, 2, 3, 5' -> '1–3, 5' for numeric step labels."""

    numbers: list[int] = []
    for label in labels:
        if not str(label).isdigit():
            return ", ".join(str(item) for item in labels)
        numbers.append(int(label))
    if not numbers:
        return ""
    numbers = sorted(set(numbers))
    parts: list[str] = []
    start = previous = numbers[0]
    for number in numbers[1:] + [None]:  # type: ignore[list-item]
        if number is not None and number == previous + 1:
            previous = number
            continue
        parts.append(f"{start}–{previous}" if previous != start else f"{start}")
        if number is not None:
            start = previous = number
    return ", ".join(parts)


def _ranges(labels: Sequence[str]) -> str:
    """'1, 2, 3, 5' -> '1–3, 5단계' for numeric step labels."""

    runs = _label_runs(labels)
    return runs + "단계" if runs else ""


# --- Where each material and piece of equipment is used (decision 3) -----------
#
# The server looks for an item's name in each source step's text, and lists
# the steps it appears in; nothing is inferred. Names are looked for in this
# order, and the first that the steps hold is used:
#   1. the name as listed, without what is in parentheses, a catalog number,
#      or what follows a comma ("Formic acid, LC-MS grade" -> "Formic acid");
#   2. if no step holds that, the name without its company and grade words:
#      capitalized words, codes with digits and grade marks at its end
#      ("Ammonium bicarbonate Merck MilliporeSigma" -> "Ammonium bicarbonate"),
#      and a capitalized word in front of a lower-case one ("Promega trypsin"
#      -> "trypsin"). A name left with one word is used only when no other
#      listed item has that word, so "Glass pipettes" is not "pipettes".
# The last word may be singular or plural. Besides its name, an item is also
# found by the short name the source calls it: one in its own parentheses
# ("Lysogeny Broth (LB)") or one a step writes right after its name
# ("ammonium bicarbonate (AMBIC)"). Where items share a first word and a
# short name ("Lysogeny Broth (LB)", "Lysogeny Agar"), the word after the
# short name tells them apart ("LB agar" is the agar).

#: Words in a listed name that say its grade or how it is listed, not what it is.
_GRADE_WORDS = frozenset({"lc-ms", "hplc", "grade", "catalog", "cat.", "model", "name",
                          "type", "brand", "sku"})
_SHORT_NAME = re.compile(r"[A-Z][A-Z0-9]{1,7}")
_ASCII_WORD = "A-Za-z0-9"


def _word_forms(word: str) -> set[str]:
    """A word and its singular or plural, compared without case."""

    lower = word.casefold()
    forms = {lower, lower + "s", lower + "es"}
    if lower.endswith("es") and len(lower) > 4:
        forms.add(lower[:-2])
    if lower.endswith("s") and len(lower) > 3:
        forms.add(lower[:-1])
    return forms


def _name_pattern(words: Sequence[str]) -> re.Pattern[str]:
    *head, last = words
    last_forms = sorted((re.escape(form) for form in _word_forms(last)), key=len, reverse=True)
    body = r"\s+".join([re.escape(word) for word in head] + [f"(?:{'|'.join(last_forms)})"])
    return re.compile(rf"(?<![{_ASCII_WORD}]){body}(?![{_ASCII_WORD}])", re.I)


#: What a protocols.io materials or equipment list writes after an item's
#: name, and the analysis copies into it (lane FX, decision 4): the labels
#: NAME, TYPE, BRAND and SKU, each after its value, and a long number code
#: at the end. Labels are upper case: "B&J Brand" is a company's name.
_LIST_LABEL = re.compile(r"(?<![A-Za-z0-9])(?:NAME|TYPE|BRAND|SKU)(?![A-Za-z0-9])")
_TRAILING_NUMBER_CODE = re.compile(r"\s+\d{6,}$")


def item_display_name(name: str) -> str:
    """A listed name as the report shows it, without the list's labels.

    "Eppendorf Thermomixer C Model 5382 NAME Thermomixer C TYPE Eppendorf
    BRAND 5382000023 SKU" -> "Eppendorf Thermomixer C Model 5382": the name
    ends before the first label after its first word, and a number code of
    six or more digits at its end is left out ("Falcon 50 mL conical tube
    352070" -> "Falcon 50 mL conical tube"). Display only: the item keeps
    the name as listed (``ReportItem.listed``).
    """

    text = " ".join(str(name).split())
    label = next((match for match in _LIST_LABEL.finditer(text) if match.start() > 0), None)
    if label is not None:
        text = text[:label.start()].rstrip()
    shorter = _TRAILING_NUMBER_CODE.sub("", text)
    return shorter if shorter.strip() else text


def _looked_for(name: str) -> str:
    """What an item's steps are looked for by: the name as shown when the
    list's labels or number code were left out of it, otherwise its first
    line, as before (lane FX, decision 4)."""

    shown = item_display_name(name)
    if shown != " ".join(str(name).split()):
        return shown
    return next((line for line in str(name).splitlines() if line.strip()), "")


def _listed_words(name: str) -> list[str]:
    """The name's words without parentheses, a catalog number, or what follows a comma."""

    first_line = next((line for line in str(name).splitlines() if line.strip()), "")
    text = re.sub(r"\([^)]*\)", " ", first_line.replace("’", "'"))
    text = re.split(r"\bCat(?:alog)?\.?\s*#|#|,", text)[0]
    return text.split()


def _core_words(words: Sequence[str]) -> list[str]:
    """The name without its company and grade words (see above)."""

    core = list(words)
    while len(core) > 1 and (
        core[-1][:1].isupper() or any(ch.isdigit() for ch in core[-1])
        or core[-1].casefold() in _GRADE_WORDS
    ):
        core.pop()
    while len(core) > 1 and core[0][:1].isupper() and core[1][:1].islower():
        core.pop(0)
    return core


def item_step_labels(
    names: Sequence[tuple[str, str]], source_steps: Sequence[Mapping[str, Any]],
) -> tuple[ReportItem, ...]:
    """Each listed (kind, name) with the source steps whose text names it."""

    texts = [(str(step["label"]), str(step.get("whole_source_text") or step.get("source_text") or "")
              .replace("’", "'")) for step in source_steps]
    looked_for = [_looked_for(name) for _, name in names]
    words_of = [_listed_words(name) for name in looked_for]
    other_words = [
        {form for j, words in enumerate(words_of) if j != i for word in words for form in _word_forms(word)}
        for i in range(len(names))
    ]
    found: list[set[str]] = [set() for _ in names]
    patterns: list[re.Pattern[str] | None] = []
    for index, words in enumerate(words_of):
        pattern = None
        core = _core_words(words)
        for candidate, shortened in ((words, False), (core, core != words)):
            if not candidate or (shortened and len(candidate) == 1
                                 and _word_forms(candidate[0]) & other_words[index]):
                continue
            compiled = _name_pattern(candidate)
            if any(compiled.search(text) for _, text in texts):
                pattern = compiled
                break
        patterns.append(pattern)
        if pattern is not None:
            found[index] |= {label for label, text in texts if pattern.search(text)}

    for index, name in enumerate(looked_for):
        first_line = next((line for line in str(name).splitlines() if line.strip()), "")
        short = {match.group(1) for match in re.finditer(r"\(\s*([^()]*?)\s*\)", first_line)
                 if _SHORT_NAME.fullmatch(match.group(1))}
        if patterns[index] is not None:
            for _, text in texts:
                for match in patterns[index].finditer(text):
                    after = re.match(r"\s*\(\s*([A-Z][A-Z0-9]{1,7})(?![A-Za-z0-9])", text[match.end():])
                    if after:
                        short.add(after.group(1))
        words = words_of[index]
        siblings = [
            (j, other[-1]) for j, other in enumerate(words_of)
            if j != index and other and words and other[0].casefold() == words[0].casefold()
        ]
        for abbreviation in short:
            use = re.compile(rf"(?<![{_ASCII_WORD}]){re.escape(abbreviation)}(?![{_ASCII_WORD}])\s*([^\s,.;:()]*)")
            for label, text in texts:
                for match in use.finditer(text):
                    owner = next((j for j, last in siblings
                                  if match.group(1).casefold() in _word_forms(last)), index)
                    found[owner].add(label)
    order = {label: position for position, (label, _) in enumerate(texts)}
    return tuple(
        ReportItem(kind=kind, name=item_display_name(name),
                   steps=tuple(sorted(found[index], key=lambda label: order.get(label, 0))),
                   listed=" ".join(str(name).split()))
        for index, (kind, name) in enumerate(names)
    )


_ENDPOINT_ANSWER = re.compile(r"^원문 종점 “(?P<quote>.*)” — 답: (?P<answer>.+)$", re.S)


def _readable_record_text(text: str) -> str:
    """An endpoint answer the rules stored as one line, in plain words."""

    match = _ENDPOINT_ANSWER.match(text.strip())
    if match is None:
        return text.strip()
    quote = " ".join(match.group("quote").split())
    if len(quote) > 90:
        quote = quote[:88].rstrip() + "…"
    answer = match.group("answer").strip()
    verdict = "충족했다고 답했다" if answer in {"네", "응", "예", "yes"} else "아직 충족하지 않았다고 답했다"
    return f"원문의 끝 조건(“{quote}”)을 {verdict}"


def _protocol_state_words(report_data: Mapping[str, Any]) -> str:
    """The run row "프로토콜 상태" (decision 7 of 2026-10-08, lane DI).

    Under the MVP rule a protocol runs when its analysis passed and the
    experimenter pressed start; that is what the row says. The curated
    development fixture is named as such: its analysis was hand-built over
    the source PDF, not produced by the analysis model.
    """

    words = "분석 통과 · 실험자가 시작함"
    if report_data.get("development_only"):
        words += " · 개발용 큐레이션 분석(원문 PDF 기준)"
    return words


def _pdf_front_matter(fixture: Any) -> tuple[str, str, str, str]:
    """(purpose sentence, keywords, DOI url, author line) from the PDF's first page."""

    try:
        pages = fixture.draft.protocol.metadata.pdf.pages
        first = pages[0].text if pages else ""
    except AttributeError:
        return "", "", "", ""
    text = first.replace("\xa0", " ")
    keywords = ""
    match = re.search(r"Keywords:\s*(.+?)(?:\n[A-Za-z ]+:|\nprotocols\.io|\Z)", text, re.S)
    if match:
        keywords = " ".join(match.group(1).split())
    doi = ""
    match = re.search(r"https?://(?:dx\.)?doi\.org/\S+", text)
    if match:
        doi = match.group(0).rstrip(".,)")
    author = ""
    match = re.search(r"Protocol Citation:\s*(.+?)(?:\n|$)", text)
    if match:
        author = " ".join(match.group(1).split())
    purpose = ""
    match = re.search(r"(?:Abstract|Purpose|Aim|Objective)s?\s*\n(.+?)(?:\n\s*\n|\Z)", text, re.S)
    if match:
        purpose = " ".join(match.group(1).split())[:600]
    return purpose, keywords, doi, author


_SECTION_DURATION = re.compile(r"\s+\d+\s*(?:m|min|h|hr|s|d)\s*$", re.I)


def build_report_facts(
    report_data: Mapping[str, Any],
    events: Sequence[Mapping[str, Any]] | None = None,
    *,
    fixture: Any = None,
    context: "GroundedReportContext | None" = None,
    zone: ZoneInfo | None = None,
    protocol_problem: str = "",
    start_check: Sequence[Mapping[str, Any]] | None = None,
    previous: "ReportFacts | None" = None,
) -> ReportFacts:
    """The report's facts from the record and the protocol the session ran.

    ``start_check`` is the answers given in the check before the start, as
    the protocol's ledger keeps them (lane RP2), or None when not read.
    ``previous`` is the facts of the finished run of the same protocol
    before this one (lane RP3), or None when there is none.
    """

    zone = zone or report_timezone()
    events = list(report_data.get("events") or ()) if events is None else list(events)
    if fixture is not None:
        source_steps = _source_steps_from_fixture(fixture)
    elif context is not None:
        source_steps = _source_steps_from_context(context)
    else:
        source_steps = []
    by_label = {step["label"]: step for step in source_steps}
    order = {step["label"]: index for index, step in enumerate(source_steps)}

    completed_at: dict[str, datetime | None] = {}
    advanced: list[str] = []
    timer_started: dict[str, tuple[datetime | None, int | None]] = {}
    early: dict[str, tuple[int, int]] = {}
    records: list[ReportRecord] = []
    pauses: list[tuple[str, datetime | None, datetime | None]] = []
    stop_label: str | None = None
    workflow_completed = False
    blocked: list[str] = []
    experimenter = ""
    #: The source's safety statements the experimenter saw before pressing
    #: start (decision of 2026-10-08): how many, from the server's record.
    safety_notices_acknowledged: int | None = None
    # Lane R7's events (decision 6): every completion with its round, every
    # timer start, a later start's skipped steps, and returns within a repeat.
    completions: dict[str, list[tuple[datetime | None, int | None]]] = {}
    timer_starts: dict[str, list[tuple[datetime | None, int]]] = {}
    skip: tuple[str, list[str], str] | None = None
    early_keys: dict[str, str] = {}
    returns: list[ReportReturn] = []
    # Lane CF, decision 3: reverts; decision 5: the file the run followed.
    reverts: list[ReportRevert] = []
    basis_words = ""
    source_disclosure_words = ""
    source_note = ""
    # Lane CB, decision 5.
    branch_answers: list[ReportBranchAnswer] = []
    repetitions: list[ReportRepetition] = []
    declined_rounds: list[tuple[str, Mapping[str, Any], str, str]] = []
    step_label_by_id = {step["step_id"]: step["label"] for step in source_steps if step.get("step_id")}

    def step_record(payload: Mapping[str, Any]) -> Mapping[str, Any]:
        record = payload.get("step_record")
        return record if isinstance(record, Mapping) else {}

    def whole_number(value: Any) -> int | None:
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    # Lane N, decision 2: corrections and withdrawals, by the record they amend.
    amendments = {
        key: (value["status"], value["text"]) for key, value in record_amendments(events).items()
    }
    # Lane N, decision 3: the values the report holds, and which were confirmed.
    review = report_review_items(events)
    # A value is "not confirmed" once the researcher was asked about it (by
    # voice, or on the screen) and has not confirmed it; before any review a
    # record makes no claim either way.
    review_offered = any(str(event.get("event_type") or "") in REVIEW_EVENT_TYPES for event in events)
    confirmed_keys = {item["source_key"] for item in review if item["confirmed"]}
    unconfirmed_keys = (
        {item["source_key"] for item in review if not item["confirmed"]} if review_offered else set()
    )
    unconfirmed_texts: list[str] = []

    def add_record(label: str, kind: str, text: str, at: datetime | None, key: str = "") -> None:
        first = _readable_record_text(text)
        status, latest = amendments.get(key, ("", "")) if key else ("", "")
        records.append(ReportRecord(
            number=len(records) + 1, step_label=label, kind=kind,
            text=latest or first, at=at.strftime("%H:%M") if at else "",
            status=status, original=first if status else "", key=key,
            confirmed=bool(key) and key in confirmed_keys,
        ))
        if key and key in unconfirmed_keys:
            unconfirmed_texts.append(latest or first)

    def point(line: str, key: str) -> str:
        """A point done differently, as corrected and confirmed (decisions 2-3)."""

        status, latest = amendments.get(key, ("", "")) if key else ("", "")
        if status == "정정됨" and latest:
            line = f"{latest} (정정됨 — 처음 기록 “{line.rstrip('.')}”)."
        if key and key in unconfirmed_keys:
            unconfirmed_texts.append(line)
        if key and key in confirmed_keys:
            line = (line[:-1] if line.endswith(".") else line) + " (실험자 확인)."
        return line

    for event in events:
        kind = str(event.get("event_type") or "")
        label = str(event.get("step_label") or "")
        at = _local(event.get("created_at"), zone)
        payload = event.get("payload") if isinstance(event.get("payload"), dict) else {}
        wording = str(event.get("user_wording") or "").strip()
        legacy = [str(item).strip() for item in payload.get("observations") or () if str(item).strip()]
        if kind == "step_completed" and label:
            completed_at.setdefault(label, at)
            record = step_record(payload)
            completions.setdefault(label, []).append((at, whole_number(record.get("round")) if (
                record.get("kind") == "repeat_round_completion") else None))
            if wording:
                add_record(label, "관찰", wording, at)
            timer = payload.get("timer") if isinstance(payload.get("timer"), dict) else {}
            if timer.get("step_exited_before_timer_elapsed") or (
                timer.get("completion_state") == "step_exited_before_timer_elapsed"
            ):
                try:
                    early[label] = (
                        int(timer.get("elapsed_seconds") or 0),
                        int(timer.get("source_duration_seconds") or timer.get("duration_seconds") or 0),
                    )
                    early_keys[label] = str(event.get("event_key") or "")
                except (TypeError, ValueError):
                    pass
        elif kind == "step_advanced" and label:
            advanced.append(label)
        elif kind == "observation" and label:
            add_record(label, RECORD_KINDS.get(str(event.get("category") or ""), "관찰"),
                       wording or str(payload.get("text") or "관찰 기록"), at,
                       str(event.get("event_key") or ""))
        elif kind == "anomaly" and label:
            add_record(label, "이상", wording or str(payload.get("text") or "이상 보고"), at,
                       str(event.get("event_key") or ""))
        elif kind == "photo_attached" and label:
            add_record(label, "사진", f"사진 첨부 — {wording}" if wording else "사진 첨부", at)
        elif kind == "timer_started" and label:
            timer = payload.get("timer") if isinstance(payload.get("timer"), dict) else {}
            seconds = timer.get("duration_seconds") or timer.get("source_duration_seconds")
            try:
                seconds = int(seconds) if seconds else None
            except (TypeError, ValueError):
                seconds = None
            if seconds:
                timer_started[label] = (at, seconds)
                timer_starts.setdefault(label, []).append((at, seconds))
        elif kind == "steps_skipped" and skip is None:
            record = step_record(payload)
            skipped = [str(item) for item in record.get("skipped_step_labels") or ()]
            if skipped:
                skip = (str(record.get("start_step") or label), skipped, str(event.get("event_key") or ""))
        elif kind == "repeat_returned":
            record = step_record(payload)
            returns.append(ReportReturn(
                from_label=str(record.get("from_step") or ""), to_label=str(record.get("to_step") or label),
                round=whole_number(record.get("round")), at=at.strftime("%H:%M") if at else "",
                key=str(event.get("event_key") or ""),
            ))
        elif kind == "step_reverted":
            record = step_record(payload)
            span = record.get("in_repeat") if isinstance(record.get("in_repeat"), Mapping) else None
            pair = tuple(str(item) for item in (span or {}).get("repeated_step_labels") or ())
            reverts.append(ReportRevert(
                from_label=str(record.get("from_step") or ""), to_label=str(record.get("to_step") or label),
                reverted_labels=tuple(str(item) for item in record.get("reverted_step_labels") or ()),
                in_repeat=(pair[0], pair[-1]) if pair else None,
                at=at.strftime("%H:%M") if at else "", key=str(event.get("event_key") or ""),
            ))
        elif kind == "protocol_basis_recorded":
            if not basis_words:
                filename = " ".join(str(payload.get("filename") or "").split())
                words = " ".join(str(payload.get("words") or "").split())
                basis_words = words.removeprefix("지금 기준:").strip() or filename
            disclosure = payload.get("source_disclosure")
            if not source_disclosure_words and isinstance(disclosure, Mapping):
                if disclosure.get("recorded") is False:
                    source_disclosure_words = "확인 기록 없음"
                elif disclosure.get("recorded") is True:
                    source_disclosure_words = {
                        "yes": "예 — 공개·시판 원문",
                        "no": "아니요 — 공개·시판 아님",
                        "unknown": "모름 — 확인 필요",
                    }.get(str(disclosure.get("status") or ""), "")
                    source_note = " ".join(
                        str(disclosure.get("note") or "").split()
                    )
        elif kind == "branch_answered":
            record = step_record(payload)
            branch_answers.append(ReportBranchAnswer(
                step_label=str(record.get("step_label") or label),
                condition=" ".join(str(record.get("condition_source_text") or "").split()),
                answer="예" if record.get("answer") == "yes" else "아니요",
                value_source=value_source_words(record.get("value_source")) + (
                    f" · {BEFORE_START_WORDS}" if record.get("asked") == "before_start" else ""),
                at=at.strftime("%H:%M") if at else "",
                skipped_labels=tuple(str(item) for item in record.get("skipped_step_labels") or ()),
                key=str(event.get("event_key") or ""),
            ))
        elif kind == "repeat_registered":
            record = step_record(payload)
            first, last = _range_pair(record)
            repetitions.append(ReportRepetition(
                first=first, last=last, asked_at=label, count=whole_number(record.get("count")),
                value_source=value_source_words(record.get("value_source")) + (
                    f" · {BEFORE_START_WORDS}" if record.get("asked") == "before_start" else ""),
                per_round=record.get("decided") == "per_round",
                at=at.strftime("%H:%M") if at else "", key=str(event.get("event_key") or ""),
            ))
        elif kind == "repeat_closed":
            record = step_record(payload)
            first, last = _range_pair(record)
            closed = ReportRepetition(
                first=first, last=last, asked_at=label, count=whole_number(record.get("count")),
                value_source=value_source_words(record.get("value_source")), per_round=True,
                at=at.strftime("%H:%M") if at else "", key=str(event.get("event_key") or ""),
            )
            earlier = next(
                (index for index, item in enumerate(repetitions)
                 if (item.first, item.last) == (first, last) and item.count is None), None,
            )
            if earlier is None:
                repetitions.append(closed)
            else:
                repetitions[earlier] = _replace_dataclass(
                    closed, asked_at=repetitions[earlier].asked_at, at=repetitions[earlier].at)
        elif kind == "repeat_rounds_declined":
            record = step_record(payload)
            declined_rounds.append((
                label, record, str(event.get("event_key") or ""),
                step_label_by_id.get(str(payload.get("post_transition_step_id") or ""), ""),
            ))
        elif kind == "workflow_paused":
            pauses.append((label, at, None))
        elif kind == "workflow_resumed" and pauses and pauses[-1][2] is None:
            pauses[-1] = (pauses[-1][0], pauses[-1][1], at)
        elif kind == "session_stopped":
            stop_label = label or None
        elif kind == "workflow_completed":
            workflow_completed = True
        elif kind == "blocked" and label:
            blocked.append(label)
        elif kind == "experimenter_recorded" and not experimenter:
            # The person who started the experiment, as the server knew them
            # (decision 5): a development account's name today, the signed-in
            # person's once sign-in is attached.
            experimenter = " ".join(str(payload.get("display_name") or "").split())[:120]
        elif kind == "safety_notices_acknowledged":
            count = payload.get("notice_count")
            safety_notices_acknowledged = count if isinstance(count, int) else 0
        for text in legacy:
            add_record(label or "—", "관찰", text, at)

    started = _local(report_data.get("started_at"), zone)
    ended = _local(report_data.get("ended_at"), zone)
    status = str(report_data.get("status") or "in_progress")
    outcome = "completed" if workflow_completed or status == "completed" else (
        "stopped" if status in {"stopped", "blocked", "incomplete"} else "in_progress")
    completed_labels = sorted(completed_at, key=lambda item: order.get(item, 10_000 + len(item)))
    total = len(source_steps) or len(completed_labels)

    def in_rounds(label: str) -> bool:
        return any(round_ is not None for _, round_ in completions.get(label, ()))

    def completion_words(label: str) -> str:
        """When the step was completed; in each round once a return was confirmed."""

        done = completions.get(label) or []
        if not in_rounds(label):
            first = done[0][0] if done else None
            return first.strftime("%H:%M") if first else ""
        return ", ".join(f"{round_ or 1}회차 {at:%H:%M}" if at else f"{round_ or 1}회차"
                         for at, round_ in done)

    def timer_actual(label: str) -> str:
        """Each completion against the last timer start before it, after the completion before."""

        starts = [at for at, _ in timer_starts.get(label, ()) if at is not None]
        parts: list[str] = []
        previous: datetime | None = None
        for done, round_ in completions.get(label, ()):
            if done is None:
                continue
            started = [at for at in starts if at <= done and (previous is None or at >= previous)]
            previous = done
            if started:
                words = _duration_words((done - started[-1]).total_seconds())
                parts.append(f"{round_ or 1}회차 {words}" if in_rounds(label) else words)
        return ", ".join(parts)

    steps: list[ReportStepFacts] = []
    shown = list(completed_labels)
    for record in records:
        if record.step_label not in shown and record.step_label != "—":
            shown.append(record.step_label)
    for back in returns:
        if back.from_label and back.from_label not in shown:
            shown.append(back.from_label)
    for answer in branch_answers:
        if answer.step_label and answer.step_label not in shown:
            shown.append(answer.step_label)
    for repetition in repetitions:
        # Lane CF, decision 4: a count given before the start is shown at its
        # step even when the run did not reach it.
        if repetition.asked_at and repetition.asked_at in order and repetition.asked_at not in shown:
            shown.append(repetition.asked_at)
    shown.sort(key=lambda item: order.get(item, 10_000))
    for label in shown:
        source = by_label.get(label, {})
        note = ""
        defined = source.get("timer")
        if label in early:
            elapsed, duration = early[label]
            note = f"원문 {_duration_words(duration or defined)} / {_duration_words(elapsed)}에 끝냄"
        elif label in timer_started:
            seconds = timer_started[label][1]
            actual = timer_actual(label)
            note = f"원문 {_duration_words(seconds)} / 실제 {actual}" if actual else f"원문 {_duration_words(seconds)}"
        elif defined:
            note = f"원문 {_duration_words(defined)} / 타이머 기록 없음"
        steps.append(ReportStepFacts(
            label=label, step_id=str(source.get("step_id") or ""),
            section=_SECTION_DURATION.sub("", str(source.get("section") or "")).strip(),
            text=str(source.get("text") or ""), source_text=str(source.get("source_text") or ""),
            translated=bool(source.get("translated")), expected=tuple(source.get("expected") or ()),
            source_timer_seconds=defined, completed=label in completed_at,
            completed_at=completion_words(label),
            timer_note=note, records=tuple(r for r in records if r.step_label == label),
            translation=str(source.get("translation") or ""),
        ))

    # What was done differently from the source, from the record only.
    deviations: list[str] = []
    if skip is not None:
        # Decision 6: a confirmed later start, as its event records it.
        deviations.append(point(f"{skip[0]}단계부터 시작({_ranges(skip[1])} 건너뜀).", skip[2]))
    elif source_steps and completed_labels:
        # No event says so: inferred from the first step recorded.
        first = min(order.get(label, 0) for label in completed_labels)
        before = [step["label"] for step in source_steps[:first]]
        if before:
            deviations.append(f"{_ranges(before)}는 이 기록에 없다({completed_labels[0]}단계부터 기록됨).")
    for back in returns:
        # The round counts returns confirmed in words, not rounds done at the bench.
        round_words = f"{back.round}회차(말로 확인한 돌아가기 기준)" if back.round else "회차 기록 없음"
        deviations.append(point(f"{back.from_label}단계에서 {back.to_label}단계로 돌아갔다 — {round_words}"
                                + (f", {back.at}." if back.at else "."), back.key))
    for back in reverts:
        # Lane CF, decision 3: the completions taken back stay in the record.
        labels = back.reverted_labels
        span = labels[0] if len(labels) == 1 else f"{labels[0]}~{labels[-1]}" if labels else ""
        taken = f"{span}단계 완료 취소(처음 완료 기록은 남김)" if labels else "완료 취소 없음"
        where = (f"반복 구간 {back.in_repeat[0]}~{back.in_repeat[1]}단계 안"
                 if back.in_repeat else "반복 구간 밖")
        deviations.append(point(
            f"{back.from_label}단계에서 {back.to_label}단계로 되돌아갔다 — {taken}, {where}"
            + (f", {back.at}." if back.at else "."), back.key))
    for label, record, key, next_label in declined_rounds:
        # Lane CB, decision 3: the rounds the source (or the person) asked
        # for and the person stopped short of; the round by confirmed returns.
        first, last = _range_pair(record)
        required = record.get("rounds_required")
        done = record.get("rounds_done")
        asked = (
            f"원문은 {first}~{last}단계를 {required}회 하라고 했으나"
            if record.get("count_source") == "source"
            else f"{first}~{last}단계를 {required}회 하기로 했으나({value_source_words(record.get('count_source'))})"
        )
        where = f" {next_label}단계로" if next_label else " 다음 단계로"
        deviations.append(point(
            f"{label}단계: {asked} {done}회차까지만 하고{where} 넘어갔다(회차는 말로 확인한 돌아가기 기준).",
            key,
        ))
    for label in completed_labels:
        if label in early:
            elapsed, duration = early[label]
            deviations.append(point(
                f"{label}단계: 원문 시간 {_duration_words(duration)} 중 {_duration_words(elapsed)}에 단계를 끝냈다(타이머를 일찍 끝냄).",
                early_keys.get(label, ""),
            ))
    for record in records:
        if record.status == "취소됨":
            continue
        if record.kind == "편차" or (record.kind == "관찰" and "반복" in record.text):
            line = f"{record.step_label}단계: 연구자 기록 — “{record.text}”."
            if record.key and record.key in unconfirmed_keys:
                unconfirmed_texts.append(line)
            deviations.append(
                (line[:-1] + " (실험자 확인).") if record.confirmed and record.kind == "편차" else line)
    for label, paused_at, resumed_at in pauses:
        if paused_at and resumed_at:
            deviations.append(
                f"{label}단계에서 {_duration_words((resumed_at - paused_at).total_seconds())} 동안 멈췄다가 다시 진행했다"
                f"({paused_at:%H:%M}–{resumed_at:%H:%M})."
            )
        elif paused_at:
            deviations.append(f"{label}단계에서 {paused_at:%H:%M}에 멈춘 뒤 다시 시작한 기록이 없다.")
    for label in advanced:
        deviations.append(f"{label}단계: 완료를 말하지 않고 다음 단계로 넘어갔다.")
    if outcome == "stopped" and source_steps:
        after = [step["label"] for step in source_steps
                 if step["label"] not in completed_at
                 and order.get(step["label"], 0) > max((order.get(l, -1) for l in completed_labels), default=-1)]
        if after:
            where = f"{stop_label}단계에서 " if stop_label else ""
            deviations.append(f"{where}실험을 끝내 {_ranges(after)}는 수행하지 않았다.")

    live = [r for r in records if r.status != "취소됨"]
    observations = [r for r in live if r.kind == "관찰"]
    anomalies = [r for r in live if r.kind == "이상"]
    photos = [r for r in live if r.kind == "사진"]
    confirmed: list[str] = []
    if completed_labels:
        confirmed.append(f"완료로 기록된 단계는 {_ranges(completed_labels)}이다.")
    for record in live:
        if record.key and record.key in unconfirmed_keys:
            # Lane N, decision 3: listed apart, under "확인되지 않은 값".
            continue
        if record.kind == "관찰" and record.text.startswith("원문의 끝 조건"):
            confirmed.append(f"{record.step_label}단계: 연구자가 {record.text}.")
        elif record.kind == "관찰":
            confirmed.append(f"{record.step_label}단계에서 연구자가 “{record.text}”라고 기록했다.")
        elif record.kind in {"측정", "메모"}:
            confirmed.append(f"{record.step_label}단계에서 연구자가 “{record.text}”라고 기록했다({record.kind}).")
    waited = [label for label in completed_labels if label in timer_started and label not in early]
    if waited:
        confirmed.append(f"타이머를 켠 {_ranges(waited)}는 타이머를 켠 뒤 완료했다(실제 걸린 시간은 3-2의 표).")
    to_check: list[str] = []
    if fixture is None:
        to_check.append(f"{protocol_problem or '프로토콜 원문을 불러오지 못했다'} — 단계 원문과 원문 조건을 이 보고서에 싣지 못했다.")
    for step in steps:
        if step.completed and step.expected and not any(r.kind == "관찰" for r in step.records):
            quote = step.expected[0] if len(step.expected[0]) <= 90 else step.expected[0][:88].rstrip() + "…"
            to_check.append(f"{step.label}단계: 원문은 끝 조건(“{quote}”)을 적고 있으나 관찰이 기록되지 않았다.")
        if step.completed and step.source_timer_seconds and step.label not in timer_started:
            to_check.append(f"{step.label}단계: 원문 시간({_duration_words(step.source_timer_seconds)})이 있으나 타이머 기록이 없다.")
        if step.label in early:
            to_check.append(f"{step.label}단계: 타이머를 일찍 끝낸 것이 결과에 영향을 주었는지 확인한다.")
    for record in anomalies:
        to_check.append(f"{record.step_label}단계 이상(“{record.text}”) 뒤에 어떻게 처리했는지 기록되지 않았다.")
    for record in photos:
        to_check.append(f"{record.step_label}단계 사진은 시스템이 해석하지 않았다 — 사진에서 본 것을 연구자가 적는다.")
    for label in blocked:
        to_check.append(f"{label}단계에서 진행이 막힌 기록이 있다.")

    if started and ended and ended.date() != started.date():
        date_words = (f"{started.year}년 {started.month}월 {started.day}일 ({_WEEKDAYS[started.weekday()]}) – "
                      f"{ended.month}월 {ended.day}일 ({_WEEKDAYS[ended.weekday()]})")
        end_words = f"{ended.month}월 {ended.day}일 {ended:%H:%M}"
    else:
        date_words = (f"{started.year}년 {started.month}월 {started.day}일 ({_WEEKDAYS[started.weekday()]})"
                      if started else "기록 없음")
        end_words = f"{ended:%H:%M}" if ended else "기록 없음 (진행 중)"
    if outcome == "completed":
        outcome_words = "끝까지 완료"
    elif outcome == "stopped":
        outcome_words = f"중단 — {stop_label}단계에서 종료" if stop_label else "중단"
    else:
        outcome_words = "진행 중"
    zone_name = "한국 시간 (Asia/Seoul)" if zone.key == "Asia/Seoul" else zone.key
    run_rows = (
        ("프로토콜", str(report_data.get("protocol_title") or "")),
        ("날짜", date_words),
        ("시작", f"{started:%H:%M}" if started else "기록 없음"),
        ("끝", end_words),
        ("시간 기준", zone_name),
        ("실험자", experimenter),
        ("걸린 시간", _duration_words((ended - started).total_seconds()) if started and ended else "진행 중"),
        ("완료 단계", f"{len(completed_labels)} / {total}" if total else str(len(completed_labels))),
        ("결과", outcome_words),
        ("기록", record_counts(records)),
        ("프로토콜 상태", _protocol_state_words(report_data)),
    ) + (
        # Lane CF, decision 5: "{파일 이름} · {올린 날짜}".
        (("기준 파일", basis_words),) if basis_words else ()
    ) + (
        (("원문 공개 범위", source_disclosure_words),)
        if source_disclosure_words else ()
    ) + (
        (("출처 메모", source_note),) if source_note else ()
    ) + (
        (("안전 주의 확인", (
            f"시작 전 화면에서 원문 안전 주의 {safety_notices_acknowledged}건을 보고 시작함"
            if safety_notices_acknowledged
            else "원문에 안전 주의가 없음 · 시작 전 화면을 보고 시작함")),)
        if safety_notices_acknowledged is not None else ()
    )

    review_version_words, review_changes = _review_version_facts(fixture)
    if review_version_words:
        run_rows += (("검토본", review_version_words),)
    purpose, keywords, doi, author = _pdf_front_matter(fixture) if fixture is not None else ("", "", "", "")
    protocol = getattr(getattr(fixture, "draft", None), "protocol", None)
    materials = tuple(" ".join(m.name_source_text.split()) for m in getattr(protocol, "materials", ()) or ())
    equipment = tuple(" ".join(e.name_source_text.split()) for e in getattr(protocol, "equipment", ()) or ())
    items = item_step_labels(
        [("재료", str(m.name_source_text)) for m in getattr(protocol, "materials", ()) or ()]
        + [("장비", str(e.name_source_text)) for e in getattr(protocol, "equipment", ()) or ()],
        source_steps,
    )
    sections: list[str] = []
    for step in source_steps:
        title = _SECTION_DURATION.sub("", step["section"]).strip()
        if title and title not in sections:
            sections.append(title)
    title = str(report_data.get("protocol_title") or "")
    reference = ReportSource(
        number=1, kind="protocol",
        title=f"{author}" if author else f"{title} (실험 PDF)",
        url=doi, note="실험에 쓴 프로토콜 원문 PDF",
    )
    facts = ReportFacts(
        protocol_title=title, protocol_loaded=bool(source_steps), total_steps=total,
        run_rows=run_rows, completed_labels=tuple(completed_labels), outcome=outcome,
        stop_label=stop_label, steps=tuple(steps), records=tuple(records),
        purpose_from_pdf=purpose, keywords=keywords, sections=tuple(sections),
        materials=materials, equipment=equipment, deviations=tuple(deviations),
        confirmed=tuple(confirmed), to_check=tuple(to_check), protocol_reference=reference,
        zone=zone, items=items, returns=tuple(returns), reverts=tuple(reverts),
        review=tuple(review) if review_offered else (),
        unconfirmed_texts=tuple(dict.fromkeys(unconfirmed_texts)),
        branch_answers=tuple(branch_answers), repetitions=tuple(repetitions),
        review_version_words=review_version_words, review_changes=review_changes,
        item_step_texts=tuple(
            (step["label"], str(step.get("whole_source_text") or step["source_text"])[:400])
            for step in source_steps if any(step["label"] in item.steps for item in items)
        ),
    )
    reviewer = " ".join(str(getattr(getattr(getattr(fixture, "draft", None), "review_version", None),
                                    "reviewer_name", "") or "").split())
    facts = _replace_dataclass(facts, glance=build_report_glance(
        facts, events, source_steps=source_steps, zone=zone, reviewer=reviewer,
        start_check=start_check))
    facts = _replace_dataclass(facts, notebook=build_lab_notebook(
        facts, events, source_steps=source_steps, zone=zone, previous=previous,
        protocol_problem=protocol_problem if fixture is None else "",
        session_id=str(report_data.get("session_id") or "")))
    return _with_notebook_source_images(facts, fixture, report_data)


# --- The report's first pages (lane RP2, decision 1, 2026-10-19) -------------
#
# Human decision: before its sections the report opens with what a reader
# takes in at a glance -- the header (protocol, review version and reviewer,
# experimenter, date, start and end, time taken), the counts, one row per
# step, one row per timer, the deviations and anomalies with the source's
# side beside each, the check before the start, and a place for the
# experimenter's name and date. All of it is read from the record (and the
# check from the protocol's ledger); no model writes any of it, a value the
# record lacks is "기록 없음", and the record is never changed.

#: The header's first rows, in the decided order; the run's other rows follow.
GLANCE_HEADER_FIRST = ("프로토콜", "검토본", "검토자", "실험자", "날짜", "시작", "끝", "걸린 시간")
GLANCE_SUMMARY_HEADER = ("단계 완료", "편차", "이상", "정정", "미확인", "연결 끊김")
GLANCE_STEP_HEADER = ("단계", "시작", "끝", "소요", "기록한 값", "편차 (원문 → 기록)", "이상")
GLANCE_TIMER_HEADER = ("단계", "원문 시간", "켠 시각", "1분 전 알림", "끝 알림", "단계 끝", "결과")
GLANCE_ISSUE_HEADER = ("번호", "단계", "종류", "원문 (쪽)", "기록", "시각", "확인")
GLANCE_PRESTART_HEADER = ("확인한 값", "단계", "원문 값", "고른 값", "결과", "시각")
GLANCE_SIGN_HEADER = ("실험자 확인", "적는 칸")
GLANCE_SIGN_ROWS = (("이름", ""), ("날짜", ""), ("서명", ""))
NO_REVIEW_VERSION_WORDS = "없음 (분석 결과 그대로 실행)"
#: Lane RP2, decision 3: what a value the experimenter has not confirmed
#: carries in the tables, and the report's line about them under its title.
UNCONFIRMED_MARK = "[미확인]"
UNCONFIRMED_WORDS = (f"실험자가 아직 확인하지 않은 값이다. 표에 {UNCONFIRMED_MARK}로 표시했다. "
                     "화면의 확인 목록에서 확인하면 빠진다.")
#: Events that move the run to a step: a step's time starts at the last one
#: before its completion.
_ARRIVALS = frozenset({
    "session_started", "step_completed", "step_advanced", "repeat_returned", "step_reverted",
    "steps_skipped", "branch_steps_skipped",
})
_STEP_DONE = frozenset({"step_completed", "step_advanced"})
_VALUE_KINDS = ("측정", "관찰", "메모", "사진")


@dataclass(frozen=True)
class ReportGlance:
    """The report's first pages (lane RP2, decision 1), each table's rows."""

    header_rows: tuple[tuple[str, str], ...]
    summary: tuple[str, ...]  # in the order of GLANCE_SUMMARY_HEADER
    connections: tuple[str, ...]
    step_rows: tuple[tuple[str, ...], ...]
    timer_rows: tuple[tuple[str, ...], ...]
    issue_rows: tuple[tuple[str, ...], ...]
    #: None when the check before the start was not read; () when it has no answer.
    prestart_rows: tuple[tuple[str, ...], ...] | None
    prestart_notes: tuple[str, ...]
    #: Lane RP2, decision 3: each value not confirmed ("2단계 측정: 세척액 450 µL").
    unconfirmed: tuple[str, ...] = ()


def _clock(moment: datetime | None) -> str:
    return f"{moment:%H:%M}" if moment else "—"


def _span_words(seconds: float | int | None) -> str:
    """"40초", "1분 30초" under ten minutes, else as ``_duration_words``."""

    if seconds is None:
        return ""
    whole = max(int(round(float(seconds))), 0)
    if whole < 60:
        return f"{whole}초"
    minutes, rest = divmod(whole, 60)
    if minutes < 10 and rest:
        return f"{minutes}분 {rest}초"
    return _duration_words(whole)


def _cell(lines: Sequence[str]) -> str:
    """A table cell holding one value per line, or "—"."""

    kept = [line for line in lines if line]
    return "\n".join(kept) if kept else "—"


def start_check_answers_between(
    events: Sequence[Any], *, since: str | None, until: str | None,
) -> list[dict[str, Any]]:
    """The answers of the check before the start recorded after ``since`` and
    by ``until``, in the order recorded (lane RP2, decision 1).

    ``events`` are the protocol ledger's events (records or mappings). The
    check is answered right before a run starts, so a run's answers are those
    after the run before it started and before it started itself.
    """

    from voiney_lab.start_check import START_CHECK_EVENT

    lower = _local(since, timezone.utc) if since else None
    upper = _local(until, timezone.utc) if until else None
    chosen: list[dict[str, Any]] = []
    for event in events:
        kind = event.get("event_type") if isinstance(event, Mapping) else getattr(event, "event_type", None)
        payload = event.get("payload") if isinstance(event, Mapping) else getattr(event, "payload", None)
        if kind != START_CHECK_EVENT or not isinstance(payload, Mapping):
            continue
        at = _local(payload.get("recorded_at") or (
            None if isinstance(event, Mapping) else getattr(event, "recorded_at", None)), timezone.utc)
        if at is None or (lower is not None and at <= lower) or (upper is not None and at > upper):
            continue
        chosen.append(dict(payload))
    return chosen


def report_start_check_answers(
    report_data: Mapping[str, Any], *, since: str | None,
) -> list[dict[str, Any]] | None:
    """The check before the start of a report's run, read from the protocol's
    ledger as the server opens it; None when it cannot be read."""

    try:
        from voiney_lab.server import _open_protocol_catalog

        _catalog, store = _open_protocol_catalog()
        try:
            events = store.list_events(str(report_data.get("protocol_id") or ""))
        finally:
            store.close()
    except Exception as exc:  # noqa: BLE001 -- said on the report, never raised
        log.warning("report start check not read error=%s", type(exc).__name__)
        return None
    owner=report_data.get("owner_principal_id")
    # Identified records never inherit anonymous or other people's answers.
    if owner is not None:
        own_events=[]
        for item in events:
            payload=item.get("payload") if isinstance(item, Mapping) else getattr(item, "payload", None)
            if isinstance(payload,Mapping) and payload.get("actor_principal_id")==owner:
                own_events.append(item)
        events=own_events
    elif report_data.get("organization_id") is not None:
        return []  # Quarantined legacy record: the experimenter is unknown.
    return start_check_answers_between(events, since=since, until=str(report_data.get("started_at") or ""))


def build_report_glance(
    facts: ReportFacts,
    events: Sequence[Mapping[str, Any]],
    *,
    source_steps: Sequence[Mapping[str, Any]],
    zone: ZoneInfo,
    reviewer: str = "",
    start_check: Sequence[Mapping[str, Any]] | None = None,
) -> ReportGlance:
    """The report's first pages from the record (lane RP2, decision 1)."""

    kinds = [str(event.get("event_type") or "") for event in events]
    labels = [str(event.get("step_label") or "") for event in events]
    keys = [str(event.get("event_key") or "") for event in events]
    moments = [_local(event.get("created_at"), zone) for event in events]
    payloads = [event.get("payload") if isinstance(event.get("payload"), Mapping) else {}
                for event in events]
    pages = {str(step["label"]): step.get("page") for step in source_steps}
    source_timers = {str(step["label"]): step.get("timer") for step in source_steps}
    order = [str(step["label"]) for step in source_steps] or [step.label for step in facts.steps]
    for step in facts.steps:
        if step.label not in order:
            order.append(step.label)
    records = {record.key: record for record in facts.records if record.key}
    review = report_review_items(events)
    confirmed_by_key = {item["source_key"]: bool(item["confirmed"]) for item in review}
    # Lane RP2, decision 3: a value the experimenter has not confirmed -- asked
    # about or not yet -- is marked wherever the first pages show it.
    unconfirmed_keys = {item["source_key"] for item in review if not item["confirmed"]}

    def marked(text: str, key: str) -> str:
        return f"{text} {UNCONFIRMED_MARK}" if key and key in unconfirmed_keys else text

    def arrival_before(index: int) -> datetime | None:
        for earlier in range(index - 1, -1, -1):
            if kinds[earlier] in _ARRIVALS:
                return moments[earlier]
        return None

    def page_words(label: str) -> str:
        page = pages.get(label)
        return f"{page}쪽" if isinstance(page, int) and not isinstance(page, bool) and page > 0 else ""

    def timer_of(index: int) -> Mapping[str, Any]:
        timer = payloads[index].get("timer")
        return timer if isinstance(timer, Mapping) else {}

    def ended_early(index: int) -> tuple[int, int] | None:
        timer = timer_of(index)
        if not (timer.get("step_exited_before_timer_elapsed")
                or timer.get("completion_state") == "step_exited_before_timer_elapsed"):
            return None
        try:
            return (int(timer.get("elapsed_seconds") or 0),
                    int(timer.get("source_duration_seconds") or timer.get("duration_seconds") or 0))
        except (TypeError, ValueError):
            return None

    def source_values(record: ReportRecord) -> str:
        payload = payloads[keys.index(record.key)] if record.key in keys else {}
        found = [str(item.get("source") or "") for item in payload.get("deviations") or ()
                 if isinstance(item, Mapping) and item.get("source")]
        return " · ".join(dict.fromkeys(found))

    def said_values(record: ReportRecord) -> list[tuple[str, str]]:
        payload = payloads[keys.index(record.key)] if record.key in keys else {}
        return [(str(item.get("source") or ""), str(item.get("said") or ""))
                for item in payload.get("deviations") or ()
                if isinstance(item, Mapping) and item.get("source") and item.get("said")]

    # The steps, each from the arrival before its completion to the completion.
    done_at: dict[str, list[int]] = {}
    seen_at: dict[str, int] = {}
    skipped: set[str] = set()
    for index, kind in enumerate(kinds):
        if labels[index]:
            seen_at.setdefault(labels[index], index)
        if kind in _STEP_DONE and labels[index]:
            done_at.setdefault(labels[index], []).append(index)
        record = payloads[index].get("step_record") if isinstance(payloads[index].get("step_record"), Mapping) else {}
        if kind in {"steps_skipped", "branch_steps_skipped", "branch_answered"}:
            skipped.update(str(item) for item in record.get("skipped_step_labels") or ())
    pauses: dict[str, list[tuple[datetime, datetime]]] = {}
    for index, kind in enumerate(kinds):
        if kind == "workflow_paused" and moments[index] is not None:
            resumed = next((later for later in range(index + 1, len(kinds))
                            if kinds[later] == "workflow_resumed"), None)
            if resumed is not None and moments[resumed] is not None:
                pauses.setdefault(labels[index], []).append((moments[index], moments[resumed]))
    by_label = {step.label: step for step in facts.steps}
    step_rows: list[tuple[str, ...]] = []
    for label in order:
        step = by_label.get(label)
        live = [record for record in (step.records if step else ()) if record.status != "취소됨"]
        values = _cell([
            marked(f"{record.kind}: {record.text}" + (" (정정됨)" if record.status == "정정됨" else ""), record.key)
            for record in live if record.kind in _VALUE_KINDS])
        deviations: list[str] = []
        for record in live:
            if record.kind != "편차":
                continue
            pairs = said_values(record)
            if pairs and record.status != "정정됨":
                shown = " · ".join(f"{source} → {said}" for source, said in pairs)
            elif record.status == "정정됨":
                shown = f"원문 {source_values(record) or '값 없음'} → “{record.text}” (정정됨)"
            else:
                shown = f"“{record.text}” (원문 값 없음)"
            deviations.append(marked(shown, record.key))
        completions = done_at.get(label, [])
        for index in completions:
            early = ended_early(index)
            if early is not None:
                deviations.append(marked(
                    f"타이머 {_duration_words(early[1])} → {_duration_words(early[0])}에 끝냄", keys[index]))
        anomalies = _cell([marked(record.text, record.key) for record in live if record.kind == "이상"])
        if completions:
            visits = [(arrival_before(index), moments[index]) for index in completions]
            start, end = visits[0][0], visits[-1][1]
            if all(arrived and done for arrived, done in visits):
                spent = sum((done - arrived).total_seconds() for arrived, done in visits)  # type: ignore[operator]
                paused = sum((back - away).total_seconds() for away, back in pauses.get(label, ())
                             if start is not None and end is not None and start <= away and back <= end)
                taken = _span_words(spent) + (f" (멈춤 {_span_words(paused)} 포함)" if paused else "")
            else:
                taken = "—"
            name = f"{label} ({len(completions)}회)" if len(completions) > 1 else label
            step_rows.append((name, _clock(start), _clock(end), taken, values, _cell(deviations), anomalies))
        elif label in seen_at:
            step_rows.append((label, _clock(arrival_before(seen_at[label])), "미완료", "—",
                              values, _cell(deviations), anomalies))
        else:
            step_rows.append((label, "—", "건너뜀" if label in skipped else "수행 안 함", "—",
                              values, _cell(deviations), anomalies))

    # The timers: each start with its notices and the completion after it, and
    # a source timer the run reached but never started.
    timer_rows: list[tuple[str, ...]] = []
    for label in order:
        starts = [index for index, kind in enumerate(kinds) if kind == "timer_started" and labels[index] == label]
        for position, index in enumerate(starts):
            until = starts[position + 1] if position + 1 < len(starts) else len(kinds)
            timer = timer_of(index)
            try:
                seconds = int(timer.get("duration_seconds") or timer.get("source_duration_seconds") or 0) or None
            except (TypeError, ValueError):
                seconds = None
            source = source_timers.get(label) or seconds

            def notice(kind: str) -> str:
                found = next((later for later in range(index + 1, until)
                              if kinds[later] == kind and labels[later] == label), None)
                if found is None:
                    return "—"
                return f"{_clock(moments[found])} ({'말함' if payloads[found].get('spoken') else '화면'})"

            done = next((later for later in done_at.get(label, []) if later > index), None)
            if done is None:
                result = "단계 끝 기록 없음"
            elif (early := ended_early(done)) is not None:
                result = f"일찍 끝냄 (원문 {_duration_words(early[1] or source)} 중 {_duration_words(early[0])})"
            elif moments[index] is not None and moments[done] is not None:
                actual = (moments[done] - moments[index]).total_seconds()
                result = (f"원문 시간 뒤 완료 (실제 {_span_words(actual)})"
                          if seconds is None or actual >= seconds - 30 else f"완료 (실제 {_span_words(actual)})")
            else:
                result = "완료 (걸린 시간 기록 없음)"
            timer_rows.append((label, _duration_words(source) or "—", _clock(moments[index]),
                               notice("timer_last_minute_notice"), notice("timer_end_notice"),
                               _clock(moments[done]) if done is not None else "—", result))
        if not starts and source_timers.get(label) and (label in done_at or label in seen_at):
            done = done_at.get(label, [None])[-1]
            timer_rows.append((label, _duration_words(source_timers[label]), "켜지 않음", "—", "—",
                               _clock(moments[done]) if done is not None else "—", "타이머 기록 없음"))

    # The deviations and anomalies, in the order recorded, with the source's side.
    issues: list[tuple[str, ...]] = []

    def checked(key: str) -> str:
        if key not in confirmed_by_key:
            return "—"
        return "확인" if confirmed_by_key[key] else "미확인"

    def add_issue(label: str, kind: str, source: str, recorded: str, index: int) -> None:
        side = " · ".join(part for part in (source, page_words(label)) if part) or "—"
        issues.append((str(len(issues) + 1), label or "—", kind, side, recorded,
                       _clock(moments[index]), checked(keys[index])))

    for index, kind in enumerate(kinds):
        record = records.get(keys[index])
        step_record = (payloads[index].get("step_record")
                       if isinstance(payloads[index].get("step_record"), Mapping) else {})
        if record is not None and record.status != "취소됨" and record.kind in {"편차", "이상"} and (
                kind in {"observation", "anomaly"}):
            recorded = record.text + (f"\n처음 기록: “{record.original}”" if record.status == "정정됨" else "")
            add_issue(record.step_label, record.kind,
                      source_values(record) if record.kind == "편차" else "", recorded, index)
        elif kind == "step_completed" and (early := ended_early(index)) is not None:
            add_issue(labels[index], "타이머 일찍 끝냄", _duration_words(early[1]),
                      f"{_duration_words(early[0])}에 단계를 끝냄", index)
        elif kind == "steps_skipped" and step_record.get("skipped_step_labels"):
            passed = [str(item) for item in step_record.get("skipped_step_labels") or ()]
            start = str(step_record.get("start_step") or labels[index])
            issues_label = passed[0] if passed else labels[index]
            add_issue(issues_label, "단계 건너뜀", _ranges(passed),
                      f"{start}단계부터 시작", index)
        elif kind == "repeat_rounds_declined":
            first, last = _range_pair(step_record)
            add_issue(labels[index], "반복 덜 함",
                      f"{first}~{last}단계 {step_record.get('rounds_required')}회",
                      f"{step_record.get('rounds_done')}회차까지만 하고 넘어감", index)
        elif kind == "step_advanced" and labels[index]:
            add_issue(labels[index], "완료 말 없이 넘어감", "", "완료를 말하지 않고 다음 단계로 넘어감", index)

    connections: list[str] = []
    for index, kind in enumerate(kinds):
        if kind != "connection_restored":
            continue
        payload = payloads[index]
        lost = _local(payload.get("lost_at"), zone)
        restored = _local(payload.get("restored_at"), zone) or moments[index]
        gap = payload.get("gap_seconds")
        gap_words = (_span_words(gap) if isinstance(gap, (int, float)) and not isinstance(gap, bool)
                     else "기록 없음")
        where = f"{labels[index]}단계에서 " if labels[index] else ""
        line = (f"{where}{_clock(lost)}에 연결이 끊겼다가 {_clock(restored)}에 다시 연결됨 — 끊긴 시간 {gap_words}"
                if lost else f"{where}연결이 끊겼다가 {_clock(restored)}에 다시 연결됨 — 끊긴 시간 {gap_words}")
        ended = payload.get("timers_ended")
        if isinstance(ended, int) and not isinstance(ended, bool) and ended:
            line += f", 그사이 끝난 타이머 {ended}개"
        connections.append(line)

    run = dict(facts.run_rows)
    header: list[tuple[str, str]] = []
    for name in GLANCE_HEADER_FIRST:
        if name == "검토본":
            header.append((name, run.get(name) or NO_REVIEW_VERSION_WORDS))
        elif name == "검토자":
            # Only a review version has a reviewer: a run of the analysis as
            # it came names none (lane DI keeps reviewer words out of it).
            if run.get("검토본") and reviewer:
                header.append((name, reviewer))
        else:
            header.append((name, run.get(name, "")))
    header += [(name, value) for name, value in facts.run_rows if name not in GLANCE_HEADER_FIRST]

    deviation_count = sum(1 for row in issues if row[2] != "이상")
    summary = (
        f"{len(facts.completed_labels)} / {facts.total_steps}" if facts.total_steps
        else str(len(facts.completed_labels)),
        f"{deviation_count}건",
        f"{len(issues) - deviation_count}건",
        f"{sum(1 for kind in kinds if kind in {'record_corrected', 'record_retracted'})}건",
        f"{sum(1 for item in review if not item['confirmed'])}개",
        f"{len(connections)}회",
    )

    from voiney_lab.start_check import KIND_KO

    prestart_rows: tuple[tuple[str, ...], ...] | None = None
    if start_check is not None:
        rows = []
        for answer in start_check:
            if answer.get("action") == "accept":
                result = "원문대로 맞다고 함"
            else:
                result = "맞음" if answer.get("correct") is True else "다르게 고름"
            rows.append((
                KIND_KO.get(str(answer.get("kind") or ""), str(answer.get("kind") or "") or "—"),
                str(answer.get("step_label") or "") or "—",
                str(answer.get("value") or "") or "—",
                str(answer.get("chosen") or "") or "—",
                result,
                _clock(_local(answer.get("recorded_at"), zone)),
            ))
        prestart_rows = tuple(rows)
    notes = []
    if run.get("안전 주의 확인"):
        notes.append(f"안전 주의: {run['안전 주의 확인']}")
    for answer in facts.branch_answers:
        if BEFORE_START_WORDS in answer.value_source:
            notes.append(f"{answer.step_label}단계 조건 “{answer.condition}” → {answer.answer} ({answer.value_source})")
    for repetition in facts.repetitions:
        if BEFORE_START_WORDS in repetition.value_source:
            count = f"{repetition.count}회" if repetition.count else "미정"
            notes.append(f"{repetition.first}~{repetition.last}단계 반복 횟수 {count} ({repetition.value_source})")

    return ReportGlance(
        header_rows=tuple(header), summary=summary, connections=tuple(connections),
        step_rows=tuple(step_rows), timer_rows=tuple(timer_rows), issue_rows=tuple(issues),
        prestart_rows=prestart_rows, prestart_notes=tuple(notes),
        unconfirmed=tuple(
            (f"{item['step_label']}단계 {item['kind']}: {item['text']}" if item["step_label"]
             else f"{item['kind']}: {item['text']}")
            for item in review if not item["confirmed"]),
    )


# --- The lab notebook (lane RP3, decision ①, 2026-10-20) -----------------------
#
# Human decision (2026-10-19/20): the report is an experiment notebook for
# the researcher who ran the experiment, their PI and the lab meeting. It
# says first what a researcher wants to know -- what was done, with which
# samples, under which conditions, what was seen, what differed from the
# source and from the run before, and what to change next -- in this order:
# the header (title: the purpose said, else the protocol's name; date and
# time; experimenter; protocol with the version run and its source), the
# samples and their conditions as said, the conditions actually used ("원문
# 값 → 실제", steps done as the source says grouped as "원문대로", only a
# difference standing out, never the source's sentences copied), the
# observations and results grouped by step and sample (values as recorded),
# what was done differently from the source (with the reason when the
# record holds one) and the problems, what differed from the run before
# (deterministic; "처음 실험" when there is none), what to change next (what
# the experimenter said, and the source's next step), the confirmation box
# (experimenter, PI -- decision ④), and a small appendix (the steps' times,
# the words as said with their corrections, one line of system events).
#
# The agent's own doings -- notices, lost connections, the check before the
# start, the router -- stay out of the body. Every number, value and fact
# comes from the record or the source; nothing in this decision is written
# by a model. What the source allows (a repeat "until", a range) is no
# difference. The old sections (purpose from keywords, a background of
# section names, a materials table with the list's labels, a conclusion of
# counts, a discussion repeating the quotations, the notices tables) are
# gone from the body.

NOTEBOOK_SECTIONS = (
    "실험 정보", "목적", "방법", "관찰·특이사항", "결과", "해석·다음 할 일", "분야별 기록", "부록",
)
NOTEBOOK_SECTION_IDS = (
    "experiment_info", "purpose", "methods", "observations", "results",
    "interpretation_next", "domain", "appendix",
)
NOTEBOOK_HEADINGS = tuple(f"## {name}" for name in NOTEBOOK_SECTIONS)
#: Decision ②'s events (written by the voice rules in that decision; read
#: here): the purpose and samples said at the start, the result and the
#: change for next time said at the end. Decision ④'s event: a PI's
#: confirmation of the notebook (name, time, one line).
PURPOSE_RECORDED = "purpose_recorded"
RESULT_RECORDED = "result_recorded"
PI_CONFIRMED = "pi_confirmed"
SAMPLE_HEADER = ("시료", "조건", "대조군")
CONDITION_HEADER = ("과정", "원문 값 → 실제")
OBSERVATION_HEADER = ("과정", "시료", "관찰·측정 (기록 그대로)", "시각")
CONFIRM_HEADER = ("확인", "이름", "날짜·시각", "서명·한 줄 의견")
APPENDIX_STEP_HEADER = ("단계", "시작", "끝", "소요", "타이머 (원문 / 실제)")
APPENDIX_RECORD_HEADER = ("시각", "단계", "종류", "말한 그대로")
METHOD_HEADER = ("단계", "원문 단계", "원문 값 → 실제", "실제 시각·타이머")
TIMELINE_HEADER = ("시각", "단계", "사건", "기록")
VALUE_BAR_HEADER = ("단계·항목", "원문", "실제")
DOMAIN_HEADER = ("항목", "원문·기록에서 찾은 내용", "근거")
SAMPLE_RECORD_HEADER = ("셀·시료 번호", "기록 그대로")
INVENTORY_HEADER = ("시약·장비 (원문)", "종류", "구매처", "카탈로그 번호", "Lot 번호", "입력 시점")
NOT_SAID_WORDS = "말하지 않음."
FIRST_RUN_WORDS = "처음 실험 — 같은 프로토콜의 앞선 기록이 없다."
SAME_AS_LAST_WORDS = "지난번과 조건·관찰이 같다."
AS_SOURCE_WORDS = "원문대로"
NO_DIFFERENCE_WORDS = "원문대로 했고, 기록된 문제가 없다."
NO_SOURCE_VALUES_WORDS = "프로토콜 원문을 불러오지 못해 원문 값을 적지 못했다."
NO_VALUED_STEPS_WORDS = "원문에 값(양·시간·온도·속도·농도)이 적힌 과정을 수행하지 않았다."
NO_OBSERVATION_WORDS = "기록된 관찰이 없다."
NO_PI_WORDS = "PI 확인: 아직 없음."
NOTEBOOK_SUBTITLE = "실험 노트"
#: Lane HELP: the event of the help given after a problem.
PROBLEM_HELP_EVENT = "problem_help"


def problem_help_line(step_label: str | None, payload: Mapping[str, Any], words: str | None, at: str) -> str:
    """The notebook's one line for a problem's help: where its lines came from."""

    found = payload.get("found") if isinstance(payload.get("found"), Mapping) else {}
    pages = sorted({
        str(item.get("page")) for item in payload.get("sources") or ()
        if isinstance(item, Mapping) and item.get("tier") == "protocol" and str(item.get("page") or "").isdigit()
    }, key=int)
    told = []
    if pages:
        told.append("원문 " + "·".join(f"p.{page}" for page in pages))
    if found.get("history"):
        told.append(f"지난 기록 {found['history']}건")
    if found.get("public"):
        told.append(f"공개 자료 {found['public']}건")
    asked = f" — 물음 “{' '.join(str(words).split())}”" if words else ""
    where = f"{step_label}단계 " if step_label else ""
    return f"도움 · {where}({at or '시각 없음'}){asked}: " + (" · ".join(told) or "원문과 지난 기록에 관련 내용 없음")


#: Lane EASY, decision 2: the events that hold today's values -- the start's
#: (on "session_started"), a change and a put-back said during the run.
RUN_VALUE_CHANGED = "run_value_changed"
RUN_VALUE_REVERTED = "run_value_reverted"
TODAY_MARK_WORDS = "오늘 바꾼 값"
#: The kinds the appendix's record table gives decision ②'s sayings.
SAYING_KINDS = {PURPOSE_RECORDED: "목적·시료", RESULT_RECORDED: "결과·다음 할 일"}


@dataclass(frozen=True)
class NotebookSample:
    """A sample as the experimenter named it at the start (decision ② a)."""

    name: str
    condition: str
    control: bool


@dataclass(frozen=True)
class NotebookValue:
    """One value the source states for a step, and what was actually used.

    ``actual`` is empty when the run kept the source's value; otherwise the
    value as recorded (a deviation note's value, a measurement that
    differs, a timer ended early), with the record's words and key.
    """

    label: str
    kind: str  # 시간 | 온도 | 양 | 농도 | 속도
    family: str  # what it is compared within: time, volume, mass, molar, or the unit itself
    written: str  # as the source writes it ("37°C", "15 min")
    actual: str = ""
    words: str = ""
    key: str = ""
    early: bool = False
    #: Lane EASY, decision 2: the value today's run chose in place of the
    #: source's ("20분"), "" when none was; ``actual`` is then that value, or
    #: what the record says happened instead (a timer ended early).
    planned: str = ""


@dataclass(frozen=True)
class NotebookObservation:
    label: str
    process: str
    sample: str
    text: str  # the latest words, without the marks
    shown: str  # as the table shows it (correction, marks)
    at: str
    kind: str


@dataclass(frozen=True)
class NotebookDifference:
    """One line of "원문과 다르게 한 점·문제", in the order recorded."""

    order: int
    text: str


@dataclass(frozen=True)
class NotebookMethodRow:
    """One source step beside the execution facts recorded for this run."""

    label: str
    source_text: str
    actual_values: str
    started_at: str
    ended_at: str
    timer: str
    source_page: int | None = None


@dataclass(frozen=True)
class NotebookTimelineItem:
    at: str
    step_label: str
    kind: str
    text: str
    source_key: str


@dataclass(frozen=True)
class NotebookValueBar:
    """Raw source/actual measures.  Renderers scale these values visually."""

    label: str
    kind: str
    source: str
    actual: str
    source_number: str
    actual_number: str
    unit: str
    source_unit: str
    actual_unit: str
    source_key: str


@dataclass(frozen=True)
class NotebookPlatePosition:
    label: str
    step_label: str
    words: str


@dataclass(frozen=True)
class NotebookSampleRow:
    identifier: str
    records: tuple[str, ...]


@dataclass(frozen=True)
class NotebookDomainField:
    label: str
    value: str
    origin: str  # source | record | calculated | blank
    source_keys: tuple[str, ...] = ()
    formula: str = ""


@dataclass(frozen=True)
class NotebookDomain:
    code: str
    label: str
    selected_by: str  # source_analysis | person
    fields: tuple[NotebookDomainField, ...]


@dataclass(frozen=True)
class NotebookInventoryItem:
    kind: str
    name: str
    vendor: str
    catalog_number: str
    lot_number: str
    stage: str  # before | after | source | blank
    source_keys: tuple[str, ...] = ()


@dataclass(frozen=True)
class NotebookSection:
    id: str
    name: str


@dataclass(frozen=True)
class NotebookImage:
    """A recorded photo or an exact crop from the source PDF."""

    kind: str  # photo | source_figure
    step_label: str
    caption: str
    url: str
    media_type: str
    asset_id: str
    source_page: int | None = None
    interpretation_status: str = ""


DEFAULT_NOTEBOOK_TEMPLATE = tuple(
    NotebookSection(section_id, name)
    for section_id, name in zip(NOTEBOOK_SECTION_IDS, NOTEBOOK_SECTIONS)
)


@dataclass(frozen=True)
class LabNotebook:
    """The notebook the exports render (decision ①), from the record alone."""

    title: str
    subtitle: str
    header_rows: tuple[tuple[str, str], ...]
    purpose: str  # the purpose as said (""), its title form
    purpose_words: str  # the whole utterance
    samples: tuple[NotebookSample, ...]
    values: tuple[NotebookValue, ...]
    condition_rows: tuple[tuple[str, str], ...]
    condition_note: str  # why there are no rows, or ""
    materials_words: str
    observations: tuple[NotebookObservation, ...]
    result: str
    result_words: str
    next_change: str
    differences: tuple[str, ...]
    previous_words: str  # "지난번: ..." or ""
    previous_lines: tuple[str, ...]
    next_lines: tuple[str, ...]
    pi_rows: tuple[tuple[str, str, str, str], ...]
    appendix_step_rows: tuple[tuple[str, ...], ...]
    appendix_record_rows: tuple[tuple[str, ...], ...]
    system_words: str
    source_words: str
    performed: tuple[str, ...]
    run_words: str  # how this run is named when it is "지난번" to the next
    #: Decision ③: each step's condition in short ("원문대로 — 100 µL",
    #: "온도 37°C → 40°C"), for the results table.
    condition_by_label: Mapping[str, str] = field(default_factory=dict)
    method_rows: tuple[NotebookMethodRow, ...] = ()
    timeline: tuple[NotebookTimelineItem, ...] = ()
    value_bars: tuple[NotebookValueBar, ...] = ()
    plate_positions: tuple[NotebookPlatePosition, ...] = ()
    sample_rows: tuple[NotebookSampleRow, ...] = ()
    domain: NotebookDomain = field(default_factory=lambda: NotebookDomain(
        "general", "공통", "source_analysis", (),
    ))
    interpretation_words: str = ""
    inventory: tuple[NotebookInventoryItem, ...] = ()
    section_template: tuple[NotebookSection, ...] = DEFAULT_NOTEBOOK_TEMPLATE
    images: tuple[NotebookImage, ...] = ()


# Samples named in words (decision ② a-b): "3번 샘플", "샘플 2번", "S2", "대조군",
# "시료 3개"; read deterministically, canonical names "{n}번 샘플", "S2", "대조군".
_SAMPLE_NOUN = r"(?:샘플|시료|시편)"
_NUMBER_LIST = r"\d+(?:\s*(?:[,·/]|~|–|-|번)\s*\d+)*"
_SAMPLES_BEFORE_NOUN = re.compile(rf"({_NUMBER_LIST})\s*번\s*({_SAMPLE_NOUN})")
_SAMPLES_AFTER_NOUN = re.compile(rf"({_SAMPLE_NOUN})\s*(\d+)\s*번?(?![\d개])")
_SAMPLE_CODE = re.compile(rf"(?:{_SAMPLE_NOUN}\s*)?(?<![A-Za-z0-9])(S\d{{1,2}})(?![A-Za-z0-9])")
_SAMPLE_COUNT = re.compile(rf"({_SAMPLE_NOUN})\s*(\d+)\s*개")
_CONTROL = re.compile(r"((?:음성|양성)\s*)?(대조군|컨트롤|(?<![A-Za-z])control(?![A-Za-z]))", re.I)
_CONTROL_COUNT = re.compile(r"((?:음성|양성)\s*)?대조군\s*(\d+)\s*개")
_BARE_NUMBERS = re.compile(rf"(?<![\d.])({_NUMBER_LIST})\s*번(?!\s*{_SAMPLE_NOUN})(?!\d)")
_CLAUSE_SPLIT = re.compile(r"[,.;!?。]\s*|\s(?:그리고|및)\s|(?<=이고)\s|(?<=하고)\s|(?<=이며)\s")
_PARTICLE_HEAD = re.compile(r"^(?:은|는|이|가|을|를|도|에는|에|으로|로|의)\s*")
_COPULA_TAIL = re.compile(r"\s*(?:이에요|예요|에요|입니다|이야|이다|야|요|임|이고|고|이며)?\s*$")
_TITLE_HEAD = re.compile(r"^(?:오늘은|오늘|이번에는|이번엔|이번 실험은|이번은)\s+")
_TITLE_TAIL = re.compile(r"\s*(?:이에요|예요|에요|입니다|이야|이다|야|요)?[.!?。]*\s*$")
_SENTENCE_END = re.compile(r"(?<=[.!?。])\s+")
#: Where "다음에 바꿀 점" begins in the one utterance said at the end (② d).
_NEXT_WORDS = re.compile(r"(?:^|(?<=\s))(?:다음(?:번)?(?:에는|엔|에|부터는|부터)|앞으로는|앞으로)\s*")


def _expand_numbers(text: str) -> list[int]:
    found: list[int] = []
    for part in re.split(r"[,·/]|\s*번\s*", text):
        part = part.strip()
        if not part:
            continue
        span = re.fullmatch(r"(\d+)\s*(?:~|–|-)\s*(\d+)", part)
        if span:
            first, last = int(span.group(1)), int(span.group(2))
            found += list(range(first, last + 1)) if first <= last else [first, last]
        elif part.isdigit():
            found.append(int(part))
    return found


def _condition_words(clause: str, spans: Sequence[tuple[int, int]]) -> str:
    """The clause without the sample mentions, without its particles."""

    kept = []
    position = 0
    for start, end in sorted(spans):
        kept.append(clause[position:start])
        position = end
    kept.append(clause[position:])
    words = " ".join("".join(kept).split())
    words = _PARTICLE_HEAD.sub("", words)
    words = _COPULA_TAIL.sub("", words)
    return words.strip(" ,·")


def samples_in_text(text: str) -> list[dict[str, Any]]:
    """The samples a sentence names, each with its condition and whether it is
    a control (decision ② a): deterministic, [] when none is named."""

    words = " ".join(str(text or "").split())
    nouns = re.findall(_SAMPLE_NOUN, words)
    default_noun = nouns[0] if nouns else "샘플"
    found: dict[str, dict[str, Any]] = {}

    def add(name: str, condition: str, control: bool) -> None:
        if name not in found:
            found[name] = {"name": name, "condition": condition, "control": control}
            return
        if condition and not found[name]["condition"]:
            found[name]["condition"] = condition
        found[name]["control"] = found[name]["control"] or control

    for clause in _CLAUSE_SPLIT.split(words):
        clause = clause.strip()
        if not clause:
            continue
        spans: list[tuple[int, int]] = []
        names: list[str] = []
        standalone_controls: list[str] = []

        def free(start: int, end: int) -> bool:
            return all(end <= a or start >= b for a, b in spans)

        for match in _SAMPLES_BEFORE_NOUN.finditer(clause):
            spans.append(match.span())
            names += [f"{n}번 {match.group(2)}" for n in _expand_numbers(match.group(1))]
        for match in _SAMPLES_AFTER_NOUN.finditer(clause):
            if free(*match.span()):
                spans.append(match.span())
                names.append(f"{match.group(2)}번 {match.group(1)}")
        for match in _SAMPLE_CODE.finditer(clause):
            if free(*match.span()):
                spans.append(match.span())
                names.append(match.group(1))
        for match in _SAMPLE_COUNT.finditer(clause):
            if free(*match.span()):
                spans.append(match.span())
                names += [f"{n}번 {match.group(1)}" for n in range(1, int(match.group(2)) + 1)]
        for match in _CONTROL_COUNT.finditer(clause):
            if free(*match.span()):
                spans.append(match.span())
                standalone_controls.append(f"{(match.group(1) or '').strip()} 대조군".strip())
        control_here = False
        for match in _CONTROL.finditer(clause):
            if free(*match.span()):
                control_here = True
        if nouns:
            for match in _BARE_NUMBERS.finditer(clause):
                if free(*match.span()):
                    spans.append(match.span())
                    names += [f"{n}번 {default_noun}" for n in _expand_numbers(match.group(1))]
        condition = _condition_words(clause, spans) if names or standalone_controls else ""
        for name in names:
            add(name, condition, control_here)
        for name in standalone_controls:
            add(name, "", True)
        if control_here and not names and not standalone_controls:
            match = next(m for m in _CONTROL.finditer(clause) if free(*m.span()))
            add(f"{(match.group(1) or '').strip()} 대조군".strip(), "", True)
    return list(found.values())


def sample_in_text(text: str) -> str:
    """The first sample a record names, in its canonical name, or "" (② b)."""

    words = " ".join(str(text or "").split())
    mentions: list[tuple[int, str]] = []
    match = _SAMPLES_BEFORE_NOUN.search(words)
    if match:
        numbers = _expand_numbers(match.group(1))
        if numbers:
            mentions.append((match.start(), f"{numbers[0]}번 {match.group(2)}"))
    match = _SAMPLES_AFTER_NOUN.search(words)
    if match:
        mentions.append((match.start(), f"{match.group(2)}번 {match.group(1)}"))
    match = _SAMPLE_CODE.search(words)
    if match:
        mentions.append((match.start(), match.group(1)))
    match = _CONTROL.search(words)
    if match:
        mentions.append((match.start(), f"{(match.group(1) or '').strip()} 대조군".strip()))
    if not mentions:
        return ""
    return min(mentions)[1]


def split_result_and_next(text: str) -> tuple[str, str]:
    """(the result, what to change next) from the one utterance said at the
    end (decision ② d): split where "다음에(는)" begins; all result without it."""

    words = " ".join(str(text or "").split())
    match = _NEXT_WORDS.search(words)
    if match is None:
        return words, ""
    return words[:match.start()].strip(), words[match.end():].strip()


def notebook_title(purpose: str) -> str:
    """The experiment's title from the purpose said: its first sentence,
    without "오늘은" and the sentence's polite ending, at most 80 characters."""

    words = " ".join(str(purpose or "").split())
    if not words:
        return ""
    first = _SENTENCE_END.split(words, 1)[0]
    first = _TITLE_HEAD.sub("", first)
    first = _TITLE_TAIL.sub("", first).strip()
    if len(first) > 80:
        first = first[:79].rstrip() + "…"
    return first


# The values a step states and what was used (decision ① 3): the units the
# notebook compares, each with the kind it is shown under and the family it
# is compared within (hours with minutes, mL with µL).
_CONDITION_KIND = {
    "h": "시간", "min": "시간", "s": "시간", "°C": "온도", "µL": "양", "mL": "양", "L": "양",
    "mg": "양", "µg": "양", "ng": "양", "g": "양", "mM": "농도", "µM": "농도", "nM": "농도",
    "M": "농도", "%": "농도", "rpm": "속도", "×g": "속도", "mm/s": "속도",
}
_CONDITION_FAMILY = {
    "h": "time", "min": "time", "s": "time", "µL": "volume", "mL": "volume", "L": "volume",
    "mg": "mass", "µg": "mass", "ng": "mass", "g": "mass", "mM": "molar", "µM": "molar", "nM": "molar",
    "M": "molar",
}
_CONDITION_BASE = {
    "h": 3600.0, "min": 60.0, "s": 1.0, "µL": 1.0, "mL": 1000.0, "L": 1e6, "mg": 1000.0, "µg": 1.0,
    "ng": 0.001, "g": 1e6, "mM": 1e6, "µM": 1e3, "nM": 1.0, "M": 1e9,
}
#: A deviation event's kind (lane RA) when its value carries no unit the table knows.
_DEVIATION_KIND_KO = {
    "temperature": "온도", "time": "시간", "volume": "양", "mass": "양", "concentration": "농도",
    "percent": "농도", "speed": "속도", "force": "속도",
}


@dataclass(frozen=True)
class _Measure:
    number: str
    unit: str
    written: str

    @property
    def kind(self) -> str:
        return _CONDITION_KIND.get(self.unit, "")

    @property
    def family(self) -> str:
        return _CONDITION_FAMILY.get(self.unit, self.unit)

    def same_as(self, other: "_Measure") -> bool:
        if self.family != other.family:
            return False
        try:
            mine, theirs = float(self.number), float(other.number)
        except ValueError:
            return self.number == other.number and self.unit == other.unit
        if self.unit == other.unit:
            return abs(mine - theirs) <= 1e-9 * max(1.0, abs(theirs))
        if self.unit in _CONDITION_BASE and other.unit in _CONDITION_BASE:
            mine, theirs = mine * _CONDITION_BASE[self.unit], theirs * _CONDITION_BASE[other.unit]
            return abs(mine - theirs) <= 1e-9 * max(1.0, abs(theirs))
        return False


def run_values(events: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Today's values of a record as they stood at its end (lane EASY, 2-3):
    the start's ("session_started"), then each change and put-back in order."""

    kept: dict[str, dict[str, Any]] = {}
    for event in events:
        kind = str(event.get("event_type") or "")
        payload = event.get("payload") if isinstance(event.get("payload"), Mapping) else {}
        record = payload.get("step_record") if isinstance(payload.get("step_record"), Mapping) else {}
        if kind == "session_started":
            for item in payload.get("run_changes") or ():
                if isinstance(item, Mapping) and item.get("item_id"):
                    kept[str(item["item_id"])] = {**item, "event_key": str(event.get("event_key") or "")}
        elif kind == RUN_VALUE_CHANGED and isinstance(record.get("change"), Mapping):
            kept[str(record["change"].get("item_id"))] = {
                **record["change"], "event_key": str(event.get("event_key") or "")}
        elif kind == RUN_VALUE_REVERTED:
            for item in record.get("reverted") or ():
                if isinstance(item, Mapping):
                    kept.pop(str(item.get("item_id")), None)
    return list(kept.values())


def _actual_words(value: NotebookValue) -> str:
    """What was used, with today's value marked: "20분 (오늘 바꾼 값)", or
    "20분(오늘 바꾼 값) · 12분에 끝냄" when its timer ended early."""

    if not value.planned:
        return value.actual
    if value.actual != value.planned:
        return f"{value.planned}({TODAY_MARK_WORDS}) · {value.actual}"
    return f"{value.actual} ({TODAY_MARK_WORDS})"


def _planned_seconds(shown: str) -> tuple[str, str, str]:
    """A Korean time as today's value shows it ("20분") as a measure's parts."""

    match = re.fullmatch(r"(\d+(?:\.\d+)?)(초|분|시간|일)", shown.strip())
    if match is None:
        return ("", "", shown)
    unit = {"초": "s", "분": "min", "시간": "h"}.get(match[2])
    if unit is None:
        return (str(float(match[1]) * 24), "h", shown)
    return (match[1], unit, shown)


def _condition_measures(text: str) -> list[_Measure]:
    """The values of the notebook's kinds a text states, in order, once each."""

    found: list[_Measure] = []
    for number, unit, written in _measures(text):
        if unit not in _CONDITION_KIND:
            continue
        measure = _Measure(number, unit, " ".join(written.split()))
        if not any(item.unit == measure.unit and item.number == measure.number for item in found):
            found.append(measure)
    return found


#: Words after an item's name that say who sells it or its grade, not what
#: it is (decision ① 3: "재료는 실험에 쓴 것만 이름 정리"). The first word is
#: always the name's ("Promega trypsin", "Eppendorf tubes").
_VENDOR_WORDS = frozenset({
    "merck", "milliporesigma", "sigma-aldrich", "sigma", "aldrich", "vwr", "thermo", "fisher",
    "scientific", "promega", "avantor", "international", "brand", "b&j", "lc-ms", "hplc", "grade",
    "catalog", "cat.", "cat", "whatman", "esslab", "kinesis", "bio-rad", "invitrogen", "gibco",
    "corning", "millipore", "roche", "qiagen", "neb", "abcam", "santa", "cruz",
})
_ITEM_TOKEN = re.compile(r"\([^)]*\)|\S+")


def _short_item_name(name: str) -> str:
    """An item's name for "쓴 재료·장비": without its catalog number, the
    company and grade words after its first word, and parentheses that hold
    neither a short name nor a measurement."""

    text = re.split(r"\s*(?:Catalog|Cat\.?)\s*#.*$|\s+#\S+.*$", item_display_name(name))[0]
    kept: list[str] = []
    for index, token in enumerate(_ITEM_TOKEN.findall(text)):
        if index and token.startswith("("):
            inside = token[1:-1].strip()
            if _SHORT_NAME.fullmatch(inside) or any(ch.isdigit() for ch in inside):
                kept.append(token)
            continue
        if index and token.casefold().strip(",;") in _VENDOR_WORDS:
            break
        kept.append(token)
    return " ".join(kept).strip(" ,;") or " ".join(text.split()) or item_display_name(name)


_REPEAT_SENTENCE = re.compile(r"[^.!?\n]*(?:[Rr]epeat|반복)[^.!?\n]*[.!?]?")


def _repeat_sentence(text: str) -> str:
    """The sentence of a step's text that states the repeat, at most 90 characters."""

    words = " ".join(str(text or "").split())
    match = _REPEAT_SENTENCE.search(words)
    chosen = " ".join(match.group(0).split()) if match else words
    return chosen if len(chosen) <= 90 else chosen[:88].rstrip() + "…"


def pi_confirmations(events: Sequence[Mapping[str, Any]]) -> list[dict[str, str]]:
    """Each PI confirmation of a notebook (decision ④): name, one line, when."""

    found = []
    for event in events:
        if str(event.get("event_type") or "") != PI_CONFIRMED:
            continue
        payload = event.get("payload") if isinstance(event.get("payload"), Mapping) else {}
        found.append({
            "name": " ".join(str(payload.get("name") or "").split()),
            "comment": " ".join(str(payload.get("comment") or "").split()),
            "at": str(event.get("created_at") or ""),
        })
    return found


def _clock_words(moment: datetime | None) -> str:
    return f"{moment.year}년 {moment.month}월 {moment.day}일 {moment:%H:%M}" if moment else ""


_DOMAIN_LABELS = {
    "general": "공통",
    "biology": "바이오",
    "chemistry": "화학 합성",
    "battery_materials": "배터리·소재",
}
_DOMAIN_FIELD_PATTERNS: dict[str, tuple[tuple[str, str], ...]] = {
    "biology": (
        ("세포주", r"\b[A-Za-z][A-Za-z0-9-]{1,24}\s+cells?\b|세포주"),
        ("계대 수", r"passage\s*(?:number|no\.?|#)?\s*\d+|계대\s*(?:수)?\s*\d+"),
        ("플레이트 배치·well 위치", r"\bwell\b|\b\d+\s*[- ]?well\b|웰|plate"),
    ),
    "chemistry": (
        ("반응식", r"reaction\s+scheme|반응식|(?:->|→)"),
        ("시약(몰·질량·당량)", r"\b(?:mmol|mol|equiv(?:alent)?s?|mg|µg|μg|kg)\b|\b(?:TEOS|HCl|ethanol)\b|시약|당량"),
        ("수득량·수율", r"yield|수득|수율"),
        ("TLC Rf", r"\bTLC\b|\bRf\b"),
        ("정제 조건", r"purif|chromatograph|column|recrystalli|정제"),
    ),
    "battery_materials": (
        ("셀·시료 번호", r"\b(?:cell|sample|electrode)\s*[A-Za-z]+[-_]?\d+\b|\b[A-Za-z]+[-_]\d+\b|셀\s*[A-Za-z0-9_-]+|시료\s*[A-Za-z0-9_-]+"),
        ("전극 로딩", r"electrode\s+loading|전극\s*로딩"),
        ("전해액 양", r"electrolyte|전해액"),
        ("조립 환경(수분·산소)", r"glovebox|H2O|O2|수분|산소"),
        ("크림핑·소성 조건", r"crimp|calcina|sinter|크림핑|소성"),
        ("사이클 조건", r"cycl(?:e|ing)|사이클"),
    ),
}

_YIELD_NUMBER_UNIT = r"(\d+(?:\.\d+)?)\s*(mmol|[µμu]mol|mol|mg|[µμu]g|ng|g)\b"
_ACTUAL_YIELD = re.compile(
    r"(?:실제\s*수득량|actual\s+yield)\s*[:：=]?\s*" + _YIELD_NUMBER_UNIT,
    re.I,
)
_THEORETICAL_YIELD = re.compile(
    r"(?:이론(?:적)?\s*수득량|theoretical\s+yield)\s*[:：=]?\s*" + _YIELD_NUMBER_UNIT,
    re.I,
)
_YIELD_UNIT = {
    "ng": ("mass", Decimal("0.000000001")),
    "µg": ("mass", Decimal("0.000001")),
    "mg": ("mass", Decimal("0.001")),
    "g": ("mass", Decimal("1")),
    "µmol": ("amount", Decimal("0.000001")),
    "mmol": ("amount", Decimal("0.001")),
    "mol": ("amount", Decimal("1")),
}


def _yield_unit(unit: str) -> str:
    return unit.replace("μ", "µ").replace("u", "µ").casefold().replace("µ", "µ")


def _decimal_words(value: Decimal) -> str:
    shown = format(value, "f")
    return shown.rstrip("0").rstrip(".") if "." in shown else shown


def _calculated_yield(
    record_evidence: Sequence[tuple[str, str, str]],
) -> NotebookDomainField | None:
    """Calculate yield only when a person's records contain both raw values."""

    actual: tuple[Decimal, str, str, str] | None = None
    theoretical: tuple[Decimal, str, str, str] | None = None
    for origin, key, text in record_evidence:
        if origin != "record":
            continue
        actual_match = _ACTUAL_YIELD.search(text)
        theoretical_match = _THEORETICAL_YIELD.search(text)
        try:
            if actual_match:
                actual = (Decimal(actual_match.group(1)), _yield_unit(actual_match.group(2)),
                          actual_match.group(1), key)
            if theoretical_match:
                theoretical = (Decimal(theoretical_match.group(1)), _yield_unit(theoretical_match.group(2)),
                               theoretical_match.group(1), key)
        except InvalidOperation:
            continue
    if actual is None or theoretical is None or theoretical[0] <= 0:
        return None
    actual_unit = _YIELD_UNIT.get(actual[1])
    theoretical_unit = _YIELD_UNIT.get(theoretical[1])
    if actual_unit is None or theoretical_unit is None or actual_unit[0] != theoretical_unit[0]:
        return None
    percent = (
        actual[0] * actual_unit[1] / (theoretical[0] * theoretical_unit[1]) * Decimal("100")
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    percent_words = _decimal_words(percent)
    actual_words = f"{actual[2]} {actual[1]}"
    theoretical_words = f"{theoretical[2]} {theoretical[1]}"
    return NotebookDomainField(
        label="수득량·수율",
        value=f"실제 {actual_words} · 이론 {theoretical_words} · 수율 {percent_words}%",
        origin="calculated",
        source_keys=tuple(dict.fromkeys((actual[3], theoretical[3]))),
        formula=f"{actual_words} ÷ {theoretical_words} × 100 = {percent_words}%",
    )


def _select_notebook_domain(
    facts: ReportFacts,
    events: Sequence[Mapping[str, Any]],
    source_steps: Sequence[Mapping[str, Any]],
) -> tuple[str, str]:
    """Select a field profile from analysed source text; the latest person
    override in the append-only report ledger wins."""

    allowed = set(_DOMAIN_LABELS)
    for event in reversed(events):
        if str(event.get("event_type") or "") != "notebook_domain_selected":
            continue
        payload = event.get("payload") if isinstance(event.get("payload"), Mapping) else {}
        code = str(payload.get("domain") or "")
        if code in allowed:
            return code, "person"
    source = "\n".join(
        [str(step.get("source_text") or "") for step in source_steps]
        + list(facts.materials) + list(facts.equipment)
    ).casefold()
    scores = {
        "biology": sum(1 for word in (
            "cell culture", " cells", "a549", "well plate", "96-well", "mtt", "passage", "세포", "계대",
        ) if word in source),
        "chemistry": sum(1 for word in (
            "reaction", "synthesis", "teos", "hcl", "ethanol", "reagent", "mmol", "equiv", "tlc", "purif", "sol-gel",
        ) if word in source),
        "battery_materials": sum(1 for word in (
            "coin cell", "cr2032", "electrode", "electrolyte", "glovebox", "crimp", "cathode", "anode", "cycl", "battery",
        ) if word in source),
    }
    selected = max(scores, key=scores.get)
    return (selected, "source_analysis") if scores[selected] else ("general", "source_analysis")


def _notebook_domain(
    facts: ReportFacts,
    events: Sequence[Mapping[str, Any]],
    source_steps: Sequence[Mapping[str, Any]],
) -> NotebookDomain:
    code, selected_by = _select_notebook_domain(facts, events, source_steps)
    source_evidence = [
        ("source", str(step.get("step_id") or step.get("label") or ""),
         " ".join(str(step.get("source_text") or "").split()))
        for step in source_steps
    ]
    record_evidence = [
        ("record", str(event.get("event_key") or ""),
         " ".join(str(event.get("user_wording") or "").split()))
        for event in events if str(event.get("user_wording") or "").strip()
    ]
    evidence = record_evidence + source_evidence
    calculated_yield = _calculated_yield(record_evidence) if code == "chemistry" else None
    fields: list[NotebookDomainField] = []
    for label, pattern in _DOMAIN_FIELD_PATTERNS.get(code, ()):
        if label == "수득량·수율" and calculated_yield is not None:
            fields.append(calculated_yield)
            continue
        matches = [(origin, key, text) for origin, key, text in evidence
                   if text and re.search(pattern, text, re.I)]
        values = tuple(dict.fromkeys(text for _, _, text in matches))
        origins = {origin for origin, _, _ in matches}
        origin = "record" if "record" in origins else "source" if origins else "blank"
        fields.append(NotebookDomainField(
            label=label,
            value=" · ".join(values[:6]),
            origin=origin,
            source_keys=tuple(dict.fromkeys(key for _, key, _ in matches if key)),
        ))
    return NotebookDomain(code, _DOMAIN_LABELS[code], selected_by, tuple(fields))


_CATALOG_VALUE = re.compile(
    r"(?:\bcatalog(?:ue)?(?:\s*(?:number|no\.?|#))?|\bcat\.?\s*(?:no\.?)?|카탈로그\s*(?:번호)?)"
    r"\s*[:：#]?\s*([A-Za-z0-9._/-]+)", re.I,
)
_LOT_VALUE = re.compile(
    r"(?:lot(?:\s*(?:number|no\.?|#|번호))?|로트\s*(?:번호)?)"
    r"\s*[:：#]?\s*([A-Za-z0-9._/-]+)",
    re.I,
)
_VENDOR_VALUE = re.compile(
    r"(?:구매처|공급사|vendor)\s*[:：]?\s*(.+?)(?=\s+(?:catalog|cat\.?|카탈로그|lot|로트)\b|[,;]|$)",
    re.I,
)


def _notebook_inventory(
    facts: ReportFacts,
    events: Sequence[Mapping[str, Any]],
) -> tuple[NotebookInventoryItem, ...]:
    """The source inventory plus optional person-entered Methods fields."""

    rows: dict[str, dict[str, Any]] = {}
    aliases: dict[str, str] = {}
    for item in facts.items:
        listed = item.listed or item.name
        short = _short_item_name(item.name)
        catalog = (_CATALOG_VALUE.search(listed).group(1) if _CATALOG_VALUE.search(listed) else "")
        vendor = ""
        before_catalog = re.split(r"\b(?:Catalog|Cat\.?)\s*#?", listed, maxsplit=1, flags=re.I)[0].strip()
        if catalog and before_catalog.casefold().startswith(short.casefold()):
            vendor = before_catalog[len(short):].strip(" ,;")
        rows[listed] = {
            # Keep the already-normalised source item name in the name
            # column; vendor/catalog source text belongs in its own columns.
            "kind": item.kind, "name": short, "vendor": vendor,
            "catalog_number": catalog, "lot_number": "",
            "stage": "source" if vendor or catalog else "blank", "source_keys": [],
        }
        for alias in {listed, item.name, short}:
            if alias:
                aliases[" ".join(alias.casefold().split())] = listed

    completed_index = next((index for index, event in enumerate(events)
                            if str(event.get("event_type") or "") == "workflow_completed"), len(events))
    for index, event in enumerate(events):
        event_type = str(event.get("event_type") or "")
        payload = event.get("payload") if isinstance(event.get("payload"), Mapping) else {}
        key = str(event.get("event_key") or "")
        if event_type == "notebook_inventory_recorded":
            named = " ".join(str(payload.get("item_name") or "").casefold().split())
            listed = aliases.get(named)
            if listed is None:
                continue
            row = rows[listed]
            for field_name in ("vendor", "catalog_number", "lot_number"):
                value = " ".join(str(payload.get(field_name) or "").split())
                if value:
                    row[field_name] = value
            stage = str(payload.get("stage") or "")
            if stage in {"before", "after"}:
                row["stage"] = stage
            row["source_keys"].append(key)
            continue
        if event_type != "observation":
            continue
        words = " ".join(str(event.get("user_wording") or "").split())
        if not words:
            continue
        listed = next((target for alias, target in aliases.items() if alias and alias in words.casefold()), None)
        if listed is None:
            continue
        vendor = _VENDOR_VALUE.search(words)
        catalog = _CATALOG_VALUE.search(words)
        lot = _LOT_VALUE.search(words)
        if not any((vendor, catalog, lot)):
            continue
        row = rows[listed]
        if vendor:
            row["vendor"] = " ".join(vendor.group(1).split())
        if catalog:
            row["catalog_number"] = catalog.group(1)
        if lot:
            row["lot_number"] = lot.group(1)
        row["stage"] = "before" if index < completed_index else "after"
        row["source_keys"].append(key)
    return tuple(NotebookInventoryItem(
        kind=str(row["kind"]), name=str(row["name"]), vendor=str(row["vendor"]),
        catalog_number=str(row["catalog_number"]), lot_number=str(row["lot_number"]),
        stage=str(row["stage"]), source_keys=tuple(row["source_keys"]),
    ) for row in rows.values())


def _notebook_photos(
    events: Sequence[Mapping[str, Any]], session_id: str,
) -> tuple[NotebookImage, ...]:
    """Recorded image metadata; the image bytes remain in workspace evidence."""

    images = []
    for event in events:
        if str(event.get("event_type") or "") != "photo_attached":
            continue
        payload = event.get("payload") if isinstance(event.get("payload"), Mapping) else {}
        evidence_id = str(payload.get("evidence_id") or "")
        media_type = str(payload.get("media_type") or "image/*")
        url = ""
        if session_id and evidence_id:
            url = (
                f"/api/workspace/experiments/{quote(session_id, safe='')}/evidence/"
                f"{quote(evidence_id, safe='')}"
            )
        images.append(NotebookImage(
            kind="photo", step_label=str(event.get("step_label") or ""),
            caption=" ".join(str(event.get("user_wording") or "사진 첨부").split()),
            url=url, media_type=media_type, asset_id=evidence_id,
            interpretation_status=str(payload.get("interpretation_status") or "not_interpreted"),
        ))
    return tuple(images)


def _with_notebook_source_images(
    facts: ReportFacts, fixture: Any, report_data: Mapping[str, Any], *, maximum: int = 24,
) -> ReportFacts:
    """Attach only deterministic crops of completed steps' source pages."""

    if facts.notebook is None or fixture is None:
        return facts
    try:
        from voiney_lab import source_figures

        steps = tuple(getattr(fixture, "steps", ()) or ())
        protocol_id = str(report_data.get("protocol_id") or getattr(fixture, "protocol_id", ""))
        revision_id = str(report_data.get("protocol_revision") or getattr(fixture, "revision_id", ""))
        completed = set(facts.completed_labels)
        found: list[NotebookImage] = []
        for index, step in enumerate(steps):
            label = str(getattr(step, "source_label", "") or (
                facts.steps[index].label if index < len(facts.steps) else index + 1
            ))
            if label not in completed:
                continue
            for figure in source_figures.figures_for_step(fixture, index):
                found.append(NotebookImage(
                    kind="source_figure", step_label=label,
                    caption=figure.caption or f"원문 그림 · PDF p.{figure.source_page}",
                    url=(
                        f"/api/protocols/{quote(protocol_id, safe='')}/revisions/"
                        f"{quote(revision_id, safe='')}/figures/{quote(figure.figure_id, safe='')}"
                    ),
                    media_type=figure.mime_type, asset_id=figure.figure_id,
                    source_page=figure.source_page,
                ))
                if len(found) >= maximum:
                    break
            if len(found) >= maximum:
                break
    except Exception as exc:  # noqa: BLE001 - a visual never fails the notebook
        log.warning("notebook source figures unavailable error=%s", type(exc).__name__)
        return facts
    book = _replace_dataclass(
        facts.notebook, images=facts.notebook.images + tuple(found),
    )
    return _replace_dataclass(facts, notebook=book)


_PLATE_POSITION = re.compile(
    r"(?<![A-Za-z0-9])(?:[A-Ha-h](?:[1-9]|1[0-2])|\d+\s*번)\s*(?:well|웰)(?![A-Za-z])",
    re.I,
)
_MATERIAL_IDENTIFIER = re.compile(
    r"(?<![A-Za-z0-9])(?:"
    r"[A-Za-z]{1,8}[-_]\d{1,6}"
    r"|\d+\s*번\s*(?:셀|cell|시료|샘플|sample)"
    r"|(?:셀|cell|시료|샘플|sample)\s*\d+\s*번"
    r")(?![A-Za-z0-9])",
    re.I,
)


def _notebook_visual_facts(
    facts: ReportFacts,
    events: Sequence[Mapping[str, Any]],
    values: Sequence[NotebookValue],
) -> tuple[
    tuple[NotebookTimelineItem, ...], tuple[NotebookValueBar, ...],
    tuple[NotebookPlatePosition, ...], tuple[NotebookSampleRow, ...],
]:
    timeline: list[NotebookTimelineItem] = []
    timeline_kinds = {
        "session_started": ("start", "실험 시작"),
        "step_completed": ("step_end", "단계 끝"),
        "timer_started": ("timer_start", "타이머 시작"),
        "timer_paused": ("timer_pause", "타이머 일시정지"),
        "timer_resumed": ("timer_resume", "타이머 다시 시작"),
        "timer_cancelled": ("timer_cancel", "타이머 취소"),
        "timer_restarted": ("timer_restart", "타이머 처음부터"),
        "timer_end_notice": ("timer_end", "타이머 끝"),
        "workflow_paused": ("pause", "실험 일시정지"),
        "workflow_resumed": ("resume", "실험 다시 시작"),
        "observation": ("observation", "관찰·특이사항"),
        "photo_attached": ("photo", "사진"),
    }
    positions: list[NotebookPlatePosition] = []
    sample_records: dict[str, list[str]] = {}
    for event in events:
        event_type = str(event.get("event_type") or "")
        words = " ".join(str(event.get("user_wording") or "").split())
        if event_type in timeline_kinds:
            kind, fallback = timeline_kinds[event_type]
            timeline.append(NotebookTimelineItem(
                at=_clock(_local(event.get("created_at"), facts.zone)),
                step_label=str(event.get("step_label") or ""),
                kind=kind,
                text=words or fallback,
                source_key=str(event.get("event_key") or ""),
            ))
        for match in _PLATE_POSITION.finditer(words):
            label = " ".join(match.group(0).split())
            positions.append(NotebookPlatePosition(
                label=label, step_label=str(event.get("step_label") or ""), words=words,
            ))
        for match in _MATERIAL_IDENTIFIER.finditer(words):
            sample_records.setdefault(match.group(0), []).append(words)
    bars: list[NotebookValueBar] = []
    for value in values:
        if not value.actual:
            continue
        source = _condition_measures(value.written)
        actual = _condition_measures(value.actual)
        pair = next(((left, right) for left in source for right in actual
                     if left.family == right.family), None)
        if pair is None:
            continue
        left, right = pair
        bars.append(NotebookValueBar(
            label=f"{value.label}단계 · {value.kind}", kind=value.kind,
            source=left.written, actual=right.written,
            source_number=left.number, actual_number=right.number,
            unit=left.unit, source_unit=left.unit, actual_unit=right.unit,
            source_key=value.key,
        ))
    sample_rows = tuple(NotebookSampleRow(identifier, tuple(dict.fromkeys(records)))
                        for identifier, records in sample_records.items())
    return tuple(timeline), tuple(bars), tuple(positions), sample_rows


def _notebook_method_rows(
    facts: ReportFacts,
    source_steps: Sequence[Mapping[str, Any]],
    condition_by_label: Mapping[str, str],
) -> tuple[NotebookMethodRow, ...]:
    glance = {str(row[0]).split(" ", 1)[0]: row for row in (
        facts.glance.step_rows if facts.glance is not None else ()
    )}
    by_label = {step.label: step for step in facts.steps}
    rows: list[NotebookMethodRow] = []
    for source in source_steps:
        label = str(source.get("label") or "")
        step = by_label.get(label)
        if step is None or not step.completed:
            continue
        timing = glance.get(label, ())
        rows.append(NotebookMethodRow(
            label=label,
            source_text=str(source.get("source_text") or ""),
            actual_values=condition_by_label.get(label, AS_SOURCE_WORDS),
            started_at=str(timing[1]) if len(timing) > 1 else "",
            ended_at=str(timing[2]) if len(timing) > 2 else step.completed_at,
            timer=step.timer_note,
            source_page=(int(source["page"]) if isinstance(source.get("page"), int) else None),
        ))
    return tuple(rows)


def build_lab_notebook(
    facts: ReportFacts,
    events: Sequence[Mapping[str, Any]],
    *,
    source_steps: Sequence[Mapping[str, Any]],
    zone: ZoneInfo,
    previous: ReportFacts | None = None,
    protocol_problem: str = "",
    session_id: str = "",
) -> LabNotebook:
    """The notebook from the record, the source and the run before (decision ①)."""

    kinds = [str(event.get("event_type") or "") for event in events]
    labels = [str(event.get("step_label") or "") for event in events]
    keys = [str(event.get("event_key") or "") for event in events]
    moments = [_local(event.get("created_at"), zone) for event in events]
    payloads = [event.get("payload") if isinstance(event.get("payload"), Mapping) else {}
                for event in events]
    wordings = [" ".join(str(event.get("user_wording") or "").split()) for event in events]
    index_by_key = {key: index for index, key in enumerate(keys) if key}
    section_of = {step.label: step.section for step in facts.steps}
    for step in source_steps:
        section_of.setdefault(str(step["label"]), _SECTION_DURATION.sub("", str(step.get("section") or "")).strip())
    source_text_of = {str(step["label"]): str(step.get("source_text") or "") for step in source_steps}
    order = {str(step["label"]): index for index, step in enumerate(source_steps)}
    # Lane RP2, decision 3 (kept): a value the experimenter has not confirmed
    # -- asked about or not yet -- is marked wherever the notebook shows it.
    unconfirmed_keys = {item["source_key"] for item in report_review_items(events) if not item["confirmed"]}

    def mark(key: str) -> str:
        return f" {UNCONFIRMED_MARK}" if key and key in unconfirmed_keys else ""

    def process_words(step_labels: Sequence[str]) -> str:
        section = section_of.get(step_labels[0], "")
        runs = _label_runs(list(step_labels))
        return f"{runs}단계 · {section}" if section else f"{runs}단계"

    # Lane N's amendments by the record they amend, with when each was made.
    history: dict[str, list[tuple[str, str, str]]] = {}
    latest_text: dict[str, str] = {}
    for index, kind in enumerate(kinds):
        if kind not in {"record_corrected", "record_retracted"}:
            continue
        fix = payloads[index].get("record_fix") if isinstance(payloads[index].get("record_fix"), Mapping) else {}
        target = str(fix.get("target_event_key") or "")
        if not target:
            continue
        if kind == "record_retracted":
            history.setdefault(target, []).append((_clock(moments[index]), "취소", ""))
        else:
            after = " ".join(str(fix.get("text_after") or "").split())
            history.setdefault(target, []).append((_clock(moments[index]), "정정", after))
            latest_text[target] = after

    # Decision ②'s sayings: the purpose and samples, the result and the next change.
    purpose_words = ""
    samples: list[NotebookSample] = []
    result_words = result = next_change = ""
    for index, kind in enumerate(kinds):
        if kind == PURPOSE_RECORDED and not purpose_words:
            amended = keys[index] in latest_text
            purpose_words = latest_text.get(keys[index]) or wordings[index]
            listed = payloads[index].get("samples") if not amended else None
            if isinstance(listed, (list, tuple)) and listed:
                samples = [NotebookSample(
                    name=" ".join(str(item.get("name") or "").split()),
                    condition=" ".join(str(item.get("condition") or "").split()),
                    control=bool(item.get("control")))
                    for item in listed if isinstance(item, Mapping) and str(item.get("name") or "").strip()]
            else:
                samples = [NotebookSample(item["name"], item["condition"], item["control"])
                           for item in samples_in_text(purpose_words)]
        elif kind == RESULT_RECORDED and not result_words:
            amended = keys[index] in latest_text
            result_words = latest_text.get(keys[index]) or wordings[index]
            payload = payloads[index]
            if not amended and ("result" in payload or "next_change" in payload):
                result = " ".join(str(payload.get("result") or "").split())
                next_change = " ".join(str(payload.get("next_change") or "").split())
            else:
                result, next_change = split_result_and_next(result_words)
    purpose = notebook_title(purpose_words)

    # The header (decision ① 1).
    run = dict(facts.run_rows)
    title = purpose or facts.protocol_title or "실험 노트"
    subtitle = f"{NOTEBOOK_SUBTITLE} · {facts.protocol_title}" if purpose and facts.protocol_title else NOTEBOOK_SUBTITLE
    if facts.outcome == "in_progress":
        when = f"{run.get('날짜', '')} {run.get('시작', '')} – 진행 중"
    else:
        when = f"{run.get('날짜', '')} {run.get('시작', '')} – {run.get('끝', '')} · {run.get('걸린 시간', '')}"
    if facts.review_version_words:
        parts = facts.review_version_words.split(" · ")
        version = f"{parts[0]} ({', '.join(parts[1:])})" if len(parts) > 1 else parts[0]
    else:
        version = "분석 결과 그대로"
    done = run.get("완료 단계", str(len(facts.completed_labels)))
    if facts.outcome == "completed":
        progress = f"끝까지 완료 · {done}단계"
    elif facts.outcome == "stopped":
        progress = (f"중단 — {facts.stop_label}단계에서 · {done}단계" if facts.stop_label
                    else f"중단 · {done}단계")
    else:
        progress = f"진행 중 · {done}단계"
    header: list[tuple[str, str]] = [
        ("날짜·시간", " ".join(when.split())),
        ("실험자", run.get("실험자") or "(직접 적어 주세요)"),
        ("프로토콜", f"{facts.protocol_title} · {version}" if facts.protocol_title else version),
    ]
    reference = facts.protocol_reference
    if reference.url or not reference.title.endswith("(실험 PDF)"):
        header.append(("출처", " ".join(part for part in (reference.title, reference.url) if part)))
    if run.get("기준 파일"):
        header.append(("기준 파일", run["기준 파일"]))
    if run.get("원문 공개 범위"):
        header.append(("원문 공개 범위", run["원문 공개 범위"]))
    if run.get("출처 메모"):
        header.append(("출처 메모", run["출처 메모"]))
    if run.get("프로토콜 상태"):
        header.append(("프로토콜 상태", run["프로토콜 상태"]))
    header.append(("진행", progress))

    # Timers ended early, by step, from the completions (as the first pages read them).
    early: dict[str, tuple[int, int, str]] = {}
    for index, kind in enumerate(kinds):
        if kind != "step_completed" or not labels[index]:
            continue
        timer = payloads[index].get("timer") if isinstance(payloads[index].get("timer"), Mapping) else {}
        if timer.get("step_exited_before_timer_elapsed") or (
                timer.get("completion_state") == "step_exited_before_timer_elapsed"):
            try:
                early[labels[index]] = (
                    int(timer.get("elapsed_seconds") or 0),
                    int(timer.get("source_duration_seconds") or timer.get("duration_seconds") or 0),
                    keys[index])
            except (TypeError, ValueError):
                continue

    # Lane EASY, decision 2: today's values -- the start's, then each change
    # and put-back said during the run, in the record's order.
    planned = {str(item["item_id"]): (item, str(item.get("event_key") or "")) for item in run_values(events)}

    # The conditions actually used (decision ① 3): each performed step's
    # source values, and for each what the record says was used instead.
    live = [record for record in facts.records if record.status != "취소됨"]
    values: list[NotebookValue] = []
    rows: list[tuple[str, str]] = []
    group: list[tuple[str, list[NotebookValue], bool]] = []
    condition_by_label: dict[str, str] = {}

    def flush_group() -> None:
        if not group:
            return
        group_labels = [label for label, _, _ in group]
        if len(group) == 1:
            written = " · ".join(dict.fromkeys(value.written for value in group[0][1]))
        else:
            written = " / ".join(
                f"{label}단계 " + " · ".join(dict.fromkeys(value.written for value in items))
                for label, items, _ in group)
        finished = all(done for _, _, done in group)
        rows.append((process_words(group_labels),
                     f"{AS_SOURCE_WORDS} — {written}" + ("" if finished else " (끝내지 않음)")))
        group.clear()

    for step in facts.steps:
        if step.label not in source_text_of:
            continue
        found = [NotebookValue(step.label, item.kind, item.family, item.written)
                 for item in _condition_measures(source_text_of[step.label])]
        measures = {id(value): item for value, item in zip(
            found, _condition_measures(source_text_of[step.label]))}

        def set_actual(position: int, actual: str, words: str, key: str, is_early: bool = False) -> None:
            if found[position].actual:
                return
            found[position] = _replace_dataclass(
                found[position], actual=actual, words=words, key=key, early=is_early)

        for record in live:
            if record.step_label != step.label or record.kind in {"이상", "사진"}:
                continue
            payload = payloads[index_by_key[record.key]] if record.key in index_by_key else {}
            deviations = [item for item in payload.get("deviations") or () if isinstance(item, Mapping)]
            if record.kind == "편차" and deviations and record.status != "정정됨":
                for item in deviations:
                    said = " ".join(str(item.get("said") or "").split())
                    said_measure = next(iter(_condition_measures(said)), None)
                    sources = [part.strip() for part in str(item.get("source") or "").split("·") if part.strip()]
                    source_measures = [m for part in sources for m in _condition_measures(part)]
                    position = next((i for i, value in enumerate(found) if any(
                        measures[id(value)].unit == m.unit and measures[id(value)].number == m.number
                        for m in source_measures)), None)
                    if position is not None:
                        set_actual(position, said, record.text, record.key)
                    else:
                        kind = said_measure.kind if said_measure else _DEVIATION_KIND_KO.get(
                            str(item.get("kind") or ""), "값")
                        found.append(NotebookValue(
                            step.label, kind, said_measure.family if said_measure else "",
                            sources[0] if sources else "값 없음", actual=said, words=record.text, key=record.key))
                continue
            for measure in _condition_measures(record.text):
                family = [i for i, value in enumerate(found) if measures.get(id(value)) is not None
                          and measures[id(value)].family == measure.family]
                if not family or any(measures[id(found[i])].same_as(measure) for i in family):
                    continue
                set_actual(family[0], measure.written, record.text, record.key)
        # Lane EASY, decision 2: a value today's run chose, where the record
        # does not say something else was done.
        for item, key in planned.values():
            if str(item.get("step_label") or "") != step.label:
                continue
            source = _Measure(str(item.get("source_number") or ""), str(item.get("source_unit") or ""),
                              str(item.get("source_written") or ""))
            shown = " ".join(str(item.get("shown_ko") or "").split())
            position = next((i for i, value in enumerate(found) if measures.get(id(value)) is not None
                             and (measures[id(value)].same_as(source) or (
                                 measures[id(value)].unit == source.unit
                                 and measures[id(value)].number == source.number))), None)
            if position is None:
                found.append(NotebookValue(
                    step.label, str(item.get("kind_ko") or "값"), source.family,
                    source.written or str(item.get("source_shown_ko") or ""),
                    actual=shown, key=key, planned=shown))
                continue
            if not found[position].actual:
                found[position] = _replace_dataclass(found[position], actual=shown, key=key, planned=shown)
        if step.label in early:
            elapsed, duration, key = early[step.label]
            ended = f"{_duration_words(elapsed)}에 끝냄"
            position = next((i for i, value in enumerate(found) if measures.get(id(value)) is not None
                             and measures[id(value)].family == "time" and measures[id(value)].same_as(
                                 _Measure(str(duration), "s", f"{duration} s"))), None)
            if position is None:
                # A timer today's run changed ran for the changed time.
                position = next((i for i, value in enumerate(found) if value.planned and value.actual == value.planned
                                 and _Measure(*_planned_seconds(value.planned)).same_as(
                                     _Measure(str(duration), "s", f"{duration} s"))), None)
                if position is not None:
                    found[position] = _replace_dataclass(found[position], actual=ended, key=key, early=True)
                    position = -1
            if position is not None and position >= 0:
                set_actual(position, ended, "", key, True)
            elif duration:
                found.append(NotebookValue(step.label, "시간", "time", _duration_words(duration),
                                           actual=ended, key=key, early=True))
        values += found
        if not found:
            continue
        differing = [value for value in found if value.actual]
        condition_by_label[step.label] = (
            " · ".join(f"{value.kind} {value.written} → {_actual_words(value)}" for value in differing)
            if differing else f"{AS_SOURCE_WORDS} — " + " · ".join(dict.fromkeys(value.written for value in found)))
        if differing:
            flush_group()
            parts = []
            for value in differing:
                part = f"{value.kind} {value.written} → {_actual_words(value)}"
                if value.early:
                    part += " (타이머 일찍 끝냄)"
                if value.words:
                    part += f" (“{value.words}”)"
                parts.append(part + mark(value.key))
            others = list(dict.fromkeys(value.written for value in found if not value.actual))
            text = " · ".join(parts)
            if others:
                text += f" · 그 밖 {AS_SOURCE_WORDS} ({' · '.join(others)})"
            if not step.completed:
                text += " (끝내지 않음)"
            rows.append((process_words([step.label]), text))
        elif group and section_of.get(group[-1][0], "") == section_of.get(step.label, ""):
            group.append((step.label, found, step.completed))
        else:
            flush_group()
            group.append((step.label, found, step.completed))
    flush_group()
    # A repeat the source states, as the run did it: rounds by confirmed returns.
    repeats: dict[tuple[str, ...], dict[str, Any]] = {}
    for index, kind in enumerate(kinds):
        if kind != "repeat_returned":
            continue
        record = payloads[index].get("step_record") if isinstance(payloads[index].get("step_record"), Mapping) else {}
        span = tuple(str(item) for item in record.get("repeated_step_labels") or ())
        if not span:
            continue
        entry = repeats.setdefault(span, {"round": 0, "stated": str(record.get("stated_at_step") or ""),
                                          "source": " ".join(str(record.get("source_text") or "").split())})
        try:
            entry["round"] = max(entry["round"], int(record.get("round") or 0))
        except (TypeError, ValueError):
            pass
    for span, entry in repeats.items():
        process = f"{span[0]}–{span[-1]}단계 반복" if len(span) > 1 else f"{span[0]}단계 반복"
        if entry["source"]:
            quote = _repeat_sentence(entry["source"])
            process += f" ({entry['stated']}단계 원문: “{quote}”)" if entry["stated"] else f" (원문: “{quote}”)"
        rows.append((process, f"{entry['round']}회차까지 함 (말로 확인한 돌아가기 기준)" if entry["round"]
                     else "돌아감 (회차 기록 없음)"))
    # Lane CB, decision 5 (kept): a source condition answered on the run and
    # a repeat's count a person gave, each with where its value came from.
    for answer in facts.branch_answers:
        skipped = f" · {_ranges(list(answer.skipped_labels))} 건너뜀" if answer.skipped_labels else ""
        rows.append((f"{answer.step_label}단계 조건 “{answer.condition}”",
                     f"{answer.answer} ({answer.value_source}){skipped}"))
    for repetition in facts.repetitions:
        count = f"{repetition.count}회" if repetition.count else "미정"
        asked = " (회차마다 물음)" if repetition.per_round else ""
        rows.append((f"{repetition.first}~{repetition.last}단계 반복 횟수",
                     f"{count} — {repetition.value_source}{asked}"))
    condition_note = ""
    if protocol_problem or not any(str(step.get("source_text") or "") for step in source_steps):
        condition_note = NO_SOURCE_VALUES_WORDS
    elif not rows:
        condition_note = NO_VALUED_STEPS_WORDS
    performed = tuple(step.label for step in facts.steps)
    used = [_short_item_name(item.name) for item in facts.items if set(item.steps) & set(performed)]
    if used:
        materials_words = "쓴 재료·장비: " + " · ".join(dict.fromkeys(used))
    elif facts.items:
        materials_words = "쓴 재료·장비: 원문 목록 중 수행한 단계 글에 나오는 것을 찾지 못했다."
    else:
        materials_words = ""

    # Observations and results (decision ① 4): by step, with the sample named.
    observations: list[NotebookObservation] = []
    for record in live:
        if record.kind not in {"측정", "관찰", "메모", "사진"}:
            continue
        payload = payloads[index_by_key[record.key]] if record.key in index_by_key else {}
        text = record.text
        if record.kind == "사진":
            text = "사진: " + text.removeprefix("사진 첨부 — ").removeprefix("사진 첨부").strip() if text != "사진 첨부" else "사진"
        shown = text
        if record.status == "정정됨":
            shown = f"{text} (정정됨 — 처음 “{record.original}”)"
        shown += mark(record.key)
        sample = " ".join(str(payload.get("sample") or "").split()) or sample_in_text(record.text) or "—"
        observations.append(NotebookObservation(
            label=record.step_label, process=process_words([record.step_label]) if record.step_label != "—" else "—",
            sample=sample, text=text, shown=shown, at=record.at or "—", kind=record.kind))

    # What was done differently from the source, and the problems (decision ① 5).
    differences: list[NotebookDifference] = []
    for index, kind in enumerate(kinds):
        record = payloads[index].get("step_record") if isinstance(payloads[index].get("step_record"), Mapping) else {}
        if kind == "steps_skipped" and record.get("skipped_step_labels"):
            skipped = [str(item) for item in record.get("skipped_step_labels") or ()]
            start = str(record.get("start_step") or labels[index])
            differences.append(NotebookDifference(-1, f"{start}단계부터 시작 — {_ranges(skipped)} 건너뜀" + mark(keys[index])))
        elif kind == "repeat_rounds_declined":
            first, last = _range_pair(record)
            differences.append(NotebookDifference(
                index, f"{first}~{last}단계 반복: {record.get('rounds_required')}회 중 "
                       f"{record.get('rounds_done')}회차까지만 하고 넘어감 (회차는 말로 확인한 돌아가기 기준)" + mark(keys[index])))
        elif kind == "step_reverted":
            reverted = [str(item) for item in record.get("reverted_step_labels") or ()]
            span = reverted[0] if len(reverted) == 1 else f"{reverted[0]}~{reverted[-1]}" if reverted else ""
            differences.append(NotebookDifference(
                index, f"{record.get('from_step') or ''}단계에서 {record.get('to_step') or labels[index]}단계로 되돌아감"
                       + (f" — {span}단계를 다시 함" if span else "")))
    if not any(kind == "steps_skipped" for kind in kinds) and source_steps and facts.completed_labels:
        first = min(order.get(label, 0) for label in facts.completed_labels)
        before = [str(step["label"]) for step in source_steps[:first]]
        if before:
            differences.append(NotebookDifference(
                -1, f"{_ranges(before)}는 이 기록에 없다 ({facts.completed_labels[0]}단계부터 기록됨)"))
    for value in values:
        if value.actual and not value.early:
            position = index_by_key.get(value.key, len(kinds))
            differences.append(NotebookDifference(
                position, f"{value.label}단계 {value.kind}: 원문 {value.written} → {_actual_words(value)}"
                          + (f" — “{value.words}”" if value.words else "") + mark(value.key)))
    for label, (elapsed, duration, key) in early.items():
        # Lane EASY, decision 2: a timer today's run changed is today's, not the source's.
        today = any(value.label == label and value.planned and value.early for value in values)
        differences.append(NotebookDifference(
            index_by_key.get(key, len(kinds)),
            f"{label}단계 타이머: {TODAY_MARK_WORDS if today else '원문'} {_duration_words(duration)} 중 "
            f"{_duration_words(elapsed)}에 끝냄" + mark(key)))
    for record in live:
        if record.kind == "이상":
            differences.append(NotebookDifference(
                index_by_key.get(record.key, len(kinds)),
                f"이상 · {record.step_label}단계 ({record.at or '시각 없음'}): “{record.text}”" + mark(record.key)))
    for index, kind in enumerate(kinds):
        # Lane HELP: what the sources said back after a problem, one line.
        if kind == PROBLEM_HELP_EVENT:
            differences.append(NotebookDifference(
                index, problem_help_line(labels[index], payloads[index], wordings[index], _clock(moments[index]))))
    if facts.outcome == "stopped" and source_steps:
        last_done = max((order.get(label, -1) for label in facts.completed_labels), default=-1)
        after = [str(step["label"]) for step in source_steps
                 if str(step["label"]) not in facts.completed_labels and order.get(str(step["label"]), 0) > last_done]
        if after:
            where = f"{facts.stop_label}단계에서 끝냄" if facts.stop_label else "실험을 끝냄"
            differences.append(NotebookDifference(len(kinds) + 1, f"{where} — {_ranges(after)}는 하지 않음"))
    differences.sort(key=lambda item: item.order)

    # What differed from the run before (decision ① 6).
    previous_words = ""
    previous_lines: list[str] = []
    if previous is not None and previous.notebook is not None:
        before = previous.notebook
        previous_words = f"지난번: {before.run_words}"
        if (purpose or before.purpose) and purpose != before.purpose:
            previous_lines.append(f"목적: 지난번 {('“' + before.purpose + '”') if before.purpose else NOT_SAID_WORDS.rstrip('.')}"
                                  f" → 이번 {('“' + purpose + '”') if purpose else NOT_SAID_WORDS.rstrip('.')}")
        mine = [sample.name for sample in samples]
        theirs = [sample.name for sample in before.samples]
        if mine != theirs and (mine or theirs):
            previous_lines.append(f"시료: 지난번 {' · '.join(theirs) or NOT_SAID_WORDS.rstrip('.')}"
                                  f" → 이번 {' · '.join(mine) or NOT_SAID_WORDS.rstrip('.')}")
        performed_before = set(before.performed)
        performed_now = set(performed)

        def value_words(value: NotebookValue | None, written: str) -> str:
            return value.actual if value is not None and value.actual else f"{AS_SOURCE_WORDS} ({written})"

        seen: list[tuple[str, str, str]] = []
        for value in list(values) + list(before.values):
            key = (value.label, value.kind, value.written)
            if key not in seen:
                seen.append(key)
        seen.sort(key=lambda key: order.get(key[0], 10_000))
        for key in seen:
            if key[0] not in performed_before or key[0] not in performed_now:
                continue
            now = next((v for v in values if (v.label, v.kind, v.written) == key), None)
            then = next((v for v in before.values if (v.label, v.kind, v.written) == key), None)
            if (now.actual if now else "") != (then.actual if then else ""):
                previous_lines.append(
                    f"{key[0]}단계 {key[1]}: 지난번 {value_words(then, key[2])} → 이번 {value_words(now, key[2])}")
        only_now = [label for label in performed if label not in performed_before]
        only_then = [label for label in before.performed if label not in performed_now]
        if only_now:
            previous_lines.append(f"이번에만 한 단계: {_ranges(only_now)}")
        if only_then:
            previous_lines.append(f"지난번에만 한 단계: {_ranges(only_then)}")
        step_labels = sorted({item.label for item in observations} | {item.label for item in before.observations},
                             key=lambda label: order.get(label, 10_000))
        for label in step_labels:
            now_texts = [f"“{item.text}”" for item in observations if item.label == label]
            then_texts = [f"“{item.text}”" for item in before.observations if item.label == label]
            if now_texts == then_texts:
                continue
            if now_texts and then_texts:
                previous_lines.append(f"{label}단계 관찰: 지난번 {' · '.join(then_texts)} → 이번 {' · '.join(now_texts)}")
            elif now_texts:
                previous_lines.append(f"{label}단계 관찰: 이번에만 기록 — {' · '.join(now_texts)}")
            else:
                previous_lines.append(f"{label}단계 관찰: 지난번에만 기록 — {' · '.join(then_texts)}")
        if (result or before.result) and result != before.result:
            previous_lines.append(f"결과: 지난번 {('“' + before.result + '”') if before.result else NOT_SAID_WORDS.rstrip('.')}"
                                  f" → 이번 {('“' + result + '”') if result else NOT_SAID_WORDS.rstrip('.')}")

    # What to change next (decision ① 7).
    next_lines: list[str] = []
    if next_change:
        next_lines.append(f"{next_change} (실험자가 말함)")

    # The confirmation box (decisions ① 8, ④).
    pi_rows: list[tuple[str, str, str, str]] = []
    for index, kind in enumerate(kinds):
        if kind != PI_CONFIRMED:
            continue
        pi_rows.append(("PI", " ".join(str(payloads[index].get("name") or "").split()),
                        _clock_words(moments[index]), " ".join(str(payloads[index].get("comment") or "").split())))

    # The appendix (decision ① 9).
    timer_words = {step.label: step.timer_note for step in facts.steps}
    step_rows: list[tuple[str, ...]] = []
    for row in (facts.glance.step_rows if facts.glance is not None else ()):
        label = row[0].split(" ")[0]
        if row[1] == "—" and row[2] in {"수행 안 함", "건너뜀"}:
            continue
        step_rows.append((row[0], row[1], row[2], row[3], timer_words.get(label) or "—"))
    record_rows: list[tuple[str, ...]] = []
    for index, kind in enumerate(kinds):
        text = wordings[index]
        if kind == "observation" and text:
            said_kind = RECORD_KINDS.get(str(events[index].get("category") or ""), "관찰")
        elif kind == "anomaly" and text:
            said_kind = "이상"
        elif kind == "photo_attached":
            said_kind, text = "사진", text or "사진 첨부"
        elif kind in SAYING_KINDS and text:
            said_kind = SAYING_KINDS[kind]
        elif kind == "step_completed" and text:
            said_kind = "관찰"
        elif kind == PROBLEM_HELP_EVENT and text:
            said_kind = "문제 질문"
        else:
            continue
        for clock, change, after in history.get(keys[index], ()):
            text += f" → {clock} {change}" + (f": {after}" if after else "")
        record_rows.append((_clock(moments[index]), labels[index] or "—", said_kind, text))

    connections = []
    for index, kind in enumerate(kinds):
        if kind != "connection_restored":
            continue
        lost = _local(payloads[index].get("lost_at"), zone)
        restored = _local(payloads[index].get("restored_at"), zone) or moments[index]
        gap = payloads[index].get("gap_seconds")
        gap_words = (_span_words(gap) if isinstance(gap, (int, float)) and not isinstance(gap, bool) else "기록 없음")
        where = f"{labels[index]}단계 " if labels[index] else ""
        connections.append(f"{where}{_clock(lost)}–{_clock(restored)}, {gap_words}")
    spoken = sum(1 for i, k in enumerate(kinds) if k in {"timer_last_minute_notice", "timer_end_notice"} and payloads[i].get("spoken"))
    shown_notices = sum(1 for i, k in enumerate(kinds) if k in {"timer_last_minute_notice", "timer_end_notice"} and not payloads[i].get("spoken"))
    system = [f"연결 끊김 {len(connections)}회" + (f" ({' · '.join(connections)})" if connections else "")]
    notices = spoken + shown_notices
    system.append(f"타이머 알림 {notices}회" + (f" (말함 {spoken} · 화면 {shown_notices})" if notices else ""))
    prestart = facts.glance.prestart_rows if facts.glance is not None else None
    if prestart is None:
        system.append("시작 전 확인 기록을 읽지 못함")
    elif not prestart:
        system.append("시작 전 확인 기록 없음")
    else:
        counts = {name: sum(1 for row in prestart if row[4] == name) for name in ("맞음", "다르게 고름", "원문대로 맞다고 함")}
        system.append(f"시작 전 확인 {len(prestart)}답 (" + " · ".join(f"{name} {count}" for name, count in counts.items()) + ")")
    safety = next((payloads[i].get("notice_count") for i, k in enumerate(kinds) if k == "safety_notices_acknowledged"), None)
    if isinstance(safety, int) and not isinstance(safety, bool):
        system.append(f"안전 주의 {safety}건 보고 시작" if safety else "원문에 안전 주의 없음 · 시작 전 화면 보고 시작")
    system.append(f"멈춤 {sum(1 for k in kinds if k == 'workflow_paused')}회")
    advanced = sum(1 for k in kinds if k == "step_advanced")
    if advanced:
        system.append(f"완료 말 없이 넘어감 {advanced}회")
    source_words = f"[{reference.number}] {reference.title}" + (f" {reference.url}" if reference.url else "") + " — 실험에 쓴 프로토콜 원문"
    run_words = " · ".join(part for part in (
        f"{run.get('날짜', '')} {run.get('시작', '')}".strip(), run.get("실험자", ""),
        f"{run.get('결과', '')} ({done}단계)") if part)
    method_rows = _notebook_method_rows(facts, source_steps, condition_by_label)
    timeline, value_bars, plate_positions, sample_rows = _notebook_visual_facts(
        facts, events, values,
    )
    domain = _notebook_domain(facts, events, source_steps)
    inventory = _notebook_inventory(facts, events)
    images = _notebook_photos(events, session_id)
    return LabNotebook(
        title=title, subtitle=subtitle, header_rows=tuple(header), purpose=purpose, purpose_words=purpose_words,
        samples=tuple(samples), values=tuple(values), condition_rows=tuple(rows), condition_note=condition_note,
        materials_words=materials_words, observations=tuple(observations), result=result, result_words=result_words,
        next_change=next_change, differences=tuple(item.text for item in differences),
        previous_words=previous_words, previous_lines=tuple(previous_lines), next_lines=tuple(next_lines),
        pi_rows=tuple(pi_rows), appendix_step_rows=tuple(step_rows), appendix_record_rows=tuple(record_rows),
        system_words="시스템 사건: " + " · ".join(system), source_words=source_words, performed=performed,
        run_words=run_words, condition_by_label=condition_by_label,
        method_rows=method_rows, timeline=timeline, value_bars=value_bars,
        plate_positions=plate_positions, sample_rows=sample_rows, domain=domain,
        interpretation_words="", inventory=inventory, images=images,
    )


def audit_notebook_number_provenance(
    notebook: LabNotebook,
    *,
    source_steps: Sequence[Mapping[str, Any]],
    source_inventory: Sequence[str] = (),
    events: Sequence[Mapping[str, Any]],
) -> tuple[str, ...]:
    """Return displayed numbers that cannot be reproduced from source or ledger.

    Besides verbatim values, the notebook displays deterministic event counts
    and elapsed times.  Those are admitted only when recomputable from the
    supplied ledger timestamps.  Identifiers used solely to link records are
    not notebook prose and are excluded from this display audit.
    """

    def normalized(value: str) -> str:
        try:
            number = float(value.replace(",", ""))
        except ValueError:
            return value
        return str(int(number)) if number.is_integer() else format(number, ".12g")

    source_blob = json.dumps(
        {"steps": list(source_steps), "inventory": list(source_inventory)},
        ensure_ascii=False, default=str,
    )
    event_blob = json.dumps(list(events), ensure_ascii=False, default=str)
    admitted = {normalized(value) for value in _NUMBER.findall(source_blob + "\n" + event_blob)}
    # Source identity (citation/DOI and protocol title) is loaded from the
    # immutable protocol revision but is not repeated in ``source_steps``.
    admitted.update(normalized(value) for value in _NUMBER.findall(
        "\n".join((notebook.title, notebook.subtitle, notebook.source_words))
    ))
    # A displayed yield is admitted only when the same deterministic
    # calculation can be rebuilt from person-recorded raw values.
    rebuilt_yield = _calculated_yield([
        ("record", str(event.get("event_key") or ""),
         " ".join(str(event.get("user_wording") or "").split()))
        for event in events if str(event.get("user_wording") or "").strip()
    ])
    shown_yield = next((field for field in notebook.domain.fields
                        if field.label == "수득량·수율"), None)
    if rebuilt_yield is not None and shown_yield == rebuilt_yield:
        admitted.update(normalized(value) for value in _NUMBER.findall(
            rebuilt_yield.value + "\n" + rebuilt_yield.formula
        ))
    moments = []
    for event in events:
        try:
            moment = datetime.fromisoformat(str(event.get("created_at") or "").replace("Z", "+00:00"))
        except ValueError:
            continue
        moments.append(moment)
        admitted.update(str(value) for value in (
            moment.year, moment.month, moment.day, moment.hour, moment.minute, moment.second,
        ))
    for left in moments:
        for right in moments:
            seconds = abs(int((right - left).total_seconds()))
            admitted.add(str(seconds))
            if seconds % 60 == 0:
                admitted.add(str(seconds // 60))
            if seconds % 3600 == 0:
                admitted.add(str(seconds // 3600))

    fragments: list[tuple[str, str]] = []

    def visit(value: Any, path: str = "notebook") -> None:
        if is_dataclass(value):
            for item in _dataclass_fields(value):
                if item.name in {"source_key", "source_keys"}:
                    continue
                visit(getattr(value, item.name), f"{path}.{item.name}")
        elif isinstance(value, Mapping):
            for key, item in value.items():
                visit(item, f"{path}.{key}")
        elif isinstance(value, (list, tuple)):
            for index, item in enumerate(value):
                visit(item, f"{path}[{index}]")
        elif isinstance(value, str):
            fragments.append((path, value))
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            fragments.append((path, str(value)))

    visit(notebook)
    problems = []
    for path, text in fragments:
        for value in _NUMBER.findall(text):
            if normalized(value) not in admitted:
                problems.append(f"{path}: {value}")
    return tuple(dict.fromkeys(problems))


# --- The server's checks on model sections (decision 5) ------------------------

_NUMBER = re.compile(r"\d+(?:\.\d+)?")
_CITATION = re.compile(r"\[(\d+(?:\s*,\s*\d+)*)\]")
_CONDITION_NUMBER = re.compile(
    r"\d+(?:[.,]\d+)?\s*(?:°\s*C|℃|도\b|분|min\b|시간|hours?\b|h\b|초|sec\b|s\b|mM\b|µM|μM|nM|M\b|"
    r"µL|μL|uL|ul\b|mL|ml\b|L\b|rpm|x\s*g\b|×\s*g\b|g\b|mg|µg|ug|ng|%|v/v|w/v|mm\b|cm\b|일\b)",
    re.I,
)
_IDENTIFIER_SHAPES = (
    ("16자 이상 16진수", re.compile(r"\b[0-9a-fA-F]{16,}\b")),
    ("ER-/fixture-/candidate- 식별자", re.compile(r"\b(?:ER|fixture|candidate)-[A-Za-z0-9-]+")),
    ("밀리초 시각", re.compile(r"\d{1,2}:\d{2}:\d{2}\.\d+")),
    ("영어 상태값·명령 이름", re.compile(
        r"\b(?:[a-z]+(?:_[a-z]+)+|stopped|completed|blocked|incomplete|in progress)\b")),
)
#: A quotation is closed by its own closing mark: a straight quote by a
#: straight quote, “ by ”, and so on. A record that itself holds curly quotes
#: (an endpoint answer quotes the source) can then sit inside straight ones.
_QUOTE_MARKS = (('"', '"'), ("“", "”"), ("‘", "’"), ("'", "'"), ("「", "」"), ("『", "』"))
_QUOTED = re.compile("|".join(
    f"{re.escape(opening)}([^{re.escape(closing)}]{{2,}}){re.escape(closing)}"
    for opening, closing in _QUOTE_MARKS
))
_SPECULATION = ("추정", "것으로 보인", "보인다", "아마", "가능성", "것 같", "듯하", "듯이", "인 듯", "한 듯")
_SENTENCE = re.compile(r"(?<=[.!?。])\s+")


#: Units a number may carry, each written the ways a protocol or a person
#: writes it, and the one form they are compared in (decision 3: "15분" is
#: "15 min", "µL" is "uL"). A number with a unit must appear with the same
#: unit in the record or the source; a bare number is checked as before.
_UNIT_FORMS = (
    (r"°\s*C|℃|도", "°C"),
    (r"시간|hours?|hrs?|h", "h"),
    (r"분|minutes?|mins?|min|m", "min"),
    (r"초|seconds?|secs?|sec|s", "s"),
    (r"[µμu][Ll]", "µL"),
    (r"m[Ll]", "mL"),
    (r"L", "L"),
    (r"mM", "mM"),
    (r"[µμu]M", "µM"),
    (r"nM", "nM"),
    (r"M", "M"),
    (r"rpm|RPM", "rpm"),
    (r"mm/s", "mm/s"),
    (r"[x×]\s*g", "×g"),
    (r"mg", "mg"),
    (r"[µμu]g", "µg"),
    (r"ng", "ng"),
    (r"g", "g"),
    (r"mm", "mm"),
    (r"cm", "cm"),
    (r"[µμ]m", "µm"),
    (r"kDa", "kDa"),
    (r"%", "%"),
)
_UNIT_PATTERNS = tuple((re.compile(pattern), unit) for pattern, unit in _UNIT_FORMS)
_MEASURE = re.compile(
    r"(?<![\d.])(\d+(?:\.\d+)?)\s*(" + "|".join(
        # Longer spellings first so "min" is not read as "m" and "mL" not as "m".
        sorted((p for pattern, _ in _UNIT_FORMS for p in pattern.split("|")), key=len, reverse=True)
    ) + r")(?![A-Za-zµμ])"
)


def _plain(text: str) -> str:
    """Text without citation marks or thousands separators."""

    return re.sub(r"(?<=\d),(?=\d{3}(?!\d))", "", _CITATION.sub(" ", text))


def _number_form(item: str) -> str:
    return item if "." in item else (item.lstrip("0") or "0")


def _numbers(text: str) -> set[str]:
    return {_number_form(item) for item in _NUMBER.findall(_plain(text))}


def _measures(text: str) -> list[tuple[str, str, str]]:
    """(number, unit, as written) for every number that carries a unit."""

    found = []
    for match in _MEASURE.finditer(_plain(text)):
        written = match.group(2)
        unit = next(unit for pattern, unit in _UNIT_PATTERNS if pattern.fullmatch(written))
        found.append((_number_form(match.group(1)), unit, match.group(0)))
    return found


def _record_texts(facts: ReportFacts) -> list[str]:
    """Every text a number in a model section may come from."""

    # Lane N, decision 3: a value the researcher did not confirm is listed
    # apart, and model prose may not state it; nor a withdrawn record.
    unconfirmed = set(facts.unconfirmed_texts)
    texts: list[str] = [facts.protocol_title, facts.keywords, facts.purpose_from_pdf]
    texts += [value for _, value in facts.run_rows]
    texts += [line for line in facts.deviations if line not in unconfirmed]
    texts += list(facts.confirmed) + list(facts.to_check)
    texts += list(facts.materials) + list(facts.equipment) + list(facts.sections)
    for step in facts.steps:
        texts += [step.label, step.text, step.source_text, step.completed_at, step.timer_note, *step.expected]
        if step.source_timer_seconds:
            texts.append(_duration_words(step.source_timer_seconds))
    # The source steps the model reads for the items' uses (decision 4) are
    # the source too, performed or not.
    texts += [text for _, text in facts.item_step_texts]
    texts += [f"{record.number} {record.text} {record.at}" for record in facts.records
              if record.text not in unconfirmed and record.status != "취소됨"]
    texts.append(str(facts.total_steps))
    texts.append(" ".join(str(n) for n in range(1, facts.total_steps + 1)))
    texts.append(" ".join(str(n) for n in range(0, len(facts.records) + 1)))
    texts.append(facts.protocol_reference.title)
    return [str(text) for text in texts]


def _allowed_numbers(facts: ReportFacts) -> set[str]:
    allowed: set[str] = set()
    for text in _record_texts(facts):
        allowed |= _numbers(text)
    return allowed


def _allowed_measures(facts: ReportFacts) -> set[tuple[str, str]]:
    return {(number, unit) for text in _record_texts(facts) for number, unit, _ in _measures(text)}


def _as_text(value: Any) -> str:
    if isinstance(value, (list, tuple)):
        parts = []
        for item in value:
            if isinstance(item, Mapping):
                parts.append(" ".join(str(v) for v in item.values()))
            else:
                parts.append(str(item))
        return "\n".join(parts)
    return "" if value is None else str(value)


#: Lane N, decision 8: a number of rounds or repetitions in a model sentence
#: ("총 2회 반복", "2회차", "두 번 반복"). Rounds are counted from returns
#: confirmed in words, and a sentence giving one must say so.
_ROUND_COUNT = re.compile(
    r"\d+\s*(?:회|번)\s*(?:째\s*)?(?:반복|차)|회차|반복\s*\d+\s*(?:회|번)"
    r"|(?:한|두|세|네|다섯)\s*(?:번|차례)\s*(?:째\s*)?(?:반복|돌|수행)"
    r"|\d+\s*회\s*(?:수행|실시|진행)"
)
ROUND_BASIS = "말로 확인한 돌아가기 기준"

#: Lane N, decision 7: the words a general-knowledge background may not hold
#: -- an instruction, a safety direction -- beyond numbers, units, citations
#: and identifiers.
_GENERAL_INSTRUCTION = re.compile(
    r"하세요|하십시오|해라|하라|할 것|해야|하지 마|마세요|마십시오|반드시|금지"
)
# --- The server's checks on AI-written notes (lane RP3, decision ③) ------------
#
# The model writes the observation sentences and one result paragraph; every
# number and unit in a sentence must be in the record or the source, a value
# the experimenter has not confirmed may not be stated, and no safety or
# procedural direction is written. A sentence that fails is dropped and the
# record's own words stand; one that passes is marked "AI 정리".

#: Not "흡입": for a pipette it is drawing liquid up (live run, headspace);
#: a warning about breathing something in still holds "주의", "위험" or "하지 마".
_USE_SAFETY = ("주의", "위험", "안전", "유해", "독성", "부식", "인화", "폭발", "화상", "응급", "보호",
               "장갑", "보안경", "고글", "마스크", "환기", "후드", "피부", "눈에", "눈을",
               "금지", "MSDS")
AI_NOTE_LABEL = "AI 정리"
#: Reasons a sentence is dropped, in the words the notebook's appendix uses.
AI_NOTE_LIMIT = 300


@dataclass(frozen=True)
class AiNote:
    """One AI-written sentence that passed its check, at the step (and sample) it is about."""

    step_label: str
    sample: str
    text: str


def _ai_sentence_problem(
    text: str, *, label: str | None, labels: set[str], allowed: set[str],
    allowed_measures: set[tuple[str, str]], unconfirmed_texts: Sequence[str],
    unconfirmed_measures: set[tuple[str, str]], record_texts: Sequence[str],
) -> str:
    """Why an AI-written sentence may not be used, or "" when it may."""

    plain = " ".join(text.split())
    if not plain:
        return "빈 문장"
    if label is not None and label not in labels:
        return f"기록에 없는 단계 {label}"
    for name, shape in _IDENTIFIER_SHAPES:
        if shape.search(_CITATION.sub(" ", plain)):
            return f"본문에 {name}"
    if any(value and value in plain for value in unconfirmed_texts) or any(
            (number, unit) in unconfirmed_measures for number, unit, _ in _measures(plain)):
        return "확인되지 않은 값"
    unknown = [written for number, unit, written in _measures(plain) if (number, unit) not in allowed_measures]
    if unknown:
        return "기록·원문에 없는 값 " + ", ".join(dict.fromkeys(unknown))
    extra = _numbers(_MEASURE.sub(" ", _plain(plain))) - allowed
    if extra:
        return "기록·원문에 없는 숫자 " + ", ".join(sorted(extra, key=lambda x: (len(x), x))[:6])
    if _ROUND_COUNT.search(plain) and ROUND_BASIS not in plain:
        # Lane N, decision 8: a round counted with no basis.
        return f"회차를 ‘{ROUND_BASIS}’ 없이 씀"
    for word in _SPECULATION:
        if word in plain:
            return f"추측 표현 “{word}”"
    for match in _QUOTED.finditer(plain):
        quoted = " ".join(next(group for group in match.groups() if group is not None).split())
        if not any(quoted in record or record in quoted for record in record_texts):
            return f"기록에 없는 관찰 “{quoted[:30]}”"
    safety = next((word for word in _USE_SAFETY if word in plain), None)
    if safety:
        return f"안전 지시 “{safety}”"
    if _GENERAL_INSTRUCTION.search(plain):
        return "절차·지시"
    if len(plain) > AI_NOTE_LIMIT:
        return f"{AI_NOTE_LIMIT}자 넘음"
    return ""


def check_ai_notes(
    reply: Mapping[str, Any], facts: ReportFacts,
) -> tuple[tuple[AiNote, ...], str, tuple[tuple[str, str], ...]]:
    """(the notes that passed, the result paragraph or "", ((sentence, why dropped), ...))."""

    allowed = _allowed_numbers(facts)
    allowed_measures = _allowed_measures(facts)
    unconfirmed_texts = [" ".join(str(text).split()) for text in facts.unconfirmed_texts]
    unconfirmed_measures = {
        (number, unit) for text in unconfirmed_texts for number, unit, _ in _measures(text)}
    labels = {step.label for step in facts.steps}
    if facts.glance is not None:
        labels |= {str(row[0]).split(" ")[0] for row in facts.glance.step_rows}
    check = dict(labels=labels, allowed=allowed, allowed_measures=allowed_measures,
                 unconfirmed_texts=unconfirmed_texts, unconfirmed_measures=unconfirmed_measures,
                 record_texts=[" ".join(text.split()) for text in facts.record_texts])
    notes: list[AiNote] = []
    rejected: list[tuple[str, str]] = []
    items = reply.get("observation_notes") if isinstance(reply, Mapping) else None
    for item in items if isinstance(items, (list, tuple)) else ():
        if not isinstance(item, Mapping):
            continue
        label = " ".join(str(item.get("단계") or "").split())
        sample = " ".join(str(item.get("시료") or "").split())
        text = " ".join(str(item.get("문장") or "").split())
        why = _ai_sentence_problem(text, label=label or None, **check)
        if why:
            rejected.append((text, why))
        else:
            notes.append(AiNote(step_label=label, sample=sample, text=text))
    result = " ".join(str((reply.get("result_summary") if isinstance(reply, Mapping) else "") or "").split())
    if result:
        why = _ai_sentence_problem(result, label=None, **check)
        if why:
            rejected.append((result, why))
            result = ""
    return tuple(notes), result, tuple(rejected)


# --- The narrative and the server's own sentences -------------------------------


@dataclass(frozen=True)
class ReportNarrative:
    """The notebook's sentences: the server's, and the AI-written notes that
    passed their check (lane RP3, decision ③)."""

    title: str
    facts: ReportFacts | None = None
    sources: tuple[ReportSource, ...] = ()
    writer: str = "서버"
    written_at: str = ""
    ai_notes: tuple[AiNote, ...] = ()
    ai_result: str = ""
    #: Each AI-written sentence dropped, with why.
    ai_rejected: tuple[tuple[str, str], ...] = ()
    failure: str = ""

    @property
    def ai_written(self) -> bool:
        return bool(self.ai_notes or self.ai_result)


def narrative_from_sections(
    facts: ReportFacts,
    model: Mapping[str, Any] | None = None,
    *,
    rejected: Mapping[str, Sequence[str]] | None = None,
    writer: str = "서버",
    written_at: str | None = None,
    failure: str = "",
) -> ReportNarrative:
    """The notebook's narrative: the server's sentences, and the model's notes
    when a reply is given and they pass (``rejected`` is kept for callers
    written before decision ③; it is not read)."""

    notes, result, dropped = check_ai_notes(model, facts) if model else ((), "", ())
    return ReportNarrative(
        title=(facts.notebook.title if facts.notebook is not None else facts.protocol_title) or "실험 노트",
        facts=facts, sources=(facts.protocol_reference,), writer=writer,
        written_at=written_at or _now(), ai_notes=notes, ai_result=result, ai_rejected=dropped,
        failure=failure,
    )


@dataclass(frozen=True)
class ReportDraftState:
    """Server-owned structured draft state for specialized report writing."""

    report_id: str
    session_id: str
    protocol_id: str
    status: str
    experiment_summary: str = ""
    materials_summary: str = ""
    equipment_summary: str = ""
    observations_narrative: str = ""
    anomalies_narrative: str = ""
    timeline_narrative: str = ""
    conclusion_narrative: str = ""
    committed_event_ids: tuple[str, ...] = ()
    last_updated_at: str = ""


@dataclass(frozen=True)
class ReportWriterSettings:
    enabled: bool = True
    model: str = ""
    timeout_seconds: float = 25.0

    @classmethod
    def from_environment(cls) -> "ReportWriterSettings":
        enabled_val = os.environ.get(
            "VOINEY_LAB_REPORT_WRITER_ENABLED", "true"
        ).strip().casefold()
        enabled = enabled_val in _TRUE
        model = os.environ.get(
            "VOINEY_LAB_REPORT_MODEL",
            os.environ.get("VOINEY_LAB_SUPPLEMENTAL_MODEL", ""),
        ).strip()
        # A role nobody set is "설정 필요" (lane CL), not xai and grok-4.6
        # (lane PL): without a provider and a model the writer stays off and
        # the report keeps the server's own sentences.
        if enabled and not (
            os.environ.get("VOINEY_LAB_REPORT_PROVIDER", "").strip() and model
        ):
            log.warning(
                "report writer off: VOINEY_LAB_REPORT_PROVIDER and "
                "VOINEY_LAB_REPORT_MODEL (or VOINEY_LAB_SUPPLEMENTAL_MODEL) "
                "must both be set (설정 필요)")
            enabled = False
        try:
            timeout = float(os.environ.get(
                "VOINEY_LAB_REPORT_WRITER_TIMEOUT_SECONDS", "25"
            ).strip())
        except ValueError:
            timeout = 25.0
        return cls(enabled=enabled, model=model, timeout_seconds=timeout)


_NARRATIVE_CACHE: dict[tuple[Any, ...], ReportNarrative] = {}
_AUTO = object()

_WRITER_INSTRUCTIONS = """너는 연구자의 실험 노트를 돕는다. 아래 JSON 은 서버가 실험 기록과 프로토콜 원문에서 뽑은 사실이다. 이 사실만으로 두 가지를 쓴다.

1. observation_notes: '기록' 의 관찰·측정을 단계(와 시료)별로 묶어 실험 노트 말투(~했다/~였다)의 한 문장씩. 항목마다 {"단계": 단계 번호, "시료": 시료 이름 또는 "", "문장": 문장}. 기록이 없는 단계는 쓰지 않는다.
2. result_summary: 결과 한 문단(1–3문장) — '기록', '실제로 한 조건', '원문과 다르게 한 점', '결과 (실험자가 말함)' 에 있는 것만으로.

지킬 것:
- 숫자와 단위는 사실 JSON 에 있는 그대로 쓴다. 바꾸거나 계산하거나 새 숫자를 만들지 않는다. 서버가 모든 숫자·단위를 기록·원문과 대조해 없는 문장은 버린다.
- 기록·원문에 없는 관찰·결과·원인·해석을 지어내지 않는다. 추측하지 않는다.
- 안전 지침·완료 기준·절차 지시("~하세요", "~해야 한다")를 쓰지 않는다.
- 반복 회차를 쓸 때는 사실 JSON 에 적힌 대로 "(말로 확인한 돌아가기 기준)" 을 붙인다. 없으면 회차 수를 쓰지 않는다.
- '확인되지 않은 값' 과 "(확인되지 않음)" 이 붙은 값은 어느 문장에도 쓰지 않는다.
- 기록 ID, 버전, 해시, 영어 상태값, 명령 이름은 쓰지 않는다.

JSON 객체 하나만 돌려준다. 키: observation_notes (목록), result_summary (문자열, 쓸 것이 없으면 "")."""


def _mark_in_words(text: str) -> str:
    """The server's "[미확인]" mark said in words, so the model knows which
    value it may not state (and copies no mark of the server's)."""

    return " ".join(text.replace(UNCONFIRMED_MARK, "(확인되지 않음)").split())


def _writer_facts(facts: ReportFacts) -> dict[str, Any]:
    """What the model reads: the notebook's facts only, no identifiers (decision ③)."""

    book = facts.notebook
    unconfirmed = set(facts.unconfirmed_texts)
    records: list[dict[str, Any]] = []
    for item in (book.observations if book is not None else ()):
        if item.text in unconfirmed:
            continue
        records.append({
            "단계": item.label, "과정": item.process, "시료": "" if item.sample == "—" else item.sample,
            "종류": item.kind, "기록": item.text, "시각": item.at,
        })
    for record in facts.records:
        if record.kind in {"편차", "이상"} and record.status != "취소됨" and record.text not in unconfirmed:
            records.append({"단계": record.step_label, "종류": record.kind, "기록": record.text, "시각": record.at})
    return {
        "프로토콜": facts.protocol_title,
        "목적": book.purpose if book is not None else "",
        "시료": [
            {"이름": sample.name, "조건": sample.condition, "대조군": sample.control}
            for sample in (book.samples if book is not None else ())
        ],
        "실제로 한 조건": [
            _mark_in_words(f"{process} — {words}")
            for process, words in (book.condition_rows if book is not None else ())],
        "기록": records,
        "원문과 다르게 한 점": [_mark_in_words(line) for line in (book.differences if book is not None else ())],
        "결과 (실험자가 말함)": book.result if book is not None else "",
        "다음에 바꿀 점 (실험자가 말함)": book.next_change if book is not None else "",
        "확인되지 않은 값": [
            f"{item['step_label']}단계 {item['kind']}: {item['text']}"
            for item in facts.review if not item["confirmed"]
        ],
    }


def _parse_reply(content: str) -> dict[str, Any]:
    text = content.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text)
    data = json.loads(text)
    if not isinstance(data, dict):
        raise ValueError("report reply is not an object")
    return data


class ReportWriterBrain:
    """Writes the report's prose sections; the server checks each before using it."""

    def __init__(
        self,
        client: Any = None,
        model: str = "",
        timeout_seconds: float = 25.0,
    ) -> None:
        self.client = client
        self.model = model
        self.timeout_seconds = timeout_seconds
        self.last_usage: dict[str, Any] = {}

    def build_deterministic_draft(
        self,
        report_data: dict[str, Any],
        events: list[dict[str, Any]],
    ) -> ReportDraftState:
        """Create a faithful, unhallucinated report draft directly from SQLite events."""
        event_ids = tuple(
            str(e.get("event_key") or e.get("event_id"))
            for e in events
            if (e.get("event_key") or e.get("event_id"))
        )
        obs_count = sum(1 for e in events if e.get("event_type") == "observation")
        anomaly_count = sum(1 for e in events if e.get("event_type") in ("anomaly", "system_anomaly"))
        completed_steps = [
            e.get("step_label") for e in events if e.get("event_type") == "step_completed"
        ]

        summary = (
            f"Protocol {report_data.get('protocol_title', 'Protocol')} "
            f"execution recorded {len(events)} total ledger events across "
            f"{len(completed_steps)} completed steps."
        )
        obs_narrative = (
            f"Operator recorded {obs_count} direct qualitative observations during the session."
            if obs_count > 0 else "해당 단계에 대해 별도의 관찰 결과가 기록되지 않았다."
        )
        anomaly_narrative = (
            f"A total of {anomaly_count} unexpected condition(s) were triaged and logged."
            if anomaly_count > 0 else "Execution proceeded without unresolved abnormal incidents."
        )
        conclusion = (
            f"Experiment concluded with status '{report_data.get('status', 'in_progress')}'. "
            f"All {len(event_ids)} event records are verified in the server SQLite ledger."
        )

        return ReportDraftState(
            report_id=str(report_data.get("report_id", "")),
            session_id=str(report_data.get("session_id", "")),
            protocol_id=str(report_data.get("protocol_id", "")),
            status=str(report_data.get("status", "in_progress")),
            experiment_summary=summary,
            observations_narrative=obs_narrative,
            anomalies_narrative=anomaly_narrative,
            conclusion_narrative=conclusion,
            committed_event_ids=event_ids,
            last_updated_at=_now(),
        )

    @staticmethod
    def facts_for(
        report_data_or_context: Mapping[str, Any] | GroundedReportContext,
        events: Sequence[Mapping[str, Any]] | None = None,
        *,
        fixture: Any = _AUTO,
        start_check: Sequence[Mapping[str, Any]] | None = None,
        previous: ReportFacts | None = None,
    ) -> ReportFacts:
        if isinstance(report_data_or_context, GroundedReportContext):
            context = report_data_or_context
            report_data = {**context.report_metadata, **context.protocol_metadata,
                           **context.session_timing}
            return build_report_facts(report_data, list(events or context.event_ledger), context=context,
                                      previous=previous)
        report_data = report_data_or_context
        problem = ""
        if fixture is _AUTO:
            fixture, problem = report_protocol_fixture(report_data)
        context = None
        if fixture is None:
            context = build_grounded_report_context(
                dict(report_data), list(events or report_data.get("events") or ()), lookup=False)
        return build_report_facts(report_data, events, fixture=fixture, context=context,
                                  protocol_problem=problem, start_check=start_check, previous=previous)

    def build_deterministic_narrative(
        self,
        report_data_or_context: dict[str, Any] | GroundedReportContext,
        events: list[dict[str, Any]] | None = None,
        context: GroundedReportContext | None = None,
        *,
        fixture: Any = _AUTO,
    ) -> ReportNarrative:
        """The report with the server's own sentence in every section."""

        facts = self.facts_for(context or report_data_or_context, events, fixture=fixture)
        return narrative_from_sections(facts)

    async def write_reply(self, facts: ReportFacts) -> tuple[dict[str, Any] | None, str]:
        """One model call for the report's prose: (reply, "") or (None, why it failed)."""

        if not self.client:
            return None, ""
        try:
            response = await asyncio.wait_for(
                self.client.chat.completions.create(
                    model=self.model,
                    messages=[
                        {"role": "system", "content": _WRITER_INSTRUCTIONS},
                        {"role": "user", "content": json.dumps(
                            _writer_facts(facts), ensure_ascii=False, indent=1)},
                    ],
                    response_format={"type": "json_object"},
                    # Room for one use per listed item (decision 4) beside the prose.
                    max_tokens=6000,
                ),
                timeout=self.timeout_seconds,
            )
            self.last_usage["writer"] = getattr(response, "usage", None)
            return _parse_reply(response.choices[0].message.content or ""), ""
        except Exception as exc:  # noqa: BLE001 -- the server's sentences stand in
            log.warning("report writer failed error=%s; using the server's sentences", type(exc).__name__)
            if isinstance(exc, (asyncio.TimeoutError, TimeoutError)):
                return None, "모델 시간 초과"
            return None, f"모델 답 실패({type(exc).__name__})"

    async def generate_narrative(
        self,
        report_data: dict[str, Any],
        events: list[dict[str, Any]],
        *,
        fixture: Any = _AUTO,
    ) -> ReportNarrative:
        """The report's prose: one model call, each section checked before use."""

        facts = self.facts_for(report_data, events, fixture=fixture)
        latest_key = str(events[-1].get("event_key") if events else "")
        cache_key = (
            str(report_data.get("report_id") or ""), int(report_data.get("finalization_version") or 0),
            latest_key, str(report_data.get("protocol_revision") or ""), self.model,
        )
        cached = _NARRATIVE_CACHE.get(cache_key)
        if cached is not None:
            return cached
        if not self.client:
            narrative = narrative_from_sections(facts)
        else:
            reply, failure = await self.write_reply(facts)
            narrative = narrative_from_reply(facts, reply, failure=failure, writer=self.model)
        _NARRATIVE_CACHE[cache_key] = narrative
        return narrative


def narrative_from_reply(
    facts: ReportFacts, reply: Mapping[str, Any] | None, *, failure: str = "", writer: str,
    written_at: str | None = None,
) -> ReportNarrative:
    """The report from one model reply, checked against the record as it is now."""

    if reply is None:
        return narrative_from_sections(
            facts, None, writer=f"{writer} (실패)", written_at=written_at, failure=failure or "모델 답 없음")
    return narrative_from_sections(facts, reply, writer=writer, written_at=written_at)


# --- Preparing the prose when the experiment ends (decision 9) --------------------

#: The prose row's writer when no model wrote it.
SERVER_WRITER = "서버"


class ReportProsePreparer:
    """Writes each finished report's prose once, in the background.

    When an experiment ends the server asks for its report's prose; the report
    role's model is called once on a background thread (or not at all when no
    model is configured) and the reply is kept with the report. Downloads use
    what was kept. A person asking for it again is the only other call.
    """

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._running: dict[tuple[str, str], threading.Event] = {}
        #: Reports whose values are being confirmed by voice (lane N,
        #: decision 3): their prose waits until the review ends.
        self._held: set[tuple[str, str]] = set()

    def hold(self, store: ExperimentReportStore, report_id: str) -> None:
        """Keep the prose back while the report's values are confirmed by voice."""

        with self._lock:
            self._held.add(self._key(store, report_id))

    def release(self, store: ExperimentReportStore, report_id: str) -> bool:
        """End a hold; True when there was one."""

        with self._lock:
            key = self._key(store, report_id)
            held = key in self._held
            self._held.discard(key)
            return held

    def held(self, store: ExperimentReportStore, report_id: str) -> bool:
        with self._lock:
            return self._key(store, report_id) in self._held

    @staticmethod
    def _key(store: ExperimentReportStore, report_id: str) -> tuple[str, str]:
        return (str(store.path), report_id)

    def start(
        self, store: ExperimentReportStore, report_id: str, make_brain: Any, *, again: bool = False,
    ) -> str:
        """"preparing", "ready" (already kept), or "not_finished"."""

        key = self._key(store, report_id)
        with self._lock:
            if key in self._running:
                return "preparing"
            if key in self._held:
                return "awaiting_review"
            report = store.get_report(report_id)
            if report["status"] == "in_progress":
                return "not_finished"
            if not again and store.get_prose(report_id) is not None:
                return "ready"
            brain = make_brain()
            if brain is None or not brain.client:
                store.save_prose(report_id, writer=SERVER_WRITER, reply=None,
                                 failure="보고서 모델이 설정되지 않음",
                                 event_count=len([e for e in report["events"]
                                                  if e.get("event_type") not in REVIEW_EVENT_TYPES]))
                return "ready"
            done = threading.Event()
            self._running[key] = done
        threading.Thread(
            target=self._prepare, args=(store, report_id, brain, key, done),
            name=f"report-prose-{report_id}", daemon=True,
        ).start()
        return "preparing"

    def _prepare(
        self, store: ExperimentReportStore, report_id: str, brain: "ReportWriterBrain",
        key: tuple[str, str], done: threading.Event,
    ) -> None:
        count = 0
        try:
            report = store.get_report(report_id)
            events = list(report["events"])
            count = len([e for e in events if e.get("event_type") not in REVIEW_EVENT_TYPES])
            reply, failure = asyncio.run(brain.write_reply(brain.facts_for(report, events)))
            store.save_prose(report_id, writer=brain.model, reply=reply, failure=failure,
                             event_count=count)
        except Exception as exc:  # noqa: BLE001 -- the server's sentences stand in
            log.warning("report prose not prepared error=%s", type(exc).__name__)
            try:
                store.save_prose(report_id, writer=brain.model, reply=None,
                                 failure=f"준비 실패({type(exc).__name__})", event_count=count)
            except Exception:  # noqa: BLE001
                log.warning("report prose failure not kept")
        finally:
            with self._lock:
                self._running.pop(key, None)
            done.set()

    def wait(self, store: ExperimentReportStore, report_id: str, timeout: float) -> bool:
        with self._lock:
            done = self._running.get(self._key(store, report_id))
        return True if done is None else done.wait(timeout)

    def status(self, store: ExperimentReportStore, report_id: str) -> dict[str, Any]:
        """What the screen shows next to the record (decision 9)."""

        with self._lock:
            preparing = self._key(store, report_id) in self._running
            held = self._key(store, report_id) in self._held
        report = store.get_report(report_id)
        prose = None if preparing else store.get_prose(report_id)
        status: dict[str, Any] = {"report_id": report_id, "ai_written": False, "new_records": 0}
        recorded = [e for e in report["events"] if e.get("event_type") not in REVIEW_EVENT_TYPES]
        if preparing:
            status.update(state="preparing", label="보고서 준비 중 — 실험 기록으로 보고서 문장을 쓰고 있습니다.")
        elif held and prose is None:
            status.update(state="awaiting_review",
                          label="보고서에 넣을 값을 확인하는 중입니다. 확인이 끝나면 보고서 문장을 준비합니다.")
        elif prose is not None:
            ai = prose["reply"] is not None
            status.update(state="ready", ai_written=ai, prepared_at=prose["prepared_at"],
                          new_records=max(0, len(recorded) - int(prose["event_count"])))
            if ai:
                label = f"보고서 준비됨 — AI({prose['writer']}) 문장 중 서버 검사를 통과한 것만 씁니다."
            elif prose["writer"] == SERVER_WRITER:
                label = "보고서 준비됨 — 서버 문장 (보고서 모델이 설정되지 않음)."
            else:
                label = f"보고서 준비됨 — 서버 문장 (AI 문장을 쓰지 못함: {prose['failure']})."
            if status["new_records"]:
                label += f" 준비한 뒤 기록이 {status['new_records']}건 늘었습니다. 표에는 들어가며, 문장은 다시 만들 수 있습니다."
            status["label"] = label
        elif report["status"] == "in_progress":
            status.update(state="not_finished", label="실험이 끝나면 보고서 문장을 준비합니다.")
        else:
            status.update(state="not_prepared", label="보고서 문장이 아직 준비되지 않았습니다.")
        return status


REPORT_PROSE = ReportProsePreparer()


# --- One document, two renderings (Word and Markdown) -----------------------------


def report_blocks(narrative: ReportNarrative) -> list[tuple[str, Any]]:
    """The grounded handoff notebook as renderer-neutral blocks."""

    facts = narrative.facts
    assert facts is not None and facts.notebook is not None
    book = facts.notebook
    glance = facts.glance
    blocks: list[tuple[str, Any]] = [("title", book.title), ("subtitle", book.subtitle)]
    if glance is not None and glance.unconfirmed:
        # Lane RP2, decision 3 (kept): the first line says how many values
        # are not confirmed and which; the notebook is still made.
        blocks += [("alert", (f"확인 안 된 값 {len(glance.unconfirmed)}개", UNCONFIRMED_WORDS)),
                   ("list", glance.unconfirmed)]
    blocks.append(("h1", NOTEBOOK_SECTIONS[0]))
    blocks.append(_table(("항목", "내용"), book.header_rows))
    if book.samples:
        blocks.append(_table(SAMPLE_HEADER, tuple(
            (sample.name, sample.condition, "예" if sample.control else "") for sample in book.samples)))

    blocks.append(("h1", NOTEBOOK_SECTIONS[1]))
    blocks.append(("p", book.purpose_words))

    blocks.append(("h1", NOTEBOOK_SECTIONS[2]))
    if book.method_rows:
        blocks.append(_table(METHOD_HEADER, tuple((
            f"{row.label}단계" + (f" · PDF p.{row.source_page}" if row.source_page else ""),
            row.source_text,
            row.actual_values,
            " · ".join(part for part in (
                f"{row.started_at}–{row.ended_at}" if row.started_at or row.ended_at else "",
                row.timer,
            ) if part),
        ) for row in book.method_rows)))
    if book.condition_note:
        blocks.append(("note", book.condition_note))
    if book.condition_rows:
        blocks += [("h2", "원문 값 → 실제 요약"), _table(CONDITION_HEADER, book.condition_rows)]
    if book.value_bars:
        blocks += [("h2", "원문 → 실제 값 비교"), _table(VALUE_BAR_HEADER, tuple(
            (item.label, item.source, item.actual) for item in book.value_bars
        ))]
        blocks.extend(("bar", (item.label, *_notebook_bar_lines(item))) for item in book.value_bars)
    if book.timeline:
        timeline_words = {
            "start": "실험 시작", "step_end": "단계 끝", "timer_start": "타이머 시작",
            "timer_pause": "타이머 일시정지", "timer_resume": "타이머 다시 시작",
            "timer_cancel": "타이머 취소", "timer_restart": "타이머 처음부터",
            "timer_end": "타이머 끝", "pause": "실험 일시정지", "resume": "실험 다시 시작",
            "observation": "관찰·특이사항", "photo": "사진",
        }
        blocks += [("h2", "단계별 타임라인"), _table(TIMELINE_HEADER, tuple(
            (item.at, item.step_label or "—", timeline_words.get(item.kind, item.kind), item.text)
            for item in book.timeline
        ))]
    if book.images:
        blocks.append(("h2", "단계 사진·원문 그림"))
        blocks.extend(("image", image) for image in sorted(
            book.images, key=lambda item: (_label_order(item.step_label), item.kind, item.asset_id)
        ))
    if book.differences:
        blocks += [("h2", "원문과 다르게 한 점·문제"), ("list", book.differences)]

    blocks.append(("h1", NOTEBOOK_SECTIONS[3]))
    if book.observations:
        blocks.append(_table(OBSERVATION_HEADER, tuple(
            (item.process, item.sample, item.shown, item.at) for item in book.observations)))
    else:
        blocks.append(("p", ""))

    blocks.append(("h1", NOTEBOOK_SECTIONS[4]))
    blocks.append(("p", book.result))

    blocks.append(("h1", NOTEBOOK_SECTIONS[5]))
    blocks.append(_table(("항목", "내용"), (
        ("해석", book.interpretation_words),
        ("다음 할 일", book.next_change),
    )))

    blocks.append(("h1", NOTEBOOK_SECTIONS[6]))
    blocks.append(("h2", book.domain.label))
    if book.domain.fields:
        blocks.append(_table(DOMAIN_HEADER, tuple(
            (item.label, item.value + (f"\n계산식: {item.formula}" if item.formula else ""),
             {"source": "원문", "record": "기록", "calculated": "계산", "blank": ""}.get(item.origin, ""))
            for item in book.domain.fields
        )))
    if book.plate_positions:
        blocks += [("h2", "플레이트 위치"), ("list", tuple(
            f"{item.label} · {item.step_label}단계 · {item.words}" for item in book.plate_positions
        ))]
    if book.domain.code == "battery_materials" and book.sample_rows:
        blocks += [("h2", "셀·시료 번호별 기록"), _table(SAMPLE_RECORD_HEADER, tuple(
            (item.identifier, " · ".join(item.records)) for item in book.sample_rows
        ))]

    blocks.append(("h1", NOTEBOOK_SECTIONS[7]))
    blocks.append(("appendix", None))
    if book.materials_words:
        blocks.append(("note", book.materials_words))
    if book.inventory:
        stage_words = {"before": "실험 전", "after": "실험 후", "source": "원문", "blank": ""}
        blocks += [("h2", "Methods 조건표"), _table(INVENTORY_HEADER, tuple(
            (item.name, item.kind, item.vendor, item.catalog_number, item.lot_number,
             stage_words.get(item.stage, "")) for item in book.inventory
        ))]
    blocks.append(("h2", "같은 프로토콜의 앞선 내 기록"))
    if book.previous_words:
        blocks.append(("note", book.previous_words))
        blocks.append(("list", book.previous_lines) if book.previous_lines else ("note", SAME_AS_LAST_WORDS))
    else:
        blocks.append(("note", FIRST_RUN_WORDS))
    blocks.append(("h2", "확인 기록"))
    confirm_rows = (("실험자", "", "", ""),) + (book.pi_rows or (("PI", "", "", ""),))
    blocks.append(("sign", (CONFIRM_HEADER, confirm_rows, TABLE_WIDTHS[CONFIRM_HEADER])))
    if not book.pi_rows:
        blocks.append(("note", NO_PI_WORDS))
    blocks.append(("h2", "단계 시각표"))
    blocks.append(_table(APPENDIX_STEP_HEADER, book.appendix_step_rows) if book.appendix_step_rows
                  else ("note", "기록된 단계가 없다."))
    blocks.append(("h2", "말한 그대로의 기록"))
    blocks.append(_table(APPENDIX_RECORD_HEADER, book.appendix_record_rows) if book.appendix_record_rows
                  else ("note", "말로 남긴 기록이 없다."))
    if facts.review_version_words:
        # Lane RV, decision g (kept, in the appendix): what the review version
        # the run followed holds apart from the source, item by item.
        blocks.append(("h2", facts.review_version_words))
        if facts.review_changes:
            blocks.append(_table(REVIEW_CHANGE_HEADER, tuple(
                (change.step_label or "—", change.kind, change.original, change.reviewed,
                 change.origin, change.reviewer)
                for change in facts.review_changes)))
        else:
            blocks.append(("note", "검토자가 원문에서 바꾼 값이 없다."))
    blocks += [("note", book.system_words), ("note", f"출처: {book.source_words}"),
               ("note", authorship_line(narrative))]
    # A lab template may rename/reorder only the eight complete sections.
    # Content and grounding rules stay fixed in these chunks.
    prefix: list[tuple[str, Any]] = []
    chunks: dict[str, list[tuple[str, Any]]] = {}
    current: str | None = None
    id_by_name = dict(zip(NOTEBOOK_SECTIONS, NOTEBOOK_SECTION_IDS))
    for block in blocks:
        if block[0] == "h1" and str(block[1]) in id_by_name:
            current = id_by_name[str(block[1])]
            chunks[current] = [block]
        elif current is None:
            prefix.append(block)
        else:
            chunks[current].append(block)
    ordered = list(prefix)
    for section in book.section_template:
        chunk = list(chunks.get(section.id, ()))
        if not chunk:
            continue
        chunk[0] = ("h1", section.name)
        ordered.extend(chunk)
    return ordered


def _label_order(label: str) -> tuple[int, str]:
    match = re.match(r"\d+", str(label))
    return (int(match.group()) if match else 10**9, str(label))


def _notebook_bar_lines(row: NotebookValueBar) -> tuple[str, str]:
    """Two deterministic text bars, scaled only from the displayed values."""

    source = _condition_measures(row.source)
    actual = _condition_measures(row.actual)
    if not source or not actual or source[0].family != actual[0].family:
        return (f"원문 {row.source}", f"실제 {row.actual}")
    left, right = source[0], actual[0]
    try:
        left_value = float(left.number) * _CONDITION_BASE.get(left.unit, 1.0)
        right_value = float(right.number) * _CONDITION_BASE.get(right.unit, 1.0)
    except ValueError:
        return (f"원문 {row.source}", f"실제 {row.actual}")
    maximum = max(left_value, right_value)
    if maximum <= 0:
        return (f"원문 {row.source}", f"실제 {row.actual}")
    width = 18
    left_bar = "█" * (0 if left_value <= 0 else max(1, round(width * left_value / maximum)))
    right_bar = "█" * (0 if right_value <= 0 else max(1, round(width * right_value / maximum)))
    return (f"원문 {row.source} {left_bar}", f"실제 {row.actual} {right_bar}")


def authorship_line(narrative: ReportNarrative) -> str:
    """Who wrote the notebook's sentences, and when (decision ①: the server)."""

    facts = narrative.facts
    written = _local(narrative.written_at, facts.zone if facts is not None else report_timezone())
    when = f"{written.year}년 {written.month}월 {written.day}일 {written:%H:%M}" if written else "시각 기록 없음"
    base = "이 실험 노트의 문장은 서버가 실험 기록과 프로토콜 원문에서 만들었다."
    # Observations stay verbatim and results stay person-authored, while
    # interpretation stays blank.  A separately prepared model suggestion is
    # therefore never part of the notebook body, even when its grounding
    # check passed.
    line = f"{base} 모델이 쓴 문장은 포함하지 않았다."
    return f"{line} 작성 {when}."


#: Each report table's column widths as shares of the text width (decision 4):
#: short values (step, kind, time, timer, item) narrow, content wide.
#: The materials and equipment tables (decisions 2-4), with and without uses.
ITEM_HEADER = ("이름 (원문 그대로)", "사용 단계")
#: Said under the tables when a name was shown without what the protocols.io
#: list writes after it (lane FX, decision 4).
ITEM_LABELS_NOTE = "이름 뒤에 붙은 protocols.io 목록 꼬리표(NAME·TYPE·BRAND·SKU)와 끝의 긴 숫자 코드는 표에서 뺐다."
TABLE_WIDTHS: dict[tuple[str, ...], tuple[float, ...]] = {
    ("항목", "내용"): (0.22, 0.78),
    ITEM_HEADER: (0.80, 0.20),
    ("단계", "원문 단계", "완료 시각", "타이머 (원문 / 실제)"): (0.07, 0.59, 0.12, 0.22),
    ("단계", "종류", "기록 내용 (연구자가 말한 그대로)", "시각"): (0.07, 0.09, 0.73, 0.11),
    REVIEW_CHANGE_HEADER: (0.07, 0.13, 0.28, 0.28, 0.13, 0.11),
    # Lane RP2, decision 1: the first pages' tables, on A4's 16.6 cm text width.
    GLANCE_SUMMARY_HEADER: (0.17, 0.17, 0.17, 0.17, 0.16, 0.16),
    GLANCE_STEP_HEADER: (0.07, 0.08, 0.11, 0.10, 0.28, 0.21, 0.15),
    GLANCE_TIMER_HEADER: (0.07, 0.11, 0.11, 0.14, 0.14, 0.11, 0.32),
    GLANCE_ISSUE_HEADER: (0.06, 0.07, 0.13, 0.18, 0.38, 0.08, 0.10),
    GLANCE_PRESTART_HEADER: (0.12, 0.08, 0.17, 0.17, 0.34, 0.12),
    GLANCE_SIGN_HEADER: (0.30, 0.70),
    # Lane RP3, decision ①: the notebook's tables.
    SAMPLE_HEADER: (0.30, 0.55, 0.15),
    CONDITION_HEADER: (0.32, 0.68),
    OBSERVATION_HEADER: (0.20, 0.12, 0.58, 0.10),
    CONFIRM_HEADER: (0.14, 0.20, 0.26, 0.40),
    APPENDIX_STEP_HEADER: (0.10, 0.14, 0.14, 0.22, 0.40),
    APPENDIX_RECORD_HEADER: (0.10, 0.08, 0.14, 0.68),
    METHOD_HEADER: (0.11, 0.42, 0.27, 0.20),
    TIMELINE_HEADER: (0.12, 0.10, 0.18, 0.60),
    VALUE_BAR_HEADER: (0.40, 0.30, 0.30),
    DOMAIN_HEADER: (0.24, 0.62, 0.14),
    SAMPLE_RECORD_HEADER: (0.22, 0.78),
    INVENTORY_HEADER: (0.30, 0.09, 0.16, 0.18, 0.15, 0.12),
}
#: The report's one font, for Korean and Latin letters alike (lane RP2,
#: decision 1): Word prints it as Malgun Gothic, and the server's LibreOffice
#: prints it as NanumGothic (the Korean names point there, lane AN), so a
#: line never mixes two fonts.
REPORT_FONT = "맑은 고딕"


def _table(header: tuple[str, ...], rows: tuple[tuple[Any, ...], ...]) -> tuple[str, Any]:
    return ("table", (header, rows, TABLE_WIDTHS[header]))


def render_markdown(narrative: ReportNarrative) -> str:
    def cell(value: Any) -> str:
        if isinstance(value, StackedCell):
            return (f"{cell(value.main)}<br><small>원문: {cell(value.below)}</small>")
        # Lane RP2: a cell holding one value per line keeps its lines.
        return "<br>".join(
            " ".join(line.split()).replace("|", "\\|") for line in str(value).split("\n")
        ).strip() or " "

    lines: list[str] = []
    for kind, content in report_blocks(narrative):
        if kind == "title":
            lines += [f"# {content}", ""]
        elif kind == "subtitle":
            lines += [f"> {content}", ""]
        elif kind == "appendix":
            continue  # Lane RP3: Word prints what follows in smaller type.
        elif kind == "h1":
            lines += [f"## {content}", ""]
        elif kind == "h2":
            lines += [f"### {content}", ""]
        elif kind == "h3":
            lines += [f"#### {content}", ""]
        elif kind == "p":
            lines += [str(content).strip(), ""]
        elif kind == "note":
            lines += [f"> {content}", ""]
        elif kind == "alert":
            lines += [f"> **{content[0]}** — {content[1]}", ""]
        elif kind == "list":
            lines += [f"- {item}" for item in content] + [""]
        elif kind == "bar":
            label, source_line, actual_line = content
            lines += [f"**{label}**", f"- {source_line}", f"- {actual_line}", ""]
        elif kind == "image":
            caption = f"{content.step_label}단계 · {content.caption}" if content.step_label else content.caption
            safe_caption = caption.replace("[", "(").replace("]", ")")
            if content.url:
                lines += [f"![{safe_caption}]({content.url})", ""]
            else:
                lines += [f"> {safe_caption}", ""]
        elif kind in {"table", "sign"}:
            header, rows, widths = content
            lines.append("| " + " | ".join(cell(h) for h in header) + " |")
            # Markdown has no column widths; the dash counts carry the
            # shares, which Pandoc reads as relative widths.
            lines.append("|" + "|".join("-" * max(3, round(share * 40)) for share in widths) + "|")
            lines += ["| " + " | ".join(cell(v) for v in row) + " |" for row in rows]
            lines.append("")
    return "\n".join(lines).rstrip() + "\n"


#: Word paragraph spacing (decision 1, 2026-10-06): points before, points
#: after, and line spacing as a multiple of one line, by block kind. The
#: template's default (10 pt after every paragraph, 1.15 lines) applied to
#: body lines, list items and table text alike and left long gaps; headings
#: keep room above them. A list's last item takes the body's space after.
DOCX_SPACING: dict[str, tuple[float, float, float]] = {
    "title": (0, 8, 1.0),
    "h1": (12, 4, 1.0),
    "h2": (8, 3, 1.0),
    "h3": (6, 2, 1.0),
    "p": (0, 4, 1.15),
    "note": (2, 4, 1.15),
    "alert": (2, 2, 1.15),
    "list": (0, 1, 1.15),
    "bar": (0, 2, 1.0),
    "image": (0, 4, 1.0),
    "table": (0, 0, 1.0),
}


def _space(paragraph: Any, kind: str, *, after: float | None = None) -> None:
    from docx.shared import Pt

    before, default_after, line = DOCX_SPACING[kind]
    form = paragraph.paragraph_format
    form.space_before = Pt(before)
    form.space_after = Pt(default_after if after is None else after)
    form.line_spacing = line


def render_docx(
    narrative: ReportNarrative,
    image_resolver: Callable[[NotebookImage], tuple[bytes, str] | None] | None = None,
) -> bytes:
    return _render_docx_blocks(report_blocks(narrative), image_resolver=image_resolver)


def _render_docx_blocks(
    blocks: Sequence[tuple[str, Any]],
    *, image_resolver: Callable[[NotebookImage], tuple[bytes, str] | None] | None = None,
) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.2))
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    # Lane RP2, decision 1: one Korean font for every script, not the theme's.
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for script in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(script), REPORT_FONT)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(theme), None)
    # Paragraphs a person adds in Word take the body's spacing, not the template's.
    body_before, body_after, body_line = DOCX_SPACING["p"]
    normal.paragraph_format.space_before = Pt(body_before)
    normal.paragraph_format.space_after = Pt(body_after)
    normal.paragraph_format.line_spacing = body_line
    text_width = section.page_width - section.left_margin - section.right_margin
    # Lane RP3, decision ① (9): the appendix is printed in smaller type.
    small = False

    for kind, content in blocks:
        if kind == "title":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(content))
            run.bold = True
            run.font.size = Pt(16)
            _space(paragraph, "title")
        elif kind == "subtitle":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(content))
            run.italic = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(90, 90, 90)
            _space(paragraph, "note")
        elif kind == "appendix":
            small = True
        elif kind in {"h1", "h2", "h3"}:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(str(content))
            run.bold = True
            run.font.size = Pt({"h1": 13, "h2": 10 if small else 11, "h3": 10.5}[kind])
            run.font.color.rgb = RGBColor(20, 50, 35)
            _space(paragraph, kind)
        elif kind == "p":
            for line in str(content).split("\n"):
                paragraph = document.add_paragraph(line)
                if small:
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
                _space(paragraph, "p")
        elif kind == "note":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(str(content))
            run.italic = True
            run.font.color.rgb = RGBColor(90, 90, 90)
            if small:
                run.font.size = Pt(8.5)
            _space(paragraph, "note")
        elif kind == "alert":
            # Lane RP2, decision 3: "확인 안 된 값 n개" in bold red, first under the title.
            paragraph = document.add_paragraph()
            run = paragraph.add_run(content[0])
            run.bold = True
            run.font.color.rgb = RGBColor(176, 32, 32)
            paragraph.add_run(f" — {content[1]}")
            _space(paragraph, "alert")
        elif kind == "list":
            for index, item in enumerate(content):
                last = index == len(content) - 1
                paragraph = document.add_paragraph(f"• {item}")
                if small:
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
                _space(paragraph, "list", after=DOCX_SPACING["p"][1] if last else None)
        elif kind == "bar":
            label, source_line, actual_line = content
            heading = document.add_paragraph()
            heading.add_run(str(label)).bold = True
            _space(heading, "bar")
            for line, color in ((source_line, RGBColor(90, 110, 100)),
                                (actual_line, RGBColor(32, 105, 72))):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(str(line))
                run.font.color.rgb = color
                run.font.name = REPORT_FONT
                _space(paragraph, "bar")
        elif kind == "image":
            caption = f"{content.step_label}단계 · {content.caption}" if content.step_label else content.caption
            added = False
            if image_resolver is not None:
                try:
                    resolved = image_resolver(content)
                    if resolved is not None and resolved[0] and len(resolved[0]) <= 32 * 1024 * 1024:
                        paragraph = document.add_paragraph()
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        paragraph.add_run().add_picture(io.BytesIO(resolved[0]), width=Cm(12.5))
                        _space(paragraph, "image")
                        added = True
                except Exception as exc:  # noqa: BLE001 - caption still survives
                    log.warning("notebook image not embedded error=%s", type(exc).__name__)
            paragraph = document.add_paragraph(caption)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.italic = True
                run.font.size = Pt(8.5 if small else 9)
                run.font.color.rgb = RGBColor(90, 90, 90)
            if not added and content.url:
                paragraph.add_run(f" · {content.url}")
            _space(paragraph, "image")
        elif kind == "pagebreak":
            # Lane RP2, decision 1: the sections start on the page after the first pages.
            paragraph = document.add_paragraph()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
            _space(paragraph, "table")
        elif kind in {"table", "sign"}:
            header, rows, widths = content
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for index, text in enumerate(header):
                table.rows[0].cells[index].text = str(text)
                for run in table.rows[0].cells[index].paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(8 if small else 9)
            # A table that runs onto the next page repeats its heading row there.
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                cells = table.add_row().cells
                for index, value in enumerate(row):
                    if isinstance(value, StackedCell):
                        # Lane N, decision 6: the source under its Korean,
                        # in the same cell, in small grey type.
                        cells[index].text = value.main
                        below = cells[index].add_paragraph().add_run(f"원문: {value.below}")
                        below.font.size = Pt(7.5)
                        below.font.color.rgb = RGBColor(110, 110, 110)
                        for run in cells[index].paragraphs[0].runs:
                            run.font.size = Pt(9)
                        continue
                    # One value per line, each its own paragraph (lane RP2).
                    first, *more = str(value).split("\n")
                    cells[index].text = first
                    for line in more:
                        cells[index].add_paragraph(line)
                    for paragraph in cells[index].paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(8 if small else 9)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        _space(paragraph, "table")
            if kind == "sign":
                # Room to write a name, a date and a signature by hand.
                for row in table.rows[1:]:
                    row.height = Cm(1.1)
            _fix_column_widths(table, [int(text_width * share) for share in widths])
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


MEETING_HEADINGS = ("목적", "시료", "핵심 결과", "지난번과 다른 점", "다음 할 일")
XLSX_HEADER = ("시료", "과정", "조건 (원문 → 실제)", "관찰", "측정값", "단위", "시각", "출처")


def meeting_blocks(narrative: ReportNarrative) -> list[tuple[str, Any]]:
    """The one-page lab-meeting summary (decision ③): purpose, samples, key
    results, what differed from the run before, what to do next."""

    facts = narrative.facts
    assert facts is not None and facts.notebook is not None
    book = facts.notebook
    run = dict(facts.run_rows)
    line = " · ".join(part for part in ("랩미팅 요약", run.get("날짜", ""), run.get("실험자", ""), facts.protocol_title) if part)
    blocks: list[tuple[str, Any]] = [("title", book.title), ("subtitle", line)]
    blocks += [("h2", "목적"), ("p", book.purpose or NOT_SAID_WORDS)]
    blocks.append(("h2", "시료"))
    if book.samples:
        blocks.append(_table(SAMPLE_HEADER, tuple(
            (sample.name, sample.condition, "예" if sample.control else "") for sample in book.samples)))
    else:
        blocks.append(("p", NOT_SAID_WORDS))
    blocks.append(("h2", "핵심 결과"))
    if narrative.ai_result:
        blocks.append(("p", f"결과 ({AI_NOTE_LABEL}): {narrative.ai_result}"))
    if book.result:
        blocks.append(("p", f"실험자: “{book.result}”"))
    if narrative.ai_notes:
        blocks.append(("list", tuple(
            f"{note.step_label}단계" + (f" · {note.sample}" if note.sample else "") + f": {note.text} ({AI_NOTE_LABEL})"
            for note in narrative.ai_notes)))
    elif book.observations:
        blocks.append(("list", tuple(
            f"{item.label}단계" + (f" · {item.sample}" if item.sample != "—" else "") + f": {item.text}"
            for item in book.observations)))
    if not (narrative.ai_result or book.result or narrative.ai_notes or book.observations):
        blocks.append(("p", NO_OBSERVATION_WORDS))
    if book.differences:
        blocks.append(("list", tuple(f"원문과 다르게: {item}" for item in book.differences)))
    blocks.append(("h2", "지난번과 다른 점"))
    if book.previous_words:
        blocks.append(("p", book.previous_words))
        blocks.append(("list", book.previous_lines) if book.previous_lines else ("p", SAME_AS_LAST_WORDS))
    else:
        blocks.append(("p", FIRST_RUN_WORDS))
    blocks.append(("h2", "다음 할 일"))
    blocks.append(("list", book.next_lines) if book.next_lines else ("p", NOT_SAID_WORDS))
    return blocks


def render_meeting_docx(narrative: ReportNarrative) -> bytes:
    """The lab-meeting summary as one Word page (decision ③)."""

    return _render_docx_blocks(meeting_blocks(narrative))


def _record_measure(text: str) -> tuple[Any, str | None]:
    """The first number with a unit in a record, as a number and its unit, else (None, None)."""

    for number, unit, _ in _measures(text):
        try:
            value: Any = float(number)
            if value.is_integer():
                value = int(value)
        except ValueError:
            value = number
        return value, unit
    return None, None


def render_xlsx(narrative: ReportNarrative) -> bytes:
    """The results table (decision ③): one row per record, each with its
    sample, its step, the condition used, the value and unit, and the time;
    the AI-written notes after them, marked; the header on a second sheet."""

    from openpyxl import Workbook

    facts = narrative.facts
    assert facts is not None and facts.notebook is not None
    book = facts.notebook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "결과"
    sheet.append(list(XLSX_HEADER))
    process_of = {item.label: item.process for item in book.observations}
    for record in facts.records:
        if record.status == "취소됨":
            continue
        payload_sample = next((item.sample for item in book.observations if item.text == record.text and item.label == record.step_label), None)
        sample = payload_sample if payload_sample and payload_sample != "—" else (sample_in_text(record.text) or "—")
        process = process_of.get(record.step_label) or (
            f"{record.step_label}단계 · {next((s.section for s in facts.steps if s.label == record.step_label), '')}".rstrip(" ·"))
        value, unit = _record_measure(record.text)
        sheet.append([sample, process, book.condition_by_label.get(record.step_label), record.text,
                      value, unit, record.at or None, "기록"])
    for note in narrative.ai_notes:
        sheet.append([note.sample or "—", process_of.get(note.step_label) or f"{note.step_label}단계", None,
                      note.text, None, None, None, AI_NOTE_LABEL])
    if narrative.ai_result:
        sheet.append(["—", "결과", None, narrative.ai_result, None, None, None, AI_NOTE_LABEL])
    for column, width in zip("ABCDEFGH", (14, 24, 30, 48, 10, 8, 8, 8)):
        sheet.column_dimensions[column].width = width
    head = workbook.create_sheet("머리")
    head.append(["실험", book.title])
    for name, value in book.header_rows:
        head.append([name, value])
    head.column_dimensions["A"].width = 14
    head.column_dimensions["B"].width = 80
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class ReportPdfError(RuntimeError):
    """The report could not be printed to PDF (lane RP2, decision 2).

    ``report_pdf_unavailable``: this server has no LibreOffice;
    ``report_pdf_failed``: LibreOffice did not print it.
    """

    def __init__(self, message: str, *, code: str) -> None:
        super().__init__(message)
        self.code = code


def report_docx_to_pdf(document: bytes, *, converter: str | None = None, run: Any = None) -> bytes:
    """The Word report as the PDF the server's LibreOffice prints (lane RP2, decision 2).

    The same report, not another rendering of it: the .docx is printed the
    way an uploaded Word protocol is (``docx_to_pdf``: its own LibreOffice
    profile, and the Korean font names pointed at Nanum).
    """

    import tempfile

    from voiney_lab.experiment_protocol_pdf import (
        ProtocolPdfError,
        ProtocolSourceConversionUnavailableError,
        docx_to_pdf,
    )

    with tempfile.TemporaryDirectory(prefix="experiment-report-") as root:
        source, target = Path(root) / "report.docx", Path(root) / "report.pdf"
        source.write_bytes(document)
        try:
            docx_to_pdf(source, target, converter=converter, run=run)
        except ProtocolSourceConversionUnavailableError as exc:
            raise ReportPdfError(
                "LibreOffice is not installed on this server.", code="report_pdf_unavailable") from exc
        except ProtocolPdfError as exc:
            raise ReportPdfError("LibreOffice did not print the report.", code="report_pdf_failed") from exc
        return target.read_bytes()


def _format_elapsed_clock(value: Any) -> str:
    if value in (None, ""):
        return ""
    try:
        total = int(round(float(value)))
    except (TypeError, ValueError):
        return ""
    if total < 0:
        total = 0
    hours, remainder = divmod(total, 3600)
    minutes, seconds = divmod(remainder, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    return f"{minutes:02d}:{seconds:02d}"


def _human_event_clock(created_at: str | None) -> str:
    if not created_at:
        return ""
    try:
        parsed = datetime.fromisoformat(str(created_at).replace("Z", "+00:00"))
    except ValueError:
        return ""
    return parsed.strftime("%H:%M:%S")


def _fix_column_widths(table: Any, widths: Sequence[int]) -> None:
    """Fixed column widths in Word: the grid, every cell, and a fixed layout."""

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table.autofit = False
    properties = table._tbl.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def new_session_id() -> str:
    return "session-" + secrets.token_hex(16)
