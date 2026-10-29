"""Explicit, source-preserving multi-Protocol catalog service.

The service is deliberately constructed only by an authorized API/session
boundary.  Importing or listing it does not contact an analysis Provider.
All writes use the separate immutable Protocol store, never ProcedureStore.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import queue
import re
import secrets
import threading
import time
import unicodedata
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from collections import OrderedDict
from dataclasses import dataclass, fields, is_dataclass, replace
from datetime import datetime, timezone
from difflib import SequenceMatcher
from enum import Enum
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Callable, Mapping, Sequence
from zoneinfo import ZoneInfo

from voiney_lab import experiment_protocol as domain
from voiney_lab import pdf_text_engine
from voiney_lab.curated_protocol import CuratedProtocolFixture
from voiney_lab.experiment_protocol_analysis import (
    AUTO_RESOLVED_OCR_TWO_ENGINES,
    AUTO_RESOLVED_REQUOTED,
    BLOCKED_FAILURE_CODE,
    MAX_PIECED_ANALYSIS_PAGES,
    MAX_SINGLE_PASS_INPUT_BYTES,
    OCR_VALUE_ENGINES_DIFFER,
    OCR_VALUE_NOT_PLACED,
    OCR_VALUE_ONE_ENGINE,
    OCR_VALUE_ONE_ENGINE_READ,
    OCR_VALUE_PAGE_DIFFERS_WIDELY,
    OCR_VALUE_PAGE_LOW_CONFIDENCE,
    OCR_VALUE_READ_DIFFERENTLY,
    AnalysisPiece,
    OpenAICompatibleProtocolAnalysisModel,
    PieceConcurrencyGate,
    ProtocolAnalysisDraft,
    ProtocolAnalysisEvidenceError,
    ProtocolAnalysisPageRangeError,
    ProtocolAnalysisIncompleteError,
    ProtocolAnalysisInputTooLargeError,
    ProtocolAnalysisModel,
    ProtocolAnalysisPieceConflictError,
    ProtocolAnalysisPieceFailedError,
    ProtocolAnalysisTooManyPagesError,
    REQUOTE_ANSWERED_BY_PERSON,
    REQUOTE_OCR_VALUES,
    REQUOTE_TIME_LIMIT_SECONDS,
    REVIEW_ITEM_KINDS,
    REVIEW_REASON_ADDRESSED_TO_PROGRAM,
    REVIEW_REASON_BLOCKED,
    REVIEW_REASON_NUMBERING_SEAM,
    REVIEW_REASON_OCR_VALUE,
    REVIEW_REASON_PIECE_FAILED,
    REVIEW_REASON_POSSIBLY_MISSED,
    REVIEW_REASON_READ_TWO_WAYS,
    REVIEW_REASON_REFERENCE_MOVED,
    REVIEW_REASON_VALUE_CHANGED,
    REVIEW_REASONS_KO,
    ReviewFlag,
    analysis_piece_concurrency,
    analysis_pieces,
    analysis_requote_enabled,
    analyze_protocol_extraction,
    analyze_protocol_piece,
    blocked_pages_ko,
    check_analysis_page_count,
    check_pieced_page_count,
    checked_page_range,
    join_piece_analyses,
    possibly_missed_flags,
    possibly_missed_places,
    possibly_missed_summary,
    prepare_protocol_analysis_request,
    program_address_flags,
    replace_at_paths,
    requote_questions,
    requote_unverified_items,
    resettle_requote,
    review_item_id,
    safety_block_provider,
    verify_step_timers,
)
from voiney_lab.experiment_protocol_pdf import (
    PDF_MEDIA_TYPE,
    ProtocolPdfExtraction,
    ProtocolPdfPage,
    SourceConversion,
    extract_protocol_pdf,
)
from voiney_lab.model_providers import ProviderRetryEvent, observe_provider_retries
from voiney_lab.experiment_protocol_store import (
    AnalysisRevisionRecord,
    ProtocolRevisionRecord,
    ProtocolStore,
    serialize_analysis,
)
# The store's own canonical encoding of domain records: a review-pending
# analysis (lane RV, decision a) is kept in a ledger event, because its
# refused quotes keep it out of the analysis table the domain validation
# guards.
from voiney_lab.experiment_protocol_store import (  # noqa: E402
    _decode_domain as _decode_stored_record,
    _encode_domain as _encode_stored_record,
)
from voiney_lab.protocol_ocr_providers import TEXT_LAYER as OCR_TEXT_LAYER_PROVIDER
# The number-and-unit tokens the two OCR engines' readings were compared by
# (lane P2/P3): an OCR value item (lane RV2, decision b) is named by the same
# token, so "both engines read it alike" means the comparison's own sense.
from voiney_lab.protocol_ocr_providers import (  # noqa: E402
    _NUMBER_WITH_UNIT as _OCR_NUMBER_WITH_UNIT,
    _UNIT_SPELLINGS as _OCR_UNIT_SPELLINGS,
    CLOVA as OCR_CLOVA,
    GOOGLE as OCR_GOOGLE,
    numeric_tokens,
)
from voiney_lab.protocol_ocr import (
    OcrResult,
    ProtocolOcrProvider,
    ocr_result_payload,
    validate_ocr_result,
)


_SAFE_FILENAME = re.compile(r"^[^/\\\x00]{1,255}\.pdf$", re.IGNORECASE)
_STABLE_PROTOCOL_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._:-]{0,127}$")
_REVISION_ID = re.compile(r"^pdf-(\d+)(?:-analysis-(\d+)|-review-(\d+))?$")


_ANALYSIS_REQUESTED_EVENT = "protocol_analysis_requested"
_ANALYSIS_STARTED_EVENT = "protocol_analysis_started"
_ANALYSIS_READY_EVENT = "protocol_analysis_ready"
_ANALYSIS_FAILED_EVENT = "protocol_analysis_failed"
_ANALYSIS_RETRY_EVENT = "protocol_analysis_retry_started"
#: Lane AN, decision 1 (2026-10-07), changing lane PA's "no automatic
#: retry": an analysis whose response broke the structure the domain
#: validation demands -- and only that failure, once per analysis request
#: (since lane AN's EV4 step, 2026-10-13, a person's "분석 다시 시도" too) --
#: is sent once more, the same request, by itself. A failed source-evidence
#: check, a time-out and a missing provider are not: sending them again costs
#: a call and settles nothing a person need not see.
#: Measured 2026-10-06 (lane PX): ANKOM ended so once and passed on other runs.
AUTOMATIC_RETRY_FAILURE_CODES = frozenset({"protocol_analysis_invalid_response"})
AUTOMATIC_RETRY_LIMIT = 1
AUTOMATIC_RETRY_AUTHORITY = "automatic_invalid_response_retry"
#: Lane AQ, decision 4 (2026-10-10): the reasoning level the one automatic
#: retry asks when the first answer was cut at the provider's output limit --
#: one level below the first request, the lowest staying where it is. DS-2
#: (2026-10-09) measured 51 of 195 Gemini 3.8 Flash "high" analysis calls
#: cut at 65,536 output tokens, a median 61k of them thinking, and every
#: same-setting retry cut again; the live server lost 12m47s on PMC8250384
#: that way on 2026-10-10 02:01. An answer that broke the structure for any
#: other reason is still sent again unchanged (lane AN).
TRUNCATED_RETRY_REASONING: dict[str, str] = {
    "xhigh": "high", "high": "medium", "medium": "low", "low": "low",
}
#: The OpenAI-style finish reason of an answer cut at the output limit. xAI
#: and OpenAI chat completions say so themselves; the Gemini adapter reports
#: MAX_TOKENS as this (model_providers). Any other or missing reason keeps
#: the retry as it was.
TRUNCATED_FINISH_REASON = "length"
_CHUNK_PLAN_EVENT = "protocol_chunk_plan_created"
_CHUNK_STARTED_EVENT = "protocol_chunk_analysis_started"
_CHUNK_COMPLETED_EVENT = "protocol_chunk_analysis_completed"
_CHUNK_FAILED_EVENT = "protocol_chunk_analysis_failed"
_MERGE_STARTED_EVENT = "protocol_chunk_merge_started"
_MERGE_CONFLICT_EVENT = "protocol_chunk_merge_conflict"
_REVIEW_REQUIRED_EVENT = "protocol_chunk_review_required"
_SINGLE_REVIEW_REQUIRED_EVENT = "protocol_review_required"
_RUN_CANCELLED_EVENT = "protocol_chunk_run_cancelled"
_DEVELOPMENT_FIXTURE_EVENT = "development_fixture_materialized"
#: An analysis that ended with items a person must check (lane RV, decision
#: a, human decision 2026-10-12): kept in this event, never as an analysis
#: revision, so nothing that runs can read it.
_REVIEW_PENDING_EVENT = "protocol_analysis_review_pending"
#: Lane RX1, decision 4: a person's own name for a protocol -- one more
#: append-only event; the latest one is the title everywhere it is read. The
#: analysis's title, the file's metadata and the file name stay as they are.
PROTOCOL_TITLE_RENAMED_EVENT = "protocol_title_renamed"
PROTOCOL_TITLE_MAX_CHARS = 120
#: The uploader's answer to the public/commercial-source question. It is a
#: declaration, not a licence check or an execution approval.
PROTOCOL_SOURCE_DISCLOSURE_EVENT = "protocol_source_disclosure_recorded"
PROTOCOL_SOURCE_NOTE_MAX_CHARS = 500
#: A person's review registered as a version that runs (lane RV, decision c):
#: the version's protocol is one more analysis revision, named
#: ``review-v<N>-...``, and this event beside it holds who registered it,
#: what it was reviewed from and every change with its original value.
_REVIEW_VERSION_EVENT = "protocol_review_version_registered"
#: Lane RP3, decision ④: a PI's "확인함" on a registered review version.
_REVIEW_VERSION_PI_CONFIRMED_EVENT = "protocol_review_version_pi_confirmed"
REVIEW_VERSION_ANALYSIS_PREFIX = "review-v"
#: Lane RV2, decision c (human decision 2026-10-16): an analysis with more
#: items for a person than ``review_max_items()`` makes no cards; it is
#: analysed once more by itself, one reasoning level up, within
#: ``review_reanalysis_max_calls()`` model calls -- this event records why,
#: at what level and what came of it. Still too many: the guidance alone.
_AUTO_REANALYSIS_EVENT = "protocol_analysis_auto_reanalysis"
#: Lane AN4, decision 1 (human decision 2026-10-19): the items the evidence
#: check refused were asked about once more; this event records how many,
#: what came of each kind of answer, and each call (no source text).
_REQUOTE_EVENT = "protocol_analysis_evidence_requote"
#: Lane AN5, decision 3 (human decision 2026-10-20): the re-ask runs after the
#: analysis is stored and its cards are shown. It started (the progress line
#: says so until its completed event above), and the cards it closed, with
#: the record the review reads from since.
_REQUOTE_STARTED_EVENT = "protocol_analysis_evidence_requote_started"
_REQUOTE_CLOSED_EVENT = "protocol_analysis_review_cards_closed"
REQUOTE_RUNNING_KO = "원문 위치를 다시 찾는 중"
REQUOTE_CLOSED_KO = "원문 위치를 다시 찾아 확인했어요"
#: A re-ask started this long ago with no end was cut off (the server
#: stopped under it); it is no longer said to be running.
_REQUOTE_STALE_SECONDS = REQUOTE_TIME_LIMIT_SECONDS + 120
TOO_MANY_REVIEW_ITEMS_CODE = "protocol_analysis_too_many_review_items"
REVIEW_MAX_ITEMS_SETTING = "VOINEY_LAB_REVIEW_MAX_ITEMS"
REVIEW_MAX_ITEMS = 12
REVIEW_REANALYSIS_MAX_CALLS_SETTING = "VOINEY_LAB_REVIEW_REANALYSIS_MAX_CALLS"
#: The default bound of the one automatic re-analysis, in model calls, for
#: one document: a single analysis is one call; a long document counts two
#: per piece (each piece may be sent once more). 16 is a document of up to
#: eight pieces (about 110 pages). Basis: lane AN2's ledger, USD 6.09 for 247
#: Gemini 3.8 Flash analysis calls at "low" (about 0.025 a call); one level
#: up spends more on thinking -- at three times that, 16 calls stay under
#: about USD 1.2 for one document.
REVIEW_REANALYSIS_MAX_CALLS = 16
#: One reasoning level up (decision c); the top level stays where it is.
ESCALATED_REASONING: dict[str, str] = {
    "none": "low", "low": "medium", "medium": "high", "high": "xhigh", "xhigh": "xhigh",
}
#: The authority an upload-time OCR acceptance is recorded under (lane PX,
#: decision 3). A person's review carries ``human_review``.
AUTOMATIC_OCR_AUTHORITY = "automatic_upload_ocr"
_OCR_REQUESTED_EVENT = "protocol_ocr_requested"
_OCR_COMPLETED_EVENT = "protocol_ocr_completed"
_OCR_FAILED_EVENT = "protocol_ocr_failed"
_OCR_REVIEWED_EVENT = "protocol_ocr_reviewed"


class ProtocolCatalogError(RuntimeError):
    code = "protocol_catalog_error"


class ProtocolRegistrationError(ProtocolCatalogError):
    code = "protocol_registration_invalid"


class ProtocolTitleInvalidError(ProtocolRegistrationError):
    """A person's name for a protocol that is empty, too long or holds a
    control character (lane RX1, decision 4)."""

    code = "protocol_title_invalid"


class ProtocolCatalogNotFoundError(ProtocolCatalogError):
    code = "protocol_catalog_not_found"


class ProtocolCatalogUnavailableError(ProtocolCatalogError):
    code = "protocol_catalog_unavailable"


class ProtocolAnalysisUnavailableError(ProtocolCatalogError):
    code = "protocol_analysis_not_configured"


class ProtocolOcrRequiredError(ProtocolCatalogError):
    code = "ocr_required"


class ProtocolReanalysisRefusedError(ProtocolCatalogError):
    """The curated development fixture is served as curated, never analysed again."""

    code = "reanalysis_not_allowed"


class ProtocolOcrReviewError(ProtocolCatalogError):
    """An OCR result that does not hold against the source (evidence, not a person's review)."""

    code = "protocol_ocr_review_invalid"


class ProtocolAnalysisTooManyItemsError(ProtocolCatalogError):
    """More items for a person than cards are made for, even after the one
    automatic re-analysis (lane RV2, decision c): the guidance alone."""

    code = TOO_MANY_REVIEW_ITEMS_CODE

    def __init__(
        self, message: str, *, items: int, limit: int, items_after: int | None,
        by_reason: dict[str, int] | None = None,
    ) -> None:
        super().__init__(message)
        self.items = items
        self.limit = limit
        self.items_after = items_after
        self.by_reason = dict(by_reason or {})


class ProtocolChunkedAnalysisRequiredError(ProtocolCatalogError):
    code = "chunked_analysis_required"


class ProtocolAnalysisPageRangeInvalidError(ProtocolCatalogError):
    """The pages chosen to analyse are not pages of the document (lane EASY, 1)."""

    code = "analysis_page_range_invalid"


def _page_range_public(page_range: tuple[int, int], page_count: int) -> dict[str, object]:
    """The chosen pages as the ledger and the screen hold them (lane EASY, 1)."""

    first, last = page_range
    return {
        "first_page": first, "last_page": last, "page_count": page_count,
        "pages_ko": f"{first}쪽" if first == last else f"{first}–{last}쪽",
    }


def _page_range_from(value: object) -> tuple[int, int] | None:
    if not isinstance(value, dict):
        return None
    first, last = value.get("first_page"), value.get("last_page")
    if (isinstance(first, int) and isinstance(last, int)
            and not isinstance(first, bool) and not isinstance(last, bool) and 1 <= first <= last):
        return first, last
    return None


def _shown_page_range(events: Sequence[Any]) -> dict[str, object] | None:
    """The chosen pages of the latest analysis that passed or waits for a
    person, as its ledger event holds them; None for one that read every page."""

    latest = None
    for event in events:
        if event.event_type in {_ANALYSIS_READY_EVENT, _REVIEW_PENDING_EVENT}:
            latest = event
    payload = latest.payload if latest is not None and isinstance(latest.payload, dict) else {}
    shown = payload.get("page_range")
    return dict(shown) if _page_range_from(shown) is not None else None


def _requote_groups(
    draft: ProtocolAnalysisDraft, pieces: Sequence[AnalysisPiece],
) -> list[tuple[int, ...]] | None:
    """The page sets a re-ask may offer: each piece's, or the chosen pages
    (lane EASY, decision 1) -- a quote is never moved off what was analysed."""

    if pieces:
        return [tuple(piece.page_numbers) for piece in pieces]
    if draft.page_range is not None:
        return [tuple(range(draft.page_range[0], draft.page_range[1] + 1))]
    return None


def _within_page_range(places: Sequence[Any], page_range: tuple[int, int] | None) -> tuple[Any, ...]:
    """The possibly-missed places on the chosen pages (lane EASY, 1): the
    check of the lines no quote covers reads only the pages analysed."""

    if page_range is None:
        return tuple(places)
    first, last = page_range
    return tuple(place for place in places if first <= place.page_number <= last)


class _RepairedAnswerWithItems(ProtocolAnalysisEvidenceError):
    """A response read after its brackets were repaired, with items the source
    does not hold (lane RV, decision a, keeping lane EV3's decision 1).

    "고쳐도 안 되면 지금처럼 자동 재시도한다": such an answer is recorded as the
    evidence failure it used to be and sent again once; should the retry not
    come back, the answer is kept for a person's review.
    """

    def __init__(self, draft: ProtocolAnalysisDraft) -> None:
        super().__init__("A repaired response holds items the source does not.")
        self.draft = draft
        self.response_repair = draft.response_repair


class ProtocolSourceStatus(str, Enum):
    YES = "yes"
    NO = "no"
    UNKNOWN = "unknown"


_PROTOCOL_SOURCE_STATUS_KO = {
    ProtocolSourceStatus.YES: "예 — 공개·시판 원문",
    ProtocolSourceStatus.NO: "아니요 — 공개·시판 아님",
    ProtocolSourceStatus.UNKNOWN: "모름 — 확인 필요",
}


@dataclass(frozen=True)
class ProtocolSourceDisclosure:
    """One person's declared source scope, retained in the catalog ledger."""

    status: ProtocolSourceStatus
    note: str = ""
    recorded_at: str | None = None

    def __post_init__(self) -> None:
        if not isinstance(self.status, ProtocolSourceStatus):
            raise ProtocolRegistrationError("Protocol source status is invalid.")
        if not isinstance(self.note, str):
            raise ProtocolRegistrationError("Protocol source note is invalid.")
        note = self.note.strip()
        try:
            note.encode("utf-8")
        except UnicodeEncodeError as exc:
            raise ProtocolRegistrationError(
                "Protocol source note is invalid."
            ) from exc
        if len(note) > PROTOCOL_SOURCE_NOTE_MAX_CHARS or any(
            unicodedata.category(character) == "Cc" and character not in "\n\r\t"
            for character in note
        ):
            raise ProtocolRegistrationError("Protocol source note is invalid.")
        if self.recorded_at is not None and (
            not isinstance(self.recorded_at, str) or not self.recorded_at.strip()
        ):
            raise ProtocolRegistrationError("Protocol source record time is invalid.")
        object.__setattr__(self, "note", note)

    def public_dict(self) -> dict[str, object]:
        return {
            "recorded": True,
            "status": self.status.value,
            "status_ko": _PROTOCOL_SOURCE_STATUS_KO[self.status],
            "note": self.note,
            "recorded_at": self.recorded_at,
        }


def _source_disclosure_public(
    disclosure: ProtocolSourceDisclosure | None,
) -> dict[str, object]:
    if disclosure is not None:
        return disclosure.public_dict()
    return {
        "recorded": False,
        "status": None,
        "status_ko": "확인 기록 없음",
        "note": "",
        "recorded_at": None,
    }


@dataclass(frozen=True)
class ProtocolCatalogEntry:
    protocol_id: str
    title: str
    source_filename: str
    source_sha256: str
    revision_id: str
    readiness_status: str
    analysis_status: str
    step_count: int
    created_at: str
    available_for_execution: bool
    lifecycle_state: str = "uploaded"
    #: The readiness reasons that keep this revision out of execution under
    #: the MVP rule (decision of 2026-10-08): nothing runnable came out of the
    #: analysis, or a page could not be read. Empty when it may run.
    execution_blocker_codes: tuple[str, ...] = ()
    #: The curated development fixture, as against an uploaded PDF's own
    #: analysis. Said on screen and in the report; it grants nothing.
    development_only: bool = False
    #: What marks the stored analysis as older than the analyser (lane RN,
    #: decision 3): ``repeat_count_kind_missing`` when a fixed repeat was
    #: stored before lane EV2 asked how the source states its count. The
    #: Korean line is advice ("다시 분석 권장"); the analysis still runs.
    analysis_outdated_codes: tuple[str, ...] = ()
    analysis_outdated_ko: str | None = None
    #: The latest analysis that ended with items for a person (lane RV,
    #: decision a), when it is newer than any analysis that passed: how many
    #: items, why, and which pages of a long document failed. None otherwise.
    review_pending: dict[str, object] | None = None
    #: The registered review versions, v1 first (lane RV, decisions c and
    #: d): number, revision id, reviewer, date and how many changes. The
    #: latest is the one that runs.
    review_versions: tuple[dict[str, object], ...] = ()
    source_disclosure: ProtocolSourceDisclosure | None = None

    def public_dict(self) -> dict[str, object]:
        return {
            "protocol_id": self.protocol_id,
            "title": self.title,
            "source_filename": self.source_filename,
            "source_sha256": self.source_sha256,
            "revision_id": self.revision_id,
            "readiness_status": self.readiness_status,
            "analysis_status": self.analysis_status,
            "step_count": self.step_count,
            "created_at": self.created_at,
            "available_for_execution": self.available_for_execution,
            "execution_blocker_codes": list(self.execution_blocker_codes),
            "development_only": self.development_only,
            "lifecycle_state": self.lifecycle_state,
            "analysis_outdated_codes": list(self.analysis_outdated_codes),
            "analysis_outdated_ko": self.analysis_outdated_ko,
            "review_pending": self.review_pending,
            "review_versions": list(self.review_versions),
            "current_review_version": (
                self.review_versions[-1] if self.review_versions else None),
            "source_disclosure": _source_disclosure_public(
                self.source_disclosure
            ),
        }


@dataclass(frozen=True)
class ProtocolRegistration:
    entry: ProtocolCatalogEntry
    deduplicated: bool


@dataclass(frozen=True)
class ProtocolDevelopmentBootstrap:
    """Result of explicitly materializing one verified development fixture."""

    entry: ProtocolCatalogEntry
    deduplicated: bool


@dataclass(frozen=True)
class ProtocolAssetResolution:
    path: Path
    sha256: str
    source_page: int
    mime_type: str = "image/svg+xml"


@dataclass(frozen=True)
class ProtocolAnalysisRunStatus:
    protocol_id: str
    candidate_revision_id: str
    analysis_run_id: str | None
    state: str
    total_chunks: int
    completed_chunks: int
    failed_chunks: int
    pending_chunks: int
    chunks: tuple[dict[str, object], ...]
    failure_code: str | None = None
    failure_detail: dict[str, object] | None = None
    merge_status: str | None = None
    restart_behavior: str = "explicit_analysis_request_only"
    lifecycle_state: str = "uploaded"
    #: When the run on screen was requested (the ledger's own time, UTC), so
    #: the screen can say how long the analysis has taken. None before any
    #: request.
    requested_at: str | None = None
    #: The run's automatic retry (lane AN, decision 1): attempt, limit, state
    #: (``in_progress``, ``passed``, ``failed``) and the failure it retried.
    #: None when the run on screen had none.
    automatic_retry: dict[str, object] | None = None
    #: Whether the run on screen was asked for over a passed analysis (lane
    #: RN, decision 2), and whether the revision may run right now -- on the
    #: analysis it has, whatever this run does.
    reanalysis: bool = False
    available_for_execution: bool = False

    def public_dict(self) -> dict[str, object]:
        return {
            "protocol_id": self.protocol_id,
            "candidate_revision_id": self.candidate_revision_id,
            "analysis_run_id": self.analysis_run_id,
            "state": self.state,
            "total_chunks": self.total_chunks,
            "completed_chunks": self.completed_chunks,
            "failed_chunks": self.failed_chunks,
            "pending_chunks": self.pending_chunks,
            "chunks": list(self.chunks),
            "failure_code": self.failure_code,
            "failure_detail": self.failure_detail,
            "merge_status": self.merge_status,
            "restart_behavior": self.restart_behavior,
            "lifecycle_state": self.lifecycle_state,
            "requested_at": self.requested_at,
            "automatic_retry": self.automatic_retry,
            "reanalysis": self.reanalysis,
            "available_for_execution": self.available_for_execution,
        }


def _failures_not_superseded(events: Sequence[Any]) -> tuple[Any, ...]:
    """The events without the analysis failures a later attempt replaced.

    The failure an automatic retry was sent for is not the run's result
    (lane AN, decision 1): while the retry runs the run is analysing, and
    when it passes nothing failed. Nor is a failure an analysis that passed
    afterwards replaced -- a person's "분석 다시 시도" that passed used to
    read as failed, because the run status still carried the earlier code.
    A failure stays in the ledger either way.
    """

    last_superseding = max(
        (index for index, event in enumerate(events)
         if event.event_type in {
             _ANALYSIS_RETRY_EVENT, _ANALYSIS_READY_EVENT, _SINGLE_REVIEW_REQUIRED_EVENT,
             _REVIEW_PENDING_EVENT,
         }),
        default=-1,
    )
    return tuple(
        event for index, event in enumerate(events)
        if not (
            index < last_superseding
            and event.event_type in {_ANALYSIS_FAILED_EVENT, _CHUNK_FAILED_EVENT}
        )
    )


#: The run state a piece run (lane AN, decision 2) ended in, by its failure
#: code; any other failure is "analysis_failed".
_PIECE_RUN_FAILED_STATES = {
    "protocol_analysis_piece_failed": "chunk_analysis_failed",
    "protocol_analysis_piece_conflict": "merge_conflict",
}


def _piece_run_end(
    run_events: Sequence[Any], revision_events: Sequence[Any],
) -> tuple[str, str | None] | None:
    """How a piece run ended, from the lifecycle events after its plan.

    (state, failure code) once the analysis passed or failed after the
    plan; None while the run is still going. A piece run records its pieces
    as chunk events and its end as the analysis' own ready or failed event.
    """

    if not run_events:
        return None
    planned_at = run_events[0].sequence_id
    after = [
        event for event in revision_events
        if event.event_type in {
            _ANALYSIS_FAILED_EVENT, _ANALYSIS_READY_EVENT, _SINGLE_REVIEW_REQUIRED_EVENT,
            _REVIEW_PENDING_EVENT,
        }
        and event.sequence_id > planned_at
    ]
    if not after:
        return None
    last = after[-1]
    if last.event_type == _REVIEW_PENDING_EVENT:
        # Lane RV, decision a: the passed pieces kept, the rest for a person.
        return "review_pending", None
    if last.event_type == _ANALYSIS_FAILED_EVENT:
        payload = last.payload if isinstance(last.payload, dict) else {}
        code = payload.get("failure_code")
        code = code if isinstance(code, str) and code else "analysis_failed"
        return _PIECE_RUN_FAILED_STATES.get(code, "analysis_failed"), code
    return "review_required", None


def _cards_closed_event(pending: Any, events: Sequence[Any]) -> Any | None:
    """The latest record of the cards the re-ask closed on this review-pending
    analysis (lane AN5, decision 3), or None."""

    return next((
        event for event in reversed(events)
        if event.event_type == _REQUOTE_CLOSED_EVENT and isinstance(event.payload, dict)
        and event.payload.get("review_pending_event") == pending.event_id
    ), None)


def _requote_running(pending: Any, events: Sequence[Any]) -> bool:
    """Whether the re-ask over this review-pending analysis is still running
    (lane AN5, decision 3): started, not ended, and not cut off."""

    started = next((
        event for event in reversed(events)
        if event.event_type == _REQUOTE_STARTED_EVENT and isinstance(event.payload, dict)
        and event.payload.get("review_pending_event") == pending.event_id
    ), None)
    if started is None or any(
        event.event_type == _REQUOTE_EVENT and event.sequence_id > started.sequence_id
        and isinstance(event.payload, dict)
        and event.payload.get("analysis_id") == started.payload.get("analysis_id")
        for event in events
    ):
        return False
    try:
        age = (datetime.now(timezone.utc) - datetime.fromisoformat(started.recorded_at)).total_seconds()
    except (TypeError, ValueError):
        return False
    return age <= _REQUOTE_STALE_SECONDS


def _review_pending_summary(event: Any, events: Sequence[Any] = ()) -> dict[str, object]:
    """What the screen says of a review-pending analysis (lane RV, decision a).

    ``needs_review_count`` items, by reason, with one Korean line: "검토 대기 ·
    확인 필요 12건(근거를 못 찾음 9 · OCR 흐림 3)", and "41–59쪽 분석 실패" for
    each failed piece of a long document.

    Lane AN5, decision 3: given the protocol's ``events``, the items left
    after the re-ask closed cards, and " · 원문 위치를 다시 찾는 중" while it runs.
    """

    payload = event.payload if isinstance(event.payload, dict) else {}
    closed = _cards_closed_event(event, events)
    record = closed.payload if closed is not None else payload
    flags = [
        ReviewFlag.from_public(item) for item in record.get("flags") or ()
        if isinstance(item, dict)
    ]
    # Lane AN3, decision A: a page's OCR values are one card, counted once.
    units = _asked_units(flags)
    items = list(units)
    by_reason: dict[str, int] = {}
    for reason in units.values():
        by_reason[reason] = by_reason.get(reason, 0) + 1
    failed_pages = [
        {"first_page": flag.first_page, "last_page": flag.last_page,
         "failure_code": flag.reason_code}
        for flag in flags if flag.first_page is not None
    ]
    words = " · ".join(
        f"{REVIEW_REASONS_KO.get(reason, reason)} {count}"
        + ("쪽" if reason == REVIEW_REASON_OCR_VALUE else "")
        for reason, count in by_reason.items()
        if reason not in {REVIEW_REASON_PIECE_FAILED, REVIEW_REASON_BLOCKED}
    )
    line = f"검토 대기 · 확인 필요 {len(items)}건"
    if words:
        line += f"({words})"
    for flag in flags:
        if flag.first_page is not None:
            # "41–59쪽 분석 실패", or lane RV2 decision f's "구글 안전 필터에 막혀 …".
            line += f" · {flag.reason_ko}"
    requoting = bool(events) and _requote_running(event, events)
    if requoting:
        line += f" · {REQUOTE_RUNNING_KO}"
    summary: dict[str, object] = {
        "analysis_id": payload.get("analysis_id"),
        "recorded_at": event.recorded_at,
        "needs_review_count": len(items),
        "by_reason": by_reason,
        "failed_pages": failed_pages,
        "message_ko": line,
    }
    if requoting:
        summary["requoting"] = True
        summary["requoting_ko"] = REQUOTE_RUNNING_KO
    if closed is not None:
        summary["closed_by_requote"] = len(closed.payload.get("closed") or ())
    return summary


def review_version_suggestion(entry: "ProtocolCatalogEntry") -> dict[str, object] | None:
    """What an upload of a file already registered says (lane RV, decision d).

    Human decision 2026-10-12: the same protocol is not uploaded again; the
    experimenter starts from its review version. None when it has none.
    """

    if not entry.review_versions:
        return None
    current = entry.review_versions[-1]
    return {
        **{key: current.get(key) for key in (
            "version", "version_ko", "revision_id", "reviewer_name", "registered_at",
            "registered_on")},
        "message_ko": (
            f"이미 올린 원문입니다. 다시 올리지 않고 {current.get('version_ko')}"
            f"({current.get('registered_on')} · {current.get('reviewer_name')})로 바로 "
            "시작할 수 있습니다."),
    }


def review_max_items(environment: Mapping[str, str] | None = None) -> int:
    """The most items an analysis may ask a person about by cards (lane RV2, c)."""

    env = os.environ if environment is None else environment
    raw = env.get(REVIEW_MAX_ITEMS_SETTING, "").strip()
    if re.fullmatch(r"\d+", raw) and 1 <= int(raw) <= 500:
        return int(raw)
    return REVIEW_MAX_ITEMS


def review_reanalysis_max_calls(environment: Mapping[str, str] | None = None) -> int:
    """The model calls the one automatic re-analysis may make for one document
    (lane RV2, c); 0 turns it off."""

    env = os.environ if environment is None else environment
    raw = env.get(REVIEW_REANALYSIS_MAX_CALLS_SETTING, "").strip()
    if re.fullmatch(r"\d+", raw) and int(raw) <= 200:
        return int(raw)
    return REVIEW_REANALYSIS_MAX_CALLS


def escalated_analysis_model(
    model: ProtocolAnalysisModel,
) -> tuple[ProtocolAnalysisModel, dict[str, object]]:
    """The analysis model one reasoning level up, and what the event records.

    Only the server's adapter has a level to raise; any other model is used
    as it is (lane RV2, decision c).
    """

    if not isinstance(model, OpenAICompatibleProtocolAnalysisModel):
        return model, {"reasoning_effort": None}
    current = model.reasoning_effort
    raised = ESCALATED_REASONING.get(current) if isinstance(current, str) else None
    if raised is None:
        return model, {"reasoning_effort": None, "previous_reasoning_effort": current}
    return (model if raised == current else replace(model, reasoning_effort=raised)), {
        "reasoning_effort": raised, "previous_reasoning_effort": current}


def _items_by_reason(draft: ProtocolAnalysisDraft) -> tuple[int, dict[str, int]]:
    """How many items the draft asks a person about, and by first reason; a
    page's OCR values count once (lane AN3, decision A3)."""

    reasons = _asked_units(draft.review_flags)
    by_reason: dict[str, int] = {}
    for reason in reasons.values():
        by_reason[reason] = by_reason.get(reason, 0) + 1
    return len(reasons), by_reason


def _korean_date(timestamp: str) -> str:
    """The date of a ledger time in Korea ("2026-10-14")."""

    try:
        moment = datetime.fromisoformat(timestamp)
    except (TypeError, ValueError):
        return str(timestamp)[:10]
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(ZoneInfo("Asia/Seoul")).date().isoformat()


def _pi_confirmations_of(events: Sequence[Any], version: object) -> list[dict[str, object]]:
    """Each PI confirmation of one review version, in the order left (lane RP3, ④)."""

    found = []
    for event in events:
        if event.event_type != _REVIEW_VERSION_PI_CONFIRMED_EVENT:
            continue
        payload = event.payload if isinstance(event.payload, dict) else {}
        if payload.get("version") != version:
            continue
        found.append({
            "name": str(payload.get("name") or ""), "comment": str(payload.get("comment") or ""),
            "at": event.recorded_at, "on": _korean_date(event.recorded_at),
        })
    return found


def _review_version_summary(event: Any, confirmations: Sequence[Any] = ()) -> dict[str, object]:
    """One registered review version as the list shows it (lane RV, c, d)."""

    payload = event.payload if isinstance(event.payload, dict) else {}
    changes = [item for item in payload.get("changes") or () if isinstance(item, dict)]
    return {
        "pi_confirmations": _pi_confirmations_of(confirmations, payload.get("version")),
        "version": payload.get("version"),
        "version_ko": f"검토본 v{payload.get('version')}",
        "revision_id": payload.get("revision_id"),
        "reviewer_name": payload.get("reviewer_name"),
        "registered_at": event.recorded_at,
        "registered_on": _korean_date(event.recorded_at),
        "base": payload.get("base"),
        "base_revision_id": payload.get("base_revision_id"),
        "change_count": sum(
            1 for item in changes if item.get("action") in {"fix", "enter", "clear"}),
        "confirmed_count": sum(1 for item in changes if item.get("action") == "confirm"),
    }


def _automatic_retry_status(events: Sequence[Any]) -> dict[str, object] | None:
    """The automatic retry of the run on screen, or None.

    A retry belongs to the run a person's request opened; a later request
    (a person pressing "분석 다시 시도") is a new run without one.
    """

    retry_index = max(
        (index for index, event in enumerate(events)
         if event.event_type == _ANALYSIS_RETRY_EVENT),
        default=-1,
    )
    if retry_index < 0 or any(
        event.event_type == _ANALYSIS_REQUESTED_EVENT
        for event in events[retry_index + 1:]
    ):
        return None
    retry = events[retry_index]
    payload = retry.payload if isinstance(retry.payload, dict) else {}
    state = "in_progress"
    for event in events[retry_index + 1:]:
        if event.event_type == _ANALYSIS_FAILED_EVENT:
            state = "failed"
        elif event.event_type in {_ANALYSIS_READY_EVENT, _SINGLE_REVIEW_REQUIRED_EVENT}:
            state = "passed"
        elif event.event_type == _REVIEW_PENDING_EVENT:
            state = "review_pending"
    status: dict[str, object] = {
        "attempt": payload.get("attempt", 1),
        "limit": payload.get("limit", AUTOMATIC_RETRY_LIMIT),
        "state": state,
        "reason_code": payload.get("reason_code"),
    }
    # A retry sent for an answer cut at the output limit says so, and what
    # reasoning level it asked (lane AQ, decision 4); any other reads as
    # lane AN wrote it.
    for key in ("truncated", "reasoning_effort", "previous_reasoning_effort"):
        if key in payload:
            status[key] = payload[key]
    return status


def _truncated_retry_ko(retry: dict[str, object]) -> str:
    """How the pipeline line says a retry was sent for a cut answer."""

    level = retry.get("reasoning_effort")
    previous = retry.get("previous_reasoning_effort")
    if level and previous and level != previous:
        return (f"첫 응답이 출력 한도에서 잘려 추론 수준을 한 단계 낮춰({previous}→{level}) "
                "같은 요청을")
    return "첫 응답이 출력 한도에서 잘렸고 추론 수준을 더 낮출 수 없어 같은 요청을 같은 설정으로"


def _response_repair_of(payload: object) -> dict[str, object] | None:
    """The structural repair an analysis attempt's event records (lane EV3,
    decision 1), or None for a response read as it came."""

    repair = payload.get("response_repaired") if isinstance(payload, dict) else None
    if not isinstance(repair, dict) or not isinstance(repair.get("bracket_count"), int):
        return None
    return repair


def _response_repair_ko(repair: dict[str, object]) -> str:
    """What the pipeline line says about a response read in a repaired shape."""

    return f"응답 형식을 고쳐 읽음(괄호 {repair['bracket_count']}개)"


class _FinishReasonRecorder:
    """An OpenAI-style client that remembers the last reply's finish reason.

    Wrapped around the analysis model's client for one analysis so the
    catalog can tell an answer cut at the output limit (finish reason
    "length") from one that broke the structure another way (lane AQ,
    decision 4). The request and the reply pass through unchanged; nothing
    here judges the reply.
    """

    def __init__(self, inner: Any) -> None:
        self._inner = inner
        self.finish_reason: str | None = None
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

    def _create(self, **request: Any) -> Any:
        self.finish_reason = None
        response = self._inner.chat.completions.create(**request)
        choices = getattr(response, "choices", None) or ()
        if choices:
            reason = getattr(choices[0], "finish_reason", None)
            self.finish_reason = reason if isinstance(reason, str) else None
        return response


@dataclass(frozen=True)
class _ObservedAnalysisModel:
    """The analysis model with its provider client observed.

    Only the server's adapter (``OpenAICompatibleProtocolAnalysisModel``)
    has a client to observe; any other ``ProtocolAnalysisModel`` is passed
    through untouched and its retry is the same object, as before.
    """

    model: ProtocolAnalysisModel
    recorder: _FinishReasonRecorder | None

    @classmethod
    def wrap(cls, model: ProtocolAnalysisModel) -> "_ObservedAnalysisModel":
        if not isinstance(model, OpenAICompatibleProtocolAnalysisModel):
            return cls(model, None)
        recorder = _FinishReasonRecorder(model.client)
        return cls(replace(model, client=recorder), recorder)

    @property
    def truncated(self) -> bool:
        """Whether the last answer was cut at the provider's output limit."""

        return (
            self.recorder is not None
            and self.recorder.finish_reason == TRUNCATED_FINISH_REASON
        )

    def retry(self) -> tuple[ProtocolAnalysisModel, dict[str, object]]:
        """The model the automatic retry calls, and what its event records.

        Unchanged unless the last answer was cut; then the same model one
        reasoning level lower (``TRUNCATED_RETRY_REASONING``), or the same
        model when its level is not known or already the lowest.
        """

        if not self.truncated:
            return self.model, {}
        current = getattr(self.model, "reasoning_effort", None)
        lowered = TRUNCATED_RETRY_REASONING.get(current) if isinstance(current, str) else None
        if lowered is None:
            return self.model, {"truncated": True}
        model = self.model if lowered == current else replace(self.model, reasoning_effort=lowered)
        return model, {
            "truncated": True,
            "reasoning_effort": lowered,
            "previous_reasoning_effort": current,
        }


#: What a person does after an analysis failure, by failure code (lane PA
#: decision 4). Every code not listed gets the default.
_ANALYSIS_RECOVERY_ACTIONS: dict[str, str] = {
    "provider_configuration_missing": (
        "분석 역할 설정이 없습니다. .env 의 VOINEY_LAB_ANALYSIS_PROVIDER·"
        "VOINEY_LAB_ANALYSIS_MODEL 과 그 공급자의 API 키를 확인하고 서버를 다시 "
        "시작한 뒤 '분석 다시 시도'를 누르세요."
    ),
    "protocol_analysis_timeout": (
        "분석 호출이 제한 시간(VOINEY_LAB_PROTOCOL_ANALYSIS_TIMEOUT_SECONDS, "
        "기본 600초) 안에 끝나지 않았습니다. '분석 다시 시도'를 누르거나, 긴 "
        "문서라면 제한 시간을 늘리세요."
    ),
    "protocol_analysis_invalid_evidence": (
        "분석 모델이 낸 근거가 원문 쪽의 글과 맞지 않아 결과를 쓰지 않았습니다. "
        "'분석 다시 시도'를 누르면 새로 분석합니다(분석 모델 호출 비용이 듭니다). "
        "원문은 바뀌지 않습니다."
    ),
    "protocol_analysis_invalid_response": (
        "분석 모델의 응답이 정해진 형식(JSON 스키마)에 맞지 않아 결과를 쓰지 않았습니다. "
        "'분석 다시 시도'를 누르면 새로 분석합니다(분석 모델 호출 비용이 듭니다). "
        "원문은 바뀌지 않습니다."
    ),
    "ocr_required": (
        "이 PDF 는 읽을 수 있는 글자가 없는 쪽이 있어 OCR 글이 먼저 필요합니다. "
        "업로드 때 OCR 이 자동으로 돌지 않았다면 OCR 공급자 설정"
        "(VOINEY_LAB_OCR_PROVIDERS 와 그 엔진의 키)을 확인하고 서버를 다시 시작한 뒤 "
        "PDF 를 다시 올리거나 'OCR 텍스트 추출'을 누르세요. OCR 이 끝나면 분석이 "
        "바로 시작됩니다."
    ),
    "protocol_pdf_too_large": (
        "PDF 가 등록 한도보다 큽니다. 더 작은 파일로 다시 올리세요."
    ),
    # Lane EV2, decision 4.
    "protocol_analysis_incomplete": (
        "원문에 번호가 붙은 단계가 분석 결과보다 훨씬 많아, 원문 일부만 읽은 분석으로 보고 "
        "쓰지 않았습니다. '분석 다시 시도'를 누르면 새로 분석합니다(분석 모델 호출 비용이 "
        "듭니다). 다시 해도 같으면 절차가 있는 쪽만 골라 PDF 로 만들어 올려 주세요."
    ),
    "protocol_analysis_too_many_pages": (
        f"분석하는 문서는 {MAX_PIECED_ANALYSIS_PAGES}쪽까지입니다. 실험에 쓸 절차가 있는 쪽만 "
        "골라 PDF 로 만들어 다시 올려 주세요(쪽 범위를 고르는 화면은 아직 없습니다). 분석 "
        "모델은 부르지 않았고 원문은 바뀌지 않습니다."
    ),
    # Lane AN, decision 2: a long document read in pieces.
    "protocol_analysis_piece_failed": (
        "실패한 조각은 자동으로 한 번 다시 보냈지만 다시 실패해, 나머지 조각의 결과도 쓰지 "
        "않았습니다(일부만 읽은 분석은 쓰지 않습니다). '분석 다시 시도'를 누르면 문서 전체를 "
        "새로 분석합니다(조각 수만큼 분석 모델 호출 비용이 듭니다). 원문은 바뀌지 않습니다."
    ),
    "protocol_analysis_piece_conflict": (
        "두 조각이 함께 읽은 쪽을 서로 다르게 읽어 결과를 쓰지 않았습니다. '분석 다시 시도'를 "
        "누르면 문서 전체를 새로 분석합니다(조각 수만큼 분석 모델 호출 비용이 듭니다). 원문은 "
        "바뀌지 않습니다."
    ),
    # Lane RV2, decision c.
    TOO_MANY_REVIEW_ITEMS_CODE: (
        "'다시 분석'을 누르거나, 실험에 쓸 절차가 있는 쪽만 골라 PDF 로 만들어 다시 올려 "
        "주세요. 원문은 바뀌지 않습니다."
    ),
    # Lane RV2, decision f.
    BLOCKED_FAILURE_CODE: (
        "모델 공급자의 안전 필터가 이 원문을 읽지 못하게 막았습니다(같은 요청은 다시 보내도 "
        "같게 막힙니다). 막힌 쪽을 빼고 실험에 쓸 절차가 있는 쪽만 골라 PDF 로 만들어 다시 "
        "올리거나, 안전 필터 설정이 바뀐 뒤 '분석 다시 시도'를 누르세요. 원문은 바뀌지 않습니다."
    ),
}
#: What a person does with an analysis waiting for review (lane RV, a).
REVIEW_PENDING_ACTION_KO = (
    "검토 화면에서 확인 필요 항목을 확인·고침·비움 중 하나로 정하고, 검토자 이름을 적어 "
    "'검토 완료하고 등록'을 누르세요. 원문이 바뀌었다면 '다시 분석'을 누르세요."
)
_DEFAULT_ANALYSIS_RECOVERY_ACTION = (
    "실패 원인을 확인한 뒤 '분석 다시 시도'를 누르세요. 다시 시도해도 원문은 "
    "바뀌지 않습니다."
)
#: What a person does after a re-analysis over a passed analysis failed
#: (lane RN, decision 2): nothing is lost, the pass before it still runs.
_REANALYSIS_RECOVERY_ACTION = (
    "'다시 분석'을 눌러 다시 할 수 있습니다. 원문과 지금 분석은 바뀌지 않습니다."
)
_NOT_RETRYABLE_FAILURE_CODES = frozenset({
    "ocr_required", "protocol_pdf_too_large", "protocol_analysis_too_many_pages",
})

#: One Korean line per readiness reason: what blocks execution and what a
#: person does about it (lane PX, decision 4 of 2026-10-06). The English
#: ``ReadinessReason.message`` the domain records is unchanged; the review
#: carries this beside it as ``message_ko``. ``{pages}`` is filled from the
#: recorded message by ``readiness_reason_korean``.
READINESS_REASON_KO: dict[str, str] = {
    "invalid_protocol": "구조화된 프로토콜이 올바르지 않습니다. '분석 다시 시도'를 누르세요.",
    "source_text_cross_check_failed": (
        "옛 추출 대조 판정입니다(2026-10-02 뒤로는 만들어지지 않음). 분석을 다시 돌리면 "
        "사라집니다."),
    "source_text_cross_check_unavailable": (
        "옛 추출 대조 판정입니다(2026-10-02 뒤로는 만들어지지 않음). 분석을 다시 돌리면 "
        "사라집니다."),
    "no_executable_steps": (
        "실행할 단계를 원문에서 찾지 못했습니다. 원문을 확인하고 '분석 다시 시도'를 누르세요."),
    "unsupported_conditional_branch": (
        "조건·선택 분기(예: '새로 만든 튜브면 36~41단계를 두 번 더 반복')가 있습니다. "
        "그 단계에 이르면 조건에 해당하는지 묻고 답대로 안내합니다."),
    "unsupported_fixed_range_repetition": (
        "정해진 범위의 단계 반복이 있습니다. 범위의 끝에서 회차를 세어 안내합니다."),
    "unsupported_operator_determined_repetition": (
        "작업자가 횟수를 정하는 반복이 있습니다. 범위의 첫 단계에서 몇 번 하실지 묻습니다."),
    "unsupported_repeat_until": (
        "조건이 될 때까지 반복하는 지시가 있습니다. 그 단계에서 관찰 결과를 묻습니다."),
    "unsupported_parallel_background_work": (
        "동시에 또는 백그라운드로 진행하는 작업은 아직 안내 기능이 없습니다. 그 단계에 이르면 "
        "한 번 알리고 원문을 읽어 드립니다."),
    "unsupported_recurring_reminder": (
        "반복 알림은 아직 안내 기능이 없습니다. 그 단계에 이르면 한 번 알리고 원문을 읽어 "
        "드립니다."),
    "unsupported_recurring_action": (
        "일정 간격으로 되풀이하는 동작은 아직 안내 기능이 없습니다. 그 단계에 이르면 한 번 "
        "알리고 원문을 읽어 드립니다."),
    "unsupported_reusable_subprocedure": (
        "다른 곳에서 다시 쓰는 하위 절차는 아직 안내 기능이 없습니다. 그 단계에 이르면 한 번 "
        "알리고 원문을 읽어 드립니다."),
    "unresolved_ambiguity": (
        "원문에 서로 다른 두 서술이 있어 어느 쪽이 맞는지 정해지지 않았습니다. 그 단계에서는 "
        "원문을 그대로 읽어 드리니 실험자가 원문을 보고 판단하세요."),
    "unresolved_execution_value_conflict": (
        "실행에 쓰는 값이 원문 안에서 서로 다릅니다. 그 단계에서 원문을 그대로 읽어 드리니 "
        "실험자가 어느 값을 쓸지 정하세요."),
    "safety_critical_conflict": (
        "안전에 중요한 값이 원문 안에서 서로 달라 실행할 수 없습니다. 원문을 확인하고 고친 "
        "PDF 를 올리세요."),
    "no_declared_safety_warnings": (
        "원문의 안전 주의는 시작 전 화면에 모아 보이고, 각 단계에서 원문 그대로 읽어 드립니다. "
        "시작을 누르면 안전 주의를 확인한 것으로 기록합니다."),
    "source_page_requires_ocr": (
        "원문 {pages}쪽은 글자 층이 없어 OCR 글이 필요합니다. 업로드 때 OCR 이 돌고 나면 "
        "다음 분석부터 이 사유는 사라집니다."),
    "source_page_not_fully_read": (
        "분석이 원문 {pages}쪽에 적힌 값을 다 다루지 못했습니다. 그 쪽의 단계에 이르면 쪽 "
        "원문을 그대로 보여 드리니 직접 확인하세요."),
    "declined_value_not_resolved": (
        "분석이 지시가 아니라고 본 원문 값이 있습니다. 원문 단계를 그대로 읽어 드리니 "
        "실험자가 확인하세요."),
    "excessive_declined_values": (
        "지시가 아니라고 본 원문 값이 너무 많습니다. 쪽을 제대로 읽지 않았을 수 있으니 "
        "원문을 함께 보며 진행하세요."),
    "source_states_an_uncaptured_repetition": (
        "원문에 적힌 반복 지시를 분석이 놓쳤습니다. 그 단계의 원문을 그대로 읽어 드리니 "
        "실험자가 반복을 확인하세요."),
    "unconfirmed_fixed_repetition": (
        "정해진 반복 횟수는 분석이 원문에서 읽은 값입니다. 범위의 끝에서 그 횟수로 회차를 "
        "세어 안내합니다."),
    "missing_execution_critical_value": (
        "실행에 꼭 필요한 값이 원문에 없습니다. 그 단계에서 원문을 그대로 읽어 드리니 "
        "실험자가 값을 정하세요."),
    # Lane MISS, decision 4; the start screen's line is possibly_missed_notices'.
    "source_text_possibly_missed": (
        "분석에 없을 수 있는 원문이 {pages}쪽에 있어요. 시작 전에 원문을 확인해 주세요."),
}
#: What a person does about each execution blocker, for the pipeline line.
_BLOCKER_ACTIONS: dict[str, str] = {
    "no_executable_steps": "원문을 확인하고 '분석 다시 시도'를 누르세요.",
    "invalid_protocol": "'분석 다시 시도'를 누르세요.",
    "source_page_requires_ocr": (
        "OCR 공급자가 설정돼 있으면 'OCR 텍스트 추출'을 누르세요. 없으면 글자 층이 있는 "
        "PDF 를 올리세요."),
    "safety_critical_conflict": "원문을 확인하고 고친 PDF 를 올리세요.",
}
_READINESS_LABEL_KO = {
    "guidance_ready": "실행 안내 준비 완료",
    "analysis_required": "시작 전에 읽을 알림이 있음",
}
_ANALYSIS_REQUIRED_LABEL_KO = "구조 분석이 아직 끝나지 않았습니다."
_ANALYSIS_REQUIRED_MESSAGE_KO = "실행에 앞서 구조 분석이 통과해야 합니다. 업로드 직후 자동으로 시작됩니다."
_PAGE_NUMBERS = re.compile(r"(?<![\w.])\d+(?![\w.])")

#: Where a document stands on its way to a start, in Korean (decision 4).
PIPELINE_STAGE_KO: dict[str, str] = {
    "extraction": "원문 읽기",
    "ocr": "OCR",
    "analysis": "분석",
    "evidence": "근거 대조",
    # Lane RV, decision a: an analysis with items a person must check.
    "review": "검토",
    "readiness": "실행 준비",
    "translation": "번역",
    "ready": "실행 가능",
}
_ANALYSIS_FAILURE_KO: dict[str, str] = {
    "provider_configuration_missing": "분석 모델 설정이 없어 분석을 시작하지 못했습니다.",
    "protocol_analysis_timeout": "분석 호출이 제한 시간 안에 끝나지 않았습니다.",
    "protocol_analysis_invalid_evidence": (
        "분석 모델이 낸 근거가 원문 쪽의 글과 맞지 않아 결과를 쓰지 않았습니다."),
    "protocol_analysis_invalid_response": (
        "분석 모델의 응답이 정해진 형식에 맞지 않아 결과를 쓰지 않았습니다."),
    "protocol_analysis_not_configured": "분석 모델 설정이 없어 분석을 시작하지 못했습니다.",
    "ocr_required": "글자 층이 없는 쪽의 OCR 글이 아직 없어 분석할 수 없습니다.",
    "protocol_pdf_too_large": "PDF 가 등록 한도보다 큽니다.",
    "protocol_analysis_incomplete": "분석이 원문 단계를 다 담지 못했어요.",
    "protocol_analysis_too_many_pages": (
        f"원문이 {MAX_PIECED_ANALYSIS_PAGES}쪽을 넘어 분석하지 않았어요. 분석할 쪽 범위를 "
        "줄여야 합니다."),
    "protocol_analysis_piece_failed": "일부 조각 분석 실패",
    BLOCKED_FAILURE_CODE: "모델 공급자의 안전 필터에 막혀 분석하지 못했어요.",
    TOO_MANY_REVIEW_ITEMS_CODE: (
        "확인할 항목이 너무 많아 검토 카드를 만들지 않았어요. 이 문서는 다시 분석하거나 "
        "절차가 있는 쪽만 올리는 편이 나아요."),
    "protocol_analysis_piece_conflict": (
        "조각을 합치다 같은 쪽을 두 조각이 서로 다르게 읽은 곳이 있어 결과를 쓰지 않았습니다."),
    "analysis_cancelled": "분석이 취소되었습니다.",
    "chunk_analysis_failed": "큰 문서의 일부 분석이 실패했습니다.",
    "merge_conflict": "큰 문서의 분석 결과를 합치다 충돌이 났습니다.",
}


def _duration_ko(seconds: int) -> str:
    """"12시간", "1시간 30분", "15분", "30초": a timer value said in Korean."""

    hours, rest = divmod(int(seconds), 3600)
    minutes, secs = divmod(rest, 60)
    parts = [
        f"{value}{unit}"
        for value, unit in ((hours, "시간"), (minutes, "분"), (secs, "초"))
        if value
    ]
    return " ".join(parts) or "0초"


def _review_timers(table: domain.StepTimerTable) -> dict[str, list[dict[str, object]]]:
    """The start screen's timer list (lane PT, decision 4).

    ``verified``: each step's timer values with the source excerpt they were
    read from; a step with more than one value is a choice the experimenter
    makes when starting the timer. ``bound`` says whether the source states
    the value as a minimum, an approximate value or a maximum, and
    ``bound_ko`` is the word the screen shows before it -- "최소", "약",
    "최대" (lane VX, decision 4). ``refused``: each time the server did not
    make a timer from, with the reason in Korean.
    """

    choices = table.choices()
    return {
        "verified": [
            {
                "step_id": timer.step_id,
                "step_label": timer.source_label,
                "source_literal": timer.literal,
                "source_excerpt": timer.excerpt,
                "source_page_number": timer.page_number,
                "seconds": list(timer.seconds),
                "value_ko": " 또는 ".join(_duration_ko(s) for s in timer.seconds),
                "choice": timer.step_id in choices,
                "source": timer.source,
                "bound": timer.bound,
                "bound_ko": domain.TIMER_BOUND_KO.get(timer.bound, ""),
            }
            for timer in table.verified
        ],
        "refused": [
            {
                "step_id": item.step_id,
                "step_label": item.source_label,
                "source_literal": item.literal,
                "source_excerpt": item.excerpt,
                "source_page_number": item.page_number,
                "reason": item.reason,
                "reason_ko": item.reason_ko,
            }
            for item in table.refused
        ],
    }


#: Lane RN, decision 3: the one line the list and the start screen show over
#: an analysis stored before the analyser improved.
ANALYSIS_OUTDATED_KO = "분석기 개선 전 분석 · 다시 분석 권장"
ANALYSIS_OUTDATED_REPEAT_KIND = "repeat_count_kind_missing"


def _analysis_outdated_codes(protocol: domain.ExperimentProtocol) -> tuple[str, ...]:
    """What marks a stored analysis as older than the analyser, if anything.

    Lane RN, decision 3: a fixed repeat stored without ``repeat_count_kind``
    was read before lane EV2 asked how the source states its count, so its
    count may be one round short (headspace: "three … in total" stored as
    2). The analysis is still served and still runs; this is the advice to
    analyse it again. An analysis with no fixed repeat, or whose repeats
    carry a kind, cannot be told from a current one and gets nothing.
    """

    if any(
        isinstance(construct, domain.FixedRangeRepetition)
        and construct.repeat_count_kind is None
        for construct in protocol.constructs
    ):
        return (ANALYSIS_OUTDATED_REPEAT_KIND,)
    return ()


def _repeat_count_checks_ko(protocol: domain.ExperimentProtocol) -> list[str]:
    """One line per fixed repeat whose count the experimenter should check.

    Lane EV2, decision 3: a count is the total number of runs. A count
    stored without saying how the source states it (every analysis before
    the decision) is kept as it was and needs checking; a count the source
    states two ways is asked before the start. A count the source states as
    a total, or as runs after the first (made a total by the server), needs
    no line.
    """

    labels = {
        step.step_id: step.source_label
        for section in protocol_with_display_step_labels(protocol).sections
        for step in section.steps
    }
    lines: list[str] = []
    for construct in protocol.constructs:
        if not isinstance(construct, domain.FixedRangeRepetition) or (
            construct.repeat_count_kind
            in {domain.REPEAT_COUNT_TOTAL, domain.REPEAT_COUNT_ADDITIONAL}
        ):
            continue
        where = (
            f"{labels.get(construct.start_step_id, construct.start_step_id)}~"
            f"{labels.get(construct.end_step_id, construct.end_step_id)}단계"
        )
        source = " ".join(construct.range_source_text.split())
        if construct.repeat_count_kind == domain.REPEAT_COUNT_AMBIGUOUS:
            how = (
                "처음을 포함한 총 횟수인지 더 하는 횟수인지 원문이 두 가지로 읽혀 "
                "시작할 때 여쭤봅니다"
            )
        elif construct.repeat_count is not None:
            how = (
                f"분석이 {construct.repeat_count}회로 읽음"
                "(처음을 포함한 총 횟수인지 원문을 확인하세요)"
            )
        else:
            how = "분석이 횟수를 읽지 않음(원문을 확인하세요)"
        lines.append(f"반복 횟수 확인 필요 · {where} · 원문 “{source}” · {how}")
    return lines


def _safety_notice_sources(protocol: Any) -> list[dict[str, object]]:
    """Every safety statement the source declares, step by step, verbatim.

    The screen the experimenter reads before pressing start (decision of
    2026-10-08): the step's warnings and its sub-actions' warnings in source
    order, each with the page it stands on. The words are the document's own
    (``evidence.source_excerpt`` where the server reconstructed one, else the
    statement text); nothing is summarised.
    """

    items: list[dict[str, object]] = []
    seen: set[tuple[str, str]] = set()
    for section in protocol.sections:
        for step in section.steps:
            statements = list(step.warnings)
            for action in step.sub_actions:
                statements.extend(action.warnings)
            for index, statement in enumerate(statements, 1):
                source = " ".join(
                    str(statement.evidence.source_excerpt
                        or statement.source_text or "").split()
                )
                if not source or (step.step_id, source) in seen:
                    continue
                seen.add((step.step_id, source))
                items.append({
                    "step_id": step.step_id,
                    "step_label": step.source_label,
                    "statement_id": statement.statement_id,
                    "warning_index": index,
                    "source_page_number": statement.evidence.source_page_number,
                    "source_text": source,
                })
    return items


def readiness_reason_korean(code: str, message: str | None = None) -> str | None:
    """The Korean line for one readiness reason, pages filled in."""

    template = READINESS_REASON_KO.get(code)
    if template is None:
        return None
    if "{pages}" not in template:
        return template
    numbers = _PAGE_NUMBERS.findall(message or "")
    if code == "source_page_not_fully_read":
        pages = numbers[0] if numbers else "일부"
    else:
        pages = ", ".join(numbers) if numbers else "일부"
    return template.format(pages=pages)


def _with_korean_reasons(readiness: dict[str, object]) -> dict[str, object]:
    status = str(readiness.get("status") or "")
    readiness["label_ko"] = _READINESS_LABEL_KO.get(status, readiness.get("label"))
    reasons = readiness.get("reasons")
    if isinstance(reasons, list):
        for reason in reasons:
            if isinstance(reason, dict):
                reason["message_ko"] = readiness_reason_korean(
                    str(reason.get("code") or ""), str(reason.get("message") or ""))
    return readiness


def _text_sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _protocol_id(checksum: str) -> str:
    return f"protocol-{checksum[:32]}"


def _revision_id(
    protocol_revision_number: int,
    analysis_revision_number: int | None = None,
) -> str:
    value = f"pdf-{protocol_revision_number}"
    if analysis_revision_number is not None:
        value += f"-analysis-{analysis_revision_number}"
    return value


def _parse_revision_id(value: str) -> tuple[int, int | None]:
    match = _REVISION_ID.fullmatch(value)
    if match is None:
        raise ProtocolCatalogNotFoundError("Protocol revision is unknown.")
    return int(match.group(1)), (
        int(match.group(2)) if match.group(2) is not None else None
    )


def _display_filename(value: str) -> str:
    if not isinstance(value, str) or not _SAFE_FILENAME.fullmatch(value):
        raise ProtocolRegistrationError(
            "Protocol filename must be one plain PDF filename."
        )
    return value


def protocol_with_display_step_labels(
    protocol: domain.ExperimentProtocol,
) -> domain.ExperimentProtocol:
    """The Protocol with its steps numbered in order where the source prints none.

    An empty source_label means the source prints no step numbers (human
    decision 2026-10-05, lane P3). The stored analysis keeps it empty -- that
    is what the source says -- and the screen and the run number the steps
    1, 2, 3 ... in their order instead. A labelled Protocol is returned as is.

    Lane AN3, decision B (human decision 2026-10-16): a long document joined
    from pieces may hold both kinds. Its printed numbers are kept, and each
    unnumbered step is shown by its place among all the steps (the 18th step
    is "18") -- a number the screen shows, never one the source printed.
    """

    steps = [step for section in protocol.sections for step in section.steps]
    if not steps or all(step.source_label.strip() for step in steps):
        return protocol
    ordinal = iter(range(1, len(steps) + 1))

    def shown(step: domain.ProtocolSourceStep, place: int) -> domain.ProtocolSourceStep:
        return step if step.source_label.strip() else replace(step, source_label=str(place))

    return replace(
        protocol,
        sections=tuple(
            replace(
                section,
                steps=tuple(shown(step, next(ordinal)) for step in section.steps),
            )
            for section in protocol.sections
        ),
    )


def _review_sections(protocol: domain.ExperimentProtocol) -> object:
    shown = protocol_with_display_step_labels(protocol)
    sections = _review_value(shown.sections)
    if shown is not protocol:
        printed = {step.step_id for section in protocol.sections for step in section.steps
                   if step.source_label.strip()}
        for section in sections:
            for step in section["steps"]:
                if step["step_id"] not in printed:
                    step["source_label_printed"] = False
    return sections


def _review_value(value: object) -> object:
    """Convert selected Protocol records into a stable, JSON-safe review view."""

    if isinstance(value, Enum):
        return value.value
    if is_dataclass(value):
        return {
            field.name: _review_value(getattr(value, field.name))
            for field in fields(value)
        }
    if isinstance(value, (tuple, list)):
        return [_review_value(item) for item in value]
    if isinstance(value, (set, frozenset)):
        return sorted(_review_value(item) for item in value)
    if isinstance(value, dict):
        return {
            str(key): _review_value(item) for key, item in value.items()
        }
    return value


@dataclass(frozen=True)
class _SourceMetadata:
    """Source facts recorded once in the existing append-only catalog ledger."""

    title: str | None
    analysis_status: str
    ocr_needed: bool

    @classmethod
    def from_extraction(cls, extraction: ProtocolPdfExtraction) -> _SourceMetadata:
        return cls(extraction.metadata.title, _analysis_state(extraction), _ocr_needed(extraction))


def _ocr_needed(extraction: ProtocolPdfExtraction | _SourceMetadata) -> bool:
    """Whether any page of this source has to come from OCR.

    Wider than ``_analysis_state(...) == "ocr_required"``, which is the whole
    document having too little text to analyse at all. A readable document
    with one page the text layer cannot supply is analysed as it stands, and
    still needs this route for that page: until the page's OCR text has been
    accepted, readiness keeps the Protocol out of execution
    (``source_page_requires_ocr``). The OCR provider reads only the marked
    pages and keeps every other page's text layer.
    """

    if isinstance(extraction, _SourceMetadata):
        return extraction.ocr_needed
    return (
        _analysis_state(extraction) == "ocr_required"
        or bool(extraction.ocr_required_page_numbers)
    )


def _analysis_state(extraction: ProtocolPdfExtraction) -> str:
    extracted_chars = sum(len(page.text.strip()) for page in extraction.pages)
    if extraction.non_empty_page_count == 0 or extracted_chars < 32:
        return "ocr_required"
    if analysis_pieces(extraction):
        # Read in pieces (lane AN, decision 2): "대형 문서 분석 요청 대기"
        # until the analysis starts, as the screen says it.
        return "chunked_analysis_required"
    try:
        prepare_protocol_analysis_request(
            extraction, max_input_bytes=MAX_SINGLE_PASS_INPUT_BYTES
        )
    except ProtocolAnalysisInputTooLargeError:
        return "chunked_analysis_required"
    return "structured_analysis_ready"


#: An analysis refused under lane EV2's 60-page bound, of a document the
#: pieces read now (lane AN, decision 2): it may be asked again.
_OLD_PAGE_BOUND_KO = (
    "이 원문은 예전 한도(한 번에 60쪽)에 걸려 분석하지 않았어요. 이제 긴 문서는 조각으로 "
    "나눠 분석합니다."
)
_OLD_PAGE_BOUND_ACTION = (
    "'분석 다시 시도'를 누르면 조각으로 나눠 새로 분석합니다(조각 수만큼 분석 모델 호출 "
    "비용이 듭니다). 원문은 바뀌지 않습니다."
)


def _refused_under_the_old_page_bound(failure_code: str | None, page_count: int) -> bool:
    return (
        failure_code == "protocol_analysis_too_many_pages"
        and page_count <= MAX_PIECED_ANALYSIS_PAGES
    )


def _failed_pieces_ko(events: Sequence[Any]) -> str | None:
    """"41–59쪽, 97–115쪽": the pieces the latest failed analysis names (lane AN, 2)."""

    failed = next(
        (
            event.payload for event in reversed(events)
            if event.event_type == _ANALYSIS_FAILED_EVENT and isinstance(event.payload, dict)
        ),
        None,
    )
    pieces = (failed or {}).get("failed_pieces")
    if not isinstance(pieces, list):
        return None
    pages = [
        f"{piece['first_page']}–{piece['last_page']}쪽" for piece in pieces
        if isinstance(piece, dict) and isinstance(piece.get("first_page"), int)
        and isinstance(piece.get("last_page"), int)
    ]
    return ", ".join(pages) or None


def _latest_failure_payload(events: Sequence[Any]) -> dict[str, object]:
    return next(
        (event.payload for event in reversed(events)
         if event.event_type == _ANALYSIS_FAILED_EVENT and isinstance(event.payload, dict)),
        {},
    )


def _unlabelled_pieces_ko(events: Sequence[Any]) -> str | None:
    """"15–29쪽·141–155쪽": pieces that read their steps without numbers while
    others numbered them (lane RV2), from the latest failure."""

    conflict = _latest_failure_payload(events).get("piece_conflict")
    pieces = conflict.get("unlabelled_pieces") if isinstance(conflict, dict) else None
    if not isinstance(pieces, list):
        return None
    return "·".join(
        f"{piece['first_page']}–{piece['last_page']}쪽" for piece in pieces
        if isinstance(piece, dict) and isinstance(piece.get("first_page"), int)
    ) or None


def _blocked_provider(events: Sequence[Any]) -> str | None:
    blocked = _latest_failure_payload(events).get("blocked")
    return str(blocked.get("provider")) if isinstance(blocked, dict) else None


def _blocked_pieces_ko(events: Sequence[Any]) -> str | None:
    """The failed pieces the safety filter blocked, said as such (lane RV2, f)."""

    pieces = _latest_failure_payload(events).get("failed_pieces")
    words = [
        blocked_pages_ko(piece["first_page"], piece["last_page"])
        for piece in pieces or () if isinstance(piece, dict)
        and piece.get("failure_code") == BLOCKED_FAILURE_CODE
        and isinstance(piece.get("first_page"), int)
    ] if isinstance(pieces, list) else []
    return " · ".join(words) or None


def _piece_failure_code(exc: BaseException) -> str:
    if safety_block_provider(exc) is not None:
        # Lane RV2, decision f: the provider's safety filter, by name.
        return BLOCKED_FAILURE_CODE
    code = getattr(exc, "code", None)
    if isinstance(code, str) and re.fullmatch(r"[a-z][a-z0-9_]{0,63}", code):
        return code
    return "protocol_analysis_error"


def _safe_evidence_failure(
    exc: BaseException,
    *,
    source_revision: str,
    source_hash: str,
) -> dict[str, object] | None:
    if not isinstance(exc, ProtocolAnalysisEvidenceError):
        return None
    del source_revision, source_hash
    return exc.diagnostic.privacy_safe_dict()


_ANALYSIS_RECORDS_LOCK = threading.Lock()
_ANALYSIS_RECORDS: OrderedDict[tuple[str, int, int, str], tuple[AnalysisRevisionRecord, int]] = OrderedDict()
_ANALYSIS_RECORD_CACHE_ENTRIES = 128
_ANALYSIS_RECORD_CACHE_BYTES = 16 * 1024 * 1024
_analysis_record_cache_bytes = 0


class ProtocolCatalog:
    """Catalog facade over immutable source and analysis records.

    The execution rule (decision of 2026-10-08, lane DI): a revision may run
    when its latest analysis passed -- which includes the source-evidence
    check every analysis goes through -- and none of its readiness reasons
    is an execution blocker (``experiment_protocol.execution_blocking_reasons``:
    nothing runnable, or a page that could not be read). The one human
    confirmation is the experimenter pressing start, recorded by the server
    on the session. There is no approval, no reviewer finding and no
    development activation here any more; the rows those paths wrote stay in
    the ledger and are no longer read.
    """

    def __init__(
        self,
        store: ProtocolStore,
        *,
        on_analysis_ready: (
            Callable[["ProtocolCatalog", str], None] | None
        ) = None,
    ) -> None:
        self.store = store
        #: Told the protocol id as soon as an analysis has passed -- the
        #: moment its sentences start being translated (lane PX, human
        #: decision 1 of 2026-10-06), so the Korean is ready by the time the
        #: experimenter presses start. It adds no authority and its failure
        #: changes nothing the catalog recorded.
        self.on_analysis_ready = on_analysis_ready

    def _analysis_ready(self, entry: ProtocolCatalogEntry) -> None:
        if self.on_analysis_ready is None or entry.analysis_status != "review_required":
            return
        try:
            self.on_analysis_ready(self, entry.protocol_id)
        except Exception as exc:  # noqa: BLE001 - see on_analysis_ready
            logging.getLogger(__name__).warning(
                "protocol.analysis_ready.hook_failed protocol_id=%s error=%s",
                entry.protocol_id, type(exc).__name__,
            )

    def _latest_protocol_revision(self, protocol_id: str) -> ProtocolRevisionRecord:
        if not _STABLE_PROTOCOL_ID.fullmatch(protocol_id):
            raise ProtocolCatalogNotFoundError("Protocol identifier is unknown.")
        revisions = self.store.list_protocol_revisions(protocol_id)
        if not revisions:
            raise ProtocolCatalogNotFoundError("Protocol identifier is unknown.")
        return revisions[-1]

    def _latest_analysis(
        self, revision: ProtocolRevisionRecord
    ) -> AnalysisRevisionRecord | None:
        """The latest analysis the analyser made; a registered review version
        (lane RV, decision c) is a person's, read by ``_review_versions``."""

        # Find the immutable reference first; listing does not need to decode
        # every historical analysis just to choose the most recent one.
        row = self.store._connection.execute(
            """SELECT analysis_revision_number FROM analysis_revisions
            WHERE experiment_id=? AND protocol_revision_number=?
              AND substr(analysis_id,1,?) != ?
            ORDER BY analysis_revision_number DESC LIMIT 1""",
            (revision.experiment_id, revision.revision_number,
             len(REVIEW_VERSION_ANALYSIS_PREFIX), REVIEW_VERSION_ANALYSIS_PREFIX),
        ).fetchone()
        return self._analysis_record(revision, row["analysis_revision_number"]) if row is not None else None

    def _analysis_record(
        self, revision: ProtocolRevisionRecord, analysis_revision_number: int,
    ) -> AnalysisRevisionRecord:
        """Reuse only bytes the store already decoded and verified.

        The source, revision, complete envelope and actual payload bytes form
        the key. Reading them again detects even out-of-band corruption after
        a cache hit. A miss always uses the store's sole decoder/validator;
        this cache owns no execution verdict. The savepoint keeps those two
        reads on one SQLite snapshot without committing a caller's work.
        """
        global _analysis_record_cache_bytes
        connection = self.store._connection
        connection.execute("SAVEPOINT catalog_analysis_read")
        try:
            row = connection.execute(
                """SELECT ar.*,ap.payload_json FROM analysis_revisions ar
                JOIN analysis_payloads ap ON ap.payload_sha256=ar.payload_sha256
                WHERE ar.experiment_id=? AND ar.protocol_revision_number=?
                  AND ar.analysis_revision_number=?""",
                (revision.experiment_id, revision.revision_number, analysis_revision_number),
            ).fetchone()
            if row is None:
                return self.store.get_analysis_revision(
                    revision.experiment_id, revision.revision_number, analysis_revision_number)
            content = json.dumps(tuple(row), ensure_ascii=False, separators=(",", ":")).encode("utf-8")
            content_hash = hashlib.sha256(content).hexdigest()
            key = (revision.pdf_checksum, revision.revision_number, analysis_revision_number, content_hash)
            with _ANALYSIS_RECORDS_LOCK:
                cached = _ANALYSIS_RECORDS.get(key)
                if cached is not None:
                    _ANALYSIS_RECORDS.move_to_end(key)
                    return cached[0]
            record = self.store.get_analysis_revision(
                revision.experiment_id, revision.revision_number, analysis_revision_number)
            # Retired coverage payloads can contain mutable dictionaries.
            # Keep their original fresh-read behavior, rather than sharing
            # one caller's mutable legacy object with another request.
            immutable = not record.page_coverage
            try:
                hash(record.protocol)
                hash(record.readiness)
            except TypeError:
                immutable = False
            # Bound retained serialized input as well as the number of
            # records. Oversized records are still verified and returned.
            if immutable and len(content) <= _ANALYSIS_RECORD_CACHE_BYTES:
                with _ANALYSIS_RECORDS_LOCK:
                    previous = _ANALYSIS_RECORDS.pop(key, None)
                    if previous is not None:
                        _analysis_record_cache_bytes -= previous[1]
                    _ANALYSIS_RECORDS[key] = (record, len(content))
                    _analysis_record_cache_bytes += len(content)
                    while (
                        len(_ANALYSIS_RECORDS) > _ANALYSIS_RECORD_CACHE_ENTRIES
                        or _analysis_record_cache_bytes > _ANALYSIS_RECORD_CACHE_BYTES
                    ):
                        _, (_, size) = _ANALYSIS_RECORDS.popitem(last=False)
                        _analysis_record_cache_bytes -= size
            return record
        finally:
            connection.execute("RELEASE catalog_analysis_read")

    def _review_version_events(self, revision: ProtocolRevisionRecord) -> tuple[Any, ...]:
        return tuple(
            event for event in self.store.list_events(revision.experiment_id)
            if event.event_type == _REVIEW_VERSION_EVENT
            and event.protocol_revision_number == revision.revision_number
            and event.analysis_revision_number is not None
        )

    def _review_versions(
        self, revision: ProtocolRevisionRecord,
    ) -> list[tuple[Any, AnalysisRevisionRecord]]:
        """Every registered review version of the revision, v1 first (lane RV, c)."""

        versions = []
        for event in self._review_version_events(revision):
            versions.append((event, self._analysis_record(revision, event.analysis_revision_number)))
        return versions

    def _reanalysis_requested(self, revision: ProtocolRevisionRecord) -> bool:
        """Whether the run on screen was asked for over a passed analysis."""

        request = next(
            (
                event for event in reversed(self.store.list_events(revision.experiment_id))
                if event.protocol_revision_number == revision.revision_number
                and event.event_type == _ANALYSIS_REQUESTED_EVENT
            ),
            None,
        )
        return (
            request is not None and isinstance(request.payload, dict)
            and request.payload.get("reanalysis") is True
        )

    @staticmethod
    def _reanalysis_status(
        events: Sequence[Any], analysis: AnalysisRevisionRecord | None,
    ) -> dict[str, object] | None:
        """The run asked for over the latest passed analysis, or None.

        Lane RN, decision 2: the lifecycle events after the pass (or after
        the development fixture's record) are that run's. Its failure is its
        own, never the executable analysis's, so the review says it apart
        from ``analysis_failure``. None while no such run was asked for.
        """

        if analysis is None:
            return None
        last_pass = max(
            (
                index for index, event in enumerate(events)
                if event.event_type in {
                    _ANALYSIS_READY_EVENT, _SINGLE_REVIEW_REQUIRED_EVENT,
                    _DEVELOPMENT_FIXTURE_EVENT,
                }
            ),
            default=-1,
        )
        after = tuple(
            event for event in events[last_pass + 1:]
            if event.event_type in {
                _ANALYSIS_REQUESTED_EVENT, _ANALYSIS_STARTED_EVENT,
                _ANALYSIS_FAILED_EVENT, _ANALYSIS_RETRY_EVENT,
                _REVIEW_PENDING_EVENT,
            }
        )
        if not after:
            return None
        state = {
            _ANALYSIS_REQUESTED_EVENT: "analysis_pending",
            _ANALYSIS_STARTED_EVENT: "analyzing",
            _ANALYSIS_FAILED_EVENT: "analysis_failed",
            _ANALYSIS_RETRY_EVENT: "analyzing",
            # Lane RV, decision a: the run ended with items for a person;
            # the analysis that passed before still runs.
            _REVIEW_PENDING_EVENT: "review_pending",
        }[after[-1].event_type]
        failure: dict[str, object] | None = None
        if state == "analysis_failed":
            payload = after[-1].payload if isinstance(after[-1].payload, dict) else {}
            code = payload.get("failure_code")
            code = code if isinstance(code, str) and code else "analysis_failed"
            detail = payload.get("evidence_failure")
            failure = {
                "code": code,
                "detail": detail if isinstance(detail, dict) else None,
                "retryable": code not in _NOT_RETRYABLE_FAILURE_CODES,
                "action": _REANALYSIS_RECOVERY_ACTION,
            }
            repair = _response_repair_of(payload)
            if repair is not None:
                # Lane EV3, decision 1.
                failure["response_repaired"] = repair
        status: dict[str, object] = {
            "state": state,
            "requested_at": next(
                (event.recorded_at for event in reversed(after)
                 if event.event_type == _ANALYSIS_REQUESTED_EVENT),
                None,
            ),
            "failure": failure,
            "automatic_retry": _automatic_retry_status(after),
        }
        if state == "review_pending":
            status["review_pending"] = _review_pending_summary(after[-1], events)
        return status

    def _latest_ocr_events(
        self, revision: ProtocolRevisionRecord
    ) -> tuple[object, ...]:
        events = tuple(
            event
            for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
            and event.event_type
            in {
                _OCR_REQUESTED_EVENT,
                _OCR_COMPLETED_EVENT,
                _OCR_FAILED_EVENT,
                _OCR_REVIEWED_EVENT,
            }
        )
        requested = tuple(
            event for event in events if event.event_type == _OCR_REQUESTED_EVENT
        )
        if not requested or not isinstance(requested[-1].payload, dict):
            return ()
        ocr_id = requested[-1].payload.get("ocr_id")
        if not isinstance(ocr_id, str):
            return ()
        return tuple(
            event
            for event in events
            if isinstance(event.payload, dict)
            and event.payload.get("ocr_id") == ocr_id
        )

    def _ocr_projection(
        self,
        revision: ProtocolRevisionRecord,
        extraction: ProtocolPdfExtraction | _SourceMetadata,
        *,
        include_text: bool,
    ) -> dict[str, object]:
        if not _ocr_needed(extraction):
            return {
                "state": "not_required",
                "review_required": False,
                "accepted_for_analysis": False,
                "executable": False,
            }
        events = self._latest_ocr_events(revision)
        if not events:
            return {
                "state": "ocr_required",
                "review_required": False,
                "accepted_for_analysis": False,
                "executable": False,
            }
        request = events[0].payload
        completed = next(
            (
                event.payload
                for event in reversed(events)
                if event.event_type == _OCR_COMPLETED_EVENT
                and isinstance(event.payload, dict)
            ),
            None,
        )
        failed = next(
            (
                event.payload
                for event in reversed(events)
                if event.event_type == _OCR_FAILED_EVENT
                and isinstance(event.payload, dict)
            ),
            None,
        )
        reviewed = next(
            (
                event.payload
                for event in reversed(events)
                if event.event_type == _OCR_REVIEWED_EVENT
                and isinstance(event.payload, dict)
            ),
            None,
        )
        if reviewed is not None:
            accepted = reviewed.get("decision") == "accepted"
            state = "accepted_for_analysis" if accepted else "rejected"
        elif completed is not None:
            accepted = False
            state = "review_required"
        elif failed is not None:
            accepted = False
            state = "failed"
        else:
            accepted = False
            state = "in_progress"
        projection: dict[str, object] = {
            "state": state,
            "ocr_id": request.get("ocr_id"),
            "source_sha256": revision.pdf_checksum,
            "review_required": state == "review_required",
            "accepted_for_analysis": accepted,
            "executable": False,
            "failure_code": (
                failed.get("failure_code") if failed is not None else None
            ),
            "review": reviewed,
        }
        if completed is not None:
            pages = completed.get("pages")
            projection.update(
                {
                    "provider": completed.get("provider"),
                    "provider_version": completed.get("provider_version"),
                    "languages": completed.get("languages", []),
                    "warnings": completed.get("warnings", []),
                    "page_count": completed.get("page_count"),
                    "engines": completed.get("engines"),
                    "pages_read_by_both": completed.get("pages_read_by_both"),
                    "pages": (
                        pages
                        if include_text
                        else [
                            {
                                key: page.get(key)
                                for key in (
                                    "source_page_number",
                                    "confidence",
                                    "text_sha256",
                                )
                            }
                            for page in pages
                            if isinstance(page, dict)
                        ]
                        if isinstance(pages, list)
                        else []
                    ),
                }
            )
        return projection

    def ocr_status(
        self, protocol_id: str, *, include_text: bool = True
    ) -> dict[str, object]:
        revision = self._latest_protocol_revision(protocol_id)
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        source = self.store.file_store.object_path(
            revision.pdf_checksum, expected_size=pdf_object.byte_size
        )
        extraction = extract_protocol_pdf(source)
        return {
            "protocol_id": protocol_id,
            "revision_id": _revision_id(revision.revision_number),
            **self._ocr_projection(
                revision, extraction, include_text=include_text
            ),
        }

    def _extraction_for_analysis(
        self,
        revision: ProtocolRevisionRecord,
        extraction: ProtocolPdfExtraction,
    ) -> ProtocolPdfExtraction:
        projection = self._ocr_projection(
            revision, extraction, include_text=True
        )
        if projection.get("accepted_for_analysis") is not True:
            return extraction
        pages = projection.get("pages")
        if not isinstance(pages, list) or len(pages) != extraction.page_count:
            raise ProtocolOcrReviewError("Accepted OCR page evidence is invalid.")
        reconstructed = []
        result_provider = projection.get("provider")
        for expected_number, page in enumerate(pages, start=1):
            if (
                not isinstance(page, dict)
                or page.get("source_page_number") != expected_number
                or not isinstance(page.get("text"), str)
                or hashlib.sha256(page["text"].encode("utf-8")).hexdigest()
                != page.get("text_sha256")
            ):
                raise ProtocolOcrReviewError(
                    "Accepted OCR page evidence failed integrity validation."
                )
            # A page the provider kept from the text layer is not OCR text;
            # a provider that names one engine for every page names it on
            # the result only.
            ocr_derived = (
                (page.get("provider") or result_provider) != OCR_TEXT_LAYER_PROVIDER
            )
            original = extraction.pages[expected_number - 1]
            # The text layer's own page keeps its geometry: the footer band
            # and the text blocks the page-boundary rule (lane PA, decision
            # 3) reads. Measured 2026-10-06 (lane PX): with them dropped for
            # every page after an OCR acceptance, ANKOM's step 21 -- "21
            # Flush procedure:" ending page 19, its sentence opening page 20
            # -- was refused, while the same response passed on the raw
            # extraction. OCR text has no geometry, so an OCR page keeps none.
            keeps_text_layer = not ocr_derived and page["text"] == original.text
            reconstructed.append(
                ProtocolPdfPage(
                    source_page_number=expected_number,
                    text=page["text"],
                    text_empty=not page["text"].strip(),
                    warning=(
                        original.warning if keeps_text_layer
                        else "Text was produced by OCR and accepted for structured review."
                    ),
                    bottom_band_offset=(
                        original.bottom_band_offset if keeps_text_layer else None
                    ),
                    blocks=original.blocks if keeps_text_layer else (),
                    ocr_derived=ocr_derived,
                )
            )
        return replace(
            extraction,
            pages=tuple(reconstructed),
            warnings=tuple(
                dict.fromkeys(
                    (
                        *extraction.warnings,
                        "OCR-derived text is read as source text; it is not an analysed protocol.",
                    )
                )
            ),
        )

    def run_ocr(
        self,
        protocol_id: str,
        provider: ProtocolOcrProvider,
        *,
        ocr_id: str,
        accepted_automatically: bool = True,
    ) -> dict[str, object]:
        """Read the pages without a usable text layer with the OCR provider.

        With ``accepted_automatically`` (lane PX, human decision 3 of
        2026-10-06: OCR runs at upload for the pages whose text layer is
        missing or unreadable) the validated result is accepted for analysis
        at once, under the authority ``automatic_upload_ocr`` written into
        the ledger, and the one confirmation a person gives before execution
        (the experimenter pressing start) covers it; there is no separate
        OCR approval step. Each OCR page keeps its provider and its
        numeric-review mark, so the screen still shows which pages came from
        OCR and where the two engines read different numbers.
        """

        if not _STABLE_PROTOCOL_ID.fullmatch(ocr_id):
            raise ProtocolOcrReviewError("OCR request identity is invalid.")
        revision = self._latest_protocol_revision(protocol_id)
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        source = self.store.file_store.object_path(
            revision.pdf_checksum, expected_size=pdf_object.byte_size
        )
        extraction = extract_protocol_pdf(source)
        if not _ocr_needed(extraction):
            raise ProtocolOcrReviewError("Protocol PDF does not require OCR.")
        current = self._ocr_projection(
            revision, extraction, include_text=False
        )
        if current["state"] in {
            "in_progress",
            "review_required",
            "accepted_for_analysis",
        }:
            return self.ocr_status(protocol_id)
        self.store.append_event(
            f"ocr-requested-{ocr_id}",
            protocol_id,
            revision.revision_number,
            _OCR_REQUESTED_EVENT,
            {
                "ocr_id": ocr_id,
                "status": "in_progress",
                "source_sha256": revision.pdf_checksum,
                "page_count": extraction.page_count,
            },
        )
        try:
            result = provider.recognize(
                source,
                source_sha256=revision.pdf_checksum,
                page_count=extraction.page_count,
            )
            validated = validate_ocr_result(
                result,
                expected_sha256=revision.pdf_checksum,
                expected_page_count=extraction.page_count,
            )
            payload = {"ocr_id": ocr_id, **ocr_result_payload(validated)}
            # Lane RV2, decision b: which engines read the document and which
            # pages both of them read, so an OCR value both read alike can be
            # settled without a person. Absent from results stored before.
            engines = ocr_engine_names(provider)
            payload["engines"] = list(engines)
            payload["pages_read_by_both"] = pages_read_by_both(validated, engines)
            self.store.append_event(
                f"ocr-completed-{ocr_id}",
                protocol_id,
                revision.revision_number,
                _OCR_COMPLETED_EVENT,
                payload,
            )
            if accepted_automatically:
                ocr_pages = [
                    page.source_page_number for page in validated.pages
                    if page.provider != OCR_TEXT_LAYER_PROVIDER
                    and (page.provider is not None
                         or validated.provider != OCR_TEXT_LAYER_PROVIDER)
                ]
                self.store.append_event(
                    f"ocr-reviewed-{ocr_id}-accepted",
                    protocol_id,
                    revision.revision_number,
                    _OCR_REVIEWED_EVENT,
                    {
                        "ocr_id": ocr_id,
                        "decision": "accepted",
                        "authority": AUTOMATIC_OCR_AUTHORITY,
                        "comment": (
                            "Upload-time OCR of the pages without a usable text "
                            "layer; covered by the one confirmation a person "
                            "gives before execution."
                        ),
                        "executable": False,
                        "ocr_page_numbers": ocr_pages,
                        "numeric_review_page_numbers": [
                            page.source_page_number for page in validated.pages
                            if page.numeric_review_required
                        ],
                    },
                )
        except Exception as exc:
            failure_code = getattr(exc, "code", "protocol_ocr_failed")
            if not isinstance(failure_code, str) or not re.fullmatch(
                r"[a-z][a-z0-9_]{0,63}", failure_code
            ):
                failure_code = "protocol_ocr_failed"
            self.store.append_event(
                f"ocr-failed-{ocr_id}",
                protocol_id,
                revision.revision_number,
                _OCR_FAILED_EVENT,
                {
                    "ocr_id": ocr_id,
                    "status": "failed",
                    "failure_code": failure_code,
                },
            )
            raise
        return self.ocr_status(protocol_id)

    def _latest_chunk_events(
        self,
        revision: ProtocolRevisionRecord,
    ) -> tuple[object, ...]:
        events = tuple(
            event
            for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
        )
        planned = tuple(
            event for event in events if event.event_type == _CHUNK_PLAN_EVENT
        )
        if not planned:
            return ()
        if any(
            event.event_type == _ANALYSIS_REQUESTED_EVENT
            and event.sequence_id > planned[-1].sequence_id
            for event in events
        ):
            # A later request opened a run of its own (lane AN, decision 2):
            # the planned run is no longer the one on screen.
            return ()
        payload = planned[-1].payload
        if not isinstance(payload, dict):
            return ()
        run_id = payload.get("analysis_run_id")
        if not isinstance(run_id, str):
            return ()
        return tuple(
            event
            for event in events
            if isinstance(event.payload, dict)
            and event.payload.get("analysis_run_id") == run_id
        )

    def analysis_run_status(
        self,
        protocol_id: str,
    ) -> ProtocolAnalysisRunStatus:
        revision = self._latest_protocol_revision(protocol_id)
        candidate_revision_id = _revision_id(revision.revision_number)
        events = self._latest_chunk_events(revision)
        entry = self._entry_for_revision(revision)
        reanalysis = self._reanalysis_requested(revision)
        if not events:
            lifecycle_events = tuple(
                event for event in self.store.list_events(revision.experiment_id)
                if event.protocol_revision_number == revision.revision_number
                and event.event_type in {
                    _ANALYSIS_REQUESTED_EVENT,
                    _ANALYSIS_STARTED_EVENT,
                    _ANALYSIS_FAILED_EVENT,
                    _ANALYSIS_RETRY_EVENT,
                    _ANALYSIS_READY_EVENT,
                    _SINGLE_REVIEW_REQUIRED_EVENT,
                    _REVIEW_PENDING_EVENT,
                }
            )
            current_events = _failures_not_superseded(lifecycle_events)
            latest_failure = next(
                (
                    event.payload.get("failure_code")
                    for event in reversed(current_events)
                    if event.event_type == _ANALYSIS_FAILED_EVENT
                    and isinstance(event.payload, dict)
                    and isinstance(event.payload.get("failure_code"), str)
                ),
                None,
            )
            latest_failure_detail = next(
                (
                    event.payload.get("evidence_failure")
                    for event in reversed(current_events)
                    if event.event_type == _ANALYSIS_FAILED_EVENT
                    and isinstance(event.payload, dict)
                    and isinstance(
                        event.payload.get("evidence_failure"), dict
                    )
                ),
                None,
            )
            analysis_run_id = next(
                (
                    event.payload.get("analysis_id")
                    for event in reversed(lifecycle_events)
                    if isinstance(event.payload, dict)
                    and isinstance(event.payload.get("analysis_id"), str)
                ),
                None,
            )
            requested_at = next(
                (
                    event.recorded_at
                    for event in reversed(lifecycle_events)
                    if event.event_type == _ANALYSIS_REQUESTED_EVENT
                ),
                None,
            )
            state = entry.analysis_status
            if lifecycle_events:
                state = {
                    _ANALYSIS_REQUESTED_EVENT: "analysis_pending",
                    _ANALYSIS_STARTED_EVENT: "analyzing",
                    _ANALYSIS_FAILED_EVENT: "analysis_failed",
                    _ANALYSIS_RETRY_EVENT: "analyzing",
                    _ANALYSIS_READY_EVENT: "analysis_ready",
                    _SINGLE_REVIEW_REQUIRED_EVENT: "review_required",
                    _REVIEW_PENDING_EVENT: "review_pending",
                }[lifecycle_events[-1].event_type]
            return ProtocolAnalysisRunStatus(
                protocol_id=protocol_id,
                candidate_revision_id=candidate_revision_id,
                analysis_run_id=analysis_run_id,
                state=state,
                total_chunks=0,
                completed_chunks=0,
                failed_chunks=0,
                pending_chunks=0,
                chunks=(),
                failure_code=latest_failure,
                failure_detail=latest_failure_detail,
                lifecycle_state=entry.lifecycle_state,
                requested_at=requested_at,
                automatic_retry=_automatic_retry_status(lifecycle_events),
                reanalysis=reanalysis,
                available_for_execution=entry.available_for_execution,
            )
        plan_event = next(
            event for event in events if event.event_type == _CHUNK_PLAN_EVENT
        )
        plan_payload = plan_event.payload
        raw_chunks = plan_payload.get("chunks", [])
        chunks: list[dict[str, object]] = []
        statuses: dict[
            str,
            tuple[str, str | None, dict[str, object] | None],
        ] = {}
        state = "chunk_planned"
        failure_code = None
        failure_detail = None
        merge_status = None
        cancelled = False
        for event in events:
            payload = event.payload
            if event.event_type == _CHUNK_STARTED_EVENT:
                state = "chunk_analysis_in_progress"
                statuses[str(payload.get("chunk_id"))] = (
                    "in_progress",
                    None,
                    None,
                )
            elif event.event_type == _CHUNK_COMPLETED_EVENT:
                statuses[str(payload.get("chunk_id"))] = (
                    "completed",
                    None,
                    None,
                )
            elif event.event_type == _CHUNK_FAILED_EVENT:
                code = payload.get("failure_code")
                failure_code = code if isinstance(code, str) else "chunk_analysis_failed"
                raw_detail = payload.get("evidence_failure")
                detail = raw_detail if isinstance(raw_detail, dict) else None
                if detail is not None:
                    failure_detail = detail
                statuses[str(payload.get("chunk_id"))] = (
                    "failed",
                    failure_code,
                    detail,
                )
                state = "chunk_analysis_failed"
            elif event.event_type == _MERGE_STARTED_EVENT:
                state = "merge_in_progress"
                merge_status = "in_progress"
            elif event.event_type == _MERGE_CONFLICT_EVENT:
                state = "merge_conflict"
                merge_status = "conflict"
                code = payload.get("failure_code")
                failure_code = code if isinstance(code, str) else "merge_conflict"
            elif event.event_type == _REVIEW_REQUIRED_EVENT:
                state = "review_required"
                merge_status = "complete"
            elif event.event_type == _RUN_CANCELLED_EVENT:
                state = "chunk_analysis_cancelled"
                failure_code = "analysis_cancelled"
                cancelled = True
        # Cancellation is a terminal fence. A worker may already be between
        # scheduling members of one batch when another connection appends the
        # cancellation event, so a later benign "started" event must never make
        # the run appear resumable or in progress again.
        if cancelled:
            state = "chunk_analysis_cancelled"
            failure_code = "analysis_cancelled"
        for raw in raw_chunks if isinstance(raw_chunks, list) else []:
            if not isinstance(raw, dict) or not isinstance(raw.get("chunk_id"), str):
                continue
            chunk_id = raw["chunk_id"]
            status, code, detail = statuses.get(
                chunk_id,
                ("pending", None, None),
            )
            chunk_status: dict[str, object] = {
                "chunk_id": chunk_id,
                "ordinal": raw.get("ordinal"),
                "source_page_start": raw.get("source_page_start"),
                "source_page_end": raw.get("source_page_end"),
                "source_page_refs": raw.get("source_page_refs", []),
                "status": status,
                "failure_code": code,
            }
            if detail is not None:
                chunk_status["failure_detail"] = detail
            chunks.append(chunk_status)
        completed = sum(chunk["status"] == "completed" for chunk in chunks)
        failed = sum(chunk["status"] == "failed" for chunk in chunks)
        end = _piece_run_end(events, tuple(
            event for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
        ))
        if end is not None:
            # A piece run (lane AN, decision 2) ends in the analysis' own
            # ready or failed event, not in a chunk event.
            state, end_code = end
            failure_code = end_code or failure_code
            merge_status = "complete" if state == "review_required" else merge_status
            if state == "review_pending":
                # Lane RV, decision a: the failed pieces are items for a
                # person, not the run's failure.
                failure_code, failure_detail, merge_status = None, None, "complete"
        return ProtocolAnalysisRunStatus(
            protocol_id=protocol_id,
            candidate_revision_id=candidate_revision_id,
            analysis_run_id=str(plan_payload.get("analysis_run_id")),
            state=state,
            total_chunks=len(chunks),
            completed_chunks=completed,
            failed_chunks=failed,
            pending_chunks=len(chunks) - completed - failed,
            chunks=tuple(chunks),
            failure_code=failure_code,
            failure_detail=failure_detail,
            merge_status=merge_status,
            restart_behavior=(
                "terminal_review_required"
                if state in {"review_required", "review_pending"}
                else "explicit_new_run_required_after_interruption"
                if state in {"chunk_planned", "chunk_analysis_in_progress", "merge_in_progress"}
                else "explicit_analysis_request_only"
            ),
            lifecycle_state=(
                state
                if state in {"review_required", "review_pending"}
                else "blocked"
                if state in {
                    "merge_conflict",
                    "chunk_analysis_failed",
                    "chunk_analysis_cancelled",
                    "analysis_failed",
                }
                else "analyzing"
            ),
            # The plan is the run's first record.
            requested_at=plan_event.recorded_at,
            reanalysis=reanalysis,
            available_for_execution=entry.available_for_execution,
        )

    def _entry_for_revision(
        self, revision: ProtocolRevisionRecord
    ) -> ProtocolCatalogEntry:
        analysis = self._latest_analysis(revision)
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        source_metadata = self._source_metadata(revision, analysis)
        analysis_status = "review_required" if analysis else source_metadata.analysis_status
        lifecycle_state = "review_required" if analysis else "uploaded"
        if analysis is None:
            if analysis_status == "ocr_required":
                ocr_state = self._ocr_projection(
                    revision, source_metadata, include_text=False
                )["state"]
                analysis_status, lifecycle_state = {
                    "ocr_required": ("ocr_required", "blocked"),
                    "in_progress": ("ocr_in_progress", "analyzing"),
                    "review_required": ("ocr_review_required", "review_required"),
                    "accepted_for_analysis": (
                        "structured_analysis_ready",
                        "uploaded",
                    ),
                    "rejected": ("ocr_rejected", "blocked"),
                    "failed": ("ocr_failed", "blocked"),
                }.get(str(ocr_state), ("ocr_required", "blocked"))
            chunk_events = self._latest_chunk_events(revision)
            if chunk_events:
                event_states = {
                    _CHUNK_PLAN_EVENT: "chunk_planned",
                    _CHUNK_STARTED_EVENT: "chunk_analysis_in_progress",
                    _CHUNK_COMPLETED_EVENT: "chunk_analysis_in_progress",
                    _CHUNK_FAILED_EVENT: "chunk_analysis_failed",
                    _MERGE_STARTED_EVENT: "merge_in_progress",
                    _MERGE_CONFLICT_EVENT: "merge_conflict",
                    _REVIEW_REQUIRED_EVENT: "review_required",
                    _RUN_CANCELLED_EVENT: "chunk_analysis_cancelled",
                }
                analysis_status = event_states.get(
                    chunk_events[-1].event_type,
                    analysis_status,
                )
            failed = next(
                (
                    event
                    for event in reversed(_failures_not_superseded(tuple(
                        event for event in self.store.list_events(revision.experiment_id)
                        if event.protocol_revision_number == revision.revision_number
                    )))
                    if event.event_type == _ANALYSIS_FAILED_EVENT
                ),
                None,
            )
            if failed is not None and not chunk_events:
                analysis_status = "analysis_failed"
            elif chunk_events:
                end = _piece_run_end(chunk_events, tuple(
                    event for event in self.store.list_events(revision.experiment_id)
                    if event.protocol_revision_number == revision.revision_number
                ))
                if end is not None and end[0] != "review_required":
                    analysis_status = end[0]
            lifecycle_events = tuple(
                event
                for event in self.store.list_events(revision.experiment_id)
                if event.protocol_revision_number == revision.revision_number
                and event.event_type in {
                    _ANALYSIS_REQUESTED_EVENT,
                    _ANALYSIS_STARTED_EVENT,
                    _ANALYSIS_FAILED_EVENT,
                    _ANALYSIS_RETRY_EVENT,
                    _REVIEW_PENDING_EVENT,
                }
            )
            if lifecycle_events:
                lifecycle_state = {
                    _ANALYSIS_REQUESTED_EVENT: "analysis_pending",
                    _ANALYSIS_STARTED_EVENT: "analyzing",
                    _ANALYSIS_FAILED_EVENT: "blocked",
                    _ANALYSIS_RETRY_EVENT: "analyzing",
                    # Lane RV, decision a: waiting for a person's review;
                    # it never runs as it is.
                    _REVIEW_PENDING_EVENT: "review_pending",
                }[lifecycle_events[-1].event_type]
                if lifecycle_events[-1].event_type == _REVIEW_PENDING_EVENT:
                    analysis_status = "review_pending"
            elif analysis_status == "ocr_required":
                lifecycle_state = "blocked"
        development_only = any(
            event.event_type == _DEVELOPMENT_FIXTURE_EVENT
            and event.protocol_revision_number == revision.revision_number
            for event in self.store.list_events(revision.experiment_id)
        )
        readiness = (
            analysis.readiness_status if analysis else "analysis_required"
        )
        # The execution rule (decision of 2026-10-08): a passed analysis whose
        # readiness carries no execution blocker may run. A missing or failed
        # analysis stays out; so does one that found no step to run or a
        # page it could not read. Every other readiness reason is a notice
        # the experimenter sees before pressing start.
        blocker_codes: tuple[str, ...] = ()
        outdated_codes: tuple[str, ...] = ()
        if analysis is not None:
            blocker_codes = tuple(
                reason.code.value
                for reason in domain.execution_blocking_reasons(analysis.readiness)
            )
            lifecycle_state = "blocked" if blocker_codes else "ready"
            outdated_codes = _analysis_outdated_codes(analysis.protocol)
        available = analysis is not None and not blocker_codes
        # An analysis whose title was emptied (lane EV2, decision 2) is named
        # as one without an analysis is: by the file's own title or name.
        title = (
            analysis.protocol.metadata.title
            if analysis is not None and analysis.protocol.metadata.title
            else source_metadata.title or Path(revision.original_filename).stem
        )
        pending = self._review_pending_event(revision)
        revision_id = _revision_id(
            revision.revision_number,
            analysis.analysis_revision_number if analysis else None,
        )
        shown = analysis
        versions = self._review_version_events(revision)
        if versions:
            # Lane RV, decision c: the latest registered review version runs,
            # whatever the analyser made after it; it was registered only
            # without an execution blocker.
            event = versions[-1]
            shown = self._analysis_record(revision, event.analysis_revision_number)
            revision_id = str(event.payload["revision_id"])
            readiness = shown.readiness_status
            blocker_codes = ()
            available = True
            lifecycle_state = "ready"
            outdated_codes = ()
            if analysis_status == "review_pending" and pending is None:
                analysis_status = "review_registered"
            if shown.protocol.metadata.title:
                title = shown.protocol.metadata.title
        renamed = self._renamed_title(revision)
        if renamed is not None:
            # Lane RX1, decision 4: a person's own name wins over every read one.
            title = renamed
        return ProtocolCatalogEntry(
            protocol_id=revision.experiment_id,
            title=title,
            source_filename=revision.original_filename,
            source_sha256=revision.pdf_checksum,
            revision_id=revision_id,
            readiness_status=readiness,
            analysis_status=analysis_status,
            step_count=(
                sum(len(section.steps) for section in shown.protocol.sections)
                if shown
                else 0
            ),
            created_at=revision.created_at,
            available_for_execution=available,
            lifecycle_state=lifecycle_state,
            execution_blocker_codes=blocker_codes,
            development_only=development_only,
            analysis_outdated_codes=outdated_codes,
            analysis_outdated_ko=ANALYSIS_OUTDATED_KO if outdated_codes else None,
            review_pending=(
                _review_pending_summary(pending, self.store.list_events(revision.experiment_id))
                if pending is not None else None),
            review_versions=tuple(
                _review_version_summary(event, self.store.list_events(revision.experiment_id))
                for event in versions),
            source_disclosure=self._source_disclosure(revision),
        )

    def record_source_disclosure(
        self,
        protocol_id: str,
        disclosure: ProtocolSourceDisclosure,
    ) -> ProtocolSourceDisclosure:
        """Append one declared source scope and return its ledger timestamp.

        Repeating exactly the same declaration is idempotent. A changed
        declaration receives a different event id, so the earlier answer is
        retained and the latest one is projected.
        """

        if not isinstance(disclosure, ProtocolSourceDisclosure):
            raise ProtocolRegistrationError("Protocol source disclosure is invalid.")
        revision = self._latest_protocol_revision(protocol_id)
        payload = {
            "status": disclosure.status.value,
            "note": disclosure.note,
        }
        latest = self._source_disclosure(revision)
        if (
            latest is not None
            and latest.status is disclosure.status
            and latest.note == disclosure.note
        ):
            return latest
        event = self.store.append_event(
            f"source-disclosure-{secrets.token_hex(12)}",
            protocol_id,
            revision.revision_number,
            PROTOCOL_SOURCE_DISCLOSURE_EVENT,
            payload,
        )
        return replace(disclosure, recorded_at=event.recorded_at)

    def _source_disclosure(
        self, revision: ProtocolRevisionRecord
    ) -> ProtocolSourceDisclosure | None:
        latest: ProtocolSourceDisclosure | None = None
        for event in self.store.list_events(revision.experiment_id):
            if (
                event.event_type != PROTOCOL_SOURCE_DISCLOSURE_EVENT
                or event.protocol_revision_number != revision.revision_number
            ):
                continue
            payload = event.payload
            try:
                if not isinstance(payload, dict) or set(payload) != {"status", "note"}:
                    raise ValueError
                latest = ProtocolSourceDisclosure(
                    status=ProtocolSourceStatus(payload["status"]),
                    note=payload["note"],
                    recorded_at=event.recorded_at,
                )
            except (ProtocolRegistrationError, TypeError, ValueError) as exc:
                raise ProtocolCatalogUnavailableError(
                    "Stored protocol source disclosure is invalid."
                ) from exc
        return latest

    def _review_pending_event(self, revision: ProtocolRevisionRecord) -> Any | None:
        """The latest review-pending analysis newer than any passed one and
        not yet registered as a review version, or None.

        Lane RV, decision a. A later failed or running request leaves it as
        it is: it is still the latest analysis a person can review.
        """

        events = [
            event for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
            and event.event_type in {_REVIEW_PENDING_EVENT, _ANALYSIS_READY_EVENT}
        ]
        if not events or events[-1].event_type != _REVIEW_PENDING_EVENT:
            return None
        pending = events[-1]
        key = {"kind": "review_pending", "event_id": pending.event_id}
        if any(
            event.event_type == _REVIEW_VERSION_EVENT
            and isinstance(event.payload, dict) and event.payload.get("base") == key
            for event in self.store.list_events(revision.experiment_id)
        ):
            return None
        return pending

    def review_pending_analysis(
        self, protocol_id: str,
    ) -> tuple[domain.ExperimentProtocol, domain.ReadinessAssessment, tuple[ReviewFlag, ...], Any]:
        """The latest review-pending analysis, read back (lane RV, decision a).

        The protocol is given its source pages again from the stored source
        and its accepted OCR; when those are not the pages the analysis read
        (an OCR read again since), it is refused: the items cannot be shown
        against pages they were not read from.

        Lane AN5, decision 3: after the re-ask closed cards, the record it
        wrote is read instead (the stored analysis itself never changes).
        """

        revision = self._latest_protocol_revision(protocol_id)
        event = self._review_pending_event(revision)
        if event is None:
            raise ProtocolCatalogNotFoundError("No analysis is waiting for review.")
        closed = _cards_closed_event(event, self.store.list_events(revision.experiment_id))
        payload = closed.payload if closed is not None else event.payload
        payload = payload if isinstance(payload, dict) else {}
        protocol = _decode_stored_record(payload.get("protocol"))
        readiness = _decode_stored_record(payload.get("readiness"))
        if not isinstance(protocol, domain.ExperimentProtocol) or not isinstance(
            readiness, domain.ReadinessAssessment
        ):
            raise ProtocolCatalogUnavailableError("The review-pending analysis is unreadable.")
        extraction = self._extraction_for_analysis(
            revision, self._source_extraction(revision))
        if (
            extraction.sha256 != protocol.metadata.pdf.sha256
            or [_text_sha256(page.text) for page in extraction.pages]
            != payload.get("source_pages")
        ):
            raise ProtocolCatalogUnavailableError(
                "The pages this analysis read are no longer the source's pages."
            )
        protocol = replace(protocol, metadata=replace(protocol.metadata, pdf=extraction))
        flags = tuple(
            ReviewFlag.from_public(item) for item in payload.get("flags") or ()
            if isinstance(item, dict)
        )
        return protocol, readiness, flags, event

    def _source_extraction(self, revision: ProtocolRevisionRecord) -> ProtocolPdfExtraction:
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        return extract_protocol_pdf(self.store.file_store.object_path(
            revision.pdf_checksum, expected_size=pdf_object.byte_size
        ))

    def _record_source_metadata(
        self, revision: ProtocolRevisionRecord, extraction: ProtocolPdfExtraction,
    ) -> None:
        metadata = _SourceMetadata.from_extraction(extraction)
        self.store.append_event(
            f"source-metadata-{revision.experiment_id}-{revision.revision_number}",
            revision.experiment_id, revision.revision_number, "protocol_source_metadata",
            {"source_sha256": revision.pdf_checksum, "title": metadata.title,
             "analysis_status": metadata.analysis_status, "ocr_needed": metadata.ocr_needed},
        )

    def _source_metadata(
        self, revision: ProtocolRevisionRecord, analysis: AnalysisRevisionRecord | None,
    ) -> _SourceMetadata:
        for event in self.store.list_events(revision.experiment_id):
            if event.event_type == "protocol_source_metadata" and event.protocol_revision_number == revision.revision_number:
                payload = event.payload
                if (
                    not isinstance(payload, dict)
                    or payload.get("source_sha256") != revision.pdf_checksum
                    or not isinstance(payload.get("analysis_status"), str)
                    or not isinstance(payload.get("ocr_needed"), bool)
                    or (payload.get("title") is not None and not isinstance(payload.get("title"), str))
                ):
                    raise ProtocolCatalogUnavailableError("Stored source metadata is invalid.")
                return _SourceMetadata(payload.get("title"), payload["analysis_status"], payload["ocr_needed"])
        if analysis is not None:
            # Old catalogs already retain exact source metadata in analyses.
            # No PDF bytes, parser or write is needed for this projection.
            return _SourceMetadata.from_extraction(analysis.protocol.metadata.pdf)
        raise ProtocolCatalogUnavailableError("Source metadata needs startup preparation.")

    def prepare_source_metadata(self) -> None:
        """Backfill pre-metadata uploads once, before serving, on a worker.

        Existing events remain unchanged. Normal list requests never extract
        PDFs or create records, including after extraction-cache eviction.
        """
        for experiment in self.store.list_experiments():
            revision = self._latest_protocol_revision(experiment.experiment_id)
            try:
                self._source_metadata(revision, None)
            except ProtocolCatalogUnavailableError as exc:
                if str(exc) != "Source metadata needs startup preparation.":
                    raise
                analysis = self._latest_analysis(revision)
                self._record_source_metadata(revision, (
                    analysis.protocol.metadata.pdf if analysis is not None
                    else self._source_extraction(revision)))

    def recover_interrupted_ocr(self, active_protocol_ids: frozenset[str] = frozenset()) -> None:
        """At process startup, end only OCR requests with no terminal event.

        This never resumes a provider or changes a protocol's analysis. A
        later upload or explicit OCR request may start a new bounded run.
        """
        for experiment in self.store.list_experiments():
            if experiment.experiment_id in active_protocol_ids:
                continue
            revision = self._latest_protocol_revision(experiment.experiment_id)
            events = self._latest_ocr_events(revision)
            if not events or any(event.event_type in {
                _OCR_COMPLETED_EVENT, _OCR_FAILED_EVENT, _OCR_REVIEWED_EVENT,
            } for event in events):
                continue
            ocr_id = events[0].payload["ocr_id"]
            self.store.append_event(
                f"ocr-failed-{ocr_id}", revision.experiment_id,
                revision.revision_number, _OCR_FAILED_EVENT,
                {"ocr_id": ocr_id, "status": "failed", "failure_code": "protocol_ocr_interrupted"},
            )

    def warm_analysis_records(self) -> None:
        """Prepare immutable records off-loop before the first list request.

        No entry or execution decision is cached. Visibility is still applied
        before constructing a request's entries, against current ledger data.
        """
        for experiment in self.store.list_experiments():
            revision = self._latest_protocol_revision(experiment.experiment_id)
            self._latest_analysis(revision)
            versions = self._review_version_events(revision)
            if versions:
                self._analysis_record(revision, versions[-1].analysis_revision_number)

    def list_entries(
        self, protocol_ids: frozenset[str] | None = None,
    ) -> tuple[ProtocolCatalogEntry, ...]:
        entries = []
        for experiment in self.store.list_experiments():
            if protocol_ids is not None and experiment.experiment_id not in protocol_ids:
                continue
            entries.append(
                self._entry_for_revision(
                    self._latest_protocol_revision(experiment.experiment_id)
                )
            )
        return tuple(entries)

    @staticmethod
    def _development_fixture_payload(
        fixture: CuratedProtocolFixture,
    ) -> dict[str, object]:
        return {
            "protocol_id": fixture.protocol_id,
            "revision_id": fixture.revision_id,
            "fixture_sha256": fixture.fixture_sha256,
            "source_sha256": fixture.source_pdf_sha256,
            "status": fixture.status,
            "development_only": True,
        }

    @staticmethod
    def _development_fixture_event_id(
        fixture: CuratedProtocolFixture,
        analysis_payload_sha256: str,
        payload: dict[str, object],
    ) -> str:
        """Name the materialization event by everything the event records.

        Keyed on the fixture alone, the event collided with its own earlier
        record as soon as the analysis changed, so the analysis digest joined
        the id. The id still could not see the event's own payload. Lane DI
        (2026-10-08) dropped ``final_approval`` from it, and every store that
        had materialized this fixture before then held a row under the same id
        with the old payload: the store refused, correctly, and from
        2026-10-08 ``scripts/run_dev.sh`` did not start.

        The digest of the payload as the store writes it is part of the id
        now (lane BT, 2026-10-09). A changed payload is a new event appended
        beside the old one, and the old row stays exactly as it was written.
        """

        event_sha256 = hashlib.sha256(
            json.dumps(
                payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
            ).encode("utf-8")
        ).hexdigest()
        return (
            f"development-fixture-{fixture.fixture_sha256[:48]}"
            f"-{analysis_payload_sha256[:16]}-{event_sha256[:16]}"
        )

    @staticmethod
    def _development_analysis_identity(
        fixture: CuratedProtocolFixture,
    ) -> tuple[str, str]:
        """Name this fixture's analysis by everything the analysis contains.

        The id used to be the fixture's own SHA-256 and nothing else, which
        held for as long as the fixture's bytes decided the analysis. They do
        not. Readiness is assessed when the fixture is loaded, against the
        capability policy, so declaring a capability changes the stored
        analysis while leaving the fixture byte-identical -- and the id could
        not see it. The store then found an id it already held whose payload
        had changed and refused, correctly, at server start, which is the
        worst moment to learn it: the pilot catalog holds
        ``curated-69517f0f...`` with payload ``47df9633...`` and the declared
        policy produces ``824e9b54...``.

        Same lesson as the chunk cache key before STEP 31. A key that cannot
        see the question reports a hit it cannot honour -- or, here, a
        collision that is really a new analysis. The payload digest is part of
        the identity now, so a changed analysis is a different analysis and is
        appended as the next revision. Nothing is overwritten and nothing is
        deleted: the earlier revisions stay exactly where they are, and so do
        the findings recorded against them, which belong to those revisions
        and not to this one.
        """

        _, payload_sha256 = serialize_analysis(
            fixture.draft.protocol,
            fixture.draft.readiness,
            fixture.draft.capability_policy_id,
        )
        return (
            f"curated-{fixture.fixture_sha256}-{payload_sha256[:16]}",
            payload_sha256,
        )

    def bootstrap_development_fixture(
        self,
        fixture: CuratedProtocolFixture,
    ) -> ProtocolDevelopmentBootstrap:
        """Idempotently materialize an already validated development fixture.

        This boundary creates immutable source and analysis records only.  It
        never appends an approval event and never changes execution readiness.
        """

        if (
            not fixture.development_only
            or fixture.status != "development_only_not_final_acceptance"
            or fixture.source_pdf_path is None
            or fixture.source_pdf_sha256
            != fixture.draft.protocol.metadata.file_checksum
            or fixture.protocol_id != fixture.draft.protocol.protocol_id
        ):
            raise ProtocolRegistrationError(
                "Development fixture is not eligible for materialization."
            )
        existing = self.store.get_experiment(fixture.protocol_id)
        analysis_id, payload_sha256 = self._development_analysis_identity(
            fixture
        )
        analysis = self.store.create_experiment_with_analysis(
            fixture.protocol_id,
            fixture.source_pdf_path,
            analysis_id,
            fixture.draft.protocol,
            fixture.draft.readiness,
            fixture.draft.capability_policy_id,
        )
        revision = self.store.get_protocol_revision(fixture.protocol_id, 1)
        if revision is None:
            raise ProtocolCatalogUnavailableError(
                "Development fixture revision is unavailable."
            )
        payload = self._development_fixture_payload(fixture)
        self.store.append_event(
            self._development_fixture_event_id(fixture, payload_sha256, payload),
            fixture.protocol_id,
            revision.revision_number,
            _DEVELOPMENT_FIXTURE_EVENT,
            payload,
            analysis_revision_number=analysis.analysis_revision_number,
        )
        return ProtocolDevelopmentBootstrap(
            entry=self._entry_for_revision(revision),
            deduplicated=existing is not None,
        )

    def development_fixture_is_materialized(
        self,
        fixture: CuratedProtocolFixture,
    ) -> bool:
        """Verify the exact development-only provenance marker in the store.

        The question is identity, not uniqueness.  Asking whether the protocol
        carried exactly one analysis revision made an ordinary event -- editing
        the fixture, whose hash then names a second analysis at the next
        bootstrap -- indistinguishable from a corrupted store, and from
        2026-09-04 the whole catalog listing failed with a 503 for that reason
        alone.  A second analysis revision is what an append-only ledger is
        supposed to accumulate.

        What has to hold is that the analysis this exact fixture materializes
        exists, still carries the fixture's own content, and is the one the
        catalog would serve.  An analysis that does not match is not tolerated
        by being ignored: if the latest analysis is not this fixture's, the
        answer is False, because then the reader and the fixture disagree about
        what this protocol is.
        """

        if not fixture.development_only:
            return False
        revisions = self.store.list_protocol_revisions(fixture.protocol_id)
        if len(revisions) != 1:
            return False
        revision = revisions[0]
        if revision.pdf_checksum != fixture.source_pdf_sha256:
            return False
        analysis = self._latest_analysis(revision)
        if analysis is None:
            return False
        if analysis.analysis_id != self._development_analysis_identity(
            fixture
        )[0]:
            return False
        if (
            analysis.protocol != fixture.draft.protocol
            or analysis.readiness != fixture.draft.readiness
            or analysis.capability_policy_id
            != fixture.draft.capability_policy_id
        ):
            return False
        expected_payload = self._development_fixture_payload(fixture)
        return any(
            event.event_type == _DEVELOPMENT_FIXTURE_EVENT
            and event.protocol_revision_number == revision.revision_number
            and event.analysis_revision_number
            == analysis.analysis_revision_number
            and event.payload == expected_payload
            for event in self.store.list_events(fixture.protocol_id)
        )

    def get_entry(self, protocol_id: str) -> ProtocolCatalogEntry:
        return self._entry_for_revision(
            self._latest_protocol_revision(protocol_id)
        )

    def review(self, protocol_id: str) -> dict[str, object]:
        """Return a source-linked, read-only view of the latest analysis draft."""

        revision = self._latest_protocol_revision(protocol_id)
        entry = self._entry_for_revision(revision)
        analysis = self._latest_analysis(revision)
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        source = self.store.file_store.object_path(
            revision.pdf_checksum, expected_size=pdf_object.byte_size
        )
        extraction = extract_protocol_pdf(source)
        revision_events = tuple(
            event
            for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
        )
        current_events = _failures_not_superseded(revision_events)
        latest_failure = next(
            (
                event.payload.get("failure_code")
                for event in reversed(current_events)
                if event.event_type in {_ANALYSIS_FAILED_EVENT, _CHUNK_FAILED_EVENT}
                and isinstance(event.payload, dict)
                and isinstance(event.payload.get("failure_code"), str)
            ),
            None,
        )
        latest_failure_detail = next(
            (
                event.payload.get("evidence_failure")
                for event in reversed(current_events)
                if event.event_type
                in {_ANALYSIS_FAILED_EVENT, _CHUNK_FAILED_EVENT}
                and isinstance(event.payload, dict)
                and isinstance(event.payload.get("evidence_failure"), dict)
            ),
            None,
        )
        reanalysis = self._reanalysis_status(revision_events, analysis)
        base: dict[str, object] = {
            **entry.public_dict(),
            "analysis_available": analysis is not None,
            "lifecycle_state": entry.lifecycle_state,
            # Lane RN decision 2: a failure after a pass belongs to the run
            # asked for over it (``reanalysis``), not to the analysis that
            # passed, which still runs.
            "analysis_failure": (
                {
                    "code": latest_failure,
                    "detail": latest_failure_detail,
                    "retryable": (
                        latest_failure not in _NOT_RETRYABLE_FAILURE_CODES
                        or _refused_under_the_old_page_bound(
                            latest_failure, extraction.page_count)
                    ),
                    # Lane PA decision 4 (2026-10-06): the recovery guidance
                    # is the server's, in Korean.
                    "action": (
                        _OLD_PAGE_BOUND_ACTION
                        if _refused_under_the_old_page_bound(
                            latest_failure, extraction.page_count)
                        else _ANALYSIS_RECOVERY_ACTIONS.get(
                            latest_failure, _DEFAULT_ANALYSIS_RECOVERY_ACTION
                        )
                    ),
                }
                if latest_failure is not None and reanalysis is None
                else None
            ),
            "reanalysis": reanalysis,
            # Lane EASY, decision 1: the pages the analysis shown here read
            # (None: the whole document).
            "analysis_page_range": _shown_page_range(revision_events),
            # Lane RV, decision a: the latest analysis waiting for a person.
            "review_pending": entry.review_pending,
            "source": {
                "filename": revision.original_filename,
                "sha256": revision.pdf_checksum,
                "media_type": extraction.media_type,
                "byte_size": extraction.byte_size,
                "page_count": extraction.page_count,
                "all_pages_inspected": extraction.all_pages_inspected,
                "extraction_warnings": list(extraction.warnings),
            },
            "ocr": self._ocr_projection(
                revision, extraction, include_text=True
            ),
            "metadata": {},
            "before_start": [],
            "materials": [],
            "equipment": [],
            "sections": [],
            "constructs": [],
            "readiness": {
                "status": domain.ReadinessStatus.ANALYSIS_REQUIRED.value,
                "label": "Structured analysis has not been completed.",
                "label_ko": _ANALYSIS_REQUIRED_LABEL_KO,
                "reasons": [
                    {
                        "code": entry.analysis_status,
                        "message": "Explicit structured analysis is required before review or execution.",
                        "message_ko": _ANALYSIS_REQUIRED_MESSAGE_KO,
                    }
                ],
            },
            "capability_policy_id": domain.P1_CAPABILITY_POLICY.profile_id,
            # The execution rule's two lists (decision of 2026-10-08): what
            # still keeps this revision from running, and what the
            # experimenter is told before starting. Empty before an analysis.
            "execution_blockers": [],
            "execution_notices": [],
            # Lane MISS, decision 4: the places of the source that read like
            # an instruction and that no quote of the analysis covers.
            "possibly_missed": None,
            "safety_notice_sources": [],
            "timers": {"verified": [], "refused": []},
            # Lane EV2, decision 2: the fields emptied because the source did
            # not confirm them, and the one line the start screen shows.
            "cleared_fields": [],
            "cleared_fields_ko": None,
            # Lane EV2, decision 3: the fixed repeat counts to check before
            # the start, one line each.
            "repeat_count_checks_ko": [],
        }
        ocr_projection = base["ocr"]
        base["pipeline"] = self._pipeline(
            entry, analysis, revision,
            ocr=ocr_projection if isinstance(ocr_projection, dict) else {},
            failure_code=latest_failure,
            automatic_retry=_automatic_retry_status(revision_events),
            reanalysis=reanalysis,
            events=revision_events,
            page_count=extraction.page_count,
        )
        versions = self._review_versions(revision)
        version_payload: dict[str, object] | None = None
        if versions:
            # Lane RV, decision c: the start screen shows the version that
            # runs -- the latest registered review version.
            version_event, analysis = versions[-1]
            version_payload = version_event.payload
            base["analysis_available"] = True
            base["review_version"] = _review_version_summary(version_event)
        if analysis is None:
            return base

        protocol = analysis.protocol
        missed = (
            version_payload.get("possibly_missed") if version_payload is not None
            else _ready_possibly_missed(revision_events, analysis))
        declared_warning_count = domain.declared_safety_warning_count(protocol)
        metadata = {
            field.name: _review_value(getattr(protocol.metadata, field.name))
            for field in fields(protocol.metadata)
            if field.name != "pdf"
        }
        base.update(
            {
                "metadata": metadata,
                "before_start": _review_value(protocol.before_start),
                "materials": _review_value(protocol.materials),
                "equipment": _review_value(protocol.equipment),
                "sections": _review_sections(protocol),
                "constructs": [
                    {
                        "construct_type": type(construct).__name__,
                        **_review_value(construct),
                    }
                    for construct in protocol.constructs
                ],
                "readiness": _with_korean_reasons(_review_value(analysis.readiness)),
                "capability_policy_id": analysis.capability_policy_id,
                "analysis_payload_sha256": analysis.payload_sha256,
                "page_coverage": [
                    dict(item) for item in analysis.page_coverage
                ],
                "declined_segment_count": sum(
                    len(item.get("declined_segment_ids") or ())
                    for item in analysis.page_coverage
                ),
                "declared_safety_warning_count": declared_warning_count,
                "execution_blockers": self._execution_reasons(
                    analysis, blocking=True
                ),
                "execution_notices": [
                    *self._execution_reasons(analysis, blocking=False),
                    *possibly_missed_notices(missed),
                ],
                "possibly_missed": dict(missed) if isinstance(missed, dict) else None,
                # The source's own safety statements, step by step, for the
                # screen the experimenter reads before pressing start. The
                # server adds the Korean beside each (it holds the
                # translation store); the words here are the document's.
                "safety_notice_sources": _safety_notice_sources(protocol),
                "cleared_fields": [
                    {
                        "field": field,
                        "label_ko": domain.cleared_field_label_ko(field, protocol),
                    }
                    for field in protocol.cleared_fields
                ],
                "cleared_fields_ko": domain.cleared_fields_line_ko(
                    protocol.cleared_fields, protocol
                ),
                "repeat_count_checks_ko": _repeat_count_checks_ko(protocol),
                # Step timers for the start screen (lane PT, decision 4):
                # each verified value with its source excerpt, and each time
                # the server could not verify, with the reason.
                "timers": _review_timers(
                    self._step_timer_table(
                        revision, extraction,
                        protocol_with_display_step_labels(protocol),
                        version=version_payload,
                    )
                ),
            }
        )
        return base

    @staticmethod
    def _execution_reasons(
        analysis: Any, *, blocking: bool
    ) -> list[dict[str, object]]:
        """The readiness reasons of one kind, each with its Korean line.

        ``blocking`` selects the reasons that keep the revision out of
        execution; otherwise the notices the experimenter reads before
        starting. ``kind`` says what sort of notice it is:
        ``no_guidance_yet`` for a construct this version cannot guide (said
        once more at the step, with the source read out) and
        ``source_note`` for a finding about the source itself.
        """

        reasons = (
            domain.execution_blocking_reasons(analysis.readiness)
            if blocking else domain.execution_notice_reasons(analysis.readiness)
        )
        return [
            {
                "code": reason.code.value,
                "message": reason.message,
                "message_ko": readiness_reason_korean(
                    reason.code.value, reason.message
                ),
                "kind": (
                    "blocking" if blocking
                    else "no_guidance_yet"
                    if reason.code in domain.NO_GUIDANCE_YET_REASON_CODES
                    else "source_note"
                ),
                "source_page_number": (
                    reason.evidence.source_page_number
                    if reason.evidence else None
                ),
                "source_excerpt": (
                    reason.evidence.source_excerpt
                    if reason.evidence else None
                ),
                "step_id": reason.step_id,
            }
            for reason in reasons
        ]

    def pipeline_status(self, protocol_id: str) -> dict[str, object]:
        """Where this document stands on the way to a start (lane PX, 4)."""

        revision = self._latest_protocol_revision(protocol_id)
        entry = self._entry_for_revision(revision)
        analysis = self._latest_analysis(revision)
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        source = self.store.file_store.object_path(
            revision.pdf_checksum, expected_size=pdf_object.byte_size
        )
        extraction = extract_protocol_pdf(source)
        revision_events = tuple(
            event for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
        )
        failure = next(
            (
                event.payload.get("failure_code")
                for event in reversed(_failures_not_superseded(revision_events))
                if event.event_type in {_ANALYSIS_FAILED_EVENT, _CHUNK_FAILED_EVENT}
                and isinstance(event.payload, dict)
                and isinstance(event.payload.get("failure_code"), str)
            ),
            None,
        )
        return self._pipeline(
            entry, analysis, revision,
            ocr=self._ocr_projection(revision, extraction, include_text=False),
            failure_code=failure,
            automatic_retry=_automatic_retry_status(revision_events),
            reanalysis=self._reanalysis_status(revision_events, analysis),
            events=revision_events,
            page_count=extraction.page_count,
        )

    def _piece_progress_ko(self, revision: ProtocolRevisionRecord) -> str | None:
        """The progress line of a piece run still going, or None (lane AN, 2).

        "긴 문서를 조각 11개로 나눠 분석 중입니다 · 조각 5·6·7·8/11 분석 중 · 끝난 조각
        4/11", a retry beside the piece it is for, and the joining once
        every piece has passed.
        """

        events = self._latest_chunk_events(revision)
        if not events or _piece_run_end(events, tuple(
            event for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
        )) is not None:
            return None
        plan = events[0].payload if isinstance(events[0].payload, dict) else {}
        chunks = plan.get("chunks") if isinstance(plan.get("chunks"), list) else []
        count = len(chunks)
        if not count:
            return None
        state: dict[str, str] = {}
        retrying: set[str] = set()
        joining = False
        for event in events[1:]:
            payload = event.payload if isinstance(event.payload, dict) else {}
            chunk_id = str(payload.get("chunk_id"))
            if event.event_type == _CHUNK_STARTED_EVENT:
                state[chunk_id] = "in_progress"
                if payload.get("attempt") == 2:
                    retrying.add(chunk_id)
            elif event.event_type == _CHUNK_COMPLETED_EVENT:
                state[chunk_id] = "completed"
                retrying.discard(chunk_id)
            elif event.event_type == _CHUNK_FAILED_EVENT:
                state[chunk_id] = "failed"
                retrying.discard(chunk_id)
            elif event.event_type == _MERGE_STARTED_EVENT:
                joining = True
        ordinal = {
            str(chunk.get("chunk_id")): chunk.get("ordinal")
            for chunk in chunks if isinstance(chunk, dict)
        }
        done = sum(1 for value in state.values() if value == "completed")
        if joining:
            return (
                f"긴 문서를 조각 {count}개로 나눠 분석했습니다 · 조각 {count}개를 합쳐 원문 "
                "전체와 다시 대조하고 있습니다."
            )
        running = sorted(
            ordinal[chunk_id] for chunk_id, value in state.items()
            if value == "in_progress" and isinstance(ordinal.get(chunk_id), int)
        )
        line = f"긴 문서를 조각 {count}개로 나눠 분석 중입니다"
        if running:
            line += f" · 조각 {'·'.join(str(index) for index in running)}/{count} 분석 중"
        line += f" · 끝난 조각 {done}/{count}"
        again = sorted(
            ordinal[chunk_id] for chunk_id in retrying
            if isinstance(ordinal.get(chunk_id), int)
        )
        if again:
            line += (
                f" · 조각 {'·'.join(str(index) for index in again)} 다시 시도 중(1/1)"
            )
        return line + "."

    def _pipeline(
        self,
        entry: ProtocolCatalogEntry,
        analysis: AnalysisRevisionRecord | None,
        revision: ProtocolRevisionRecord,
        *,
        ocr: dict[str, object],
        failure_code: str | None,
        automatic_retry: dict[str, object] | None = None,
        reanalysis: dict[str, object] | None = None,
        events: Sequence[Any] = (),
        page_count: int | None = None,
    ) -> dict[str, object]:
        """One Korean line: the stage a document is at, whether it is stuck
        there, why, and what a person does (lane PX, decision 4).

        Stages, in order: 원문 읽기 → OCR → 분석 → 근거 대조 → 실행 준비 →
        사람 확인 → 실행 가능. Translation is added by the server, which
        holds the translation store. A document that reached this method
        was read, so 원문 읽기 never blocks here. An automatic retry (lane
        AN, decision 1) is said while it runs and after it failed. A run
        asked for over a passed analysis (lane RN, decision 2) is said beside
        the pass, which still runs whatever that run does. A response read
        after its brackets were repaired (lane EV3, decision 1) is said on the
        line of the attempt it belongs to.
        """

        passed_repair = next(
            (
                _response_repair_of(event.payload) for event in events
                if analysis is not None and event.event_type == _ANALYSIS_READY_EVENT
                and event.analysis_revision_number == analysis.analysis_revision_number
            ),
            None,
        )
        failed_repair = next(
            (
                _response_repair_of(event.payload)
                for event in reversed(_failures_not_superseded(events))
                if event.event_type in {_ANALYSIS_FAILED_EVENT, _CHUNK_FAILED_EVENT}
                and isinstance(event.payload, dict)
                and isinstance(event.payload.get("failure_code"), str)
            ),
            None,
        )
        passed_clause = (
            f" {_response_repair_ko(passed_repair)}: 고친 응답도 형식·원문 근거·경계 검사를 "
            "모두 통과했습니다." if passed_repair else "")
        passed_extra: dict[str, object] = (
            {"response_repair": passed_repair} if passed_repair else {})
        retry = automatic_retry or {}
        again = reanalysis or {}
        again_failure = again.get("failure") if isinstance(again.get("failure"), dict) else None
        # Lane EV3, decision 5: the re-analysis's own automatic retry.
        again_retry = again.get("automatic_retry") if isinstance(
            again.get("automatic_retry"), dict) else {}
        again_retry_count = (f"{again_retry.get('attempt', 1)}/"
                             f"{again_retry.get('limit', AUTOMATIC_RETRY_LIMIT)}")
        if again.get("state") == "review_pending":
            # Lane RV, decision a: the run over a passed analysis ended with
            # items for a person; the passed analysis still runs.
            summary = again.get("review_pending") if isinstance(
                again.get("review_pending"), dict) else {}
            again_clause = (
                f" 다시 분석은 {summary.get('message_ko') or '검토 대기'}입니다 · "
                "지금 분석으로 계속 실행할 수 있습니다.")
        elif again.get("state") in {"analysis_pending", "analyzing"}:
            again_clause = (
                (f" 다시 분석 중입니다 · 다시 시도 중({again_retry_count}) · 첫 응답이 정해진 "
                 "형식에 맞지 않아 같은 요청을 자동으로 한 번 다시 보냈습니다 · "
                 if again_retry.get("state") == "in_progress"
                 else " 다시 분석 중입니다 · 원문 근거를 확인하고 있습니다 · ")
                + "지금 분석으로 계속 실행할 수 있습니다.")
        elif again_failure is not None:
            again_repair = _response_repair_of(again_failure)
            again_clause = (
                " 다시 분석이 실패했습니다: "
                f"{_ANALYSIS_FAILURE_KO.get(str(again_failure.get('code')), '분석이 실패했습니다.')} "
                + (f"같은 요청을 자동으로 한 번 다시 보냈지만 자동 재시도({again_retry_count})도 "
                   "실패했습니다. " if again_retry.get("state") == "failed" else "")
                + (f"{_response_repair_ko(again_repair)}. " if again_repair else "")
                + "지금까지의 통과 분석으로 계속 실행할 수 있습니다.")
        else:
            again_clause = ""
        again_extra: dict[str, object] = (
            {"reanalysis": again, **({"failure_code": again_failure["code"]} if again_failure else {})}
            if again else {}
        )
        retry_count = f"{retry.get('attempt', 1)}/{retry.get('limit', AUTOMATIC_RETRY_LIMIT)}"
        truncated_how = _truncated_retry_ko(retry) if retry.get("truncated") else None

        def result(stage: str, *, blocked: bool, message: str, action: str | None = None,
                   **extra: object) -> dict[str, object]:
            return {
                "stage": stage, "stage_ko": PIPELINE_STAGE_KO[stage], "blocked": blocked,
                "message": message, "action": action, **extra,
            }

        ocr_state = str(ocr.get("state") or "not_required")
        ocr_pages = [
            page["source_page_number"] for page in (ocr.get("pages") or ())
            if isinstance(page, dict) and page.get("provider") != OCR_TEXT_LAYER_PROVIDER
            and isinstance(page.get("source_page_number"), int)
        ] if isinstance(ocr.get("pages"), list) else []
        if entry.available_for_execution:
            return result(
                "ready", blocked=False,
                message="분석을 통과했습니다. 실행할 수 있습니다." + passed_clause + again_clause,
                action=(_REANALYSIS_RECOVERY_ACTION if again_failure is not None
                        else "시작 전 화면의 안전 주의를 읽고 '이 프로토콜로 시작'을 누르세요."),
                ocr_page_numbers=ocr_pages, **again_extra, **passed_extra)
        if analysis is None:
            if entry.analysis_status == "review_pending" and entry.review_pending:
                # Lane RV, decision a (human decision 2026-10-12): kept for a
                # person instead of thrown away; it does not run as it is.
                return result(
                    "review", blocked=True,
                    message=(
                        f"{entry.review_pending['message_ko']}. 분석은 끝났지만 원문에서 "
                        "확인하지 못한 항목이 있어 이대로는 실행하지 않습니다."
                        + passed_clause),
                    action=REVIEW_PENDING_ACTION_KO,
                    review_pending=entry.review_pending,
                    ocr_page_numbers=ocr_pages)
            if ocr_state in {"queued", "in_progress"}:
                return result(
                    "ocr", blocked=False,
                    message="글자 층이 없는 쪽을 OCR 로 읽는 중입니다. 끝나면 분석이 바로 시작됩니다.")
            if ocr_state in {"failed", "rejected"}:
                code = str(ocr.get("failure_code") or "protocol_ocr_failed")
                return result(
                    "ocr", blocked=True,
                    message=("OCR 결과를 거절했습니다." if ocr_state == "rejected"
                             else "글자 층이 없는 쪽의 OCR 이 실패했습니다."),
                    action="'OCR 텍스트 추출'을 눌러 다시 시도하세요. 원문은 바뀌지 않습니다.",
                    failure_code=code)
            if ocr_state == "ocr_required" and entry.analysis_status in {"ocr_required"}:
                return result(
                    "ocr", blocked=True,
                    message="글자 층이 없는 쪽이 있어 OCR 글이 먼저 필요합니다.",
                    action=("OCR 공급자가 설정돼 있으면 업로드 때 자동으로 돕니다. 돌지 않았다면 "
                            "서버의 OCR 설정(VOINEY_LAB_OCR_PROVIDERS 와 엔진 키)을 확인하고 "
                            "'OCR 텍스트 추출'을 누르세요."))
            if failure_code is not None and entry.lifecycle_state not in {
                "analysis_pending", "analyzing",
            }:
                stage = ("evidence" if failure_code == "protocol_analysis_invalid_evidence"
                         else "analysis")
                message = _ANALYSIS_FAILURE_KO.get(failure_code, "분석이 실패했습니다.")
                if page_count is not None and _refused_under_the_old_page_bound(
                    failure_code, page_count
                ):
                    # Lane AN, decision 2: refused under lane EV2's 60-page
                    # bound, which the pieces replaced.
                    return result(
                        "analysis", blocked=True, message=_OLD_PAGE_BOUND_KO,
                        action=_OLD_PAGE_BOUND_ACTION, failure_code=failure_code)
                failed_pieces = _failed_pieces_ko(events)
                if failure_code == "protocol_analysis_piece_failed" and failed_pieces:
                    # Lane AN, decision 2: "일부 조각 분석 실패 · 41–59쪽".
                    message += f" · {failed_pieces}"
                    blocked = _blocked_pieces_ko(events)
                    if blocked:
                        message += f" · {blocked}"
                if failure_code == BLOCKED_FAILURE_CODE and page_count is not None:
                    # Lane RV2, decision f: "구글 안전 필터에 막혀 1–12쪽을 분석하지 못했어요."
                    message = blocked_pages_ko(1, page_count, _blocked_provider(events)) + "."
                too_many = _latest_failure_payload(events).get("review_items")
                if failure_code == TOO_MANY_REVIEW_ITEMS_CODE and isinstance(too_many, dict):
                    # Lane RV2, decision c: the counts behind the guidance.
                    after = too_many.get("items_after")
                    message = (
                        f"확인할 항목이 {too_many.get('items')}건이라(카드는 "
                        f"{too_many.get('limit')}건까지) 검토 카드를 만들지 않았어요"
                        + (f". 한 번 더 분석해도 {after}건이었어요" if isinstance(after, int)
                           else "")
                        + ". 이 문서는 다시 분석하거나 절차가 있는 쪽만 올리는 편이 나아요.")
                    by_reason = too_many.get("by_reason")
                    ocr_values = (by_reason or {}).get(REVIEW_REASON_OCR_VALUE, 0) \
                        if isinstance(by_reason, dict) else 0
                    if isinstance(ocr_values, int) and ocr_values * 2 >= int(too_many.get("items") or 0) > 0:
                        # Most of them are OCR values no second engine confirmed:
                        # reading again does not change those; say what does.
                        # (Lane AN3, decision A3: a page's OCR values are one.)
                        message += (
                            f" 그중 {ocr_values}건은 OCR 엔진 하나만 읽었거나 두 엔진이 다르게 읽은 "
                            "숫자가 있는 쪽이에요(쪽마다 한 건) — 두 OCR 엔진"
                            "(VOINEY_LAB_OCR_PROVIDERS=clova,google)이 모두 읽거나 절차가 있는 쪽만 "
                            "올리면 줄어요.")
                unlabelled = _unlabelled_pieces_ko(events)
                if failure_code == "protocol_analysis_piece_conflict" and unlabelled:
                    message = (
                        f"조각마다 단계 번호를 다르게 읽어 합치지 못했어요({unlabelled}은 "
                        "번호 없이, 나머지는 번호를 붙여 읽음). 원문이 부분마다 번호를 다르게 "
                        "붙였다면 한 문서에서 둘을 함께 쓸지는 사람이 정할 일이에요.")
                if retry.get("state") == "failed" and truncated_how:
                    message += (f" {truncated_how} 자동으로 한 번 다시 보냈지만 "
                                f"자동 재시도({retry_count})도 실패했습니다.")
                elif retry.get("state") == "failed":
                    message += f" 같은 요청을 자동으로 한 번 다시 보냈지만 자동 재시도({retry_count})도 실패했습니다."
                if failed_repair:
                    message += f" {_response_repair_ko(failed_repair)}."
                return result(
                    stage, blocked=True,
                    message=message,
                    action=_ANALYSIS_RECOVERY_ACTIONS.get(
                        failure_code, _DEFAULT_ANALYSIS_RECOVERY_ACTION),
                    failure_code=failure_code,
                    **({"response_repair": failed_repair} if failed_repair else {}))
            pieces_line = self._piece_progress_ko(revision)
            if pieces_line is not None and entry.lifecycle_state in {
                "analysis_pending", "analyzing",
            }:
                # Lane AN, decision 2: "조각 k/n 분석 중".
                return result("analysis", blocked=False, message=pieces_line)
            if entry.lifecycle_state in {"analysis_pending", "analyzing"} or (
                entry.analysis_status in {
                    "chunk_planned", "chunk_analysis_in_progress", "merge_in_progress",
                }
            ):
                if retry.get("state") == "in_progress" and truncated_how:
                    return result(
                        "analysis", blocked=False,
                        message=(f"분석 중입니다 · 다시 시도 중({retry_count}). {truncated_how} "
                                 "자동으로 한 번 다시 보냈습니다."),
                        automatic_retry=retry)
                if retry.get("state") == "in_progress":
                    return result(
                        "analysis", blocked=False,
                        message=(f"분석 중입니다 · 다시 시도 중({retry_count}). 첫 응답이 정해진 형식에 "
                                 "맞지 않아 같은 요청을 자동으로 한 번 다시 보냈습니다."),
                        automatic_retry=retry)
                return result(
                    "analysis", blocked=False,
                    message="분석 중입니다. 원문 근거를 확인하고 있습니다.")
            if entry.analysis_status in {"chunk_analysis_failed", "merge_conflict",
                                         "chunk_analysis_cancelled"}:
                return result(
                    "analysis", blocked=True,
                    message=_ANALYSIS_FAILURE_KO.get(entry.analysis_status, "분석이 실패했습니다."),
                    action=_DEFAULT_ANALYSIS_RECOVERY_ACTION)
            return result(
                "analysis", blocked=False,
                message="분석 대기 중입니다. 업로드 직후 자동으로 시작되며, 시작되지 않았다면 "
                        "'분석 다시 시도'를 누르세요.",
                ocr_page_numbers=ocr_pages)
        blockers = list(domain.execution_blocking_reasons(analysis.readiness))
        if not blockers:
            # Cannot happen while available_for_execution reads the same
            # list, but a pipeline line must still say something true.
            return result(
                "ready", blocked=False,
                message="분석을 통과했습니다. 실행할 수 있습니다." + passed_clause,
                action="시작 전 화면의 안전 주의를 읽고 '이 프로토콜로 시작'을 누르세요.",
                ocr_page_numbers=ocr_pages, **passed_extra)
        first = readiness_reason_korean(
            blockers[0].code.value, blockers[0].message) or "실행을 막는 사유가 있습니다."
        return result(
            "readiness", blocked=True,
            message=f"{first}" + (f" (외 {len(blockers) - 1}건)" if len(blockers) > 1 else "")
            + passed_clause
            + again_clause.replace("지금 분석으로 계속 실행할 수 있습니다.", "지금 분석은 그대로입니다.")
            .replace("지금까지의 통과 분석으로 계속 실행할 수 있습니다.", "지금 분석은 그대로입니다."),
            action=_BLOCKER_ACTIONS.get(
                blockers[0].code.value,
                "원문을 확인하고 고친 PDF 를 올리거나 '다시 분석'을 누르세요."),
            remaining_reason_codes=[reason.code.value for reason in blockers],
            ocr_page_numbers=ocr_pages, **again_extra, **passed_extra)

    def register(
        self,
        source_pdf: str | Path,
        *,
        source_filename: str,
        media_type: str,
        converted_from: SourceConversion | None = None,
        organization_id: str | None = None,
        visible_protocol_ids: frozenset[str] | None = None,
    ) -> ProtocolRegistration:
        """Register one source PDF, or find the one already registered.

        ``converted_from`` names the Word document or photo the PDF was made
        from (lane AN, FM); it is kept in the registration event, where
        ``find_converted_upload`` finds it.
        """

        filename = _display_filename(source_filename)
        if media_type.casefold().split(";", 1)[0].strip() != PDF_MEDIA_TYPE:
            raise ProtocolRegistrationError(
                "Protocol registration requires application/pdf."
            )
        source = Path(source_pdf)
        extraction = extract_protocol_pdf(source)
        registration_key=extraction.sha256
        if organization_id is None:
            existing = self.store.find_protocol_revision_by_checksum(extraction.sha256)
        else:
            # Only registration is scoped here. Source bytes remain addressed
            # by their original checksum; analyses/events belong to this lab.
            registration_key=hashlib.sha256(f"{organization_id}\x1f{extraction.sha256}".encode()).hexdigest()
            candidates=set(visible_protocol_ids or ()) | {_protocol_id(registration_key)}
            existing=next((revision for identifier in sorted(candidates)
                for revision in self.store.list_protocol_revisions(identifier)
                if revision.pdf_checksum==extraction.sha256),None)
        if existing is not None:
            return ProtocolRegistration(
                self._entry_for_revision(existing), deduplicated=True
            )
        protocol_id = _protocol_id(registration_key)
        creation = self.store.create_experiment(
            protocol_id, source, original_filename=filename
        )
        self._record_source_metadata(creation.protocol_revision, extraction)
        status = _analysis_state(extraction)
        self.store.append_event(
            f"registered-{registration_key[:40]}",
            protocol_id,
            creation.protocol_revision.revision_number,
            "protocol_registered",
            {
                "source_filename": filename,
                "source_sha256": extraction.sha256,
                "analysis_status": status,
                **(
                    {"converted_from": converted_from.payload()}
                    if converted_from is not None else {}
                ),
            },
        )
        return ProtocolRegistration(
            self._entry_for_revision(creation.protocol_revision),
            deduplicated=False,
        )

    def attach_reagent_table(
        self, protocol_id: str, source: str | Path, *, filename: str,
    ) -> tuple[dict[str, object], bool]:
        """Attach an .xlsx workbook to a protocol as its reagent table (FM b).

        Append-only: one event per workbook, named by its checksum; the same
        workbook attached again returns the one already there. Returns the
        table and whether it was already attached.
        """

        if not isinstance(filename, str) or not _SAFE_TABLE_FILENAME.fullmatch(filename):
            raise ReagentTableError("Reagent table filename must be one plain .xlsx filename.")
        revision = self._latest_protocol_revision(protocol_id)
        digest = hashlib.sha256(Path(source).read_bytes()).hexdigest()
        for table in self.reagent_tables(protocol_id):
            if table.get("sha256") == digest:
                return table, True
        workbook = read_reagent_workbook(source)
        table_id = f"reagent-table-{digest[:24]}"
        payload = {
            "table_id": table_id,
            "filename": filename,
            "media_type": REAGENT_TABLE_MEDIA_TYPE,
            "sha256": digest,
            "byte_size": Path(source).stat().st_size,
            "authority": "person_upload",
            **workbook,
        }
        self.store.append_event(
            f"{table_id}-{revision.pdf_checksum[:16]}",
            protocol_id,
            revision.revision_number,
            _REAGENT_TABLE_EVENT,
            payload,
        )
        return payload, False

    def reagent_tables(self, protocol_id: str) -> list[dict[str, object]]:
        """The reagent tables attached to a protocol, oldest first."""

        revision = self._latest_protocol_revision(protocol_id)
        return [
            event.payload for event in self.store.list_events(revision.experiment_id)
            if event.event_type == _REAGENT_TABLE_EVENT and isinstance(event.payload, dict)
        ]

    def find_converted_upload(
        self, sha256: str, *, visible_protocol_ids: frozenset[str] | None = None,
    ) -> ProtocolRegistration | None:
        """The protocol a Word document or photo with this checksum made, if any.

        Lane AN, FM: LibreOffice writes the time into each PDF it prints, so
        the same .docx uploaded again gives other PDF bytes; it is found by
        its own checksum before it is converted again.
        """

        revision = self._revision_converted_from(sha256,visible_protocol_ids=visible_protocol_ids)
        if revision is None:
            return None
        return ProtocolRegistration(self._entry_for_revision(revision), deduplicated=True)

    def _revision_converted_from(
        self, sha256: str, *, visible_protocol_ids: frozenset[str] | None = None,
    ) -> ProtocolRevisionRecord | None:
        """The revision registered from the upload with this checksum, if any."""

        for experiment in self.store.list_experiments():
            if visible_protocol_ids is not None and experiment.experiment_id not in visible_protocol_ids:
                continue
            for event in self.store.list_events(experiment.experiment_id):
                if event.event_type != "protocol_registered" or not isinstance(
                    event.payload, dict
                ):
                    continue
                origin = event.payload.get("converted_from")
                if isinstance(origin, dict) and origin.get("sha256") == sha256:
                    return self.store.get_protocol_revision(
                        experiment.experiment_id, event.protocol_revision_number
                    )
        return None

    def analyze(
        self,
        protocol_id: str,
        model: ProtocolAnalysisModel,
        *,
        analysis_id: str,
    ) -> ProtocolCatalogEntry:
        """Run analysis only at this explicit caller-owned boundary."""

        revision = self._latest_protocol_revision(protocol_id)
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        source = self.store.file_store.object_path(
            revision.pdf_checksum, expected_size=pdf_object.byte_size
        )
        extraction = extract_protocol_pdf(source)
        extraction = self._extraction_for_analysis(revision, extraction)
        # Lane EASY, decision 1: the pages the request chose, or the whole.
        page_range = self._requested_page_range(revision, analysis_id)
        pieces = analysis_pieces(extraction, page_range=page_range)
        try:
            # Lane EV2, decision 4, changed by lane AN, decision 2: a document
            # longer than one analysis reads is read in pieces; one longer
            # than the pieces read is refused here, before any model call.
            if pieces:
                check_pieced_page_count(extraction, page_range=page_range)
            else:
                check_analysis_page_count(extraction, page_range=page_range)
        except ProtocolAnalysisTooManyPagesError as exc:
            self._record_analysis_failure(revision, analysis_id, exc)
            raise
        status = _analysis_state(extraction)
        if status == "ocr_required":
            raise ProtocolOcrRequiredError(
                "Protocol requires reviewed OCR before structured analysis."
            )
        self.store.append_event(
            f"analysis-started-{analysis_id}",
            protocol_id,
            revision.revision_number,
            _ANALYSIS_STARTED_EVENT,
            {"status": "analyzing", "analysis_id": analysis_id,
             **({"page_range": _page_range_public(page_range, extraction.page_count)}
                if page_range is not None else {})},
        )
        if pieces:
            try:
                draft = self._analyze_in_pieces(
                    revision, extraction, model, pieces, analysis_id=analysis_id)
            except Exception as exc:
                self._record_analysis_failure(revision, analysis_id, exc)
                raise
            if page_range is not None:
                draft = replace(draft, page_range=page_range)
            return self._store_analysis_result(
                revision, draft, analysis_id, pieces=pieces, model=model)
        if status == "chunked_analysis_required" and page_range is None:
            raise ProtocolChunkedAnalysisRequiredError(
                "Protocol text is larger than one analysis reads."
            )
        observed = _ObservedAnalysisModel.wrap(model)
        # A repaired answer with items for a person, kept should its retry
        # not come back better (lane RV, decision a).
        kept: ProtocolAnalysisDraft | None = None
        try:
            draft = analyze_protocol_extraction(
                extraction, observed.model, review_unverified=True, page_range=page_range)
            if draft.review_flags and draft.response_repair is not None:
                raise _RepairedAnswerWithItems(draft)
        except Exception as exc:
            if isinstance(exc, _RepairedAnswerWithItems):
                kept = exc.draft
            failure_code = self._record_analysis_failure(
                revision, analysis_id, exc, truncated=observed.truncated)
            repair = getattr(exc, "response_repair", None)
            if not self._automatic_retry_allowed(
                revision, failure_code, repaired=repair is not None
            ):
                if kept is not None:
                    return self._store_analysis_result(
                        revision, kept, analysis_id, model=model)
                raise
            # The same request again: the same extraction to the same model,
            # nothing about the failure added (lane AN, decision 1) -- except
            # that an answer cut at the output limit is asked for one
            # reasoning level lower (lane AQ, decision 4).
            retry_model, retry_payload = observed.retry()
            retry_of, analysis_id = analysis_id, f"{analysis_id}-retry-1"
            self.store.append_event(
                f"analysis-retry-{analysis_id}",
                protocol_id,
                revision.revision_number,
                _ANALYSIS_RETRY_EVENT,
                {
                    "status": "analyzing",
                    "analysis_id": analysis_id,
                    "retry_of": retry_of,
                    "attempt": 1,
                    "limit": AUTOMATIC_RETRY_LIMIT,
                    "reason_code": failure_code,
                    "authority": AUTOMATIC_RETRY_AUTHORITY,
                    **retry_payload,
                    **(
                        {"response_repaired": repair.public_dict()}
                        if repair is not None else {}
                    ),
                },
            )
            logging.getLogger(__name__).info(
                "protocol.analysis.automatic_retry protocol_id=%s attempt=1 limit=%d "
                "reason_code=%s truncated=%s reasoning_effort=%s",
                protocol_id, AUTOMATIC_RETRY_LIMIT, failure_code,
                bool(retry_payload.get("truncated")), retry_payload.get("reasoning_effort"),
            )
            try:
                draft = analyze_protocol_extraction(
                    extraction, retry_model, review_unverified=True, page_range=page_range)
            except Exception as retry_exc:
                self._record_analysis_failure(
                    revision, analysis_id, retry_exc, truncated=observed.truncated)
                if kept is not None:
                    # Lane RV, decision a: the first answer's items wait for a
                    # person rather than the analysis being thrown away.
                    return self._store_analysis_result(revision, kept, retry_of, model=model)
                raise
        return self._store_analysis_result(revision, draft, analysis_id, model=model)

    def _store_analysis_result(
        self,
        revision: ProtocolRevisionRecord,
        draft: ProtocolAnalysisDraft,
        analysis_id: str,
        *,
        pieces: tuple[AnalysisPiece, ...] = (),
        model: ProtocolAnalysisModel | None = None,
    ) -> ProtocolCatalogEntry:
        """A passed analysis as the next analysis revision; one with items a
        person must check as a review-pending ledger event (lane RV, a).

        Lane RV2, decision b: the values read by OCR are settled or made
        items first, so an analysis whose OCR values are not confirmed does
        not run as it is. Decision c: with ``model``, an analysis with more
        items than ``review_max_items()`` is analysed once more instead --
        except a document read in pieces (lane AN3, round 2, decision ④),
        which may ask about ``review_max_items()`` plus one item per piece and
        otherwise ends with the guidance.

        Lane AN4, decision 1: with ``model``, the items the evidence check
        refused are asked about once more. Lane AN5, decision 3 (human
        decision 2026-10-20): the analysis is stored and its cards shown
        first, and the re-ask runs after (``_requote_after_storing``) --
        unless the items are more than the cards may be, where the re-ask
        comes first as before (``_requoted``), since what it keeps decides
        whether the document is read again or ends with the guidance."""

        groups = _requote_groups(draft, pieces)
        asks = (
            model is not None and bool(draft.review_flags) and analysis_requote_enabled()
            and bool(requote_questions(draft, groups=groups)))
        first = draft
        draft = self._with_ocr_value_flags(revision, first)
        limit = review_max_items() + len(pieces) if pieces else review_max_items()
        if asks and _items_by_reason(draft)[0] > limit:
            draft = self._with_ocr_value_flags(revision, self._requoted(
                revision, first, analysis_id, model=model, pieces=pieces))
            asks = False
        if pieces:
            # Lane AN3, round 2, decision ④ (human decision 2026-10-18): a long
            # document is not analysed again by itself -- the one live
            # re-analysis cost USD 0.91 and changed nothing -- and may ask about
            # one more item per piece (a seam or an unread page is one press).
            count, by_reason = _items_by_reason(draft)
            if count > limit:
                too_many = ProtocolAnalysisTooManyItemsError(
                    "The analysis asks a person about too many items.",
                    items=count, limit=limit, items_after=None, by_reason=by_reason)
                self._record_analysis_failure(
                    revision, f"{analysis_id}-too-many-items", too_many)
                raise too_many
        elif model is not None and _items_by_reason(draft)[0] > review_max_items():
            draft, analysis_id = self._fewer_items(revision, draft, analysis_id, model=model)
        draft = self._with_program_address_flags(revision, draft)
        if draft.review_flags:
            draft = self._with_possibly_missed_flags(revision, draft)
            entry = self._store_review_pending(revision, draft, analysis_id, pieces=pieces)
            if asks:
                self._requote_after_storing(
                    revision, draft, analysis_id, model=model, pieces=pieces)
                entry = self.get_entry(revision.experiment_id)
            return entry
        return self._store_passed_analysis(revision, draft, analysis_id, pieces=pieces)

    def _requoted(
        self,
        revision: ProtocolRevisionRecord,
        draft: ProtocolAnalysisDraft,
        analysis_id: str,
        *,
        model: ProtocolAnalysisModel | None,
        pieces: tuple[AnalysisPiece, ...] = (),
    ) -> ProtocolAnalysisDraft:
        """The draft after its refused items were asked about once more.

        Lane AN4, decision 1 (human decision 2026-10-19). Once per analysis --
        a long document once per piece, ``analysis_piece_concurrency()`` at a
        time -- the items the evidence check refused (원문에 없음, 근거를 못
        찾음, a quote that does not open with its step number) go in one small
        request with their pages' text only. What comes back passes the same
        checks as any answer or the item stays a card as it was
        (``requote_unverified_items``); a kept item is recorded beside the
        server's other settlements as "근거를 다시 찾아 확인". Off with
        ``VOINEY_LAB_ANALYSIS_REQUOTE=false``. The automatic re-analysis of
        decision c, should it follow, is not asked about again. Lane AN5,
        decision 3: here only for an analysis with more items than its cards
        may be; the calls are waited for ``REQUOTE_TIME_LIMIT_SECONDS``.
        """

        if model is None or not draft.review_flags or not analysis_requote_enabled():
            return draft
        result = requote_unverified_items(
            draft, model,
            groups=_requote_groups(draft, pieces),
            concurrency=analysis_piece_concurrency(), time_limit=REQUOTE_TIME_LIMIT_SECONDS)
        if not result.questions:
            return draft
        self._record_requote(revision, analysis_id, result)
        return result.draft

    def _record_requote(
        self, revision: ProtocolRevisionRecord, analysis_id: str, result: Any,
        **extra: object,
    ) -> None:
        summary = result.summary()
        self.store.append_event(
            f"analysis-requote-{analysis_id}",
            revision.experiment_id,
            revision.revision_number,
            _REQUOTE_EVENT,
            {"status": "completed", "analysis_id": analysis_id, **summary, **extra},
        )
        logging.getLogger(__name__).info(
            "protocol.analysis.evidence_requote protocol_id=%s questions=%d items=%d "
            "items_requoted=%d calls=%d failed_calls=%d",
            revision.experiment_id, summary["questions"], summary["items"],
            summary["items_requoted"], len(result.calls),
            sum(1 for call in result.calls if call.get("status") != "answered"),
        )

    def _requote_after_storing(
        self,
        revision: ProtocolRevisionRecord,
        draft: ProtocolAnalysisDraft,
        analysis_id: str,
        *,
        model: ProtocolAnalysisModel,
        pieces: tuple[AnalysisPiece, ...] = (),
    ) -> None:
        """Ask about the refused items once more after the cards are shown.

        Lane AN5, decision 3 (human decision 2026-10-20). The analysis waits
        for a person already; while this runs its line says "원문 위치를 다시
        찾는 중". The calls are waited for ``REQUOTE_TIME_LIMIT_SECONDS`` at
        most (a call not back leaves its cards as they are). Then
        (``_close_requoted_cards``) each item kept closes its card as "원문
        위치를 다시 찾아 확인했어요" -- one more ledger event; nothing stored
        before changes -- except a card a person answered meanwhile, which
        stays theirs, and one whose values on an OCR page would ask on a card
        of their own. When that closes every card and nobody has touched
        them, the analysis passes as it would have with the items found the
        first time.
        """

        pending = self._review_pending_event(revision)
        if pending is None:
            return
        self.store.append_event(
            f"analysis-requote-started-{analysis_id}",
            revision.experiment_id,
            revision.revision_number,
            _REQUOTE_STARTED_EVENT,
            {"status": "running", "analysis_id": analysis_id,
             "review_pending_event": pending.event_id,
             "time_limit_seconds": REQUOTE_TIME_LIMIT_SECONDS},
        )
        try:
            result = requote_unverified_items(
                draft, model,
                groups=_requote_groups(draft, pieces),
                concurrency=analysis_piece_concurrency(), time_limit=REQUOTE_TIME_LIMIT_SECONDS)
        except Exception as exc:  # noqa: BLE001 - the cards stay as they are
            self.store.append_event(
                f"analysis-requote-{analysis_id}",
                revision.experiment_id,
                revision.revision_number,
                _REQUOTE_EVENT,
                {"status": "failed", "analysis_id": analysis_id,
                 "review_pending_event": pending.event_id, "failure_code": type(exc).__name__},
            )
            logging.getLogger(__name__).warning(
                "protocol.analysis.evidence_requote_failed protocol_id=%s error=%s",
                revision.experiment_id, type(exc).__name__)
            return
        self._close_requoted_cards(revision, pending, draft, result, analysis_id, pieces=pieces)

    def _close_requoted_cards(
        self,
        revision: ProtocolRevisionRecord,
        pending: Any,
        draft: ProtocolAnalysisDraft,
        result: Any,
        analysis_id: str,
        *,
        pieces: tuple[AnalysisPiece, ...] = (),
    ) -> None:
        """Close the cards of the items the re-ask kept (lane AN5, decision 3)."""

        extra: dict[str, object] = {"review_pending_event": pending.event_id}
        current = self._review_pending_event(revision)
        if current is None or current.event_id != pending.event_id:
            # Registered, analysed again or passed meanwhile: nothing to close.
            self._record_requote(revision, analysis_id, result, **extra, closed="superseded")
            return
        target = self._review_target(revision)
        decisions, memos = (
            self._review_decisions(revision, target) if target is not None else ({}, {}))
        answered = sorted(set(result.kept_items) & set(decisions))
        if answered:
            # "사람이 이미 답한 카드는 그대로 둔다".
            result = resettle_requote(result, answered, outcome=REQUOTE_ANSWERED_BY_PERSON)
        result, closed_draft = self._with_requoted_ocr_values(revision, draft, result)
        kept = result.kept_items
        if answered:
            extra["answered_by_person"] = answered
        if not kept:
            self._record_requote(revision, analysis_id, result, **extra)
            return
        closed = [
            {"item": entry["item"], "field": entry["field"], "page": entry["page"],
             "reason": entry["reason"], "reason_ko": REQUOTE_CLOSED_KO}
            for entry in closed_draft.auto_resolved
            if entry.get("reason") == AUTO_RESOLVED_REQUOTED
        ]
        if not any(
            flag.reason != REVIEW_REASON_POSSIBLY_MISSED for flag in closed_draft.review_flags
        ) and not decisions and not memos:
            # Lane MISS, decision 4: what is left is only lines no quote
            # covers; the analysis passes, and they are its start notice.
            self._record_requote(revision, analysis_id, result, **extra, closed="all")
            self._store_passed_analysis(
                revision, replace(closed_draft, review_flags=()), analysis_id, pieces=pieces)
            return
        self.store.append_event(
            f"analysis-cards-closed-{analysis_id}",
            revision.experiment_id,
            revision.revision_number,
            _REQUOTE_CLOSED_EVENT,
            {"review_pending_event": pending.event_id, "analysis_id": analysis_id,
             "closed": closed, **self._review_pending_record(closed_draft, pieces)},
        )
        self._record_requote(revision, analysis_id, result, **extra, closed="some")
        logging.getLogger(__name__).info(
            "protocol.analysis.review_cards_closed protocol_id=%s closed=%d left=%d",
            revision.experiment_id, len(closed),
            len({flag.item_id for flag in closed_draft.review_flags}),
        )

    def _with_requoted_ocr_values(
        self, revision: ProtocolRevisionRecord, draft: ProtocolAnalysisDraft, result: Any,
    ) -> tuple[Any, ProtocolAnalysisDraft]:
        """The re-ask's result, and its draft with the OCR values read again.

        A quote kept on an OCR page now states values the first pass left to
        its card (lane RV2, decision b). Where reading the values again with
        it gives the same OCR cards, its own values settled by both engines
        are recorded as such; where it would change a card or open one, the
        items kept on OCR pages are left to their cards.
        """

        before = [flag.public_dict() for flag in draft.review_flags
                  if flag.reason == REVIEW_REASON_OCR_VALUE]
        ocr_pages = {page.source_page_number for page in draft.extraction.pages if page.ocr_derived}
        for _ in range(3):
            if not result.kept_items:
                return result, result.draft
            base = replace(
                result.draft,
                review_flags=tuple(flag for flag in result.draft.review_flags
                                   if flag.reason != REVIEW_REASON_OCR_VALUE),
                auto_resolved=tuple(entry for entry in result.draft.auto_resolved
                                    if entry.get("kind") != "ocr_value"),
            )
            checked = self._with_ocr_value_flags(revision, base)
            if [flag.public_dict() for flag in checked.review_flags
                    if flag.reason == REVIEW_REASON_OCR_VALUE] == before:
                return result, checked
            on_ocr = {
                question.item_id for question in result.questions
                if result.pages.get(question.question_id) in ocr_pages
            } & set(result.kept_items)
            result = resettle_requote(
                result, on_ocr or set(result.kept_items), outcome=REQUOTE_OCR_VALUES)
        return resettle_requote(result, set(result.kept_items), outcome=REQUOTE_OCR_VALUES), draft

    def _fewer_items(
        self,
        revision: ProtocolRevisionRecord,
        first: ProtocolAnalysisDraft,
        analysis_id: str,
        *,
        model: ProtocolAnalysisModel,
    ) -> tuple[ProtocolAnalysisDraft, str]:
        """Too many items for cards: one more analysis, or the guidance alone.

        Lane RV2, decision c (human decision 2026-10-16). The documents with
        the most items are mostly steps the model rewrote rather than quoted
        (lane RV replay: 64 items, nearly all "원문에 없음"), where reading
        again does better than a person checking each. The same extraction
        is analysed once more, one reasoning level up, provided the calls it
        may take fit ``review_reanalysis_max_calls()``; the event says why,
        at what level and what came of it. A result within the bound is
        kept (with cards or none); otherwise the analysis fails with the
        guidance and makes no cards. Returns the draft to keep and its id.
        A document read in pieces is never analysed again this way (lane AN3,
        round 2, decision ④): its call is one analysis.
        """

        protocol_id = revision.experiment_id
        limit = review_max_items()
        count, by_reason = _items_by_reason(first)
        planned = 1
        cap = review_reanalysis_max_calls()
        escalated, reasoning = escalated_analysis_model(model)
        again_id = f"{analysis_id}-more-1"
        event: dict[str, object] = {
            "analysis_id": again_id, "reanalysis_of": analysis_id,
            "reason": "too_many_review_items", "items": count, "limit": limit,
            "by_reason": by_reason, "planned_calls": planned, "call_cap": cap, **reasoning,
        }
        after: tuple[int, dict[str, int]] | None = None
        failure_code: str | None = None
        if planned > cap:
            self.store.append_event(
                f"auto-reanalysis-{again_id}", protocol_id, revision.revision_number,
                _AUTO_REANALYSIS_EVENT, {**event, "status": "skipped_over_call_cap"})
        else:
            self.store.append_event(
                f"auto-reanalysis-{again_id}", protocol_id, revision.revision_number,
                _AUTO_REANALYSIS_EVENT, {**event, "status": "analyzing"})
            logging.getLogger(__name__).info(
                "protocol.analysis.auto_reanalysis protocol_id=%s items=%d limit=%d "
                "reasoning_effort=%s planned_calls=%d",
                protocol_id, count, limit, reasoning.get("reasoning_effort"), planned,
            )
            try:
                again = analyze_protocol_extraction(
                    first.extraction, escalated, review_unverified=True,
                    page_range=first.page_range)
                again = self._with_ocr_value_flags(revision, again)
                after = _items_by_reason(again)
            except Exception as exc:  # noqa: BLE001 - the first result's guidance stands
                failure_code = _piece_failure_code(exc)
            self.store.append_event(
                f"auto-reanalysis-ended-{again_id}", protocol_id, revision.revision_number,
                _AUTO_REANALYSIS_EVENT, {
                    **event, "status": "failed" if after is None else "completed",
                    **({"failure_code": failure_code} if failure_code else {}),
                    **({"items_after": after[0], "by_reason_after": after[1]} if after else {}),
                })
            if after is not None and after[0] <= limit:
                return again, again_id
        too_many = ProtocolAnalysisTooManyItemsError(
            "The analysis asks a person about too many items.",
            items=count, limit=limit, items_after=None if after is None else after[0],
            by_reason=by_reason)
        # Its own id: a kept repaired answer is stored under the first
        # attempt's id, whose failure the ledger already holds.
        self._record_analysis_failure(revision, f"{analysis_id}-too-many-items", too_many)
        raise too_many

    def _with_program_address_flags(
        self, revision: ProtocolRevisionRecord, draft: ProtocolAnalysisDraft,
    ) -> ProtocolAnalysisDraft:
        """The draft with an item for each record whose words read as written
        to the program (lane TIDY, item 3; fake-user test F6).

        A step, an action or a statement beside them in one of the shapes of
        ``program_address_flags`` is a card, "이 문장은 실험 지시가 아니라
        프로그램에 하는 말처럼 보여요. 단계에서 뺄까요?", so an analysis that
        would have passed waits for a person. Like the possibly-missed lines,
        these items are made after the bound on cards was applied: they never
        send a document to be read again or refuse it.
        """

        asked = {flag.item_id for flag in draft.review_flags
                 if flag.reason == REVIEW_REASON_ADDRESSED_TO_PROGRAM}
        flags = tuple(flag for flag in program_address_flags(draft.protocol)
                      if flag.item_id not in asked)
        if not flags:
            return draft
        logging.getLogger(__name__).info(
            "protocol.analysis.addressed_to_program protocol_id=%s items=%d",
            revision.experiment_id, len(flags))
        return replace(draft, review_flags=(*draft.review_flags, *flags))

    def _with_possibly_missed_flags(
        self, revision: ProtocolRevisionRecord, draft: ProtocolAnalysisDraft,
    ) -> ProtocolAnalysisDraft:
        """The draft with an item for the lines no quote covers (lane MISS, 3).

        Human decision 2026-10-23: in an analysis that waits for review, each
        place of the source that reads like an instruction and that no quote
        covers (``possibly_missed_places``) is a card, "분석에 없는 원문일 수
        있어요", whose only buttons are [단계가 아니에요] and [다시 분석할게요].
        The lines another card already asks about -- a failed piece's pages,
        where a refused quote's item may stand -- are left to that card
        (``_lines_asked_about``). These items are made after the bound on
        cards was applied (they
        never send a document to be read again or refuse it); their own number
        is held to ``review_max_items()`` by putting places of one page, then
        of a run of pages, on one card (``possibly_missed_flags``).
        """

        places = _within_page_range(possibly_missed_places(
            draft.protocol, draft.extraction, asked_lines=_lines_asked_about(draft)),
            draft.page_range)
        flags = possibly_missed_flags(places, limit=review_max_items())
        if not flags:
            return draft
        logging.getLogger(__name__).info(
            "protocol.analysis.possibly_missed protocol_id=%s places=%d items=%d",
            revision.experiment_id, len(places), len(flags))
        return replace(draft, review_flags=(*draft.review_flags, *flags))

    def _store_review_pending(
        self,
        revision: ProtocolRevisionRecord,
        draft: ProtocolAnalysisDraft,
        analysis_id: str,
        *,
        pieces: tuple[AnalysisPiece, ...] = (),
    ) -> ProtocolCatalogEntry:
        """Keep an analysis with items for a person, and say so (lane RV, a).

        Human decision 2026-10-12/14: the analysis is kept as the model read
        it -- the source and this record are never changed afterwards -- with
        each item the server could not find in the source flagged and why. It
        is not an analysis revision, so nothing that runs a protocol can read
        it; a person's review and registration make the version that runs.
        The source's pages are not copied into the record: their checksums
        are, and the pages are read again from the stored source.
        """

        protocol_id = revision.experiment_id
        payload: dict[str, object] = {
            "status": "review_pending",
            "analysis_id": analysis_id,
            **self._review_pending_record(draft, pieces, protocol_id=protocol_id),
        }
        self.store.append_event(
            f"analysis-review-pending-{analysis_id}",
            protocol_id,
            revision.revision_number,
            _REVIEW_PENDING_EVENT,
            payload,
        )
        logging.getLogger(__name__).info(
            "protocol.analysis.review_pending protocol_id=%s items=%d flags=%d",
            protocol_id, payload["needs_review_count"], len(draft.review_flags),
        )
        return self.get_entry(protocol_id)

    @staticmethod
    def _review_pending_record(
        draft: ProtocolAnalysisDraft,
        pieces: tuple[AnalysisPiece, ...] = (),
        *,
        protocol_id: str | None = None,
    ) -> dict[str, object]:
        """What a review-pending analysis is read back from: its protocol
        without the pages, its items, readiness and settlements (lane RV, a;
        lane AN5, decision 3 records it again with the cards it closed)."""

        protocol = draft.protocol
        if protocol_id is not None and protocol.protocol_id != protocol_id:
            protocol = replace(protocol, protocol_id=protocol_id)
        pdf = protocol.metadata.pdf
        stored = replace(protocol, metadata=replace(
            protocol.metadata, pdf=replace(pdf, pages=())))
        flags = draft.review_flags
        record: dict[str, object] = {
            "flags": [flag.public_dict() for flag in flags],
            "needs_review_count": len({flag.item_id for flag in flags}),
            "verified_evidence_count": draft.verified_evidence_count,
            "capability_policy_id": draft.capability_policy_id,
            "source_pages": [_text_sha256(page.text) for page in pdf.pages],
            "protocol": _encode_stored_record(stored),
            "readiness": _encode_stored_record(draft.readiness),
        }
        if protocol.cleared_fields:
            record["cleared_fields"] = list(protocol.cleared_fields)
        if pieces:
            record["pieces"] = [piece.public_dict() for piece in pieces]
        if draft.page_range is not None:
            # Lane EASY, decision 1: the pages this analysis read.
            record["page_range"] = _page_range_public(draft.page_range, pdf.page_count)
        if draft.response_repair is not None:
            record["response_repaired"] = draft.response_repair.public_dict()
        if draft.auto_resolved:
            # Lane RV2, decision b: what the server settled itself, and why.
            record["auto_resolved"] = [dict(entry) for entry in draft.auto_resolved]
        return record

    def _requested_page_range(
        self, revision: ProtocolRevisionRecord, analysis_id: str,
    ) -> tuple[int, int] | None:
        """The pages the request of ``analysis_id`` chose (lane EASY, 1), or None."""

        return self._analysis_page_range(revision, analysis_id)

    def _analysis_page_range(
        self, revision: ProtocolRevisionRecord, analysis_id: str | None,
    ) -> tuple[int, int] | None:
        """The pages an analysis read, from its request in the ledger.

        Lane EASY, decision 1. ``analysis_id`` may be a run's later attempt
        ("…-retry-1", "…-more-1"): the request it answers is the one whose id
        it begins with. None for a request that chose no pages.
        """

        if not analysis_id:
            return None
        found: tuple[int, int] | None = None
        for event in self.store.list_events(revision.experiment_id):
            if (event.event_type != _ANALYSIS_REQUESTED_EVENT
                    or event.protocol_revision_number != revision.revision_number):
                continue
            payload = event.payload if isinstance(event.payload, dict) else {}
            requested = str(payload.get("analysis_id") or "")
            if requested and (analysis_id == requested or analysis_id.startswith(requested + "-")):
                found = _page_range_from(payload.get("page_range"))
        return found

    def _analyze_in_pieces(
        self,
        revision: ProtocolRevisionRecord,
        extraction: ProtocolPdfExtraction,
        model: ProtocolAnalysisModel,
        pieces: tuple[AnalysisPiece, ...],
        *,
        analysis_id: str,
    ) -> ProtocolAnalysisDraft:
        """Analyse a long document in pieces, ``analysis_piece_concurrency()`` at once.

        Lane AN, decision 2. The plan and each piece's start, pass and
        failure are ledger events of this request's run (``analysis_id``),
        which the run status and the pipeline line read ("조각 k/n 분석 중").
        A piece that fails is sent once more by itself -- the same request,
        or one reasoning level lower when its answer was cut at the output
        limit (lane AQ). Lane RV, decision a (human decision 2026-10-12,
        changing lane AN's "nothing passes on part of it"): a piece that fails
        again no longer fails the document. The pieces that passed are joined
        and the joined analysis is checked against the whole document; each
        failed piece's pages become an item a person checks, so the analysis
        is one waiting for review and does not run as it is. Only when every
        piece failed does the document fail.

        Lane AN2, decision 3: a 429 the provider answered a piece (the
        adapter may have sent it again and got through) halves the number of
        pieces still to start, for this document only; the 429s and the cut
        are written into the piece's event.
        """

        protocol_id = revision.experiment_id
        run = {"analysis_run_id": analysis_id}

        def append(event_id: str, event_type: str, payload: dict[str, object]) -> None:
            self.store.append_event(
                event_id, protocol_id, revision.revision_number, event_type,
                {**run, **payload},
            )

        append(f"pieces-planned-{analysis_id}", _CHUNK_PLAN_EVENT, {
            "status": "planned",
            "chunks": [
                {"chunk_id": f"piece-{piece.index}", "ordinal": piece.index,
                 "source_page_start": piece.first_page,
                 "source_page_end": piece.last_page}
                for piece in pieces
            ],
        })
        # The workers only call the model and check its answer. The ledger
        # connection belongs to this thread, so their starts come back here
        # through ``outbox`` and are written in order.
        outbox: queue.SimpleQueue[tuple[str, str, dict[str, object]]] = queue.SimpleQueue()

        concurrency = analysis_piece_concurrency()
        gate = PieceConcurrencyGate(concurrency)
        #: What the provider answered each piece besides an answer, kept for
        #: its completed or failed event (written by the worker, read here).
        notes: dict[int, dict[str, object]] = {}

        def one(piece: AnalysisPiece) -> ProtocolAnalysisDraft:
            chunk = {"chunk_id": f"piece-{piece.index}", "ordinal": piece.index}
            retries: list[dict[str, object]] = []
            lowered: list[int] = []

            def seen(event: ProviderRetryEvent) -> None:
                retries.append(event.public_dict())
                if event.reason != "rate_limited":
                    return
                cut = gate.rate_limited()
                if cut is not None:
                    lowered.append(cut)
                    logging.getLogger(__name__).info(
                        "protocol.analysis.piece_concurrency_lowered protocol_id=%s piece=%d to=%d",
                        protocol_id, piece.index, cut,
                    )

            try:
                with gate, observe_provider_retries(seen):
                    observed = _ObservedAnalysisModel.wrap(model)
                    outbox.put((f"piece-started-{analysis_id}-{piece.index}", _CHUNK_STARTED_EVENT,
                                {**chunk, "status": "in_progress", "attempt": 1}))
                    try:
                        return analyze_protocol_piece(
                            extraction, piece, observed.model, review_unverified=True)
                    except Exception as exc:  # noqa: BLE001 - sent once more by itself
                        if safety_block_provider(exc) is not None:
                            # Lane RV2, decision f: a prompt the safety filter
                            # blocks is not sent again here. The adapter has
                            # already sent it once more (lane AN3, decision C:
                            # lane GM saw a block answered on another send);
                            # blocked twice, its pages are an item.
                            raise
                        retry_model, retry_payload = observed.retry()
                        outbox.put((f"piece-retry-{analysis_id}-{piece.index}", _CHUNK_STARTED_EVENT, {
                            **chunk, "status": "in_progress", "attempt": 2,
                            "retry_of_failure_code": _piece_failure_code(exc), **retry_payload,
                        }))
                    return analyze_protocol_piece(
                        extraction, piece, retry_model, review_unverified=True)
            finally:
                if retries:
                    notes[piece.index] = {
                        "provider_retries": retries,
                        **({"concurrency_lowered_to": lowered[-1]} if lowered else {}),
                    }

        def drain() -> None:
            while True:
                try:
                    event_id, event_type, payload = outbox.get_nowait()
                except queue.Empty:
                    return
                append(event_id, event_type, payload)

        drafts: dict[int, ProtocolAnalysisDraft] = {}
        failures: dict[int, BaseException] = {}
        with ThreadPoolExecutor(
            max_workers=concurrency,
            thread_name_prefix="protocol-piece-analysis",
        ) as pool:
            futures = {pool.submit(one, piece): piece for piece in pieces}
            pending = set(futures)
            while pending:
                done, pending = wait(pending, timeout=0.5, return_when=FIRST_COMPLETED)
                drain()
                for future in sorted(done, key=lambda item: futures[item].index):
                    piece = futures[future]
                    chunk = {"chunk_id": f"piece-{piece.index}", "ordinal": piece.index}
                    try:
                        draft = future.result()
                    except Exception as exc:  # noqa: BLE001 - the piece's result
                        failures[piece.index] = exc
                        failure: dict[str, object] = {
                            **chunk, "status": "failed", "attempt": 2,
                            "failure_code": _piece_failure_code(exc),
                        }
                        if isinstance(exc, ProtocolAnalysisEvidenceError):
                            failure["evidence_failure"] = exc.diagnostic.privacy_safe_dict()
                        failure.update(notes.get(piece.index, {}))
                        append(f"piece-failed-{analysis_id}-{piece.index}",
                               _CHUNK_FAILED_EVENT, failure)
                        continue
                    drafts[piece.index] = draft
                    append(f"piece-completed-{analysis_id}-{piece.index}",
                           _CHUNK_COMPLETED_EVENT, {
                               **chunk, "status": "completed",
                               "steps": sum(
                                   len(section.steps) for section in draft.protocol.sections),
                               **notes.get(piece.index, {}),
                           })
        failed = tuple(piece for piece in pieces if piece.index in failures)
        if len(failed) == len(pieces):
            raise ProtocolAnalysisPieceFailedError(
                "A piece of the document failed analysis after its retry.",
                pieces=failed,
                failure_codes=tuple(
                    _piece_failure_code(failures[piece.index]) for piece in failed),
            )
        # Lane RV, decision a: the pieces that passed are kept and joined;
        # each failed piece's pages wait for a person ("41–59쪽 분석 실패").
        passed = tuple(piece for piece in pieces if piece.index in drafts)
        append(f"pieces-joining-{analysis_id}", _MERGE_STARTED_EVENT, {"status": "in_progress"})
        try:
            return join_piece_analyses(
                extraction, passed, [drafts[piece.index] for piece in passed],
                review_unverified=True,
                failed_pieces=tuple(
                    (piece, _piece_failure_code(failures[piece.index]),
                     safety_block_provider(failures[piece.index]))
                    for piece in failed),
            )
        except ProtocolAnalysisPieceConflictError as exc:
            append(f"pieces-conflict-{analysis_id}", _MERGE_CONFLICT_EVENT, {
                "status": "conflict", "failure_code": exc.code, "reason": exc.reason,
                "source_page_start": exc.first_page, "source_page_end": exc.last_page,
            })
            raise

    def _store_passed_analysis(
        self,
        revision: ProtocolRevisionRecord,
        draft: ProtocolAnalysisDraft,
        analysis_id: str,
        *,
        pieces: tuple[AnalysisPiece, ...] = (),
    ) -> ProtocolCatalogEntry:
        """Store a passed analysis as the revision's next analysis, and say so."""

        protocol_id = revision.experiment_id
        if draft.protocol.protocol_id != protocol_id:
            assigned_protocol = replace(draft.protocol, protocol_id=protocol_id)
            domain.validate_protocol(assigned_protocol)
            draft = replace(draft, protocol=assigned_protocol)
        analysis = self.store.append_analysis_revision(
            protocol_id,
            revision.revision_number,
            analysis_id,
            draft.protocol,
            draft.readiness,
            draft.capability_policy_id,
        )
        ready_payload: dict[str, object] = {
            "status": "analysis_ready",
            "analysis_payload_sha256": analysis.payload_sha256,
        }
        if draft.protocol.cleared_fields:
            # Lane EV2, decision 2: the fields emptied because the source
            # did not confirm them, kept in the ledger beside the analysis.
            ready_payload["cleared_fields"] = list(draft.protocol.cleared_fields)
        if pieces:
            # Lane AN, decision 2: read in these pieces and joined.
            ready_payload["pieces"] = [piece.public_dict() for piece in pieces]
        if draft.auto_resolved:
            # Lane RV2, decision b: what the server settled itself, and why.
            ready_payload["auto_resolved"] = [dict(entry) for entry in draft.auto_resolved]
        places = _within_page_range(
            possibly_missed_places(draft.protocol, draft.extraction), draft.page_range)
        if draft.page_range is not None:
            # Lane EASY, decision 1: the pages this analysis read.
            ready_payload["page_range"] = _page_range_public(
                draft.page_range, draft.extraction.page_count)
        if places:
            # Lane MISS, decision 4: lines no quote covers that read like an
            # instruction; told before the start, never a card or a blocker.
            ready_payload["possibly_missed"] = possibly_missed_summary(places)
            logging.getLogger(__name__).info(
                "protocol.analysis.possibly_missed protocol_id=%s places=%d pages=%d",
                protocol_id, len(places), len({place.page_number for place in places}))
        if draft.response_repair is not None:
            # Lane EV3, decision 1: the response was read after its brackets
            # were repaired in one of the two fixed shapes; every check above
            # ran on the repaired response like on any other.
            ready_payload["response_repaired"] = draft.response_repair.public_dict()
            logging.getLogger(__name__).info(
                "protocol.analysis.response_repaired protocol_id=%s shapes=%s brackets=%d",
                protocol_id, ",".join(draft.response_repair.shapes),
                draft.response_repair.bracket_count,
            )
        self.store.append_event(
            f"analysis-ready-{analysis_id}",
            protocol_id,
            revision.revision_number,
            _ANALYSIS_READY_EVENT,
            ready_payload,
            analysis_revision_number=analysis.analysis_revision_number,
        )
        self.store.append_event(
            f"review-required-{analysis_id}",
            protocol_id,
            revision.revision_number,
            _SINGLE_REVIEW_REQUIRED_EVENT,
            {
                "status": "review_required",
                "analysis_payload_sha256": analysis.payload_sha256,
            },
            analysis_revision_number=analysis.analysis_revision_number,
        )
        entry = self.get_entry(protocol_id)
        self._analysis_ready(entry)
        return entry

    def _record_analysis_failure(
        self,
        revision: ProtocolRevisionRecord,
        analysis_id: str,
        exc: BaseException,
        *,
        truncated: bool = False,
    ) -> str:
        """Persist one attempt's bounded failure code; return the code.

        ``truncated`` marks an answer the provider cut at its output limit
        (lane AQ, decision 4), so the ledger tells that failure from one
        the model wrote whole.
        """

        failure_code = getattr(exc, "code", "analysis_failed")
        if not isinstance(failure_code, str) or not failure_code:
            failure_code = "analysis_failed"
        blocked_by = safety_block_provider(exc)
        if blocked_by is not None:
            # Lane RV2, decision f: named for what it is.
            failure_code = BLOCKED_FAILURE_CODE
        failure_digest = hashlib.sha256(
            analysis_id.encode("utf-8")
        ).hexdigest()[:24]
        failure_payload: dict[str, object] = {
            "status": "failed",
            "failure_code": failure_code,
        }
        if truncated:
            failure_payload["truncated"] = True
        if blocked_by is not None:
            failure_payload["blocked"] = {"provider": blocked_by}
        repair = getattr(exc, "response_repair", None)
        if repair is not None:
            # Lane EV3, decision 1: this attempt was read after a repair of
            # its brackets and then failed a check.
            failure_payload["response_repaired"] = repair.public_dict()
        if isinstance(exc, ProtocolAnalysisPieceFailedError):
            # Lane AN, decision 2: the pieces that failed after their retry.
            failure_payload["failed_pieces"] = [
                {**piece.public_dict(), "failure_code": code}
                for piece, code in zip(exc.pieces, exc.failure_codes)
            ]
        if isinstance(exc, ProtocolAnalysisPieceConflictError):
            failure_payload["piece_conflict"] = {
                "reason": exc.reason, "source_page_start": exc.first_page,
                "source_page_end": exc.last_page,
            }
            if exc.unlabelled_pieces:
                # Lane RV2: the pieces that read their steps without numbers.
                failure_payload["piece_conflict"]["unlabelled_pieces"] = [
                    piece.public_dict() for piece in exc.unlabelled_pieces]
                failure_payload["piece_conflict"]["labelled_pieces"] = [
                    piece.public_dict() for piece in exc.labelled_pieces]
        if isinstance(exc, ProtocolAnalysisTooManyItemsError):
            # Lane RV2, decision c: the counts the guidance rests on.
            failure_payload["review_items"] = {
                "items": exc.items, "limit": exc.limit, "items_after": exc.items_after,
                "by_reason": exc.by_reason}
        if isinstance(exc, ProtocolAnalysisIncompleteError):
            # Lane EV2, decision 4: the two counts the refusal compared.
            failure_payload["step_counts"] = {
                "source_numbered_steps": exc.source_numbered_steps,
                "analysis_steps": exc.analysis_steps,
            }
        evidence_failure = _safe_evidence_failure(
            exc,
            source_revision=_revision_id(revision.revision_number),
            source_hash=revision.pdf_checksum,
        )
        if evidence_failure is not None:
            failure_payload["evidence_failure"] = evidence_failure
        self.store.append_event(
            f"analysis-failed-{failure_digest}",
            revision.experiment_id,
            revision.revision_number,
            _ANALYSIS_FAILED_EVENT,
            failure_payload,
        )
        return failure_code

    def _automatic_retry_allowed(
        self,
        revision: ProtocolRevisionRecord,
        failure_code: str,
        *,
        repaired: bool = False,
    ) -> bool:
        """Whether this failure is sent again by itself (lane AN, decision 1).

        Mechanical: the failure is a response that broke the structure and no
        automatic retry was sent after the revision's latest analysis request
        (after none, when a caller analyses without recording one).

        Lane EV3, decision 5 (human decision 2026-10-12): a person's "다시
        분석" over a passed analysis (lane RN) is a run of its own and gets
        the same one retry.

        Lane AN (2026-10-13), EV4: so does a person's "분석 다시 시도" after a
        failed analysis -- each request is one run with one automatic retry.
        Before, only the revision's first request was retried, and the
        person's request whose first answer did not close as JSON had to be
        pressed again by hand.

        Lane EV3, decision 1: an answer read only after its brackets were
        repaired (``repaired``) and then refused by any check is sent again
        like the broken answer it was ("고쳐도 안 되면 지금처럼 자동
        재시도한다"); refused whole, an evidence failure still is not.
        """

        if failure_code not in AUTOMATIC_RETRY_FAILURE_CODES and not repaired:
            return False
        events = [
            event for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
        ]
        latest = max(
            (index for index, event in enumerate(events)
             if event.event_type == _ANALYSIS_REQUESTED_EVENT),
            default=-1,
        )
        return not any(
            event.event_type == _ANALYSIS_RETRY_EVENT for event in events[latest + 1:]
        )

    def request_analysis(
        self, protocol_id: str, analysis_id: str, *, reanalysis: bool = False,
        page_range: tuple[int, int] | None = None,
    ) -> ProtocolCatalogEntry:
        """Persist an explicit analysis request without contacting a provider.

        A revision that holds a passed analysis is not asked again by a plain
        request (lane PA, decision 2: an upload never calls a passed analysis
        again). ``reanalysis`` is the person's "다시 분석" (lane RN, decision
        2): one more request event, naming the analysis it is asked over; a
        pass then becomes one more analysis revision beside the old one, and
        a failure leaves the old one as it is. The curated development
        fixture is never analysed again.

        ``page_range`` (lane EASY, decision 1, human decision 2026-10-24) is
        the pages a person chose to analyse -- a paper's Methods pages -- kept
        in the request event; the analysis then reads those pages alone. None
        is the whole document, as before.
        """

        revision = self._latest_protocol_revision(protocol_id)
        previous = self._latest_analysis(revision)
        if previous is not None and not reanalysis:
            return self.get_entry(protocol_id)
        if not reanalysis and (
            self._review_pending_event(revision) is not None
            or self._review_versions(revision)
        ):
            # Lane RV, decisions a and d: an analysis waiting for a person, or
            # a registered review version, is a result; a plain request (an
            # upload of the same file) does not ask again. "다시 분석" does.
            return self.get_entry(protocol_id)
        if previous is not None and any(
            event.event_type == _DEVELOPMENT_FIXTURE_EVENT
            and event.protocol_revision_number == revision.revision_number
            for event in self.store.list_events(revision.experiment_id)
        ):
            raise ProtocolReanalysisRefusedError(
                "The curated development fixture is not analysed again."
            )
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        source = self.store.file_store.object_path(
            revision.pdf_checksum, expected_size=pdf_object.byte_size
        )
        extraction = extract_protocol_pdf(source)
        try:
            page_range = checked_page_range(extraction, page_range)
        except ProtocolAnalysisPageRangeError as exc:
            raise ProtocolAnalysisPageRangeInvalidError(str(exc)) from exc
        if (
            _analysis_state(extraction) == "ocr_required"
            and self._ocr_projection(
                revision, extraction, include_text=False
            ).get("accepted_for_analysis")
            is not True
        ):
            raise ProtocolOcrRequiredError(
                "Protocol requires reviewed OCR before structured analysis."
            )
        request_payload: dict[str, object] = {
            "status": "analysis_pending", "analysis_id": analysis_id,
        }
        if page_range is not None:
            request_payload["page_range"] = _page_range_public(page_range, extraction.page_count)
        if previous is not None:
            request_payload["reanalysis"] = True
            request_payload["previous_analysis_revision_number"] = (
                previous.analysis_revision_number)
            request_payload["previous_analysis_payload_sha256"] = previous.payload_sha256
        self.store.append_event(
            f"analysis-requested-{analysis_id}",
            protocol_id,
            revision.revision_number,
            _ANALYSIS_REQUESTED_EVENT,
            request_payload,
        )
        return self.get_entry(protocol_id)

    def check_page_range(
        self, protocol_id: str, page_range: tuple[int, int] | None,
    ) -> tuple[int, int] | None:
        """``page_range`` when it names pages of the protocol's source (lane EASY, 1)."""

        if page_range is None:
            return None
        revision = self._latest_protocol_revision(protocol_id)
        try:
            return checked_page_range(self._source_extraction(revision), page_range)
        except ProtocolAnalysisPageRangeError as exc:
            raise ProtocolAnalysisPageRangeInvalidError(str(exc)) from exc

    def fail_analysis_request(
        self,
        protocol_id: str,
        analysis_id: str,
        *,
        failure_code: str,
    ) -> ProtocolCatalogEntry:
        """Persist only a bounded failure code for an analysis that did not start."""

        if not re.fullmatch(r"[a-z][a-z0-9_]{0,63}", failure_code):
            failure_code = "analysis_failed"
        revision = self._latest_protocol_revision(protocol_id)
        self.store.append_event(
            f"analysis-failed-{analysis_id}",
            protocol_id,
            revision.revision_number,
            _ANALYSIS_FAILED_EVENT,
            {"status": "failed", "failure_code": failure_code},
        )
        return self.get_entry(protocol_id)

    def load_executable_fixture(
        self, protocol_id: str
    ) -> CuratedProtocolFixture:
        revision = self._latest_protocol_revision(protocol_id)
        entry = self._entry_for_revision(revision)
        versions = self._review_versions(revision)
        if versions:
            # Lane RV, decision c: the latest registered review version runs.
            event, analysis = versions[-1]
            return self._fixture_for_analysis(
                revision, analysis, entry, status="analysis_passed",
                version=_version_payload(event))
        analysis = self._latest_analysis(revision)
        if analysis is None or not entry.available_for_execution:
            raise ProtocolCatalogUnavailableError(
                "Protocol analysis has not passed, or a page could not be read."
            )
        return self._fixture_for_analysis(
            revision, analysis, entry, status="analysis_passed"
        )

    def load_fixture_for_revision(
        self, protocol_id: str, revision_id: str,
    ) -> CuratedProtocolFixture:
        """The fixture of the exact revision an experiment started on (lane RV, f).

        Lane RN found that after an analysis passed again, continuing an
        experiment started on the older one was refused: the session loaded
        the latest fixture and the experiment named the older revision. An
        experiment is continued on the revision it started on -- an analysis
        that passed with no execution blocker, or a registered review
        version -- never on one that waits for review or did not pass.
        """

        revision = self._latest_protocol_revision(protocol_id)
        match = _REVISION_ID.fullmatch(revision_id) if isinstance(revision_id, str) else None
        if match is None or int(match[1]) != revision.revision_number:
            raise ProtocolCatalogNotFoundError("Protocol revision is unknown.")
        entry = self._entry_for_revision(revision)
        if match[3] is not None:
            for event, analysis in self._review_versions(revision):
                if event.payload.get("revision_id") == revision_id:
                    return self._fixture_for_analysis(
                        revision, analysis, entry, status="analysis_passed",
                        version=_version_payload(event))
            raise ProtocolCatalogNotFoundError("Protocol revision is unknown.")
        if match[2] is None:
            raise ProtocolCatalogUnavailableError("This revision has no analysis that runs.")
        analysis = next(
            (
                item for item in self.store.list_analysis_revisions(
                    protocol_id, revision.revision_number)
                if item.analysis_revision_number == int(match[2])
                and not item.analysis_id.startswith(REVIEW_VERSION_ANALYSIS_PREFIX)
            ),
            None,
        )
        if analysis is None:
            raise ProtocolCatalogNotFoundError("Protocol revision is unknown.")
        if domain.execution_blocking_reasons(analysis.readiness):
            raise ProtocolCatalogUnavailableError("This revision's analysis does not run.")
        return self._fixture_for_analysis(
            revision, analysis, entry, status="analysis_passed", revision_id=revision_id)

    def load_analysis_fixture(self, protocol_id: str) -> CuratedProtocolFixture:
        """The fixture a session *would* run for the latest passed analysis.

        For translation only (lane PX, decision 1): the same steps, facts,
        ``revision_id`` and ``fixture_sha256`` the executable fixture will
        carry, so the Korean made from it is the Korean the session shows --
        but marked ``analysis_draft``, and never handed to a session: no
        authority is read or granted here. Refuses while no analysis has
        passed.
        """

        revision = self._latest_protocol_revision(protocol_id)
        versions = self._review_versions(revision)
        if versions:
            # Lane RV, decision c: the version a session will run.
            event, analysis = versions[-1]
            return self._fixture_for_analysis(
                revision, analysis, self._entry_for_revision(revision),
                status="analysis_passed", version=_version_payload(event))
        analysis = self._latest_analysis(revision)
        if analysis is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol analysis is required before translation."
            )
        entry = self._entry_for_revision(revision)
        return self._fixture_for_analysis(
            revision, analysis, entry,
            status="analysis_passed" if entry.available_for_execution
            else "analysis_draft",
        )

    def _fixture_for_analysis(
        self,
        revision: ProtocolRevisionRecord,
        analysis: AnalysisRevisionRecord,
        entry: ProtocolCatalogEntry,
        *,
        status: str,
        version: dict[str, object] | None = None,
        revision_id: str | None = None,
    ) -> CuratedProtocolFixture:
        protocol_id = revision.experiment_id
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol source object is unavailable."
            )
        source = self.store.file_store.object_path(
            revision.pdf_checksum, expected_size=pdf_object.byte_size
        )
        extraction = extract_protocol_pdf(source)
        draft_fields = dict(
            extraction=extraction,
            protocol=protocol_with_display_step_labels(analysis.protocol),
            readiness=analysis.readiness,
            capability_policy=domain.P1_CAPABILITY_POLICY,
            analysis_schema_version=analysis.analysis_schema_version,
            verified_evidence_count=0,
        )
        draft = (
            ReviewedAnalysisDraft(**draft_fields, review_version=_review_version_info(version))
            if version else ProtocolAnalysisDraft(**draft_fields))
        labels = tuple(
            step.source_label
            for section in draft.protocol.sections
            for step in section.steps
        )
        timers = self._step_timer_table(
            revision, extraction, draft.protocol, version=version)
        return CuratedProtocolFixture(
            draft=draft,
            status=status,
            ordered_step_labels=labels,
            fixture_sha256=analysis.payload_sha256,
            revision_id=(
                str(version["revision_id"]) if version
                else revision_id or entry.revision_id),
            development_only=entry.development_only,
            source_pdf_path=source,
            source_pdf_sha256=revision.pdf_checksum,
            source_filename=revision.original_filename,
            # The timers this analysis's own source supports (lane PT,
            # decision 1): the same step_id -> seconds shape the in-gel
            # sidecar loads into, plus the steps whose timer is a choice.
            timer_manifest=timers.manifest(),
            timer_choices=timers.choices() or None,
            timer_table=timers,
        )

    def _step_timer_table(
        self,
        revision: ProtocolRevisionRecord,
        extraction: ProtocolPdfExtraction,
        protocol: domain.ExperimentProtocol,
        *,
        version: dict[str, object] | None = None,
    ) -> domain.StepTimerTable:
        """Every time the analysis attached to a step, verified or refused.

        Checked against the text the analysis itself was checked against --
        the accepted OCR for a page that had none -- so a scanned page's time
        is found where its claims were found. A registered review version's
        timers (lane RV, decision c) replace a step's own: a time a reviewer
        entered or fixed runs as the reviewer's, and a step whose timer the
        reviewer left out has none.
        """

        try:
            source = self._extraction_for_analysis(revision, extraction)
        except ProtocolOcrReviewError:
            source = extraction
        table = verify_step_timers(protocol, source)
        timers = (version or {}).get("timers")
        if isinstance(timers, dict) and timers:
            table = _with_reviewer_timers(table, protocol, timers)
        return table

    # --- The review tab (lane RV, decision b) ---------------------------------

    def _review_target(self, revision: ProtocolRevisionRecord) -> _ReviewTarget | None:
        """The analysis a person reviews now: the latest of the analysis that
        passed, the one waiting for review, and (decision c) the latest
        registered review version. None before any, and for the curated
        development fixture, which is never reviewed."""

        events = tuple(
            event for event in self.store.list_events(revision.experiment_id)
            if event.protocol_revision_number == revision.revision_number
        )
        if any(event.event_type == _DEVELOPMENT_FIXTURE_EVENT for event in events):
            return None
        latest = None
        for event in events:
            if event.event_type in {
                _ANALYSIS_READY_EVENT, _REVIEW_PENDING_EVENT, _REVIEW_VERSION_EVENT,
            }:
                latest = event
        if latest is None:
            return None
        if latest.event_type == _REVIEW_VERSION_EVENT:
            # Reviewed again: "다시 고치면 v2 가 되고 v1 기록은 남는다".
            payload = latest.payload if isinstance(latest.payload, dict) else {}
            analysis = self.store.get_analysis_revision(
                revision.experiment_id, revision.revision_number,
                latest.analysis_revision_number,
            )
            return _ReviewTarget(
                kind="review_version",
                key={"kind": "review_version", "version": payload.get("version")},
                protocol=analysis.protocol, readiness=analysis.readiness, flags=(),
                recorded_at=latest.recorded_at,
                revision_id=str(payload.get("revision_id")),
                analysis_id=analysis.analysis_id,
                origins=dict(payload.get("origins") or {}),
                timers=dict(payload.get("timers") or {}),
            )
        if latest.event_type == _REVIEW_PENDING_EVENT:
            protocol, readiness, flags, event = self.review_pending_analysis(
                revision.experiment_id)
            closed = _cards_closed_event(event, events)
            return _ReviewTarget(
                kind="review_pending",
                key={"kind": "review_pending", "event_id": event.event_id},
                protocol=protocol, readiness=readiness, flags=flags,
                recorded_at=event.recorded_at,
                analysis_id=str((event.payload or {}).get("analysis_id") or ""),
                closed={
                    str(entry["item"]): {**entry, "recorded_at": closed.recorded_at}
                    for entry in closed.payload.get("closed") or () if isinstance(entry, dict)
                } if closed is not None else {},
            )
        analysis = self.store.get_analysis_revision(
            revision.experiment_id, revision.revision_number,
            latest.analysis_revision_number,
        )
        return _ReviewTarget(
            kind="analysis",
            key={"kind": "analysis",
                 "analysis_revision_number": analysis.analysis_revision_number},
            protocol=analysis.protocol, readiness=analysis.readiness, flags=(),
            recorded_at=latest.recorded_at,
            revision_id=_revision_id(
                revision.revision_number, analysis.analysis_revision_number),
            analysis_id=analysis.analysis_id,
        )

    def _review_decisions(
        self, revision: ProtocolRevisionRecord, target: _ReviewTarget,
    ) -> tuple[dict[str, dict[str, object]], dict[str, dict[str, object]]]:
        """The draft decisions on the target, latest per item, and the memos."""

        decisions: dict[str, dict[str, object]] = {}
        memos: dict[str, dict[str, object]] = {}
        for event in self.store.list_events(revision.experiment_id):
            if (
                event.event_type != _REVIEW_DECISION_EVENT
                or event.protocol_revision_number != revision.revision_number
                or not isinstance(event.payload, dict)
                or event.payload.get("target") != target.key
            ):
                continue
            decision = {**event.payload, "recorded_at": event.recorded_at,
                        "event_id": event.event_id, "sequence_id": event.sequence_id}
            item_id = str(decision.get("item_id"))
            if decision.get("action") == "memo":
                memos[item_id] = decision
            elif decision.get("action") == "reset":
                decisions.pop(item_id, None)
            else:
                decisions[item_id] = decision
        return decisions, memos

    def _value_change_links(
        self,
        revision: ProtocolRevisionRecord,
        target: _ReviewTarget,
        decisions: dict[str, dict[str, object]],
    ) -> tuple[dict[str, list[dict[str, object]]], dict[str, dict[str, object]]]:
        """The items a fixed value makes wait for a person again (lane RV2, d).

        Human decision 2026-10-16: a value a person fixed is said by the
        step sentence, not by the value record (lane RV judgement 11), so
        every item that states the old value -- the step sentence, its
        sub-actions, and the step's timer; for an OCR value, every item on
        its page that states it -- is asked again, with the sentence the
        server makes by putting the new value in the old one's place when
        the old value stands in it exactly once (otherwise a person writes
        it). A decision on such an item made before the fix no longer
        counts. Returns the links by item and the decisions that still
        count. A safety statement is never linked: its words never change.
        """

        records = _review_records(target.protocol)
        by_id = {record.item_id: record for record in records}
        ocr_flags = {flag.item_id: flag for flag in target.flags
                     if flag.reason == REVIEW_REASON_OCR_VALUE}
        table = self._review_timer_table(revision, target)
        links: dict[str, list[dict[str, object]]] = {}
        for item_id, decision in decisions.items():
            value = decision.get("value") if isinstance(decision.get("value"), dict) else {}
            original = decision.get("original") if isinstance(decision.get("original"), dict) else {}
            new, old = value.get("text"), original.get("text")
            if decision.get("action") != "fix" or not isinstance(new, str) or not isinstance(old, str):
                continue
            if item_id in ocr_flags:
                extra = ocr_flags[item_id].extra or {}
                linked = [
                    other for other in extra.get("items") or ()
                    if other in by_id and by_id[other].kind in _FIXABLE_KINDS]
            elif item_id in by_id and by_id[item_id].kind in {"value", "duration"}:
                record = by_id[item_id]
                linked = [
                    other.item_id for other in records
                    if other.item_id != item_id
                    and other.kind in _FIXABLE_KINDS - {"value", "duration"}
                    and _same_scope(other, record)]
            else:
                continue
            old_tokens = [token for token, _ in ocr_value_tokens(old)]
            new_values = ocr_value_tokens(new)
            removed = [token for token in old_tokens
                       if token not in {token for token, _ in new_values}]
            added = [words for token, words in new_values if token not in old_tokens]
            if not removed:
                continue
            sequence = decision.get("sequence_id")
            stated: list[tuple[str, str, int | None]] = []
            for other in linked:
                record = by_id[other]
                text = _record_words(record) or ""
                spans = [span for span in _ocr_value_spans(text) if span[0] in removed]
                if not spans:
                    continue
                candidate = None
                if len(spans) == 1 and len(removed) == 1 and len(added) == 1:
                    _, start, end = spans[0]
                    candidate = text[:start] + added[0] + text[end:]
                stated.append((other, candidate, None))
                step = record.step
                if step is not None and record.kind in {"step", "sub_action"}:
                    timers = table.for_step(step.step_id)
                    if any(token in removed for timer in timers
                           for token, _ in ocr_value_tokens(timer.literal)):
                        stated.append((
                            f"timer:{step.step_id}", None,
                            _duration_seconds(added[0]) if len(added) == 1 else None))
            for other, candidate, seconds in stated:
                link = {"source_item": item_id, "from": old, "to": new,
                        "candidate": candidate, "candidate_seconds": seconds}
                if link not in links.setdefault(other, []):
                    links[other].append(link)
                earlier = decisions.get(other)
                if (
                    earlier is not None and isinstance(sequence, int)
                    and isinstance(earlier.get("sequence_id"), int)
                    and earlier["sequence_id"] < sequence
                ):
                    earlier["stale"] = True
        for item_links in links.values():
            if len(item_links) > 1:
                for link in item_links:
                    link["candidate"] = None
        return links, {
            item_id: decision for item_id, decision in decisions.items()
            if not decision.get("stale")
        }

    def _review_timer_table(
        self, revision: ProtocolRevisionRecord, target: _ReviewTarget,
    ) -> domain.StepTimerTable:
        extraction = target.protocol.metadata.pdf
        try:
            table = verify_step_timers(
                _safe_page_view(target.protocol, extraction), extraction)
        except (IndexError, ValueError):
            return domain.StepTimerTable()
        if target.timers:
            table = _with_reviewer_timers(table, target.protocol, target.timers)
        return table

    def _review_items(
        self,
        revision: ProtocolRevisionRecord,
        target: _ReviewTarget,
        decisions: dict[str, dict[str, object]],
        memos: dict[str, dict[str, object]],
        *,
        timer_steps: Sequence[str] = (),
        links: dict[str, list[dict[str, object]]] | None = None,
    ) -> list[dict[str, object]]:
        """Every item of the target as the review screen shows it.

        A step's timer is an item where the step has a time (verified or
        refused) or a decision; ``timer_steps`` adds more steps' timers, so a
        reviewer may enter one where the source prints none.
        """

        protocol = target.protocol
        extraction = protocol.metadata.pdf
        flags_by_item: dict[str, list[ReviewFlag]] = {}
        for flag in target.flags:
            flags_by_item.setdefault(flag.item_id, []).append(flag)
        for item_id, item_links in (links or {}).items():
            # Lane RV2, decision d: an item that states a value a person
            # fixed waits for a person again.
            for link in item_links:
                flags_by_item.setdefault(item_id, []).append(ReviewFlag(
                    item_id=item_id, reason=REVIEW_REASON_VALUE_CHANGED,
                    field="text", reason_code=REVIEW_REASON_VALUE_CHANGED, extra=link))
        origins = target.origins or {}
        items: list[dict[str, object]] = []

        def view(item_id: str, kind: str, text: object, *, step=None, section=None,
                 evidence: domain.SourceEvidence | None = None,
                 evidence_item: str | None = None, extra: dict[str, object] | None = None,
                 ) -> dict[str, object]:
            own = flags_by_item.get(item_id, [])
            quote_refused = evidence_item is not None and any(
                flag.field == "evidence" for flag in flags_by_item.get(evidence_item, []))
            needs_review = bool(own)
            covered_by = (
                evidence_item if quote_refused and not needs_review and evidence_item != item_id
                else None)
            # Lane TIDY, item 3: words asked about only because they read as
            # written to the program are still the source's.
            doubted = any(flag.reason != REVIEW_REASON_ADDRESSED_TO_PROGRAM for flag in own)
            origin = origins.get(item_id) or (
                REVIEW_ORIGIN_AI if doubted or covered_by or (needs_review and quote_refused)
                else REVIEW_ORIGIN_SOURCE)
            source: dict[str, object] | None = None
            if evidence is not None:
                lines = None if quote_refused else _quote_lines(evidence, extraction)
                source = {
                    "page": evidence.source_page_number,
                    "excerpt": evidence.source_excerpt,
                    "continued_page": evidence.continued_on_page_number,
                    "continued_excerpt": evidence.continued_excerpt,
                    "verified": not quote_refused,
                    "lines": (
                        {"page": evidence.source_page_number, "first_line": lines[0],
                         "last_line": lines[1]} if lines else None),
                    "suggested_lines": (
                        _suggested_lines(
                            evidence.source_excerpt or str(text or ""), extraction,
                            (evidence.source_page_number,
                             *(page for flag in own for page in flag.matching_source_pages)))
                        if quote_refused and evidence_item == item_id else None),
                }
            allowed = ["memo", "reset"]
            if kind in _FIXABLE_KINDS:
                allowed[:0] = ["confirm", "fix"]
            elif kind in {"warning", "section", "metadata", "construct", "reminder",
                          "timer_spec", "pages"}:
                allowed[:0] = ["confirm"]
            elif kind == "repeat_count":
                allowed[:0] = (["confirm", "fix", "clear"] if text is not None
                               else ["enter"])
            elif kind == "missing_value":
                allowed[:0] = ["confirm", "enter"] if needs_review else ["enter"]
            elif kind == "timer":
                allowed[:0] = ["fix", "clear"] if text is not None else ["enter"]
            elif kind == "ocr_value":
                # Lane RV2, decision b: the value as read, or as the image shows it.
                allowed[:0] = ["confirm", "fix"]
                if (extra or {}).get("in_safety_statement") is True:
                    # Lane FX4, decision 5: no reading is the source's.
                    allowed.insert(2, "source_differs")
            elif kind in {"seam", "missed"}:
                # Lane AN3, decision B: [이어져요]; [아니요] is a re-analysis.
                # Lane MISS, decision 3: [단계가 아니에요]; [다시 분석할게요]
                # is a re-analysis, and no button writes a step.
                allowed[:0] = ["confirm"]
            if any(flag.reason == REVIEW_REASON_READ_TWO_WAYS for flag in own):
                # Lane RV2, decision e: this reading or the other piece's.
                allowed = ["choose", *(action for action in allowed if action != "confirm")]
            addressed = [flag for flag in own if flag.reason == REVIEW_REASON_ADDRESSED_TO_PROGRAM]
            if needs_review and kind in _CLEARABLE_KINDS and not (
                    kind == "warning" and len(addressed) == len(own)):
                allowed.insert(allowed.index("memo"), "clear")
            elif addressed and kind in {"step", "sub_action"}:
                # Lane TIDY, item 3: words written to the program are taken out
                # of the steps when a person says so ([빼 주세요]). A safety
                # statement the source prints is not (AGENTS.md rule 5).
                allowed.insert(allowed.index("memo"), "clear")
            decision = decisions.get(item_id)
            current: object = text
            origin_after = origin
            if decision is not None:
                action = decision.get("action")
                value = decision.get("value") if isinstance(decision.get("value"), dict) else {}
                if action in {"fix", "enter"}:
                    current = value.get("text")
                    origin_after = REVIEW_ORIGIN_REVIEWER
                elif action == "clear":
                    current = None
                elif action == "choose":
                    current = value.get("text")
                elif action == "source_differs":
                    # Lane FX4, decision 5: the value stays as read, for the source.
                    extra = {**(extra or {}), "source_check": {
                        "state_ko": SOURCE_CHECK_STATE_KO,
                        "guidance_ko": SOURCE_CHECK_GUIDANCE_KO}}
            reading_flag = next(
                (flag for flag in own if flag.reason == REVIEW_REASON_READ_TWO_WAYS), None)
            if reading_flag is not None:
                # Lane RV2, decision e: both pieces' readings, each with the
                # source lines its quotes stand on.
                readings = (reading_flag.extra or {}).get("readings") or []
                extra = {**(extra or {}), "readings": [
                    [{**step, "lines": _reading_lines(step, extraction)} for step in group]
                    for group in readings]}
            closed = (target.closed or {}).get(item_id)
            if closed is not None and not needs_review:
                # Lane AN5, decision 3: the re-ask found its source after the
                # card was shown, and closed it.
                extra = {**(extra or {}), "closed": {
                    "by": "server", "reason": closed.get("reason"), "label_ko": REQUOTE_CLOSED_KO,
                    "page": closed.get("page"), "recorded_at": closed.get("recorded_at")}}
            return {
                "item_id": item_id,
                "kind": kind,
                "kind_ko": REVIEW_ITEM_KIND_KO.get(kind, kind),
                "step_id": step.step_id if step is not None else None,
                "step_label": step.source_label if step is not None else None,
                "section_id": section.section_id if section is not None else None,
                "text": text,
                "value_kind_ko": (
                    value_kind_ko(str(text)) if kind in {"value", "duration", "timer"}
                    and text is not None else None),
                "source": source,
                "origin": origin,
                "origin_ko": REVIEW_ORIGIN_KO[origin],
                "needs_review": needs_review,
                "covered_by": covered_by,
                "reasons": [flag.public_dict() for flag in own],
                "allowed_actions": allowed,
                "decision": (
                    {key: decision.get(key) for key in (
                        "action", "value", "source_lines", "reviewer_name", "recorded_at")}
                    if decision is not None else None),
                "memo": memos[item_id].get("memo") if item_id in memos else None,
                "current": current,
                "origin_after": origin_after,
                "origin_after_ko": REVIEW_ORIGIN_KO[origin_after],
                "differs_from_original": bool(
                    decision is not None and decision.get("action") in {"fix", "enter", "clear"}),
                **(extra or {}),
            }

        listed: set[str] = set()
        if "metadata" in flags_by_item:
            items.append(view(
                "metadata", "metadata", protocol.metadata.title,
                evidence=protocol.metadata.evidence, evidence_item="metadata"))
            listed.add("metadata")
        for record in _review_records(protocol):
            value = record.record
            if record.kind == "repeat_count":
                text: object = value.repeat_count
            elif record.kind == "timer_spec":
                text = value.duration.source_text if value.duration is not None else None
            else:
                text = getattr(value, _REVIEW_TEXT_FIELD[type(value)])
            items.append(view(
                record.item_id, record.kind, text, step=record.step, section=record.section,
                evidence=record.evidence, evidence_item=record.evidence_item_id))
            listed.add(record.item_id)
        steps = {step.step_id: step for section in protocol.sections for step in section.steps}
        table = self._review_timer_table(revision, target)
        timer_steps = list(dict.fromkeys(
            [timer.step_id for timer in table.verified]
            + [time.step_id for time in table.refused if time.step_id]
            + [str(item_id)[len("timer:"):] for item_id in decisions
               if str(item_id).startswith("timer:")]
            + list(timer_steps)
        ))
        for step_id in timer_steps:
            step = steps.get(step_id)
            if step is None:
                continue
            verified = table.for_step(step_id)
            seconds = sorted({value for timer in verified for value in timer.seconds})
            items.append(view(
                f"timer:{step_id}", "timer",
                " · ".join(dict.fromkeys(timer.literal for timer in verified)) or None,
                step=step, evidence=step.evidence, evidence_item=None,
                extra={
                    "seconds": seconds,
                    "candidates": [
                        {"literal": time.literal, "excerpt": time.excerpt,
                         "reason": time.reason, "reason_ko": time.reason_ko,
                         "page": time.page_number}
                        for time in table.refused if time.step_id == step_id
                    ],
                }))
            listed.add(f"timer:{step_id}")
        for item_id, flags in flags_by_item.items():
            if item_id in listed:
                continue
            flag = flags[0]
            if flag.reason == REVIEW_REASON_OCR_VALUE:
                items.append(self._ocr_value_view(view, flag, extraction))
                continue
            if flag.reason == REVIEW_REASON_POSSIBLY_MISSED:
                items.append(self._possibly_missed_view(view, flag, extraction))
                continue
            if flag.reason == REVIEW_REASON_NUMBERING_SEAM:
                # Lane AN3, decision B: the two steps that meet where the
                # numbering changes, with their source lines and the numbers
                # the screen and the run show them by.
                shown_labels = {
                    step.step_id: step.source_label
                    for section in protocol_with_display_step_labels(protocol).sections
                    for step in section.steps}
                # No words of its own: the card's reason says what it is and
                # its two steps are the content (lane AN3, stage ③ found the
                # reason drawn twice).
                items.append(view(item_id, "seam", None, extra={"steps": [
                    {**step, "lines": _reading_lines(step, extraction),
                     "label_shown": shown_labels.get(str(step.get("step_id"))),
                     "label_printed": bool(str(step.get("label") or "").strip())}
                    for step in (flag.extra or {}).get("steps") or () if isinstance(step, dict)
                ]}))
                continue
            items.append(view(
                item_id, "pages" if flag.field == "pages" else "construct",
                flag.reason_ko,
                extra={"first_page": flag.first_page, "last_page": flag.last_page}))
        return items

    @staticmethod
    def _ocr_value_view(
        view: Callable[..., dict[str, object]], flag: ReviewFlag,
        extraction: ProtocolPdfExtraction,
    ) -> dict[str, object]:
        """An OCR value item (lane RV2, decision b): the value as read, the
        items that state it, its page's lines that print it and the OCR
        value boxes to show it beside its piece of the image."""

        extra = flag.extra or {}
        page = flag.page_number
        token = str(extra.get("token") or "")
        lines: list[int] = []
        if isinstance(page, int) and 1 <= page <= extraction.page_count:
            lines = [
                number for number, (_, _, line) in enumerate(
                    _page_lines(extraction.pages[page - 1].text), start=1)
                if token in {key for key, _ in ocr_value_tokens(line)}
            ]
        item = view(flag.item_id, "ocr_value", extra.get("value"), extra={
            "page": page, "token": token, "linked_items": list(extra.get("items") or ()),
            "ocr_state": flag.reason_code,
            # Lane AN3, decision A: each engine's reading in the value's place,
            # and the page's card the value is asked on.
            "readings": [dict(reading) for reading in extra.get("readings") or ()
                         if isinstance(reading, dict)],
            "page_card": f"{OCR_PAGE_CARD_PREFIX}{page}",
            "in_safety_statement": extra.get("in_safety_statement") is True,
        })
        item["source"] = {
            "page": page, "excerpt": None, "continued_page": None, "continued_excerpt": None,
            "verified": True,
            "lines": ({"page": page, "first_line": lines[0], "last_line": lines[0]}
                      if lines else None),
            "value_lines": lines,
            "value_boxes": list(extra.get("value_boxes") or ()),
            "suggested_lines": None,
        }
        return item

    @staticmethod
    def _possibly_missed_view(
        view: Callable[..., dict[str, object]], flag: ReviewFlag,
        extraction: ProtocolPdfExtraction,
    ) -> dict[str, object]:
        """A possibly-missed item (lane MISS, decision 3): the lines no quote
        covers, as the page prints them, with where they stand."""

        extra = flag.extra or {}
        places = [place for place in extra.get("places") or () if isinstance(place, dict)]
        texts: list[str] = []
        for place in places:
            page = place.get("page")
            if not isinstance(page, int) or not 1 <= page <= extraction.page_count:
                continue
            lines = extraction.pages[page - 1].text.split("\n")
            first, last = int(place.get("first_line") or 1), int(place.get("last_line") or 1)
            texts.append("\n".join(" ".join(line.split()) for line in lines[first - 1:last]
                                   if line.strip()))
        shown = "\n…\n".join(text for text in texts if text)
        if len(shown) > _MAX_REVIEW_TEXT:
            shown = shown[:_MAX_REVIEW_TEXT].rstrip() + "…"
        first_page, last_page = extra.get("first_page"), extra.get("last_page")
        item = view(flag.item_id, "missed", shown, extra={
            "places": places, "first_page": first_page, "last_page": last_page,
            "missed_scope": flag.reason_code})
        one = places[0] if len(places) == 1 else None
        item["source"] = {
            "page": first_page, "excerpt": shown if one is not None else None,
            "continued_page": None, "continued_excerpt": None, "verified": True,
            "lines": ({"page": one["page"], "first_line": one["first_line"],
                       "last_line": one["last_line"]} if one is not None else None),
            "suggested_lines": None,
        }
        return item

    def review_draft(self, protocol_id: str) -> dict[str, object]:
        """What the review screen shows (lane RV, decision b): the target, its
        items with their source lines, origins and draft decisions, and
        whether it can be registered now."""

        revision = self._latest_protocol_revision(protocol_id)
        entry = self._entry_for_revision(revision)
        base: dict[str, object] = {
            "protocol_id": protocol_id,
            "title": entry.title,
            "source_filename": entry.source_filename,
            "analysis_status": entry.analysis_status,
            "lifecycle_state": entry.lifecycle_state,
            "available_for_execution": entry.available_for_execution,
            "review_pending": entry.review_pending,
            "target": None,
            "items": [],
            "cards": [],
            "next_card": None,
            "needs_review": [],
            "counts": {"items": 0, "needs_review": 0, "decided": 0, "undecided": 0,
                       "changed": 0},
            "pages": [],
            "registrable": False,
            "registration_blockers": [],
        }
        target = self._review_target(revision)
        if target is None:
            base["message_ko"] = (
                "개발용 고정 프로토콜은 검토하지 않습니다."
                if entry.development_only else "검토할 분석이 아직 없습니다.")
            return base
        decisions, memos = self._review_decisions(revision, target)
        links, decisions = self._value_change_links(revision, target, decisions)
        items = self._review_items(revision, target, decisions, memos, links=links)
        pending = [
            item["item_id"] for item in items
            if item["needs_review"] and item["item_id"] not in decisions
        ]
        cards = review_cards(items)
        held = _source_check_items(items)
        base.update({
            "target": {
                **target.key, "recorded_at": target.recorded_at,
                "revision_id": target.revision_id, "analysis_id": target.analysis_id,
            },
            "items": items,
            # Lane RV2, decision a: one card per item that waits for a person,
            # with the only choices the screen may offer for it.
            "cards": cards,
            "next_card": next(
                (index for index, card in enumerate(cards) if not card["decided"]), None),
            "needs_review": [item["item_id"] for item in items if item["needs_review"]],
            "counts": {
                "items": len(items),
                "needs_review": sum(1 for item in items if item["needs_review"]),
                "decided": len(decisions),
                "undecided": len(pending),
                "changed": sum(1 for item in items if item["differs_from_original"]),
            },
            "pages": self._review_pages(revision, target),
            "registration_blockers": [
                {"item_id": item_id, "message_ko": "확인 필요 항목을 정하지 않았습니다."}
                for item_id in pending
            ] + [
                # Lane FX4, decision 5.
                {"item_id": item_id,
                 "message_ko": f"{SOURCE_CHECK_STATE_KO}: {SOURCE_CHECK_GUIDANCE_KO}",
                 "source_check": True}
                for item_id in held
            ],
            "registrable": not pending and not held,
        })
        return base

    def _review_pages(
        self, revision: ProtocolRevisionRecord, target: _ReviewTarget,
    ) -> list[dict[str, object]]:
        extraction = target.protocol.metadata.pdf
        ocr_pages = self._ocr_pages(revision)
        pages = []
        for page in extraction.pages:
            ocr = ocr_pages.get(page.source_page_number) or {}
            lines = self._line_views(page, ocr)
            pages.append({
                "page": page.source_page_number,
                "ocr": page.ocr_derived,
                "ocr_unclear": _ocr_page_unclear(page, ocr),
                "confidence": ocr.get("confidence") if page.ocr_derived else None,
                "line_count": len(lines),
                "value_line_count": sum(1 for line in lines if line["has_value"]),
                "unclear_line_count": sum(1 for line in lines if line["unclear"]),
                "value_box_count": len(ocr.get("value_boxes") or ()) if page.ocr_derived else 0,
            })
        return pages

    def _ocr_pages(self, revision: ProtocolRevisionRecord) -> dict[int, dict[str, object]]:
        projection = self._ocr_projection(
            revision, self._source_extraction(revision), include_text=True)
        if projection.get("accepted_for_analysis") is not True:
            return {}
        return {
            page["source_page_number"]: page for page in projection.get("pages") or ()
            if isinstance(page, dict) and isinstance(page.get("source_page_number"), int)
        }

    def _ocr_page_states(
        self, revision: ProtocolRevisionRecord, extraction: ProtocolPdfExtraction,
    ) -> tuple[dict[int, str], dict[int, dict[str, object]]]:
        """Each OCR page's state for its values, and the OCR pages (lane RV2, b).

        ``agree``: both engines read the page and their numbers and units
        were the same; ``differ``: both read it and they were not;
        ``one_engine``: one engine read it -- one configured, the other
        failed on the page, or an OCR result stored before the engines were
        recorded.
        """

        ocr_numbers = [page.source_page_number for page in extraction.pages if page.ocr_derived]
        if not ocr_numbers:
            return {}, {}
        projection = self._ocr_projection(
            revision, self._source_extraction(revision), include_text=True)
        pages = {
            page["source_page_number"]: page for page in projection.get("pages") or ()
            if isinstance(page, dict) and isinstance(page.get("source_page_number"), int)
        }
        both = projection.get("pages_read_by_both")
        both = {page for page in both if isinstance(page, int)} if isinstance(both, list) else set()
        states = {}
        for number in ocr_numbers:
            if (pages.get(number) or {}).get("numeric_review_required"):
                states[number] = OCR_PAGE_DIFFER
            elif number in both:
                states[number] = OCR_PAGE_AGREE
            else:
                states[number] = OCR_PAGE_ONE_ENGINE
        return states, pages

    def _with_ocr_value_flags(
        self, revision: ProtocolRevisionRecord, draft: ProtocolAnalysisDraft,
    ) -> ProtocolAnalysisDraft:
        """The draft with its OCR values settled or made items (lane RV2, b).

        Every value an item used in execution states on an OCR page is named
        by its page and token (lane AN3, decision A4: only those). Where both
        engines read it alike in its place (decision A1; by page for a result
        stored before) and the item's quote passed, it is settled and
        recorded with that reason; otherwise it is one item a person checks,
        naming the items that state it and what each engine read. Items
        already waiting for a person are left to their card.
        """

        extraction = draft.extraction
        states, pages = self._ocr_page_states(revision, extraction)
        if not states:
            return draft
        flagged = {flag.item_id for flag in draft.review_flags}
        quote_refused = {flag.item_id for flag in draft.review_flags if flag.field == "evidence"}
        values: dict[tuple[int, str], dict[str, object]] = {}
        page_tokens: dict[int, object] = {}
        for record in _review_records(draft.protocol):
            evidence = record.evidence
            if (
                record.kind not in _OCR_VALUE_EXECUTION_KINDS
                or isinstance(record.record, _OCR_VALUE_NOT_EXECUTED) or evidence is None
                or record.item_id in flagged or record.evidence_item_id in quote_refused
            ):
                continue
            quoted = [evidence.source_page_number]
            if evidence.continued_on_page_number is not None:
                quoted.append(evidence.continued_on_page_number)
            if not any(page in states for page in quoted):
                continue
            text = _record_words(record)
            if not text:
                continue
            label = record.step.source_label if record.kind == "step" and record.step else None
            for token, words in ocr_value_tokens(text, label=label):
                page = next(
                    (number for number in quoted if 1 <= number <= extraction.page_count
                     and token in page_tokens.setdefault(
                         number, numeric_tokens(extraction.pages[number - 1].text))),
                    quoted[0])
                if page not in states:
                    continue
                entry = values.setdefault((page, token), {"value": words, "items": []})
                if record.item_id not in entry["items"]:
                    entry["items"].append(record.item_id)
                if record.kind == _OCR_SAFETY_KIND:
                    entry["in_safety_statement"] = True
        flags: list[ReviewFlag] = []
        settled: list[dict[str, object]] = []
        for (page, token), entry in values.items():
            reason_code, readings = _ocr_value_state(states[page], pages.get(page) or {}, token)
            if reason_code is None:
                settled.append({
                    "kind": "ocr_value", "page": page, "token": token,
                    "value": entry["value"], "items": list(entry["items"]),
                    "reason": AUTO_RESOLVED_OCR_TWO_ENGINES,
                    **({"by": "place"} if readings else {}),
                })
                continue
            boxes = (pages.get(page) or {}).get("value_boxes") or ()
            flags.append(ReviewFlag(
                item_id=f"ocr_value:{page}:{token}",
                reason=REVIEW_REASON_OCR_VALUE,
                field="value",
                reason_code=reason_code,
                page_number=page,
                extra={
                    "value": entry["value"], "token": token, "items": list(entry["items"]),
                    "value_boxes": [
                        index for index, box in enumerate(boxes)
                        if isinstance(box, dict) and token in {
                            key for key, _ in ocr_value_tokens(str(box.get("text") or ""))}
                    ],
                    # Lane AN3, decision A: what each engine read in its place.
                    "readings": readings,
                    # Round 2, decision ③: stated by a safety statement.
                    **({"in_safety_statement": True} if entry.get("in_safety_statement") else {}),
                },
            ))
        if not flags and not settled:
            return draft
        return replace(
            draft,
            review_flags=(*draft.review_flags, *flags),
            auto_resolved=(*draft.auto_resolved, *settled),
        )

    @staticmethod
    def _line_views(page: ProtocolPdfPage, ocr: dict[str, object]) -> list[dict[str, object]]:
        unclear_page = _ocr_page_unclear(page, ocr)
        views = []
        boxes = _line_boxes(page)
        for number, (_, _, text) in enumerate(_page_lines(page.text), start=1):
            has_value = bool(re.search(r"[0-9]", text))
            views.append({
                "number": number, "text": text, "has_value": has_value,
                "unclear": unclear_page and has_value,
                # Lane RX1, decision 2: where the line stands on the page, in
                # points ([x0, y0, x1, y1], y down), or None when the engine
                # kept no block for it.
                "box": boxes[number - 1] if number - 1 < len(boxes) else None,
            })
        return views

    def _renamed_title(self, revision: ProtocolRevisionRecord) -> str | None:
        """The latest name a person gave this protocol, or None (lane RX1, 4)."""

        title = None
        for event in self.store.list_events(revision.experiment_id):
            if event.event_type == PROTOCOL_TITLE_RENAMED_EVENT and isinstance(event.payload, dict):
                candidate = event.payload.get("title")
                if isinstance(candidate, str) and candidate.strip():
                    title = candidate.strip()
        return title

    def rename(
        self, protocol_id: str, title: object, *, actor_principal_id: str | None = None,
    ) -> ProtocolCatalogEntry:
        """A person's own name for a protocol (lane RX1, decision 4).

        Human decision 2026-10-25: when the analysis's title is a cover line
        ("표지 최종 하리"), the screen shows the file name instead and lets a
        person name the protocol. The name is one more append-only ledger
        event with who gave it and the title it replaces; the source, its
        metadata and every analysis stay as they are, and the latest name
        is the title wherever it is read. One line of printable text, at
        most ``PROTOCOL_TITLE_MAX_CHARS`` characters.
        """

        if not isinstance(title, str):
            raise ProtocolTitleInvalidError("A protocol's name is one line of text.")
        name = " ".join(title.strip().split(" "))
        if (not name or len(name) > PROTOCOL_TITLE_MAX_CHARS
                or any(not char.isprintable() for char in name)):
            raise ProtocolTitleInvalidError(
                f"A protocol's name is 1 to {PROTOCOL_TITLE_MAX_CHARS} printable characters.")
        revision = self._latest_protocol_revision(protocol_id)
        before = self._entry_for_revision(revision)
        self.store.append_event(
            f"title-renamed-{secrets.token_hex(12)}", protocol_id, revision.revision_number,
            PROTOCOL_TITLE_RENAMED_EVENT,
            {"title": name, "title_before": before.title,
             "actor_principal_id": actor_principal_id,
             "recorded_at": datetime.now(timezone.utc).isoformat(timespec="seconds")},
        )
        return self.get_entry(protocol_id)

    def source_page_count(self, protocol_id: str) -> int:
        """How many pages the latest revision's source has, before or after an
        analysis (lane RX1, decision 1: the page picker)."""

        revision = self._latest_protocol_revision(protocol_id)
        return self._source_extraction(revision).page_count

    def source_pdf(self, protocol_id: str) -> tuple[Path, str]:
        """The stored source PDF of the latest protocol revision and its checksum."""

        revision = self._latest_protocol_revision(protocol_id)
        pdf_object = self.store.get_pdf_object(revision.pdf_checksum)
        if pdf_object is None:
            raise ProtocolCatalogUnavailableError("Protocol source object is unavailable.")
        return (
            self.store.file_store.object_path(
                revision.pdf_checksum, expected_size=pdf_object.byte_size),
            revision.pdf_checksum,
        )

    def review_source_page(self, protocol_id: str, page_number: int) -> dict[str, object]:
        """One page of the source the target was read from, line by line, with
        line numbers, and the OCR's value boxes (lane RV, decision b; UI2's
        "원문 쪽 전체 줄과 줄 번호", "OCR 흐린 줄 수")."""

        revision = self._latest_protocol_revision(protocol_id)
        target = self._review_target(revision)
        if target is not None:
            extraction = target.protocol.metadata.pdf
        else:
            extraction = self._extraction_for_analysis(
                revision, self._source_extraction(revision))
        if not isinstance(page_number, int) or not 1 <= page_number <= extraction.page_count:
            raise ProtocolCatalogNotFoundError("Protocol source page is unknown.")
        page = extraction.pages[page_number - 1]
        ocr = self._ocr_pages(revision).get(page_number) or {}
        lines = self._line_views(page, ocr)
        boxes = ocr.get("value_boxes") if page.ocr_derived else None
        return {
            "protocol_id": protocol_id,
            "page": page_number,
            "page_count": extraction.page_count,
            "ocr": page.ocr_derived,
            "ocr_unclear": _ocr_page_unclear(page, ocr),
            "confidence": ocr.get("confidence") if page.ocr_derived else None,
            "numeric_review_required": bool(ocr.get("numeric_review_required"))
            if page.ocr_derived else False,
            "lines": lines,
            "unclear_line_count": sum(1 for line in lines if line["unclear"]),
            "render_dpi": ocr.get("render_dpi") if page.ocr_derived else None,
            "value_boxes": [
                {"index": index, "text": box.get("text"), "box": box.get("box")}
                for index, box in enumerate(boxes or ())
                if isinstance(box, dict)
            ],
        }

    def review_value_box_image(
        self, protocol_id: str, page_number: int, box_index: int,
    ) -> bytes:
        """The piece of the page image an OCR value was read from (lane RV, e).

        Lane AN kept, for a photo or handwriting read by OCR, each value it
        read with its box on the page rendered at ``render_dpi``. The review
        shows that piece of the original image beside the value, a little
        wider than the box, rendered from the stored source by the one PDF
        engine and kept in memory only.
        """

        revision = self._latest_protocol_revision(protocol_id)
        ocr = self._ocr_pages(revision).get(page_number)
        if not isinstance(ocr, dict) or ocr.get("provider") == OCR_TEXT_LAYER_PROVIDER:
            raise ProtocolCatalogNotFoundError("The page has no OCR value boxes.")
        boxes = ocr.get("value_boxes")
        dpi = ocr.get("render_dpi")
        if (
            not isinstance(boxes, list) or not isinstance(dpi, int) or isinstance(dpi, bool)
            or dpi <= 0 or not isinstance(box_index, int) or not 0 <= box_index < len(boxes)
        ):
            raise ProtocolCatalogNotFoundError("The page has no such value box.")
        box = boxes[box_index].get("box") if isinstance(boxes[box_index], dict) else None
        if not isinstance(box, list) or len(box) != 4 or not all(
            isinstance(edge, (int, float)) and not isinstance(edge, bool) for edge in box
        ):
            raise ProtocolCatalogNotFoundError("The value box is unreadable.")
        x0, y0, x1, y1 = (float(edge) * 72.0 / dpi for edge in box)
        margin = max(6.0, (y1 - y0) * 0.6)
        path, _ = self.source_pdf(protocol_id)
        (png,) = pdf_text_engine.render_clips_png(
            path, [(page_number, (max(0.0, x0 - margin * 2), max(0.0, y0 - margin),
                                  x1 + margin * 2, y1 + margin))],
            max_side_px=640,
        )
        return png

    def record_review_decision(
        self, protocol_id: str, request: dict[str, object],
    ) -> dict[str, object]:
        """Record one decision on one item of the target (lane RV, decision b).

        One more ledger event: the item, the action, the value, the source
        lines a person pointed at, the original value and origin, and the
        reviewer's name when given. Nothing is changed but the draft; the
        registration (decision c) makes the version that runs.
        """

        if not isinstance(request, dict):
            raise ProtocolReviewError("A review decision must be an object.")
        allowed_keys = {"item_id", "action", "text", "seconds", "count", "source_lines",
                        "memo", "reviewer_name", "target", "reading"}
        if set(request) - allowed_keys:
            raise ProtocolReviewError("A review decision holds an unknown field.")
        revision = self._latest_protocol_revision(protocol_id)
        target = self._review_target(revision)
        if target is None:
            raise ProtocolReviewNotReadyError("There is no analysis to review.")
        if request.get("target") is not None and not _same_target(request["target"], target.key):
            raise ProtocolReviewTargetChangedError(
                "The analysis under review changed; read the review again.")
        decisions, memos = self._review_decisions(revision, target)
        links, decisions = self._value_change_links(revision, target, decisions)
        item_id = request.get("item_id")
        items = {item["item_id"]: item for item in self._review_items(
            revision, target, decisions, memos, links=links,
            timer_steps=(
                (item_id[len("timer:"):],)
                if isinstance(item_id, str) and item_id.startswith("timer:") else ()),
        )}
        if isinstance(item_id, str) and item_id.startswith(OCR_PAGE_CARD_PREFIX):
            return self._confirm_ocr_page(protocol_id, revision, target, request, items)
        if not isinstance(item_id, str) or item_id not in items:
            raise ProtocolReviewError("The review has no such item.")
        item = items[item_id]
        action = request.get("action")
        if action not in REVIEW_ACTIONS or action not in item["allowed_actions"]:
            raise ProtocolReviewError("This action is not allowed on this item.")
        value: dict[str, object] = {}
        if action == "choose":
            reading = request.get("reading")
            readings = item.get("readings") if isinstance(item.get("readings"), list) else []
            if reading not in {"a", "b"} or len(readings) != 2:
                raise ProtocolReviewError("A choice names reading a or b.")
            chosen = readings[0 if reading == "a" else 1]
            value = {"reading": reading,
                     "text": " / ".join(str(step.get("text")) for step in chosen)}
        elif request.get("reading") is not None:
            raise ProtocolReviewError("A reading is named only by a choice.")
        if action == "source_differs" and any(
                request.get(key) is not None for key in ("text", "seconds", "count")):
            # Lane FX4, decision 5: it marks the value; it writes nothing.
            raise ProtocolReviewError("[원문과 달라요] carries no value.")
        if action in {"fix", "enter"}:
            if item["kind"] == "timer":
                seconds = request.get("seconds")
                if (not isinstance(seconds, int) or isinstance(seconds, bool)
                        or not 1 <= seconds <= _MAX_TIMER_SECONDS):
                    raise ProtocolReviewError("A timer needs whole seconds.")
                value = {"seconds": seconds, "text": _duration_words(seconds)}
            elif item["kind"] == "repeat_count":
                count = request.get("count")
                if (not isinstance(count, int) or isinstance(count, bool)
                        or not 1 <= count <= _MAX_REPEAT_COUNT):
                    raise ProtocolReviewError("A repeat count needs a whole number.")
                value = {"count": count, "text": f"{count}회"}
            else:
                limit = (_MAX_REVIEW_VALUE
                         if item["kind"] in {"value", "duration", "missing_value", "ocr_value"}
                         else _MAX_REVIEW_TEXT)
                value = {"text": _clean_review_text(request.get("text"), limit=limit,
                                                    label="Reviewed text")}
                if action == "fix" and value["text"] == item["text"]:
                    raise ProtocolReviewError("A fix must change the text; confirm it instead.")
                if item["kind"] == "ocr_value" and item.get("in_safety_statement") is True and (
                    value["text"] not in {
                        str(reading.get("text")) for reading in item.get("readings") or ()
                        if isinstance(reading, dict) and reading.get("token") != item.get("token")}
                ):
                    # Round 2, decision ③: the other engine's reading, or nothing.
                    raise ProtocolReviewError(
                        "A number in a safety statement takes one of the engines' readings.")
        source_lines = None
        source = item.get("source") if isinstance(item.get("source"), dict) else None
        if action in {"confirm", "fix"} and item["needs_review"] and source is not None \
                and source.get("verified") is False:
            # The quote was not found: a person points at where the item is.
            evidence = source_lines_excerpt(target.protocol.metadata.pdf,
                                            request.get("source_lines"))
            source_lines = {**request["source_lines"], "excerpt": evidence.source_excerpt}
            suggested = source.get("suggested_lines")
            if isinstance(suggested, dict) and suggested.get("unique") and all(
                source_lines.get(key, source_lines.get("first_line")) == suggested.get(key)
                for key in ("page", "first_line", "last_line")
            ):
                # Lane RV2, decision b: the lines the server narrowed it to,
                # which a person said are right -- recorded as such.
                source_lines["suggested_by_server"] = True
        elif request.get("source_lines") is not None:
            raise ProtocolReviewError("Source lines are asked only for an item whose quote was not found.")
        memo = None
        if action == "memo":
            memo = _clean_review_text(request.get("memo"), limit=_MAX_REVIEW_MEMO, label="Memo")
        elif request.get("memo") is not None:
            raise ProtocolReviewError("A memo is its own action.")
        name = request.get("reviewer_name")
        payload: dict[str, object] = {
            "target": target.key,
            "item_id": item_id,
            "kind": item["kind"],
            "action": action,
            "value": value or None,
            "source_lines": source_lines,
            "memo": memo,
            "reviewer_name": reviewer_name(name) if name is not None else None,
            "original": {
                "text": item["text"],
                "origin": item["origin"],
                **({"seconds": item.get("seconds")} if item["kind"] == "timer" else {}),
            },
        }
        self.store.append_event(
            f"review-decision-{secrets.token_hex(16)}",
            protocol_id,
            revision.revision_number,
            _REVIEW_DECISION_EVENT,
            payload,
        )
        draft = self.review_draft(protocol_id)
        return {
            "item": next(entry for entry in draft["items"] if entry["item_id"] == item_id),
            "counts": draft["counts"],
            "registrable": draft["registrable"],
            "target": draft["target"],
            # Lane RV2, decision a: the cards after this one, so a button
            # moves straight on to the next card.
            "cards": draft["cards"],
            "next_card": draft["next_card"],
        }

    def _confirm_ocr_page(
        self,
        protocol_id: str,
        revision: ProtocolRevisionRecord,
        target: _ReviewTarget,
        request: dict[str, object],
        items: dict[str, dict[str, object]],
    ) -> dict[str, object]:
        """[모두 맞아요] on a page's OCR card (lane AN3, decision A2).

        Every value of the page not decided yet is confirmed as the analysis
        read it -- one decision event each, naming the card -- and a value a
        person already decided is left as it was.
        """

        card_id = str(request["item_id"])
        if set(request) - {"item_id", "action", "reviewer_name", "target"}:
            raise ProtocolReviewError("A page's OCR card takes [모두 맞아요] only.")
        if request.get("action") != "confirm":
            raise ProtocolReviewError("This action is not allowed on this item.")
        values = [item for item in items.values()
                  if item["kind"] == "ocr_value" and item.get("page_card") == card_id
                  and item["needs_review"]]
        if not values:
            raise ProtocolReviewError("The review has no such item.")
        name = request.get("reviewer_name")
        for item in values:
            if item.get("decision") is not None:
                continue
            self.store.append_event(
                f"review-decision-{secrets.token_hex(16)}",
                protocol_id,
                revision.revision_number,
                _REVIEW_DECISION_EVENT,
                {
                    "target": target.key,
                    "item_id": item["item_id"],
                    "kind": item["kind"],
                    "action": "confirm",
                    "value": None,
                    "source_lines": None,
                    "memo": None,
                    "reviewer_name": reviewer_name(name) if name is not None else None,
                    "original": {"text": item["text"], "origin": item["origin"]},
                    "page_card": card_id,
                },
            )
        draft = self.review_draft(protocol_id)
        decided = {item["item_id"] for item in values}
        return {
            "item": None,
            "items": [entry for entry in draft["items"] if entry["item_id"] in decided],
            "counts": draft["counts"],
            "registrable": draft["registrable"],
            "target": draft["target"],
            "cards": draft["cards"],
            "next_card": draft["next_card"],
        }

    def register_review(
        self, protocol_id: str, request: dict[str, object],
    ) -> dict[str, object]:
        """"검토 완료하고 등록": the review becomes version N (lane RV, decision c).

        Human decision 2026-10-12: the reviewer's name is required; every
        item that waited for a person must be decided; the version is one
        more analysis revision -- the analysis it was reviewed from, the
        versions before it and the source stay as they are -- and the
        registration event beside it holds who registered it, what it was
        reviewed from and every change with its original value. A version
        whose protocol would carry an execution blocker is not registered.
        The latest version is the one that runs (decision d).
        """

        if not isinstance(request, dict) or set(request) - {"reviewer_name", "target"}:
            raise ProtocolReviewError("A registration names the reviewer and the target.")
        name = reviewer_name(request.get("reviewer_name"))
        revision = self._latest_protocol_revision(protocol_id)
        target = self._review_target(revision)
        if target is None:
            raise ProtocolReviewNotReadyError("There is no analysis to review.")
        if request.get("target") is not None and not _same_target(request["target"], target.key):
            raise ProtocolReviewTargetChangedError(
                "The analysis under review changed; read the review again.")
        decisions, memos = self._review_decisions(revision, target)
        links, decisions = self._value_change_links(revision, target, decisions)
        items = self._review_items(revision, target, decisions, memos, links=links)
        pending = [
            item["item_id"] for item in items
            if item["needs_review"] and item["item_id"] not in decisions
        ]
        if pending:
            raise ProtocolReviewItemsPendingError(
                "Every item that waits for a person must be decided first.",
                item_ids=pending)
        held = _source_check_items(items)
        if held:
            raise ProtocolReviewSourceCheckError(
                "A safety number differs from the source; read the page again.",
                item_ids=held)
        protocol, changes, timers, origins = _apply_review(target, items, decisions, name)
        try:
            domain.validate_protocol(protocol)
        except domain.ProtocolValidationError as exc:
            raise ProtocolReviewNotReadyError(
                "The reviewed protocol does not hold against its source.") from exc
        readiness = domain.assess_readiness(protocol, capability_policy=domain.P1_CAPABILITY_POLICY)
        blockers = domain.execution_blocking_reasons(readiness)
        if blockers:
            raise ProtocolReviewBlockedError(
                "The reviewed protocol would not run.",
                codes=[reason.code.value for reason in blockers])
        version = len(self._review_versions(revision)) + 1
        analysis_id = f"{REVIEW_VERSION_ANALYSIS_PREFIX}{version}-{secrets.token_hex(12)}"
        analysis = self.store.append_analysis_revision(
            protocol_id, revision.revision_number, analysis_id, protocol, readiness,
            domain.P1_CAPABILITY_POLICY.profile_id,
        )
        revision_id = f"pdf-{revision.revision_number}-review-{version}"
        missed = self._version_possibly_missed(revision, target, protocol, items, decisions)
        payload: dict[str, object] = {
            "status": "review_registered",
            "version": version,
            "revision_id": revision_id,
            "reviewer_name": name,
            "base": target.key,
            "base_revision_id": target.revision_id,
            "base_analysis_id": target.analysis_id,
            "analysis_id": analysis_id,
            "analysis_payload_sha256": analysis.payload_sha256,
            "changes": [change.public_dict() for change in changes],
            "timers": timers,
            "origins": origins,
            "memos": {
                item_id: memo.get("memo") for item_id, memo in memos.items()
                if isinstance(memo.get("memo"), str)
            },
            "decision_event_ids": [str(decision["event_id"]) for decision in decisions.values()],
        }
        if missed:
            payload["possibly_missed"] = possibly_missed_summary(missed)
        event = self.store.append_event(
            f"review-version-{analysis_id}",
            protocol_id,
            revision.revision_number,
            _REVIEW_VERSION_EVENT,
            payload,
            analysis_revision_number=analysis.analysis_revision_number,
        )
        logging.getLogger(__name__).info(
            "protocol.review.registered protocol_id=%s version=%d changes=%d",
            protocol_id, version, len(changes),
        )
        if self.on_analysis_ready is not None:
            # The version's Korean, as for a passed analysis (lane PX, 1).
            try:
                self.on_analysis_ready(self, protocol_id)
            except Exception as exc:  # noqa: BLE001 - see on_analysis_ready
                logging.getLogger(__name__).warning(
                    "protocol.review.translation_start_failed error=%s", type(exc).__name__)
        return {
            "protocol": self.get_entry(protocol_id).public_dict(),
            "review_version": {
                **_review_version_summary(event),
                "changes": payload["changes"],
            },
        }

    def _version_possibly_missed(
        self, revision: ProtocolRevisionRecord, target: _ReviewTarget,
        protocol: domain.ExperimentProtocol, items: Sequence[dict[str, object]],
        decisions: dict[str, dict[str, object]],
    ) -> list[Any]:
        """The possibly-missed places a review version still tells before the
        start (lane MISS, decision 4): those of its protocol that no person
        answered on a card -- reviewed from a passed analysis, all of them;
        from an earlier version, those it still told."""

        places = _possibly_missed_answered(_within_page_range(
            possibly_missed_places(protocol, protocol.metadata.pdf),
            self._analysis_page_range(revision, target.analysis_id)), items, decisions)
        if target.kind == "review_version":
            versions = self._review_versions(revision)
            told = (versions[-1][0].payload or {}).get("possibly_missed") if versions else None
            seen = [place for place in (told or {}).get("places") or () if isinstance(place, dict)]
            places = [
                place for place in places
                if any(item.get("page") == place.page_number
                       and int(item.get("first_line") or 0) <= place.last_line
                       and int(item.get("last_line") or 0) >= place.first_line for item in seen)
            ]
        return places

    def review_version(self, protocol_id: str, version: int) -> dict[str, object]:
        """One registered version with its changes against what it was
        reviewed from (lane RV, decision c; the review tab's v1→v2)."""

        revision = self._latest_protocol_revision(protocol_id)
        for event, _ in self._review_versions(revision):
            if event.payload.get("version") == version:
                return {
                    **_review_version_summary(event, self.store.list_events(revision.experiment_id)),
                    "changes": list(event.payload.get("changes") or ()),
                    "memos": dict(event.payload.get("memos") or {}),
                    "timers": dict(event.payload.get("timers") or {}),
                }
        raise ProtocolCatalogNotFoundError("The protocol has no such review version.")

    def confirm_review_version(
        self, protocol_id: str, version: int, *, name: str, comment: str = "",
    ) -> dict[str, object]:
        """A PI's "확인함" on a registered review version (lane RP3, decision ④):
        one append-only ledger event with the name, the line and the time.
        It holds nothing back and changes no version; the version's record
        shows every confirmation left."""

        revision = self._latest_protocol_revision(protocol_id)
        if not any(event.payload.get("version") == version for event, _ in self._review_versions(revision)):
            raise ProtocolCatalogNotFoundError("The protocol has no such review version.")
        recorded_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
        self.store.append_event(
            "pi-confirmed-" + secrets.token_hex(12), protocol_id, revision.revision_number,
            _REVIEW_VERSION_PI_CONFIRMED_EVENT,
            {"version": version, "name": " ".join(str(name or "").split())[:120],
             "comment": " ".join(str(comment or "").split())[:300], "recorded_at": recorded_at},
        )
        return self.review_version(protocol_id, version)

    def resolve_asset(
        self,
        protocol_id: str,
        revision_id: str,
        asset_id: str,
    ) -> ProtocolAssetResolution:
        fixture = self.load_executable_fixture(protocol_id)
        if fixture.revision_id != revision_id:
            raise ProtocolCatalogNotFoundError("Protocol revision is unknown.")
        matching = tuple(
            asset
            for index in range(len(fixture.steps))
            if (asset := fixture.visual_for_step(index)) is not None
            and asset.asset_id == asset_id
        )
        if not matching:
            raise ProtocolCatalogNotFoundError("Protocol visual asset is unknown.")
        if fixture.source_pdf_path is None or fixture.source_pdf_sha256 is None:
            raise ProtocolCatalogUnavailableError(
                "Protocol visual source is unavailable."
            )
        return ProtocolAssetResolution(
            path=fixture.source_pdf_path,
            sha256=fixture.source_pdf_sha256,
            source_page=matching[0].source_page,
        )


# ---------------------------------------------------------------------------
# The review tab (lane RV, decisions b-d; human decisions 2026-10-12 and the
# review scope recommended 2026-10-14)
# ---------------------------------------------------------------------------
#
# A person reviews one analysis -- one waiting for review, one that passed, or
# the latest registered review version -- item by item: confirm it, fix it (a
# value, a unit, a typo), leave it empty because the source does not hold it,
# or enter a value the source does not print. Steps are not added or removed
# and a safety statement's words are not changed (a memo is); a value a
# person wrote is shown as "검토자 입력", one the analysis or the OCR's second
# reading produced as "AI 제안", and an AI suggestion becomes something that
# runs only when a person confirms it. Every decision is one more ledger
# event; the source and the analysis are never changed.

REVIEW_ORIGIN_SOURCE = "source"
REVIEW_ORIGIN_AI = "ai_suggestion"
REVIEW_ORIGIN_REVIEWER = "reviewer_input"
REVIEW_ORIGIN_KO: dict[str, str] = {
    REVIEW_ORIGIN_SOURCE: "원문",
    REVIEW_ORIGIN_AI: "AI 제안",
    REVIEW_ORIGIN_REVIEWER: "검토자 입력",
}
REVIEW_ITEM_KIND_KO: dict[str, str] = {
    "metadata": "문서 정보",
    "section": "절 제목",
    "step": "단계",
    "sub_action": "세부 동작",
    "warning": "안전 주의",
    "note": "참고",
    "tip": "팁",
    "expected_result": "완료 기준",
    "observation": "완료 기준(관찰)",
    "condition": "조건",
    "material": "재료",
    "equipment": "장비",
    "prerequisite": "시작 전 준비",
    "value": "값",
    "duration": "시간",
    "timer_spec": "타이머 정보",
    "reminder": "알림",
    "missing_value": "원문에 없는 값",
    "construct": "구조",
    "repeat_count": "반복 횟수",
    "timer": "단계 타이머",
    "pages": "분석 실패 쪽",
    "ocr_value": "OCR 값",
    "seam": "조각 경계",
    # Lane MISS, decision 3.
    "missed": "분석에 없는 원문",
}
#: The review's actions: confirm, fix, leave empty (원문에 없음), enter a value
#: the source does not print, a memo beside an item, and taking a draft
#: decision back.
REVIEW_ACTIONS: tuple[str, ...] = (
    "confirm", "fix", "clear", "enter", "memo", "reset", "choose", "source_differs")
REVIEW_ACTION_KO: dict[str, str] = {
    "confirm": "확인", "fix": "고침", "clear": "원문에 없음으로 비워 둠",
    "enter": "입력", "memo": "메모", "reset": "되돌리기",
    # Lane RV2, decision e: one of two pieces' readings of a passage.
    "choose": "두 읽기 중 고름",
    # Lane FX4, decision 5: a safety number no engine's reading gets right.
    "source_differs": "원문과 달라요(원문 확인 필요)",
}
#: Lane FX4, decision 5 (human decision 2026-10-19): a number in a safety
#: statement that differs from the source in every engine's reading. It stays
#: as read, marked for the source, and the version is not registered; the
#: page is photographed again more clearly or the document analysed again.
SOURCE_CHECK_STATE_KO = "원문 확인 필요"
SOURCE_CHECK_GUIDANCE_KO = "이 쪽을 더 선명하게 다시 찍어 올리거나 다시 분석해 주세요."
_STATEMENT_KINDS = {
    "warnings": "warning", "notes": "note", "tips": "tip",
    "expected_results": "expected_result", "conditions": "condition",
}
#: The text field a person reads (and may fix) on each record.
_REVIEW_TEXT_FIELD: dict[type, str] = {
    domain.ProtocolSection: "title_source_text",
    domain.ProtocolSourceStep: "instruction_source_text",
    domain.ProtocolSubAction: "instruction_source_text",
    domain.SourceStatement: "source_text",
    domain.RequiredObservation: "source_text",
    domain.Material: "name_source_text",
    domain.Equipment: "name_source_text",
    domain.BeforeStartPrerequisite: "source_text",
    domain.ScientificValue: "source_text",
    domain.EstimatedDuration: "source_text",
    domain.MissingExecutionValue: "description",
    domain.ProcessTimerSpecification: "timer_id",
    domain.OneTimeReminder: "message_source_text",
    domain.RecurringReminder: "message_source_text",
    domain.ConditionalBranch: "condition_source_text",
    domain.FixedRangeRepetition: "range_source_text",
    domain.OperatorDeterminedRepetition: "range_source_text",
    domain.RepeatUntil: "condition_source_text",
    domain.ParallelWork: "source_text",
    domain.RecurringAction: "source_text",
    domain.ReusableSubprocedure: "source_text",
    domain.SourceAmbiguity: "source_text",
    domain.ProtocolConflict: "source_text",
}
#: The kinds whose words a person may fix: values, units and typos in the
#: step sentences and the statements beside them (review scope, 2026-10-14).
#: A safety statement's words are not changed, nor a structure's.
_FIXABLE_KINDS = frozenset({
    "step", "sub_action", "note", "tip", "expected_result", "observation",
    "condition", "material", "equipment", "prerequisite", "value", "duration",
})
#: The kinds a person may leave empty because the source does not hold them --
#: only while flagged: a quote the server found is the source's own.
_CLEARABLE_KINDS = frozenset({
    "warning", "note", "tip", "expected_result", "observation", "condition",
    "material", "equipment", "prerequisite", "value", "duration", "reminder",
})
_MAX_REVIEW_TEXT = 2000
_MAX_REVIEW_VALUE = 200
_MAX_REVIEW_MEMO = 1000
_MAX_REVIEWER_NAME = 80
_MAX_TIMER_SECONDS = 7 * 24 * 3600
_MAX_REPEAT_COUNT = 999

#: What kind of quantity a value is, read from its unit (UI2's "수치 종류 이름").
_VALUE_KINDS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("온도", re.compile(r"°\s*[CF]\b|℃|(?<![A-Za-z])[0-9.]+\s*C\b|도\b", re.I)),
    ("원심·회전", re.compile(r"\brpm\b|[x×]\s*g\b|\bRCF\b", re.I)),
    ("시간", re.compile(
        r"\b(?:h|hr|hrs|hours?|min|mins|minutes?|s|sec|secs|seconds?|days?|overnight)\b"
        r"|분|시간|초|일\b|밤새|[0-9]{1,2}:[0-9]{2}", re.I)),
    ("부피", re.compile(r"(?:[µμu]|m|n)?L\b|(?:[µμu]|m)l\b|ml\b", re.I)),
    ("농도", re.compile(r"(?<![A-Za-z])(?:[µμunm]?M)\b|%|\bw/v\b|\bv/v\b|/\s*m?L\b|\bU/", re.I)),
    ("질량", re.compile(r"(?<![A-Za-z])(?:[µμunmk]?g)\b|\bkDa\b", re.I)),
    ("횟수", re.compile(r"\b(?:times?|cycles?|rounds?|washes)\b|회\b|번\b|[0-9]\s*[x×]\b", re.I)),
    ("길이", re.compile(r"(?<![A-Za-z])(?:[µμn]?m|mm|cm)\b", re.I)),
    ("전압·전류", re.compile(r"(?<![A-Za-z])(?:[mk]?V|[mµμ]?A)\b")),
)


def value_kind_ko(text: str | None) -> str | None:
    """"온도", "시간", "부피", ... for a value's words, or None when there is no number."""

    if not isinstance(text, str) or not re.search(r"[0-9]", text):
        return None
    for name, pattern in _VALUE_KINDS:
        if pattern.search(text):
            return name
    return "수치"


def _clean_review_text(value: object, *, limit: int, label: str) -> str:
    if not isinstance(value, str):
        raise ProtocolReviewError(f"{label} must be text.")
    cleaned = " ".join(value.split()) if "\n" not in value else "\n".join(
        " ".join(line.split()) for line in value.splitlines()).strip()
    if not cleaned or len(cleaned) > limit or any(
        ord(character) < 32 and character != "\n" for character in cleaned
    ):
        raise ProtocolReviewError(f"{label} is empty, too long or holds control characters.")
    return cleaned


def reviewer_name(value: object) -> str:
    """The reviewer's name as written, required (human decision 2026-10-12)."""

    if not isinstance(value, str) or not value.strip():
        raise ProtocolReviewNameRequiredError("A reviewer name is required.")
    return _clean_review_text(value, limit=_MAX_REVIEWER_NAME, label="Reviewer name")


def _duration_words(seconds: int) -> str:
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    parts = []
    if hours:
        parts.append(f"{hours}시간")
    if minutes:
        parts.append(f"{minutes}분")
    if secs or not parts:
        parts.append(f"{secs}초")
    return " ".join(parts)


@dataclass(frozen=True)
class _ReviewRecord:
    """One reviewable record found walking a protocol."""

    item_id: str
    kind: str
    record: Any
    path: str
    step: domain.ProtocolSourceStep | None
    section: domain.ProtocolSection | None
    evidence: domain.SourceEvidence | None
    evidence_item_id: str | None


def _review_records(protocol: domain.ExperimentProtocol) -> list[_ReviewRecord]:
    """Every record a person reviews, in document order."""

    found: list[_ReviewRecord] = []
    item_ids: dict[int, str] = {}

    def visit(value: Any, path: str, *, field_name: str, step, section, evidence, evidence_item):
        if isinstance(value, tuple):
            for index, item in enumerate(value):
                visit(item, f"{path}[{index}]", field_name=field_name, step=step,
                      section=section, evidence=evidence, evidence_item=evidence_item)
            return
        if not is_dataclass(value) or isinstance(
            value, (domain.SourceEvidence, ProtocolPdfExtraction, domain.DependencyTarget,
                    domain.ActualElapsedTime)
        ):
            return
        kind_name = REVIEW_ITEM_KINDS.get(type(value))
        if type(value) in _REVIEW_TEXT_FIELD:
            item_id = review_item_id(protocol, f"{path}.x")
            if isinstance(value, domain.ProtocolSourceStep):
                step = value
                kind = "step"
            elif isinstance(value, domain.ProtocolSection):
                section = value
                kind = "section"
            elif isinstance(value, domain.ProtocolSubAction):
                kind = "sub_action"
            elif isinstance(value, domain.SourceStatement):
                kind = _STATEMENT_KINDS.get(field_name, "note")
            elif isinstance(value, domain.ScientificValue):
                kind = "value"
            elif isinstance(value, domain.EstimatedDuration):
                kind = "duration"
            else:
                kind = kind_name[0] if kind_name else "construct"
            own = getattr(value, "evidence", None)
            if isinstance(own, domain.SourceEvidence):
                evidence, evidence_item = own, item_id
            found.append(_ReviewRecord(
                item_id, kind, value, path, step, section, evidence, evidence_item))
            if isinstance(value, domain.FixedRangeRepetition):
                found.append(_ReviewRecord(
                    f"{item_id}/repeat_count", "repeat_count", value, path, step, section,
                    evidence, evidence_item))
        for field in fields(value):
            visit(getattr(value, field.name), f"{path}.{field.name}", field_name=field.name,
                  step=step, section=section, evidence=evidence, evidence_item=evidence_item)

    for name in ("before_start", "materials", "equipment", "sections", "constructs"):
        visit(getattr(protocol, name), f"protocol.{name}", field_name=name, step=None,
              section=None, evidence=None, evidence_item=None)
    del item_ids
    return found


def _squeezed(text: str) -> str:
    return " ".join(text.split())


def _line_boxes(page: ProtocolPdfPage) -> list[list[float] | None]:
    """Each line's box on the page, in points, from the engine's text blocks
    (lane RX1, decision 2: the confirmation screen's highlight band).

    A block carries one box for all its lines; a line inside it takes an
    equal share of the block's height, top to bottom. Lines are matched to
    the blocks' lines by their words, in page order, so a line the blocks do
    not hold (an empty one, a merged one) is None. Geometry only; nothing
    here reads a word for its meaning.
    """

    block_lines: list[tuple[str, list[float]]] = []
    for block in page.blocks:
        # The engine's PdfTextBlock, or the plain tuple a stored analysis
        # keeps it as: (x0, y0, x1, y1, font size, bold, text).
        if not isinstance(block, (tuple, list)) or len(block) < 7:
            continue
        x0, y0, x1, y1, _, _, words = block[:7]
        try:
            x0, y0, x1, y1 = float(x0), float(y0), float(x1), float(y1)
        except (TypeError, ValueError):
            continue
        parts = [part for part in str(words or "").split("\n") if part.strip()]
        if not parts or y1 <= y0:
            continue
        height = (y1 - y0) / len(parts)
        for index, part in enumerate(parts):
            top = y0 + index * height
            block_lines.append((_squeezed(part), [
                round(x0, 2), round(top, 2), round(x1, 2), round(top + height, 2)]))
    boxes: list[list[float] | None] = []
    cursor = 0
    for _, _, text in _page_lines(page.text):
        wanted = _squeezed(text)
        found = None
        if wanted:
            for offset in range(cursor, len(block_lines)):
                words, box = block_lines[offset]
                if words == wanted or (len(wanted) >= 8 and (wanted in words or words in wanted)):
                    found = box
                    cursor = offset + 1
                    break
        boxes.append(found)
    return boxes


def _page_lines(text: str) -> list[tuple[int, int, str]]:
    """(start, end, text) of each line of a page, numbered from 1 by position."""

    lines: list[tuple[int, int, str]] = []
    start = 0
    for line in text.split("\n"):
        lines.append((start, start + len(line), line))
        start += len(line) + 1
    return lines


def _lines_of_span(text: str, start: int, end: int) -> tuple[int, int]:
    lines = _page_lines(text)
    first = next(index for index, (s, e, _) in enumerate(lines, 1) if e >= start)
    last = next(index for index, (s, e, _) in enumerate(lines, 1) if e >= max(start, end - 1))
    return first, last


def _quote_lines(evidence: domain.SourceEvidence, extraction: ProtocolPdfExtraction) -> tuple[int, int] | None:
    page_number = evidence.source_page_number
    if not isinstance(page_number, int) or not 1 <= page_number <= extraction.page_count:
        return None
    text = extraction.pages[page_number - 1].text
    found = text.find(evidence.source_excerpt)
    if found < 0:
        return None
    return _lines_of_span(text, found, found + len(evidence.source_excerpt))


#: Lane RV2, decision b: a suggested source line is the only candidate when
#: it is this alike to the item's words and no other place on the pages
#: looked at is within the margin of it. A person still presses [맞아요].
SUGGESTION_UNIQUE_MIN = 0.8
SUGGESTION_UNIQUE_MARGIN = 0.1


def _suggested_lines(
    words: str, extraction: ProtocolPdfExtraction, pages: Sequence[int],
) -> dict[str, object] | None:
    """The source lines most like ``words`` on these pages: an AI suggestion
    of where an item comes from, for a person to accept or change.

    ``unique`` (lane RV2, decision b) says the lines narrow to this one
    place: at least ``SUGGESTION_UNIQUE_MIN`` alike, and no place that does
    not overlap them within ``SUGGESTION_UNIQUE_MARGIN`` of it.
    """

    target = " ".join(words.split()).casefold()
    if not target:
        return None
    candidates: list[tuple[float, int, int, int]] = []
    for page_number in dict.fromkeys(pages):
        if not isinstance(page_number, int) or not 1 <= page_number <= extraction.page_count:
            continue
        lines = _page_lines(extraction.pages[page_number - 1].text)
        for first in range(len(lines)):
            joined = ""
            for last in range(first, min(first + 4, len(lines))):
                joined = f"{joined} {lines[last][2]}".strip()
                if not joined:
                    continue
                ratio = SequenceMatcher(None, target, " ".join(joined.split()).casefold(),
                                        autojunk=False).ratio()
                candidates.append((ratio, page_number, first + 1, last + 1))
    if not candidates:
        return None
    best = max(candidates, key=lambda candidate: candidate[0])
    if best[0] < 0.6:
        return None
    rivals = [
        candidate for candidate in candidates
        if candidate[0] >= best[0] - SUGGESTION_UNIQUE_MARGIN
        and (candidate[1] != best[1] or candidate[3] < best[2] or candidate[2] > best[3])
    ]
    return {"page": best[1], "first_line": best[2], "last_line": best[3],
            "similarity": round(best[0], 2),
            "unique": best[0] >= SUGGESTION_UNIQUE_MIN and not rivals}


def source_lines_excerpt(
    extraction: ProtocolPdfExtraction, source_lines: object,
) -> domain.SourceEvidence:
    """The quote a person pointed at: page and lines, as the page prints them."""

    if not isinstance(source_lines, dict):
        raise ProtocolReviewError("Source lines must name a page and its lines.")
    page = source_lines.get("page")
    first = source_lines.get("first_line")
    last = source_lines.get("last_line", first)
    if not all(isinstance(value, int) and not isinstance(value, bool)
               for value in (page, first, last)):
        raise ProtocolReviewError("Source lines must be numbers.")
    if not 1 <= page <= extraction.page_count:
        raise ProtocolReviewError("The source has no such page.")
    text = extraction.pages[page - 1].text
    lines = _page_lines(text)
    if not 1 <= first <= last <= len(lines) or last - first > 20:
        raise ProtocolReviewError("The page has no such lines.")
    start, end = lines[first - 1][0], lines[last - 1][1]
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    if start >= end:
        raise ProtocolReviewError("The lines pointed at hold no text.")
    return domain.SourceEvidence(source_page_number=page, source_excerpt=text[start:end])


class ProtocolReviewError(ProtocolCatalogError):
    """A review decision the server does not take (lane RV, decision b)."""

    code = "review_decision_invalid"


class ProtocolReviewNameRequiredError(ProtocolReviewError):
    code = "reviewer_name_required"


class ProtocolReviewNotReadyError(ProtocolCatalogError):
    """Nothing to review yet, or the items are not all decided (lane RV, c)."""

    code = "review_not_ready"


class ProtocolReviewItemsPendingError(ProtocolReviewNotReadyError):
    code = "review_items_pending"

    def __init__(self, message: str, *, item_ids: Sequence[str]) -> None:
        super().__init__(message)
        self.item_ids = tuple(item_ids)


class ProtocolReviewSourceCheckError(ProtocolReviewNotReadyError):
    """Lane FX4, decision 5: a safety number marked [원문과 달라요] holds the version."""

    code = "review_source_check_needed"

    def __init__(self, message: str, *, item_ids: Sequence[str]) -> None:
        super().__init__(message)
        self.item_ids = tuple(item_ids)


def _source_check_items(items: Sequence[dict[str, object]]) -> list[str]:
    """The items whose decision is [원문과 달라요] (lane FX4, decision 5)."""

    return [str(item["item_id"]) for item in items
            if (item.get("decision") or {}).get("action") == "source_differs"]


_REVIEW_DECISION_EVENT = "protocol_review_decision_recorded"


@dataclass(frozen=True)
class _ReviewTarget:
    """The analysis a person reviews now (lane RV, decision b)."""

    kind: str
    key: dict[str, object]
    protocol: domain.ExperimentProtocol
    readiness: domain.ReadinessAssessment
    flags: tuple[ReviewFlag, ...]
    recorded_at: str
    revision_id: str | None = None
    analysis_id: str | None = None
    #: Item id -> origin carried from the review version reviewed again.
    origins: dict[str, str] | None = None
    #: Step id -> the reviewer's timer seconds (None: no timer) of that version.
    timers: dict[str, int | None] | None = None
    #: Item id -> the card the re-ask closed (lane AN5, decision 3).
    closed: dict[str, dict[str, object]] | None = None


def _safe_page_view(
    protocol: domain.ExperimentProtocol, extraction: ProtocolPdfExtraction,
) -> domain.ExperimentProtocol:
    """``protocol`` with each quote that cites no page of the source pointed
    at page 1, so a step's timers can be read; for reading only."""

    replacements = {
        path: replace(evidence, source_page_number=1, continued_on_page_number=None,
                      continued_excerpt=None)
        for path, evidence in _evidence_paths(protocol)
        if not isinstance(evidence.source_page_number, int)
        or not 1 <= evidence.source_page_number <= extraction.page_count
        or (evidence.continued_on_page_number is not None
            and not 1 <= evidence.continued_on_page_number <= extraction.page_count)
    }
    return replace_at_paths(protocol, replacements) if replacements else protocol


def _evidence_paths(value: Any, path: str = "protocol"):
    if isinstance(value, domain.SourceEvidence):
        yield path, value
        return
    if isinstance(value, (ProtocolPdfExtraction, Enum, str)) or value is None:
        return
    if isinstance(value, tuple):
        for index, item in enumerate(value):
            yield from _evidence_paths(item, f"{path}[{index}]")
        return
    if is_dataclass(value):
        for field in fields(value):
            yield from _evidence_paths(getattr(value, field.name), f"{path}.{field.name}")


class ProtocolReviewTargetChangedError(ProtocolReviewError):
    code = "review_target_changed"


def _same_target(seen: object, key: dict[str, object]) -> bool:
    """Whether the target a screen read is the one under review now."""

    return isinstance(seen, dict) and all(seen.get(name) == value for name, value in key.items())


#: An OCR page whose reading a person should look at: the two engines read
#: different numbers or units, or the engine was not sure of the page.
OCR_UNCLEAR_CONFIDENCE = 0.8


def _ocr_page_unclear(page: ProtocolPdfPage, ocr: dict[str, object]) -> bool:
    if not page.ocr_derived:
        return False
    confidence = ocr.get("confidence")
    return bool(ocr.get("numeric_review_required")) or (
        isinstance(confidence, (int, float)) and not isinstance(confidence, bool)
        and confidence < OCR_UNCLEAR_CONFIDENCE
    )


# --- OCR values (lane RV2, decision b; lane AN3, decision A) ----------------
#
# A number OCR read is executed only when it is confirmed (human decisions
# 2026-10-16): both engines read the same number and unit in its place on
# the page (``protocol_ocr_providers.readings_by_place`` says what "in its
# place" is), and the item's quote passed the source check -- then it is
# settled without a person, the reason recorded. Otherwise the value is
# asked: the other engine read another number there, or none, or only one
# engine read the page. The values a page asks about are one card for the
# page, and that card counts once under the limit of items (decision A3).
#
# An OCR result stored before lane AN3 has no number-by-number reading; for
# it the comparison stays lane RV2's, by page: a page both engines read with
# the same numbers and units settles its values, any other page asks them.

#: Page states an OCR page can be in for its values (by page: lane RV2).
OCR_PAGE_AGREE = "agree"
OCR_PAGE_DIFFER = "differ"
OCR_PAGE_ONE_ENGINE = "one_engine"
#: Lane AN3, decision A4: the kinds whose numbers are used in execution --
#: the steps and their sub-actions, amounts and durations (with the step
#: timers read from them), the conditions and constructs that decide what
#: runs (branches, repeats), the reminders the run sets, the preparation
#: before the start and the completion criteria. Numbers anywhere else are
#: not asked: notes and tips (background), material and equipment names
#: (catalog, lot and model numbers), section headings and the document's
#: metadata, a value the source leaves out, and a timer spec and a repeat
#: count (their numbers are other records' words). Round 2, decision ③ (human
#: decision 2026-10-18): a safety statement's numbers are compared too; one
#: read differently takes one of the engines' readings, never a person's own
#: words (``_OCR_SAFETY_KIND``).
_OCR_VALUE_EXECUTION_KINDS = frozenset({
    "step", "sub_action", "value", "duration", "condition", "construct", "reminder",
    "prerequisite", "expected_result", "observation", "warning",
})
_OCR_SAFETY_KIND = "warning"
#: Constructs that describe the source rather than what runs.
_OCR_VALUE_NOT_EXECUTED = (domain.SourceAmbiguity, domain.ProtocolConflict)
_ENGINE_FAILED_WARNING = re.compile(r"^Page ([0-9]+): .+ failed;")
#: The engines' names as a card shows them.
_OCR_ENGINE_KO = {OCR_CLOVA: "CLOVA", OCR_GOOGLE: "Google"}
#: Lane AN3, decision A2: the card of one page's OCR values ("ocr_page:3"),
#: which [모두 맞아요] names as a whole.
OCR_PAGE_CARD_PREFIX = "ocr_page:"


def _asked_units(flags: Sequence[ReviewFlag]) -> dict[str, str]:
    """What a person is asked about, each with its first reason: an item, or
    one page's OCR values as one (lane AN3, decision A3: under the limit of
    items, a page's OCR card counts once)."""

    units: dict[str, str] = {}
    for flag in flags:
        key = (f"{OCR_PAGE_CARD_PREFIX}{flag.page_number}"
               if flag.reason == REVIEW_REASON_OCR_VALUE and flag.page_number is not None
               else flag.item_id)
        units.setdefault(key, flag.reason)
    return units


#: Lane AN3, round 2, decision ② (human decision 2026-10-18): a page whose
#: numbers the engines read differently, or not at all in place, at least this
#: often, or that either engine read below this page confidence, settles no
#: value by the engines' agreement. Live measurement (lane AN3, stage ②): the
#: four values both engines read alike and wrongly stood on pages where 17-34%
#: of the numbers differed and the confidences were 0.80-0.85; on the other 30
#: runs these lines cost no pass.
OCR_PAGE_MAX_DIFFERING_SHARE = 0.15
OCR_PAGE_MIN_CONFIDENCE = 0.90


def _ocr_page_gate(page: dict[str, object]) -> str | None:
    """Why a page compared place by place settles no value, or None."""

    readings = [reading for reading in page.get("value_readings") or ()
                if isinstance(reading, dict)]
    if readings and sum(
        1 for reading in readings if reading.get("other_token") != reading.get("token")
    ) >= OCR_PAGE_MAX_DIFFERING_SHARE * len(readings):
        return OCR_VALUE_PAGE_DIFFERS_WIDELY
    confidences = [value for value in (page.get("confidence"), page.get("other_confidence"))
                   if isinstance(value, (int, float)) and not isinstance(value, bool)]
    if confidences and min(confidences) < OCR_PAGE_MIN_CONFIDENCE:
        return OCR_VALUE_PAGE_LOW_CONFIDENCE
    return None


def _ocr_value_state(
    page_state: str, page: dict[str, object], token: str,
) -> tuple[str | None, list[dict[str, object]]]:
    """Whether one value on an OCR page is settled, and the engines' readings
    of it (lane AN3, decision A).

    Returns (None, readings) when both engines read ``token`` in every place
    the chosen engine read it; otherwise the reason code it is asked for and
    the readings to show: the chosen engine's, then the other engine's
    distinct numbers in those places. A page compared only as a whole (a
    result stored before lane AN3, or engines without word boxes) gives the
    page's state with no readings. A value read alike on a page read poorly
    (``_ocr_page_gate``, round 2, decision ②) is asked with the gate's reason.
    """

    readings = page.get("value_readings")
    if not isinstance(readings, list):
        if page_state == OCR_PAGE_AGREE:
            # Round 2, decision ②: the confidence line holds here too (the
            # result kept the chosen engine's page confidence only).
            return _ocr_page_gate(page), []
        return (OCR_VALUE_ENGINES_DIFFER if page_state == OCR_PAGE_DIFFER
                else OCR_VALUE_ONE_ENGINE), []
    places = [reading for reading in readings
              if isinstance(reading, dict) and reading.get("token") == token]
    if not places:
        return OCR_VALUE_NOT_PLACED, []
    chosen = _OCR_ENGINE_KO.get(str(page.get("provider")), str(page.get("provider")))
    other = _OCR_ENGINE_KO.get(str(page.get("other_provider")), str(page.get("other_provider")))
    shown = [{"engine": chosen, "text": places[0].get("text"), "token": token}]
    for place in places:
        other_token = place.get("other_token")
        if other_token is not None and other_token != token and other_token not in {
                item["token"] for item in shown}:
            shown.append({"engine": other, "text": place.get("other_text"), "token": other_token})
    if all(place.get("other_token") == token for place in places):
        # Lane AN3, round 2, decision ②: alike, but on a page read poorly.
        return _ocr_page_gate(page), shown
    if len(shown) > 1:
        return OCR_VALUE_READ_DIFFERENTLY, shown
    return OCR_VALUE_ONE_ENGINE_READ, shown


def _record_words(record: "_ReviewRecord") -> str | None:
    """The words of a reviewable record a person reads (and the run says)."""

    field = _REVIEW_TEXT_FIELD.get(type(record.record))
    words = getattr(record.record, field, None) if field else None
    return words if isinstance(words, str) else None


def _reading_lines(step: dict[str, object], extraction: ProtocolPdfExtraction) -> dict[str, int] | None:
    """The source lines one piece's reading of a step quotes (lane RV2, e)."""

    page, excerpt = step.get("page"), step.get("excerpt")
    if not isinstance(page, int) or not isinstance(excerpt, str):
        return None
    lines = _quote_lines(
        domain.SourceEvidence(source_page_number=page, source_excerpt=excerpt), extraction)
    return {"page": page, "first_line": lines[0], "last_line": lines[1]} if lines else None


def _same_scope(record: "_ReviewRecord", value: "_ReviewRecord") -> bool:
    """Whether ``record`` is one of the records a value's sentence lives in:
    its step's records, or, for a value outside any step, its owner."""

    if value.step is not None:
        return record.step is not None and record.step.step_id == value.step.step_id
    return value.path.startswith(record.path + ".")


def _ocr_value_spans(text: str) -> list[tuple[str, int, int]]:
    """Each value ``text`` states as (token, start, end) in ``text``; empty
    when the comparison form does not keep the text's positions."""

    normalized = unicodedata.normalize("NFKC", text).casefold().replace("μ", "µ")
    if len(normalized) != len(text):
        return []
    return [
        (match[1] + _OCR_UNIT_SPELLINGS.get(match[2] or "", match[2] or ""),
         match.start(), match.end())
        for match in _OCR_NUMBER_WITH_UNIT.finditer(normalized)
    ]


def _with_value_copied(text: str, token: str, reading: str) -> str:
    """``text`` with every place it states the value ``token`` written as
    ``reading`` (lane AN3, round 2, decision ③). Positions are kept through
    the comparison form character by character, so a text the form lengthens
    or shortens is still placed exactly."""

    owners: list[int] = []
    forms: list[str] = []
    for index, character in enumerate(text):
        form = unicodedata.normalize("NFKC", character).casefold().replace("μ", "µ")
        forms.append(form)
        owners.extend([index] * len(form))
    normalized = "".join(forms)
    spans = [
        (owners[match.start()], owners[match.end() - 1] + 1)
        for match in _OCR_NUMBER_WITH_UNIT.finditer(normalized)
        if match[1] + _OCR_UNIT_SPELLINGS.get(match[2] or "", match[2] or "") == token
    ]
    for start, end in reversed(spans):
        text = text[:start] + reading + text[end:]
    return text


_DURATION_WORDS = re.compile(
    r"^\s*([0-9]+(?:\.[0-9]+)?)\s*(s|sec|secs|seconds?|min|mins|minutes?|h|hr|hrs|hours?"
    r"|초|분|시간)\s*$", re.IGNORECASE)
_DURATION_UNIT_SECONDS = {"s": 1, "sec": 1, "secs": 1, "second": 1, "seconds": 1, "초": 1,
                          "min": 60, "mins": 60, "minute": 60, "minutes": 60, "분": 60,
                          "h": 3600, "hr": 3600, "hrs": 3600, "hour": 3600, "hours": 3600,
                          "시간": 3600}


def _duration_seconds(words: str) -> int | None:
    """Whole seconds of a value that is one duration ("10 min", "90 s"), or None."""

    match = _DURATION_WORDS.match(words)
    if match is None:
        return None
    seconds = float(match[1]) * _DURATION_UNIT_SECONDS[match[2].casefold()]
    return int(seconds) if seconds.is_integer() and 1 <= seconds <= _MAX_TIMER_SECONDS else None


def ocr_engine_names(provider: object) -> tuple[str, ...]:
    """The OCR engines a provider calls, as far as it says (lane RV2, b)."""

    names = getattr(provider, "engine_names", ())
    if not isinstance(names, (tuple, list)):
        return ()
    return tuple(name for name in names if name in {OCR_CLOVA, OCR_GOOGLE})


def pages_read_by_both(result: OcrResult, engines: Sequence[str]) -> list[int]:
    """The pages both engines read: OCR pages, with two engines configured,
    on which neither engine failed (the provider warns of each failure)."""

    if set(engines) != {OCR_CLOVA, OCR_GOOGLE}:
        return []
    failed = {
        int(match[1]) for warning in result.warnings
        if (match := _ENGINE_FAILED_WARNING.match(warning)) is not None
    }
    return [
        page.source_page_number for page in result.pages
        if page.provider in {OCR_CLOVA, OCR_GOOGLE}
        and page.source_page_number not in failed
    ]


def ocr_value_tokens(text: str, *, label: str | None = None) -> list[tuple[str, str]]:
    """The values ``text`` states, as (token, words): the token the two OCR
    engines were compared by ("4000rpm", "37°c"), and the words as written.

    A step's own number opening its sentence ("3.", "(2)", or its label) is
    not a value.
    """

    normalized = unicodedata.normalize("NFKC", text).casefold().replace("μ", "µ")
    shown = text if len(normalized) == len(text) else normalized
    start = 0
    opening = _LEADING_LABEL.match(normalized)
    if opening is not None:
        start = opening.end()
    elif label and label.strip():
        own = re.match(rf"\s*\(?{re.escape(label.strip().casefold())}[.)]?\s+", normalized)
        if own is not None:
            start = own.end()
    found: list[tuple[str, str]] = []
    for match in _OCR_NUMBER_WITH_UNIT.finditer(normalized, start):
        unit = match[2] or ""
        found.append((match[1] + _OCR_UNIT_SPELLINGS.get(unit, unit),
                      shown[match.start():match.end()]))
    return found


# --- Cards (lane RV2, decision a; human decision 2026-10-16) ----------------
#
# A person is asked item by item, one card at a time, by buttons: the
# server builds each card's question and the only choices the screen may
# draw for it (each a review decision the server takes, or "다시 분석"),
# so the screen decides nothing. An analysis whose every executed item
# passed has no card and runs as before.


def _choice(choice_id: str, label_ko: str, request: dict[str, object] | None, *,
            needs: str | None = None, route: str | None = None,
            primary: bool = False) -> dict[str, object]:
    """One button of a card: the decision it sends, and what it still needs
    from the person (lines pointed at, a text, seconds)."""

    return {"choice_id": choice_id, "label_ko": label_ko, "request": request,
            "needs": needs, "route": route, "primary": primary}


def _subject(word: object) -> str:
    """``word`` with the subject particle its last syllable takes: "단계가", "값이"."""

    text = str(word)
    last = text[-1] if text else ""
    if "가" <= last <= "힣" and (ord(last) - ord("가")) % 28:
        return f"{text}이"
    return f"{text}가"


def _card_choices(
    item: dict[str, object],
) -> tuple[str, str, list[dict[str, object]], dict[str, object]]:
    """The card's type, its question, its choices and what else it shows."""

    item_id = str(item["item_id"])
    kind = str(item["kind"])
    allowed = list(item["allowed_actions"])
    source = item.get("source") if isinstance(item.get("source"), dict) else None
    reasons = [reason for reason in item.get("reasons") or () if isinstance(reason, dict)]
    request = lambda action, **extra: {"item_id": item_id, "action": action, **extra}
    choices: list[dict[str, object]] = []
    changed = [
        reason["extra"] for reason in reasons
        if reason.get("reason") == REVIEW_REASON_VALUE_CHANGED and isinstance(reason.get("extra"), dict)
    ]
    if changed:
        # Decision d: a value a person fixed is stated here too; the server's
        # sentence with the new value, or a person writes it.
        link = changed[0]
        candidate, seconds = link.get("candidate"), link.get("candidate_seconds")
        lines = source is not None and source.get("verified") is False
        change = {"from": link.get("from"), "to": link.get("to"), "source_item": link.get("source_item")}
        question = (f"{link.get('from')} → {link.get('to')}(으)로 고쳤어요. "
                    f"이 {item['kind_ko']}에도 그 값이 있어요. 이렇게 바꿀까요?")
        if kind == "timer":
            if isinstance(seconds, int):
                choices.append(_choice("fix_proposed", "이렇게 바꿔요",
                                       request("fix", seconds=seconds), primary=True))
            choices.append(_choice("fix", "다시 고칠게요", request("fix"), needs="seconds"))
            return "value_changed", question, choices, {
                "change": change, "proposal": seconds and _duration_words(seconds)}
        if isinstance(candidate, str):
            choices.append(_choice("fix_proposed", "이렇게 바꿔요", request("fix", text=candidate),
                                   needs="source_lines" if lines else None, primary=True))
        choices.append(_choice("fix", "다시 고칠게요", request("fix"),
                               needs="source_lines_and_text" if lines else "text"))
        return "value_changed", question, choices, {"change": change, "proposal": candidate}
    two_ways = next((reason for reason in reasons
                     if reason.get("reason") == REVIEW_REASON_READ_TWO_WAYS), None)
    if two_ways is not None:
        # Decision e: two pieces read this passage differently; a person
        # picks one reading, with the page's lines beside both.
        page = two_ways.get("page_number")
        return "two_readings", f"두 조각이 {page}쪽을 다르게 읽었어요. 어느 쪽이 맞나요?", [
            _choice("choose_a", "이쪽으로 읽기", request("choose", reading="a"), primary=True),
            _choice("choose_b", "저쪽으로 읽기", request("choose", reading="b")),
            _choice("fix", "직접 고칠게요", request("fix"), needs="text"),
        ], {"readings": item.get("readings")}
    if any(reason.get("reason") == REVIEW_REASON_REFERENCE_MOVED for reason in reasons):
        # Lane AN3, round 2, decision ⑤: the join pointed a construct at the
        # section its step stands in; a person says it holds, or the document
        # is read again.
        return "reference_moved", (
            "조각을 합친 뒤 이 구조가 가리키는 단계가 다른 절에 있어요. "
            "원문대로 이 단계를 가리키면 [맞아요], 아니면 다시 분석해요."), [
            _choice("confirm", "맞아요", request("confirm"), primary=True),
            _choice("reanalysis", "다시 분석", None, route="reanalysis"),
        ], {}
    if kind == "pages":
        blocked = any(reason.get("reason") == REVIEW_REASON_BLOCKED for reason in reasons)
        return "blocked_pages" if blocked else "failed_pages", (
            f"{reasons[0].get('reason_ko') if reasons else '분석하지 못한 쪽'}. "
            "이 쪽에 실험할 단계가 없나요?"), [
            _choice("confirm", "이 쪽엔 단계 없음", request("confirm"), primary=True),
            _choice("reanalysis", "다시 분석", None, route="reanalysis"),
        ], {}
    if kind == "missed":
        # Lane MISS, decision 3 (human decision 2026-10-23): lines no quote
        # covers that read like an instruction. A person says they are no
        # step, or the document is read again; no button puts them in as a
        # step, since that would be a person writing values and safety words.
        first, last = item.get("first_page"), item.get("last_page")
        pages = (f"{first}쪽" if first == last else f"{first}–{last}쪽")
        lines = "이 줄들이" if len(item.get("places") or ()) > 1 or "\n" in str(
            item.get("text") or "") else "이 줄이"
        return "possibly_missed", f"{pages}의 {lines} 실험할 단계인가요?", [
            _choice("confirm", "단계가 아니에요", request("confirm")),
            _choice("reanalysis", "다시 분석할게요", None, route="reanalysis"),
        ], {}
    if kind == "seam":
        # Lane AN3, decision B (human decision 2026-10-16, "순서 확인"): where
        # the numbering changes between pieces, does the second step follow
        # the first? [아니요] does not reorder anything: the document is read
        # again.
        return "numbering_seam", (
            "번호를 붙인 부분과 붙이지 않은 부분이 만나는 곳이에요. "
            "앞 단계 다음에 뒤 단계가 이어지나요?"), [
            _choice("confirm", "이어져요", request("confirm"), primary=True),
            _choice("reanalysis", "아니요", None, route="reanalysis"),
        ], {"steps": item.get("steps")}
    if kind == "ocr_value":
        # Decision b: a number OCR read that the two engines did not both
        # read alike; the image piece is beside it on the card.
        value = item.get("text")
        return "ocr_value", f"원문 그림에서 이 값이 {value} 맞나요?", [
            _choice("confirm", f"{value} 맞아요", request("confirm"), primary=True),
            _choice("fix", "직접 입력", request("fix"), needs="text"),
        ], {"proposal": value}
    quote_refused = source is not None and source.get("verified") is False
    if any(reason.get("reason") == REVIEW_REASON_ADDRESSED_TO_PROGRAM for reason in reasons):
        # Lane TIDY, item 3 (fake-user test F6): words that read as written to
        # the program, not as an experiment's. [빼 주세요] takes them out of
        # the steps; [실험 지시예요] keeps them (pointing at the source lines
        # when the quote was not found). A safety statement is never taken out
        # (AGENTS.md rule 5): it is the source's, or the document is read again.
        keep_needs = "source_lines" if quote_refused else None
        if kind == "warning":
            choices = [_choice("confirm", "맞아요", request("confirm"), needs=keep_needs,
                               primary=True)]
            if "clear" in allowed:
                choices.append(_choice("clear", "원문에 없어요", request("clear")))
            choices.append(_choice("reanalysis", "다시 분석할게요", None, route="reanalysis"))
            return "addressed_to_program", (
                "이 안전 주의는 프로그램에 하는 말처럼 보여요. 원문의 안전 주의가 맞나요?"), choices, {}
        choices = []
        if "clear" in allowed:
            choices.append(_choice("clear", "빼 주세요", request("clear"), primary=True))
        choices.append(_choice("confirm", "실험 지시예요", request("confirm"), needs=keep_needs))
        return "addressed_to_program", (
            "이 문장은 실험 지시가 아니라 프로그램에 하는 말처럼 보여요. 단계에서 뺄까요?"), choices, {}
    suggested = source.get("suggested_lines") if source is not None else None
    if quote_refused and isinstance(suggested, dict) and suggested.get("unique"):
        # Decision b: the lines narrow to one place; the server proposes it
        # and a person only says it is right.
        lines = {key: suggested[key] for key in ("page", "first_line", "last_line")}
        choices.append(_choice(
            "confirm_suggested", "맞아요", request("confirm", source_lines=lines), primary=True))
        choices.append(_choice(
            "confirm_lines", "다른 줄이에요", request("confirm"), needs="source_lines"))
        question = (f"이 {item['kind_ko']}의 근거는 원문 {lines['page']}쪽 "
                    f"{lines['first_line']}번째 줄"
                    + (f"부터 {lines['last_line']}번째 줄" if lines["last_line"] != lines["first_line"]
                       else "") + "인가요?")
        card_type = "evidence_suggested"
    elif quote_refused:
        choices.append(_choice(
            "confirm_lines", "고른 줄이 맞아요", request("confirm"), needs="source_lines",
            primary=True))
        question = f"이 {item['kind_ko']}의 근거를 원문에서 찾지 못했어요. 원문의 어느 줄인가요?"
        card_type = "evidence"
    else:
        if "confirm" in allowed:
            choices.append(_choice("confirm", "맞아요", request("confirm"), primary=True))
        question = f"이 {_subject(item['kind_ko'])} 원문과 같나요?"
        card_type = "claim"
    if "fix" in allowed and card_type != "evidence_suggested":
        choices.append(_choice(
            "fix", "직접 고칠게요", request("fix"),
            needs="source_lines_and_text" if quote_refused else "text"))
    if "clear" in allowed:
        choices.append(_choice("clear", "원문에 없어요", request("clear")))
    elif kind == "step":
        # A step is never left empty (lane RV): one the source does not hold
        # is read again.
        choices.append(_choice("reanalysis", "다시 분석", None, route="reanalysis"))
    return card_type, question, choices, {}


def _ocr_value_entry(item: dict[str, object]) -> dict[str, object]:
    """One value of a page's OCR card: what each engine read in its place and
    its buttons -- each engine's reading (the analysis's own one confirms it,
    another one fixes it to that reading) and [직접 입력]."""

    item_id = str(item["item_id"])
    request = lambda action, **extra: {"item_id": item_id, "action": action, **extra}
    readings = [reading for reading in item.get("readings") or () if isinstance(reading, dict)]
    choices: list[dict[str, object]] = []
    for index, reading in enumerate(readings):
        text = str(reading.get("text"))
        same = reading.get("token") == item.get("token")
        choices.append({
            **_choice(f"reading_{index}", text,
                      request("confirm") if same else request("fix", text=text),
                      primary=index == 0),
            "engine": reading.get("engine"),
        })
    if not readings:
        choices.append(_choice("confirm", f"{item.get('text')} 맞아요", request("confirm"),
                               primary=True))
    if item.get("in_safety_statement") is not True:
        # Round 2, decision ③: a safety statement's number is one of the
        # engines' readings, never words of a person's own.
        choices.append(_choice("fix", "직접 입력", request("fix"), needs="text"))
    elif "source_differs" in item.get("allowed_actions", ()):
        # Lane FX4, decision 5: when no reading is the source's, the value
        # waits for the source (원문 확인 필요) and the version is not registered.
        choices.append(_choice("source_differs", "원문과 달라요", request("source_differs")))
    source = item.get("source") if isinstance(item.get("source"), dict) else {}
    reason = next((reason for reason in item.get("reasons") or ()
                   if isinstance(reason, dict) and reason.get("reason") == REVIEW_REASON_OCR_VALUE),
                  {})
    return {
        "item_id": item_id,
        "token": item.get("token"),
        "shown": item.get("text"),
        "reason_code": reason.get("reason_code"),
        "reason_ko": reason.get("reason_ko"),
        "readings": readings,
        "linked_items": item.get("linked_items"),
        "value_boxes": source.get("value_boxes"),
        "lines": source.get("lines"),
        "choices": choices,
        "decided": item.get("decision") is not None,
        "decision": item.get("decision"),
        "current": item.get("current"),
        **({"source_check": item["source_check"]} if item.get("source_check") else {}),
    }


def _ocr_page_card(values: Sequence[dict[str, object]]) -> dict[str, object]:
    """One page's OCR values on one card (lane AN3, decision A2): the page's
    piece (the values' boxes on it), the values with what each engine read,
    [모두 맞아요] for the page, and each value's own buttons."""

    page = values[0].get("page")
    card_id = str(values[0].get("page_card"))
    entries = [_ocr_value_entry(value) for value in values]
    boxes = list(dict.fromkeys(
        index for entry in entries for index in entry["value_boxes"] or ()))
    lines = sorted({line for value in values
                    for line in ((value.get("source") or {}).get("value_lines") or ())})
    return {
        "card_id": card_id,
        "item_id": card_id,
        "type": "ocr_page",
        "kind": "ocr_page",
        "kind_ko": REVIEW_ITEM_KIND_KO["ocr_value"],
        "step_label": None,
        "page": page,
        "question_ko": (
            f"원문 {page}쪽 그림의 값 {len(entries)}개를 확인해 주세요. "
            "두 OCR 엔진이 다르게 읽었거나 한 엔진만 읽은 값이에요."),
        # Each value says its own reason; the card does not repeat them (lane
        # AN3, stage ③).
        "reasons_ko": [],
        "shown": None,
        "source": {
            "page": page, "excerpt": None, "continued_page": None, "continued_excerpt": None,
            "verified": True,
            "lines": ({"page": page, "first_line": lines[0], "last_line": lines[-1]}
                      if lines else None),
            "value_lines": lines,
            "value_boxes": boxes,
            "suggested_lines": None,
        },
        "values": entries,
        "choices": [_choice("confirm_all", "모두 맞아요",
                            {"item_id": card_id, "action": "confirm"}, primary=True)],
        "proposal": None,
        "change": None,
        "readings": None,
        "steps": None,
        "memo_allowed": False,
        "decided": all(entry["decided"] for entry in entries),
        "decision": None,
        "current": None,
    }


def review_cards(items: Sequence[dict[str, object]]) -> list[dict[str, object]]:
    """One card per item waiting for a person, in document order.

    Each card says what is asked, shows the item as the analysis read it
    and where it stands in the source, and lists the buttons the screen may
    draw -- each the review decision it sends. Decided cards stay in place
    (marked), so "← 이전" can go back to them.
    """

    # Lane AN3, decision A2: a page's OCR values are asked on one card, in
    # the place of the first of them.
    asked: list[dict[str, object] | list[dict[str, object]]] = []
    pages: dict[str, list[dict[str, object]]] = {}
    for item in items:
        if not item.get("needs_review") and not item.get("closed"):
            continue
        page_card = item.get("page_card") if item["kind"] == "ocr_value" else None
        if isinstance(page_card, str):
            if page_card not in pages:
                pages[page_card] = []
                asked.append(pages[page_card])
            pages[page_card].append(item)
            continue
        asked.append(item)
    cards = []
    for position, item in enumerate(asked, start=1):
        if isinstance(item, list):
            cards.append({
                **_ocr_page_card(item),
                "position": position, "total": len(asked),
                "progress_ko": f"{position} / {len(asked)}",
            })
            continue
        source = item.get("source") if isinstance(item.get("source"), dict) else None
        closed = item.get("closed") if not item.get("needs_review") else None
        if isinstance(closed, dict):
            # Lane AN5, decision 3: closed by the re-ask -- in its place, marked
            # decided, nothing to press; a person's own decision on it shows.
            cards.append({
                "card_id": item["item_id"], "item_id": item["item_id"],
                "position": position, "total": len(asked),
                "progress_ko": f"{position} / {len(asked)}",
                "type": "requoted", "kind": item["kind"], "kind_ko": item["kind_ko"],
                "step_label": item.get("step_label"),
                "question_ko": REQUOTE_CLOSED_KO,
                "reasons_ko": [], "shown": item.get("text"), "source": source, "choices": [],
                "proposal": None, "change": None, "readings": None, "steps": None,
                "memo_allowed": "memo" in item["allowed_actions"],
                "decided": True,
                "decision": item.get("decision") or {
                    "action": "requoted", "value": None, "source_lines": None,
                    "reviewer_name": None, "recorded_at": closed.get("recorded_at")},
                "current": item.get("current"),
                "closed_ko": closed.get("label_ko"),
            })
            continue
        card_type, question, choices, shown = _card_choices(item)
        cards.append({
            "card_id": item["item_id"],
            "item_id": item["item_id"],
            "position": position,
            "total": len(asked),
            "progress_ko": f"{position} / {len(asked)}",
            "type": card_type,
            "kind": item["kind"],
            "kind_ko": item["kind_ko"],
            "step_label": item.get("step_label"),
            "question_ko": question,
            "reasons_ko": [reason.get("reason_ko") for reason in item.get("reasons") or ()
                           if isinstance(reason, dict)],
            "shown": item.get("text"),
            "source": source,
            "choices": choices,
            "proposal": shown.get("proposal"),
            "change": shown.get("change"),
            "readings": shown.get("readings"),
            "steps": shown.get("steps"),
            "memo_allowed": "memo" in item["allowed_actions"],
            "decided": item.get("decision") is not None,
            "decision": item.get("decision"),
            "current": item.get("current"),
        })
    return cards


def _lines_asked_about(draft: ProtocolAnalysisDraft) -> dict[int, set[int]]:
    """Lane MISS, decision 3: the source lines another card asks about -- a
    failed piece's pages, and where the server suggests a refused quote's
    item stands (``_suggested_lines``) -- so they are not asked about twice."""

    asked: dict[int, set[int]] = {}
    for flag in draft.review_flags:
        if flag.field == "pages" and flag.first_page is not None:
            # A failed or blocked piece's pages are asked about whole on
            # their own card ("이 쪽엔 단계 없음").
            for page in range(flag.first_page, (flag.last_page or flag.first_page) + 1):
                if 1 <= page <= draft.extraction.page_count:
                    asked.setdefault(page, set()).update(range(
                        1, draft.extraction.pages[page - 1].text.count("\n") + 2))
            continue
        if flag.field != "evidence" or not flag.path:
            continue
        try:
            evidence = _value_at_path(draft.protocol, flag.path)
        except (AttributeError, IndexError, AssertionError):
            continue
        if not isinstance(evidence, domain.SourceEvidence):
            continue
        suggested = _suggested_lines(
            str(evidence.source_excerpt or ""), draft.extraction,
            (evidence.source_page_number, *flag.matching_source_pages))
        if suggested is not None:
            asked.setdefault(int(suggested["page"]), set()).update(
                range(int(suggested["first_line"]), int(suggested["last_line"]) + 1))
    return asked


#: Lane MISS, decision 4: the code of the start notice for lines no quote
#: covers. A notice; never in ``EXECUTION_BLOCKING_REASON_CODES``.
POSSIBLY_MISSED_NOTICE_CODE = domain.ReadinessReasonCode.SOURCE_TEXT_POSSIBLY_MISSED.value


def _pages_ko(pages: Sequence[int]) -> str:
    return "·".join(str(page) for page in pages)


def _ready_possibly_missed(events: Sequence[Any], analysis: Any) -> object:
    """What the ledger recorded beside this passed analysis (lane MISS, 4)."""

    number = getattr(analysis, "analysis_revision_number", None)
    for event in reversed(events):
        if (event.event_type == _ANALYSIS_READY_EVENT and number is not None
                and event.analysis_revision_number == number
                and isinstance(event.payload, dict)):
            return event.payload.get("possibly_missed")
    return None


def possibly_missed_notices(summary: object) -> list[dict[str, object]]:
    """The start notice of a passed analysis's possibly-missed places (lane
    MISS, decision 4): one line, "분석에 없을 수 있는 원문이 3곳 있어요(2·5쪽)",
    with the places for the screen; nothing for no places."""

    if not isinstance(summary, dict) or not summary.get("count"):
        return []
    count = int(summary["count"])
    pages = [int(page) for page in summary.get("pages") or () if isinstance(page, int)]
    return [{
        "code": POSSIBLY_MISSED_NOTICE_CODE,
        "message": (
            f"{count} source passage(s) on page(s) {', '.join(map(str, pages))} read like "
            "instructions and no analysis quote covers them."),
        "message_ko": (
            f"분석에 없을 수 있는 원문이 {count}곳 있어요({_pages_ko(pages)}쪽). "
            "시작 전에 원문을 확인해 주세요."),
        "kind": "source_note",
        "source_page_number": pages[0] if pages else None,
        "source_excerpt": None,
        "step_id": None,
        "pages": pages,
        "places": [dict(place) for place in summary.get("places") or () if isinstance(place, dict)],
    }]


def _possibly_missed_answered(
    places: Sequence[Any], items: Sequence[dict[str, object]],
    decisions: dict[str, dict[str, object]],
) -> list[Any]:
    """The places a person has not answered on a card (lane MISS): a place
    lies on an answered card when the card's places hold its lines, or on a
    failed piece's pages a person said hold no step."""

    answered = [
        place for item in items if item["kind"] == "missed" and item["item_id"] in decisions
        for place in item.get("places") or () if isinstance(place, dict)
    ]
    no_step_pages = {
        page for item in items if item["kind"] == "pages" and item["item_id"] in decisions
        and isinstance(item.get("first_page"), int)
        for page in range(int(item["first_page"]), int(item.get("last_page") or item["first_page"]) + 1)
    }
    return [
        place for place in places
        if place.page_number not in no_step_pages and not any(
            seen.get("page") == place.page_number
            and int(seen.get("first_line") or 0) <= place.last_line
            and int(seen.get("last_line") or 0) >= place.first_line
            for seen in answered)
    ]


@dataclass(frozen=True)
class ReviewedChange:
    """One item a reviewer changed or confirmed in a review version (lane RV, c).

    ``action`` is ``fix``, ``enter``, ``clear`` or ``confirm`` (an item that
    waited for a person, confirmed); ``original`` is the value before -- the
    source's or the AI's reading -- and ``reviewed`` the value the version
    holds (None when left empty). ``origin`` is what the version shows the
    value as: ``reviewer_input`` ("검토자 입력") or ``ai_suggestion`` ("AI 제안",
    confirmed by a person). ``spoken`` is the value words said at the step the
    first time it is read (decision g), None for a change nothing speaks.
    """

    item_id: str
    kind: str
    kind_ko: str
    step_id: str | None
    step_label: str | None
    action: str
    original: str | None
    reviewed: str | None
    origin: str
    spoken: str | None = None
    decided_by: str | None = None
    decided_at: str | None = None

    def public_dict(self) -> dict[str, object]:
        return {
            "item_id": self.item_id, "kind": self.kind, "kind_ko": self.kind_ko,
            "step_id": self.step_id, "step_label": self.step_label,
            "action": self.action, "original": self.original, "reviewed": self.reviewed,
            "origin": self.origin, "spoken": self.spoken,
            "decided_by": self.decided_by, "decided_at": self.decided_at,
        }

    @classmethod
    def from_public(cls, payload: dict[str, object]) -> "ReviewedChange":
        text = lambda key: payload.get(key) if isinstance(payload.get(key), str) else None
        return cls(
            item_id=str(payload["item_id"]), kind=str(payload["kind"]),
            kind_ko=str(payload.get("kind_ko") or payload["kind"]),
            step_id=text("step_id"), step_label=text("step_label"),
            action=str(payload["action"]), original=text("original"),
            reviewed=text("reviewed"), origin=str(payload["origin"]),
            spoken=text("spoken"), decided_by=text("decided_by"),
            decided_at=text("decided_at"),
        )


@dataclass(frozen=True)
class ReviewVersionInfo:
    """A registered review version, as the run and its report read it (lane RV, c)."""

    version: int
    revision_id: str
    reviewer_name: str
    registered_at: str
    changes: tuple[ReviewedChange, ...] = ()
    base_revision_id: str | None = None
    #: The registration's date in Korea ("2026-10-14").
    registered_on: str = ""

    def spoken_for_step(self, step_id: str) -> tuple[str, ...]:
        """The reviewer's values said at this step, in the order they were decided."""

        return tuple(dict.fromkeys(
            change.spoken for change in self.changes
            if change.step_id == step_id and change.spoken
            and change.origin == "reviewer_input"
        ))


@dataclass(frozen=True)
class ReviewedAnalysisDraft(ProtocolAnalysisDraft):
    """The draft of a registered review version (lane RV, decision c): the
    analysis draft with the version it is and what its reviewer changed, for
    the run (decision g's notice) and its report to read."""

    review_version: ReviewVersionInfo | None = None


def _with_reviewer_timers(
    table: domain.StepTimerTable,
    protocol: domain.ExperimentProtocol,
    timers: dict[str, object],
) -> domain.StepTimerTable:
    """``table`` with a review version's timers in place of the steps' own.

    A step a reviewer gave a time runs that time, said as the reviewer's
    (``source="reviewer_input"``); a step whose timer the reviewer left out
    has none. The source's refused times at those steps are dropped: the
    reviewer has decided them.
    """

    steps = {step.step_id: step for section in protocol.sections for step in section.steps}
    decided = {step_id for step_id in timers if step_id in steps}
    verified = [timer for timer in table.verified if timer.step_id not in decided]
    for step_id in decided:
        seconds = timers[step_id]
        if isinstance(seconds, int) and not isinstance(seconds, bool) and seconds > 0:
            step = steps[step_id]
            words = _duration_words(seconds)
            verified.append(domain.VerifiedStepTimer(
                step_id, step.source_label, None, f"검토자 입력: {words}", words,
                (seconds,), step.evidence.source_page_number, source=REVIEW_ORIGIN_REVIEWER,
            ))
    order = {step_id: index for index, step_id in enumerate(steps)}
    verified.sort(key=lambda timer: order.get(timer.step_id, len(order)))
    return domain.StepTimerTable(
        tuple(verified),
        tuple(time for time in table.refused if time.step_id not in decided),
    )


def _version_payload(event: Any) -> dict[str, object]:
    """A registration event's payload with the time it was recorded."""

    return {
        **(event.payload if isinstance(event.payload, dict) else {}),
        "registered_at": event.recorded_at,
        "registered_on": _korean_date(event.recorded_at),
    }


def _review_version_info(payload: dict[str, object]) -> ReviewVersionInfo:
    return ReviewVersionInfo(
        version=int(payload["version"]),
        revision_id=str(payload["revision_id"]),
        reviewer_name=str(payload["reviewer_name"]),
        registered_at=str(payload.get("registered_at") or ""),
        registered_on=str(payload.get("registered_on") or ""),
        changes=tuple(
            ReviewedChange.from_public(item) for item in payload.get("changes") or ()
            if isinstance(item, dict)),
        base_revision_id=(
            str(payload["base_revision_id"]) if payload.get("base_revision_id") else None),
    )


def _path_depth(path: str) -> int:
    return path.count(".") + path.count("[")


def _removed_at(protocol: domain.ExperimentProtocol, path: str) -> domain.ExperimentProtocol:
    """``protocol`` without the record at ``path``: out of its tuple, or None."""

    match = re.fullmatch(r"(.+)\[([0-9]+)\]", path)
    if match is not None:
        parent, index = match[1], int(match[2])
        items = _value_at_path(protocol, parent)
        return replace_at_paths(protocol, {parent: items[:index] + items[index + 1:]})
    return replace_at_paths(protocol, {path: None})


def _value_at_path(protocol: domain.ExperimentProtocol, path: str) -> Any:
    current: Any = protocol
    for part in path.split(".")[1:]:
        match = re.fullmatch(r"([A-Za-z_][A-Za-z0-9_]*)(?:\[([0-9]+)\])?", part)
        assert match is not None
        current = getattr(current, match[1])
        if match[2] is not None:
            current = current[int(match[2])]
    return current


def _shifted_item_ids(
    origins: dict[str, str], removed: Sequence[str],
) -> dict[str, str]:
    """Origins keyed by item id after values were taken out of their tuples.

    A value is named by its place ("action:a1/quantities/2"); once the value
    before it is left empty it is the second ("…/1").
    """

    gone: dict[str, set[int]] = {}
    for item_id in removed:
        head, _, index = item_id.rpartition("/")
        if index.isdigit():
            gone.setdefault(head, set()).add(int(index))
    shifted: dict[str, str] = {}
    for item_id, origin in origins.items():
        head, _, index = item_id.rpartition("/")
        if index.isdigit() and head in gone:
            number = int(index)
            if number in gone[head]:
                continue
            item_id = f"{head}/{number - sum(1 for value in gone[head] if value < number)}"
        elif item_id in removed:
            continue
        shifted[item_id] = origin
    return shifted


def _other_reading(
    protocol: domain.ExperimentProtocol, item: dict[str, object],
) -> domain.ExperimentProtocol:
    """``protocol`` with the later piece's reading in place of the kept one
    (lane RV2, decision e): the kept steps are replaced, where the first of
    them stood, by the other piece's steps, the first taking its id."""

    flag = next((reason for reason in item.get("reasons") or ()
                 if isinstance(reason, dict)
                 and reason.get("reason") == REVIEW_REASON_READ_TWO_WAYS), None)
    extra = (flag or {}).get("extra") or {}
    kept = [str(step_id) for step_id in extra.get("kept") or ()]
    others = [_decode_stored_record(step) for step in extra.get("other_steps") or ()]
    if not kept or not others or not all(
            isinstance(step, domain.ProtocolSourceStep) for step in others):
        raise ProtocolReviewNotReadyError("The other reading of this passage is unreadable.")
    others[0] = replace(others[0], step_id=kept[0])
    sections = []
    for section in protocol.sections:
        steps: list[domain.ProtocolSourceStep] = []
        for step in section.steps:
            if step.step_id == kept[0]:
                steps.extend(others)
            elif step.step_id not in kept:
                steps.append(step)
        sections.append(replace(section, steps=tuple(steps)))
    return replace(protocol, sections=tuple(sections))


#: A step sentence's own number ("3.", "(2)") before its words.
_LEADING_LABEL = re.compile(r"^\s*\(?[0-9]+(?:\.[0-9]+)*[.)]\s+")
#: A number with the unit printed after it, as a step sentence states a value.
_STATED_VALUE = re.compile(
    r"[0-9]+(?:[.,][0-9]+)*(?:\s*(?:°\s*[CF]|℃|C\b|%|[µμunm]?[LlMg]\b|rpm\b|[x×]\s*g\b|"
    r"min\b|h\b|hr\b|hrs\b|s\b|sec\b|분|시간|초|회|번))?")


def _numbers_of(text: str | None) -> list[str]:
    """The values a sentence states, its own step number left out."""

    return [" ".join(value.split()) for value in _STATED_VALUE.findall(
        _LEADING_LABEL.sub("", text or "", count=1))]


#: The kinds the analysis may not word for the run (human decision
#: 2026-10-12): a safety statement and a completion criterion, confirmed,
#: run in the source's own words.
_SOURCE_WORDS_ONLY = frozenset({"warning", "expected_result", "observation"})


class ProtocolReviewBlockedError(ProtocolReviewNotReadyError):
    code = "review_version_blocked"

    def __init__(self, message: str, *, codes: Sequence[str]) -> None:
        super().__init__(message)
        self.codes = tuple(codes)


def _apply_review(
    target: "_ReviewTarget",
    items: list[dict[str, object]],
    decisions: dict[str, dict[str, object]],
    name: str,
) -> tuple[domain.ExperimentProtocol, list[ReviewedChange], dict[str, int | None], dict[str, str]]:
    """The protocol a review version holds, its changes, timers and origins.

    Applied to the target's protocol, deepest record first so a record's
    change and its value's change both hold: a fix writes the reviewer's
    words, a confirmed or fixed item whose quote was not found takes the
    lines a person pointed at as its quote, and an item left empty leaves
    the protocol. Timers, a repeat count and a value the source does not
    print are the reviewer's own values beside it.
    """

    protocol = target.protocol
    by_id = {item["item_id"]: item for item in items}
    records = {record.item_id: record for record in _review_records(protocol)}
    edits: list[tuple[str, Callable[[Any], Any]]] = []
    removals: list[tuple[str, str]] = []
    changes: list[ReviewedChange] = []
    timers: dict[str, int | None] = dict(target.timers or {})
    origins: dict[str, str] = dict(target.origins or {})
    confirmed_owners: set[str] = set()
    for item_id, decision in decisions.items():
        item = by_id.get(item_id)
        if item is None:
            continue
        action = str(decision.get("action"))
        kind = str(item["kind"])
        value = decision.get("value") if isinstance(decision.get("value"), dict) else {}
        original = item.get("text")
        reviewed: object = original
        spoken: str | None = None
        if kind == "timer":
            step_id = item_id[len("timer:"):]
            timers[step_id] = int(value["seconds"]) if action in {"fix", "enter"} else None
            reviewed = value.get("text") if action in {"fix", "enter"} else None
            spoken = value.get("text") if action in {"fix", "enter"} else None
        elif kind == "missed":
            # Lane MISS, decision 3: lines a person said are no step change
            # nothing in the version; the decision events record the answer
            # (``decision_event_ids``).
            continue
        elif kind in {"pages", "ocr_value", "seam"}:
            # A check, not a record: a fixed OCR value changes the items that
            # state it through their own decisions (decision d); a seam a
            # person said follows on runs in the joined order (lane AN3, B).
            pass
        elif kind == "missing_value" and action == "enter":
            reviewed = spoken = value.get("text")
        else:
            record = records.get(item_id)
            path = "protocol.metadata" if kind == "metadata" else record.path if record else None
            if path is None:
                continue
            if action == "clear":
                reviewed = None
                if kind == "repeat_count":
                    edits.append((path, lambda current: replace(current, repeat_count=None)))
                else:
                    removals.append((path, item_id))
            elif action in {"fix", "enter"}:
                if kind == "repeat_count":
                    count = int(value["count"])
                    edits.append((path, lambda current, count=count: replace(
                        current, repeat_count=count,
                        repeat_count_kind=domain.REPEAT_COUNT_TOTAL)))
                    reviewed = spoken = value.get("text")
                else:
                    text = str(value["text"])
                    reviewed = text

                    def fixed(current, text=text):
                        changed: dict[str, object] = {_REVIEW_TEXT_FIELD[type(current)]: text}
                        if isinstance(current, domain.ScientificValue):
                            changed.update(parsed_value=None, normalized_unit=None)
                        if isinstance(current, domain.EstimatedDuration):
                            changed["parsed_seconds"] = None
                        return replace(current, **changed)

                    edits.append((path, fixed))
                    if kind in {"step", "sub_action"}:
                        before = _numbers_of(str(original or ""))
                        spoken = " · ".join(
                            number for number in _numbers_of(text) if number not in before
                        ) or None
            lines = decision.get("source_lines")
            if isinstance(lines, dict) and action in {"confirm", "fix"}:
                quote = domain.SourceEvidence(
                    source_page_number=int(lines["page"]),
                    source_excerpt=str(lines["excerpt"]),
                )
                edits.append((path, lambda current, quote=quote: replace(current, evidence=quote)))
                confirmed_owners.add(item_id)
            if action == "confirm" and item["needs_review"] and kind in _SOURCE_WORDS_ONLY:
                # "안전 주의와 완료 기준은 AI 가 새로 써 넣지 않는다": confirmed,
                # a safety statement or a completion criterion the source
                # did not print as read runs in the source's own words -- the
                # lines a person pointed at, or the quote the server found.
                words = (
                    str(lines["excerpt"]) if isinstance(lines, dict)
                    else str((item.get("source") or {}).get("excerpt") or original))
                field = _REVIEW_TEXT_FIELD[type(record.record)] if record else None
                if field is not None:
                    edits.append((path, lambda current, field=field, words=words: replace(
                        current, **{field: words})))
                    reviewed = words
        if action == "choose":
            # Lane RV2, decision e: the reading a person chose.
            reviewed = value.get("text")
        if action == "confirm" and not item["needs_review"]:
            continue
        only_addressed = all(
            isinstance(reason, dict) and reason.get("reason") == REVIEW_REASON_ADDRESSED_TO_PROGRAM
            for reason in item.get("reasons") or ())
        origin = (
            REVIEW_ORIGIN_REVIEWER if action in {"fix", "enter"}
            else REVIEW_ORIGIN_SOURCE if kind in _SOURCE_WORDS_ONLY and action == "confirm"
            # Lane TIDY, item 3: kept as the source's own words.
            else str(item["origin"]) if item["needs_review"] and only_addressed
            else REVIEW_ORIGIN_AI if item["needs_review"] else str(item["origin"])
        )
        if action in {"fix", "enter", "choose"} or (
                action == "confirm" and item["needs_review"]):
            origins[item_id] = origin
        changes.append(ReviewedChange(
            item_id=item_id, kind=kind, kind_ko=str(item["kind_ko"]),
            step_id=item.get("step_id"), step_label=item.get("step_label"),
            action=action,
            original=None if original is None else str(original),
            reviewed=None if reviewed is None else str(reviewed),
            origin=origin, spoken=spoken,
            decided_by=str(decision.get("reviewer_name") or name),
            decided_at=str(decision.get("recorded_at") or ""),
        ))
    for item_id, decision in decisions.items():
        # Lane AN3, round 2, decision ③ (human decision 2026-10-18): a number of
        # a safety statement the two engines read differently takes the reading
        # a person chose -- the source's characters as that engine read them,
        # copied in its place; the statement is not otherwise edited.
        item = by_id.get(item_id)
        value = decision.get("value") if isinstance(decision.get("value"), dict) else {}
        if (item is None or item["kind"] != "ocr_value" or item.get("in_safety_statement") is not True
                or decision.get("action") != "fix" or not isinstance(value.get("text"), str)):
            continue
        for linked in item.get("linked_items") or ():
            record = records.get(linked)
            if record is None or record.kind != _OCR_SAFETY_KIND:
                continue
            words = _record_words(record) or ""
            copied = _with_value_copied(words, str(item.get("token") or ""), value["text"])
            if copied == words:
                continue
            edits.append((record.path, lambda current, copied=copied: replace(
                current, source_text=copied)))
            shown = by_id.get(linked) or {}
            changes.append(ReviewedChange(
                item_id=linked, kind=record.kind, kind_ko=str(shown.get("kind_ko") or record.kind),
                step_id=shown.get("step_id"), step_label=shown.get("step_label"),
                action="fix", original=words, reviewed=copied, origin=REVIEW_ORIGIN_SOURCE,
                decided_by=str(decision.get("reviewer_name") or name),
                decided_at=str(decision.get("recorded_at") or ""),
            ))
    for item in items:
        if item.get("covered_by") in confirmed_owners:
            origins.setdefault(str(item["item_id"]), REVIEW_ORIGIN_AI)
    for path, edit in sorted(edits, key=lambda pair: _path_depth(pair[0]), reverse=True):
        protocol = replace_at_paths(protocol, {path: edit(_value_at_path(protocol, path))})
    for path, _ in sorted(
        removals, key=lambda pair: (_path_depth(pair[0]), pair[0]), reverse=True,
    ):
        protocol = _removed_at(protocol, path)
    for item_id, decision in decisions.items():
        value = decision.get("value") if isinstance(decision.get("value"), dict) else {}
        if decision.get("action") == "choose" and value.get("reading") == "b" and item_id in by_id:
            protocol = _other_reading(protocol, by_id[item_id])
    origins = _shifted_item_ids(origins, [item_id for _, item_id in removals])
    changes.sort(key=lambda change: list(by_id).index(change.item_id))
    return protocol, changes, timers, origins


# ---------------------------------------------------------------------------
# Reagent tables (lane AN, FM b; human decision 2026-10-13)
# ---------------------------------------------------------------------------
#
# An Excel workbook is attached to a protocol as its "시약 표": every sheet
# and every non-empty cell kept where the workbook prints it (sheet name, cell
# reference, value, number format), so a question about an amount can be
# answered with the table position and the cell value. It is not a protocol
# source: it is never analysed, never advances a step, and the protocol's
# readiness does not depend on it. Connecting it to voice is a later lane.

REAGENT_TABLE_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
_REAGENT_TABLE_EVENT = "protocol_reagent_table_attached"
#: The largest workbook read: a reagent list is small, and openpyxl holds the
#: whole workbook in memory to read its merged cells.
MAX_REAGENT_TABLE_BYTES = 5 * 1024 * 1024
#: The most non-empty cells kept; a larger workbook is refused, not cut.
MAX_REAGENT_TABLE_CELLS = 20_000
_SAFE_TABLE_FILENAME = re.compile(r"^[^/\\\x00]{1,255}\.xlsx$", re.IGNORECASE)


class ReagentTableError(ProtocolCatalogError):
    code = "invalid_reagent_table"


class ReagentTableTooLargeError(ReagentTableError):
    code = "reagent_table_too_large"


def _cell_value(value: object) -> object:
    """A cell value as JSON keeps it: numbers and text as they are."""

    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if value == value and value not in (float("inf"), float("-inf")) else str(value)
    isoformat = getattr(value, "isoformat", None)
    return isoformat() if callable(isoformat) else str(value)


def _cell_display(value: object, number_format: str | None) -> str:
    """How the cell reads: the number format's percent and decimals applied.

    Only those two; any other format reads as the plain value, and the
    format itself is kept beside it.
    """

    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return "" if value is None else str(value)
    fmt = (number_format or "General").split(";", 1)[0]
    decimals_match = re.search(r"0\.(0+)", fmt)
    decimals = len(decimals_match[1]) if decimals_match else None
    if "%" in fmt:
        return f"{value * 100:.{decimals or 0}f}%"
    if decimals is not None:
        return f"{value:.{decimals}f}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_reagent_workbook(source: str | Path) -> dict[str, object]:
    """Every sheet and non-empty cell of an .xlsx workbook, where it prints them.

    A formula cell keeps the value the workbook saved for it; one saved
    without a value (written by a program that does not calculate) is kept
    with its formula and no value, and counted in ``warnings``.
    """

    try:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:  # pragma: no cover - openpyxl is a dependency
        raise ReagentTableError("openpyxl is not installed.") from exc
    import zipfile

    path = Path(source)
    if path.stat().st_size > MAX_REAGENT_TABLE_BYTES:
        raise ReagentTableTooLargeError("The workbook is larger than a reagent table reads.")
    try:
        values = openpyxl.load_workbook(path, data_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ReagentTableError("The file is not an .xlsx workbook.") from exc
    sheets: list[dict[str, object]] = []
    count = 0
    unsaved = 0
    for sheet in values.worksheets:
        written = formulas[sheet.title]
        cells: list[dict[str, object]] = []
        for row in sheet.iter_rows():
            for cell in row:
                formula = written[cell.coordinate].value
                has_formula = isinstance(formula, str) and formula.startswith("=")
                if cell.value is None and not has_formula:
                    continue
                count += 1
                if count > MAX_REAGENT_TABLE_CELLS:
                    raise ReagentTableTooLargeError(
                        "The workbook has more cells than a reagent table keeps.")
                value = _cell_value(cell.value)
                item: dict[str, object] = {
                    "ref": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "value": value,
                    "display": _cell_display(value, cell.number_format),
                }
                if cell.number_format and cell.number_format != "General":
                    item["number_format"] = cell.number_format
                if has_formula:
                    item["formula"] = formula
                    if cell.value is None:
                        unsaved += 1
                cells.append(item)
        sheets.append({
            "name": sheet.title,
            "state": sheet.sheet_state,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged": sorted(str(merged) for merged in sheet.merged_cells.ranges),
            "cells": cells,
        })
    warnings = []
    if unsaved:
        warnings.append(
            f"계산값이 저장되지 않은 수식 셀 {unsaved}개 — 엑셀에서 열어 저장한 뒤 다시 올리면 값이 보입니다."
        )
    return {"sheets": sheets, "cell_count": count, "warnings": warnings}


def _table_key(text: object) -> str:
    """Text as a lookup compares it: NFKC, case and spacing/punctuation aside."""

    import unicodedata

    folded = unicodedata.normalize("NFKC", str(text)).casefold()
    return "".join(character for character in folded if character.isalnum())


def find_reagent_cells(
    tables: Sequence[dict[str, object]], query: str, *, limit: int = 20,
) -> list[dict[str, object]]:
    """The cells whose text holds ``query``, each with its row and header row.

    Matching sets case, spacing and punctuation aside ("tris hcl" finds
    "Tris-HCl"). The header row is the sheet's first row of two or more
    cells -- a title line above the table ("S1 Table. Key resources") is
    not it -- or its first row when no row has two. The answer carries
    positions and values only -- what a person or the voice lane reads from
    them is theirs.
    """

    wanted = _table_key(query)
    if not wanted:
        return []
    found: list[dict[str, object]] = []
    for table in tables:
        for sheet in table.get("sheets", []) if isinstance(table, dict) else []:
            cells = [cell for cell in sheet.get("cells", []) if isinstance(cell, dict)]
            if not cells:
                continue
            by_row: dict[int, list[dict[str, object]]] = {}
            for cell in cells:
                by_row.setdefault(cell["row"], []).append(cell)
            header_row = min(
                (row for row, items in by_row.items() if len(items) >= 2),
                default=min(by_row),
            )
            header = {cell["column"]: cell for cell in by_row.get(header_row, [])}
            for cell in cells:
                if wanted not in _table_key(cell.get("display") or cell.get("value") or ""):
                    continue
                found.append({
                    "table_id": table.get("table_id"),
                    "filename": table.get("filename"),
                    "sheet": sheet.get("name"),
                    "ref": cell["ref"],
                    "value": cell.get("value"),
                    "display": cell.get("display"),
                    "row": [
                        {
                            "ref": item["ref"],
                            "display": item.get("display"),
                            "header": (
                                header[item["column"]].get("display")
                                if item["column"] in header and cell["row"] > header_row
                                else None
                            ),
                        }
                        for item in sorted(by_row[cell["row"]], key=lambda item: item["column"])
                    ],
                })
                if len(found) >= limit:
                    return found
    return found


def reagent_table_summary(payload: dict[str, object]) -> dict[str, object]:
    return {
        key: payload.get(key)
        for key in ("table_id", "filename", "sha256", "byte_size", "cell_count", "warnings")
    } | {
        "sheets": [
            {"name": sheet.get("name"), "cells": len(sheet.get("cells", []))}
            for sheet in payload.get("sheets", []) if isinstance(sheet, dict)
        ],
    }
